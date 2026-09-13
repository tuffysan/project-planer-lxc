from flask import Flask, render_template, request, redirect, url_for, flash, send_file, session, abort, jsonify
import sqlite3, os, json, secrets, hmac, hashlib, shutil, platform, re
from pathlib import Path
from datetime import datetime, date, timedelta
from io import BytesIO
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Border, Side, Protection
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.table import Table, TableStyleInfo

APP_VERSION = "15.2.9"
BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "projectplan.db"
SECRET_FILE = DATA_DIR / ".secret_key"

app = Flask(__name__)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("PROJECT_PLAN_HTTPS", "0") == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(minutes=int(os.getenv("PROJECT_PLAN_SESSION_MINUTES", "60"))),
)

def get_secret():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    env = os.getenv("PROJECT_PLAN_SECRET")
    if env:
        return env
    if SECRET_FILE.exists():
        return SECRET_FILE.read_text(encoding="utf-8").strip()
    value = secrets.token_hex(32)
    SECRET_FILE.write_text(value, encoding="utf-8")
    return value

app.secret_key = get_secret()

STATUSES = ["Ej startad", "Pågår", "Blockerad", "Klar", "Pausad"]
PRIORITIES = ["Låg", "Normal", "Hög", "Kritisk"]
RISK_STATUSES = ["Öppen", "Bevakas", "Åtgärdas", "Stängd"]
ROLES = ["admin", "pm", "member", "viewer"]

PROJECT_ROLES = ["pm", "member", "viewer"]

LANGUAGES = {
    "sv": {"name": "Svenska", "flag": "SV"},
    "en": {"name": "English", "flag": "EN"},
    "de": {"name": "Deutsch", "flag": "DE"},
    "no": {"name": "Norsk", "flag": "NO"},
    "da": {"name": "Dansk", "flag": "DA"},
    "fi": {"name": "Suomi", "flag": "FI"},
}

TRANSLATIONS = {
    "sv": {
        "dashboard":"Dashboard","admin":"Admin","password":"Lösenord","logout":"Logga ut","api":"API",
        "notifications":"Notiser","pmo":"PMO","resources":"Resurser","my_work":"Mina uppgifter",
        "language":"Språk","projects":"Projekt","new_project":"Nytt projekt","search":"Sök",
        "welcome":"Välkommen tillbaka!","active_projects":"Aktiva projekt","activities":"Aktiviteter",
        "milestones":"Milstolpar","risks":"Risker","project_status":"Projektstatus","upcoming_milestones":"Kommande milstolpar",
        "gantt":"Gantt","kanban":"Kanban","calendar":"Kalender","time_reporting":"Tidrapportering",
        "documents":"Dokument","meetings":"Möten","reports":"Rapporter","settings":"Inställningar",
        "users":"Användare","system":"System","open":"Öppna","save":"Spara","cancel":"Avbryt",
        "login":"Logga in","username":"Användarnamn","display_name":"Visningsnamn","current_password":"Nuvarande lösenord",
        "new_password":"Nytt lösenord","change_password":"Byt lösenord","project_portfolio":"Mina projekt",
        "only_assigned":"Du ser bara projekt där du är medlem.","admin_all_projects":"Administratörsvy: alla projekt.",
        "brand_tagline":"PLAN · GENOMFÖR · LYCKAS"
    },
    "en": {
        "dashboard":"Dashboard","admin":"Admin","password":"Password","logout":"Sign out","api":"API",
        "notifications":"Notifications","pmo":"PMO","resources":"Resources","my_work":"My Work",
        "language":"Language","projects":"Projects","new_project":"New project","search":"Search",
        "welcome":"Welcome back!","active_projects":"Active projects","activities":"Activities",
        "milestones":"Milestones","risks":"Risks","project_status":"Project status","upcoming_milestones":"Upcoming milestones",
        "gantt":"Gantt","kanban":"Kanban","calendar":"Calendar","time_reporting":"Time reporting",
        "documents":"Documents","meetings":"Meetings","reports":"Reports","settings":"Settings",
        "users":"Users","system":"System","open":"Open","save":"Save","cancel":"Cancel",
        "login":"Sign in","username":"Username","display_name":"Display name","current_password":"Current password",
        "new_password":"New password","change_password":"Change password","project_portfolio":"My projects",
        "only_assigned":"You only see projects where you are a member.","admin_all_projects":"Administrator view: all projects.",
        "brand_tagline":"PLAN · EXECUTE · SUCCEED"
    },
    "de": {
        "dashboard":"Dashboard","admin":"Admin","password":"Passwort","logout":"Abmelden","api":"API",
        "notifications":"Benachrichtigungen","pmo":"PMO","resources":"Ressourcen","my_work":"Meine Aufgaben",
        "language":"Sprache","projects":"Projekte","new_project":"Neues Projekt","search":"Suchen",
        "welcome":"Willkommen zurück!","active_projects":"Aktive Projekte","activities":"Aufgaben",
        "milestones":"Meilensteine","risks":"Risiken","project_status":"Projektstatus","upcoming_milestones":"Kommende Meilensteine",
        "gantt":"Gantt","kanban":"Kanban","calendar":"Kalender","time_reporting":"Zeiterfassung",
        "documents":"Dokumente","meetings":"Meetings","reports":"Berichte","settings":"Einstellungen",
        "users":"Benutzer","system":"System","open":"Öffnen","save":"Speichern","cancel":"Abbrechen",
        "login":"Anmelden","username":"Benutzername","display_name":"Anzeigename","current_password":"Aktuelles Passwort",
        "new_password":"Neues Passwort","change_password":"Passwort ändern","project_portfolio":"Meine Projekte",
        "only_assigned":"Sie sehen nur Projekte, denen Sie zugewiesen sind.","admin_all_projects":"Administratoransicht: alle Projekte.",
        "brand_tagline":"PLANEN · UMSETZEN · ERFOLG"
    },
    "no": {
        "dashboard":"Dashboard","admin":"Admin","password":"Passord","logout":"Logg ut","api":"API",
        "notifications":"Varsler","pmo":"PMO","resources":"Ressurser","my_work":"Mine oppgaver",
        "language":"Språk","projects":"Prosjekter","new_project":"Nytt prosjekt","search":"Søk",
        "welcome":"Velkommen tilbake!","active_projects":"Aktive prosjekter","activities":"Aktiviteter",
        "milestones":"Milepæler","risks":"Risikoer","project_status":"Prosjektstatus","upcoming_milestones":"Kommende milepæler",
        "gantt":"Gantt","kanban":"Kanban","calendar":"Kalender","time_reporting":"Timeregistrering",
        "documents":"Dokumenter","meetings":"Møter","reports":"Rapporter","settings":"Innstillinger",
        "users":"Brukere","system":"System","open":"Åpne","save":"Lagre","cancel":"Avbryt",
        "login":"Logg inn","username":"Brukernavn","display_name":"Visningsnavn","current_password":"Nåværende passord",
        "new_password":"Nytt passord","change_password":"Endre passord","project_portfolio":"Mine prosjekter",
        "only_assigned":"Du ser bare prosjekter du er medlem av.","admin_all_projects":"Administratorvisning: alle prosjekter.",
        "brand_tagline":"PLANLEGG · GJENNOMFØR · LYKKES"
    },
    "da": {
        "dashboard":"Dashboard","admin":"Admin","password":"Adgangskode","logout":"Log ud","api":"API",
        "notifications":"Notifikationer","pmo":"PMO","resources":"Ressourcer","my_work":"Mine opgaver",
        "language":"Sprog","projects":"Projekter","new_project":"Nyt projekt","search":"Søg",
        "welcome":"Velkommen tilbage!","active_projects":"Aktive projekter","activities":"Aktiviteter",
        "milestones":"Milepæle","risks":"Risici","project_status":"Projektstatus","upcoming_milestones":"Kommende milepæle",
        "gantt":"Gantt","kanban":"Kanban","calendar":"Kalender","time_reporting":"Tidsregistrering",
        "documents":"Dokumenter","meetings":"Møder","reports":"Rapporter","settings":"Indstillinger",
        "users":"Brugere","system":"System","open":"Åbn","save":"Gem","cancel":"Annuller",
        "login":"Log ind","username":"Brugernavn","display_name":"Visningsnavn","current_password":"Nuværende adgangskode",
        "new_password":"Ny adgangskode","change_password":"Skift adgangskode","project_portfolio":"Mine projekter",
        "only_assigned":"Du ser kun projekter, hvor du er medlem.","admin_all_projects":"Administratorvisning: alle projekter.",
        "brand_tagline":"PLANLÆG · UDFØR · LYKKES"
    },
    "fi": {
        "dashboard":"Kojelauta","admin":"Ylläpito","password":"Salasana","logout":"Kirjaudu ulos","api":"API",
        "notifications":"Ilmoitukset","pmo":"PMO","resources":"Resurssit","my_work":"Omat tehtävät",
        "language":"Kieli","projects":"Projektit","new_project":"Uusi projekti","search":"Haku",
        "welcome":"Tervetuloa takaisin!","active_projects":"Aktiiviset projektit","activities":"Tehtävät",
        "milestones":"Välitavoitteet","risks":"Riskit","project_status":"Projektin tila","upcoming_milestones":"Tulevat välitavoitteet",
        "gantt":"Gantt","kanban":"Kanban","calendar":"Kalenteri","time_reporting":"Työajanseuranta",
        "documents":"Dokumentit","meetings":"Kokoukset","reports":"Raportit","settings":"Asetukset",
        "users":"Käyttäjät","system":"Järjestelmä","open":"Avaa","save":"Tallenna","cancel":"Peruuta",
        "login":"Kirjaudu sisään","username":"Käyttäjänimi","display_name":"Näyttönimi","current_password":"Nykyinen salasana",
        "new_password":"Uusi salasana","change_password":"Vaihda salasana","project_portfolio":"Omat projektit",
        "only_assigned":"Näet vain projektit, joissa olet jäsenenä.","admin_all_projects":"Ylläpitäjän näkymä: kaikki projektit.",
        "brand_tagline":"SUUNNITTELE · TOTEUTA · ONNISTU"
    },
}


def db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA busy_timeout=5000")
    try:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA synchronous=NORMAL")
    except sqlite3.OperationalError:
        pass
    return conn

def ensure_column(conn, table, column, definition):
    cols = {r["name"] for r in conn.execute(f"PRAGMA table_info({table})")}
    if column not in cols:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")

def init_db():
    with db() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS schema_migrations (
            version TEXT PRIMARY KEY,
            applied_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            display_name TEXT NOT NULL DEFAULT '',
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'member',
            active INTEGER NOT NULL DEFAULT 1,
            force_password_change INTEGER NOT NULL DEFAULT 0,
            failed_logins INTEGER NOT NULL DEFAULT 0,
            locked_until TEXT DEFAULT '',
            last_login TEXT DEFAULT '',
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS projects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            customer TEXT DEFAULT '',
            project_manager TEXT DEFAULT '',
            description TEXT DEFAULT '',
            start_date TEXT DEFAULT '',
            end_date TEXT DEFAULT '',
            template_name TEXT DEFAULT '',
            created_by INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL
        );

        CREATE TABLE IF NOT EXISTS project_members (
            project_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            project_role TEXT NOT NULL DEFAULT 'viewer',
            added_at TEXT NOT NULL DEFAULT '',
            added_by INTEGER,
            PRIMARY KEY(project_id, user_id),
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
            FOREIGN KEY(added_by) REFERENCES users(id) ON DELETE SET NULL
        );

        CREATE TABLE IF NOT EXISTS tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            wbs TEXT DEFAULT '',
            title TEXT NOT NULL,
            owner TEXT DEFAULT '',
            start_date TEXT DEFAULT '',
            end_date TEXT DEFAULT '',
            actual_start TEXT DEFAULT '',
            actual_end TEXT DEFAULT '',
            status TEXT DEFAULT 'Ej startad',
            priority TEXT DEFAULT 'Normal',
            progress INTEGER DEFAULT 0,
            milestone INTEGER DEFAULT 0,
            dependencies TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            sort_order INTEGER DEFAULT 0,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS task_dependencies (
            task_id INTEGER NOT NULL,
            depends_on_task_id INTEGER NOT NULL,
            PRIMARY KEY(task_id, depends_on_task_id),
            FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE CASCADE,
            FOREIGN KEY(depends_on_task_id) REFERENCES tasks(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS risks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            kind TEXT NOT NULL DEFAULT 'Risk',
            title TEXT NOT NULL,
            description TEXT DEFAULT '',
            probability INTEGER DEFAULT 3,
            impact INTEGER DEFAULT 3,
            owner TEXT DEFAULT '',
            action TEXT DEFAULT '',
            status TEXT DEFAULT 'Öppen',
            due_date TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS comments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            task_id INTEGER,
            user_id INTEGER,
            body TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE CASCADE,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
        );

        CREATE TABLE IF NOT EXISTS audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            user_id INTEGER,
            entity_type TEXT NOT NULL,
            entity_id INTEGER,
            action TEXT NOT NULL,
            details TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            ip_address TEXT DEFAULT '',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
        );

        CREATE TABLE IF NOT EXISTS baselines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            created_at TEXT NOT NULL,
            created_by INTEGER,
            snapshot_json TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL
        );

        CREATE TABLE IF NOT EXISTS project_templates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            description TEXT DEFAULT '',
            tasks_json TEXT NOT NULL,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS system_settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS api_keys (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            key_hash TEXT NOT NULL UNIQUE,
            prefix TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            last_used TEXT DEFAULT '',
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS webhooks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            name TEXT NOT NULL,
            target_url TEXT NOT NULL,
            event_type TEXT NOT NULL DEFAULT 'project.updated',
            active INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS custom_field_definitions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            entity_type TEXT NOT NULL,
            name TEXT NOT NULL,
            field_type TEXT NOT NULL DEFAULT 'text',
            options_json TEXT NOT NULL DEFAULT '[]',
            required INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS custom_field_values (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            field_id INTEGER NOT NULL,
            entity_type TEXT NOT NULL,
            entity_id INTEGER NOT NULL,
            value TEXT DEFAULT '',
            UNIQUE(field_id,entity_type,entity_id),
            FOREIGN KEY(field_id) REFERENCES custom_field_definitions(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS dashboard_preferences (
            user_id INTEGER PRIMARY KEY,
            layout_json TEXT NOT NULL DEFAULT '{}',
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            body TEXT NOT NULL DEFAULT '',
            updated_by INTEGER,
            updated_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            entity_type TEXT NOT NULL DEFAULT 'project',
            entity_id INTEGER,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            uploaded_by INTEGER,
            uploaded_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS meetings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            meeting_date TEXT NOT NULL,
            attendees TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS action_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            meeting_id INTEGER,
            title TEXT NOT NULL,
            owner TEXT DEFAULT '',
            due_date TEXT DEFAULT '',
            status TEXT NOT NULL DEFAULT 'Open',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(meeting_id) REFERENCES meetings(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS notifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            project_id INTEGER,
            message TEXT NOT NULL,
            link TEXT DEFAULT '',
            is_read INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS automation_rules (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER,
            name TEXT NOT NULL,
            trigger_type TEXT NOT NULL,
            action_type TEXT NOT NULL,
            config_json TEXT NOT NULL DEFAULT '{}',
            active INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS portfolios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            description TEXT DEFAULT '',
            owner TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS portfolio_projects (
            portfolio_id INTEGER NOT NULL,
            project_id INTEGER NOT NULL,
            PRIMARY KEY(portfolio_id,project_id),
            FOREIGN KEY(portfolio_id) REFERENCES portfolios(id) ON DELETE CASCADE,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS status_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            report_date TEXT NOT NULL,
            overall_rag TEXT NOT NULL DEFAULT 'Green',
            scope_rag TEXT NOT NULL DEFAULT 'Green',
            schedule_rag TEXT NOT NULL DEFAULT 'Green',
            budget_rag TEXT NOT NULL DEFAULT 'Green',
            resources_rag TEXT NOT NULL DEFAULT 'Green',
            summary TEXT DEFAULT '',
            achievements TEXT DEFAULT '',
            next_steps TEXT DEFAULT '',
            created_by INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS raid_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            item_type TEXT NOT NULL,
            title TEXT NOT NULL,
            owner TEXT DEFAULT '',
            status TEXT DEFAULT 'Open',
            due_date TEXT DEFAULT '',
            details TEXT DEFAULT '',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS decisions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            decision TEXT NOT NULL,
            decided_by TEXT DEFAULT '',
            decision_date TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS change_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            description TEXT DEFAULT '',
            impact_scope TEXT DEFAULT '',
            impact_days INTEGER NOT NULL DEFAULT 0,
            impact_cost REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'Proposed',
            requested_by INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS approvals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            entity_type TEXT NOT NULL,
            entity_id INTEGER NOT NULL,
            requested_by INTEGER,
            approved_by INTEGER,
            status TEXT NOT NULL DEFAULT 'Pending',
            comment TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            decided_at TEXT DEFAULT '',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS resource_profiles (
            user_id INTEGER PRIMARY KEY,
            weekly_capacity REAL NOT NULL DEFAULT 40,
            hourly_cost REAL NOT NULL DEFAULT 0,
            hourly_rate REAL NOT NULL DEFAULT 0,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS time_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            task_id INTEGER,
            user_id INTEGER NOT NULL,
            work_date TEXT NOT NULL,
            hours REAL NOT NULL,
            note TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE SET NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS project_costs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            category TEXT NOT NULL DEFAULT 'External',
            description TEXT NOT NULL,
            planned REAL NOT NULL DEFAULT 0,
            actual REAL NOT NULL DEFAULT 0,
            cost_date TEXT DEFAULT '',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS project_finance (
            project_id INTEGER PRIMARY KEY,
            budget REAL NOT NULL DEFAULT 0,
            currency TEXT NOT NULL DEFAULT 'SEK',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS task_links (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            predecessor_id INTEGER NOT NULL,
            successor_id INTEGER NOT NULL,
            link_type TEXT NOT NULL DEFAULT 'FS',
            lag_days INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(predecessor_id) REFERENCES tasks(id) ON DELETE CASCADE,
            FOREIGN KEY(successor_id) REFERENCES tasks(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS saved_views (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            project_id INTEGER,
            name TEXT NOT NULL,
            view_type TEXT NOT NULL DEFAULT 'list',
            filter_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );
        """)

        conn.executescript("""

        CREATE TABLE IF NOT EXISTS resource_allocations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            user_id INTEGER,
            resource_name TEXT NOT NULL,
            week_start TEXT NOT NULL,
            allocation_pct INTEGER NOT NULL DEFAULT 0,
            planned_hours REAL NOT NULL DEFAULT 0,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
        );

        """)

        conn.executescript("""

        CREATE TABLE IF NOT EXISTS programs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            owner TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS program_projects (
            program_id INTEGER NOT NULL,
            project_id INTEGER NOT NULL,
            PRIMARY KEY(program_id,project_id),
            FOREIGN KEY(program_id) REFERENCES programs(id) ON DELETE CASCADE,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        """)

        conn.executescript("""

        CREATE TABLE IF NOT EXISTS document_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            version_no INTEGER NOT NULL,
            content TEXT NOT NULL DEFAULT '',
            created_by INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE,
            FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS mentions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            mentioned_user_id INTEGER NOT NULL,
            source_type TEXT NOT NULL,
            source_id INTEGER,
            message TEXT NOT NULL,
            created_at TEXT NOT NULL,
            read_at TEXT NOT NULL DEFAULT '',
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(mentioned_user_id) REFERENCES users(id) ON DELETE CASCADE
        );

        """)

        conn.executescript("""

        CREATE TABLE IF NOT EXISTS report_definitions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            report_type TEXT NOT NULL DEFAULT 'portfolio',
            config_json TEXT NOT NULL DEFAULT '{}',
            owner_user_id INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(owner_user_id) REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS report_export_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            format TEXT NOT NULL,
            report_type TEXT NOT NULL DEFAULT 'executive',
            file_name TEXT NOT NULL DEFAULT '',
            exported_by INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(exported_by) REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE INDEX IF NOT EXISTS idx_report_export_project_created
            ON report_export_log(project_id, created_at DESC);
        CREATE TABLE IF NOT EXISTS oidc_settings (
            id INTEGER PRIMARY KEY CHECK (id=1),
            enabled INTEGER NOT NULL DEFAULT 0,
            issuer TEXT NOT NULL DEFAULT '',
            client_id TEXT NOT NULL DEFAULT '',
            scopes TEXT NOT NULL DEFAULT 'openid profile email'
        );
        INSERT OR IGNORE INTO oidc_settings(id) VALUES(1);

        """)

        conn.executescript("""

        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS favorites (
    user_id INTEGER NOT NULL,
    entity_type TEXT NOT NULL,
    entity_id INTEGER NOT NULL,
    created_at TEXT NOT NULL,
    PRIMARY KEY(user_id,entity_type,entity_id),
    FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS recent_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id INTEGER NOT NULL,
    entity_type TEXT NOT NULL,
    entity_id INTEGER NOT NULL,
    label TEXT NOT NULL,
    url TEXT NOT NULL,
    viewed_at TEXT NOT NULL,
    FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS task_constraints (
    task_id INTEGER PRIMARY KEY,
    constraint_type TEXT NOT NULL DEFAULT 'ASAP',
    constraint_date TEXT DEFAULT '',
    calendar_name TEXT DEFAULT 'Standard',
    FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS planning_scenarios (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    project_id INTEGER NOT NULL,
    name TEXT NOT NULL,
    snapshot_json TEXT NOT NULL,
    created_by INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS automation_executions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    rule_id INTEGER,
    project_id INTEGER,
    status TEXT NOT NULL,
    details TEXT DEFAULT '',
    started_at TEXT NOT NULL,
    finished_at TEXT DEFAULT '',
    FOREIGN KEY(rule_id) REFERENCES automation_rules(id) ON DELETE SET NULL
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS scheduled_reports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    project_id INTEGER,
    name TEXT NOT NULL,
    format TEXT NOT NULL DEFAULT 'pdf',
    cadence TEXT NOT NULL DEFAULT 'weekly',
    recipient TEXT DEFAULT '',
    active INTEGER NOT NULL DEFAULT 1,
    created_by INTEGER,
    created_at TEXT NOT NULL
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS integration_settings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    provider TEXT NOT NULL UNIQUE,
    enabled INTEGER NOT NULL DEFAULT 0,
    base_url TEXT DEFAULT '',
    config_json TEXT NOT NULL DEFAULT '{}',
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS webhook_deliveries (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    webhook_id INTEGER,
    event_type TEXT NOT NULL,
    payload_json TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'Pending',
    attempts INTEGER NOT NULL DEFAULT 0,
    last_error TEXT DEFAULT '',
    created_at TEXT NOT NULL,
    FOREIGN KEY(webhook_id) REFERENCES webhooks(id) ON DELETE CASCADE
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS user_sessions (
    id TEXT PRIMARY KEY,
    user_id INTEGER NOT NULL,
    created_at TEXT NOT NULL,
    last_seen TEXT NOT NULL,
    ip_address TEXT DEFAULT '',
    user_agent TEXT DEFAULT '',
    revoked_at TEXT DEFAULT '',
    FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS login_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id INTEGER,
    username TEXT NOT NULL,
    success INTEGER NOT NULL,
    ip_address TEXT DEFAULT '',
    user_agent TEXT DEFAULT '',
    created_at TEXT NOT NULL,
    FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
);
CREATE TABLE IF NOT EXISTS security_events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    severity TEXT NOT NULL,
    event_type TEXT NOT NULL,
    user_id INTEGER,
    details TEXT DEFAULT '',
    ip_address TEXT DEFAULT '',
    created_at TEXT NOT NULL
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS project_environments (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, name TEXT NOT NULL, base_url TEXT DEFAULT '', owner TEXT DEFAULT '', status TEXT DEFAULT 'Ready', notes TEXT DEFAULT '', UNIQUE(project_id,name), FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS deliverables (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, name TEXT NOT NULL, owner TEXT DEFAULT '', due_date TEXT DEFAULT '', status TEXT DEFAULT 'Planned', acceptance_criteria TEXT DEFAULT '', FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS interface_register (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, name TEXT NOT NULL, source_system TEXT DEFAULT '', target_system TEXT DEFAULT '', protocol TEXT DEFAULT '', owner TEXT DEFAULT '', status TEXT DEFAULT 'Design', notes TEXT DEFAULT '', FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS test_cycles (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, name TEXT NOT NULL, test_type TEXT NOT NULL DEFAULT 'SIT', start_date TEXT DEFAULT '', end_date TEXT DEFAULT '', status TEXT DEFAULT 'Planned', passed INTEGER NOT NULL DEFAULT 0, failed INTEGER NOT NULL DEFAULT 0, FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS requirements_traceability (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, requirement_id TEXT NOT NULL, requirement TEXT NOT NULL, deliverable_id INTEGER, test_cycle_id INTEGER, status TEXT DEFAULT 'Open', UNIQUE(project_id,requirement_id), FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS cutover_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, sequence_no INTEGER NOT NULL DEFAULT 0, title TEXT NOT NULL, owner TEXT DEFAULT '', planned_at TEXT DEFAULT '', status TEXT DEFAULT 'Planned', rollback_step TEXT DEFAULT '', FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS project_health_snapshots (
    id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, snapshot_date TEXT NOT NULL, progress INTEGER NOT NULL, overdue_count INTEGER NOT NULL, blocked_count INTEGER NOT NULL, high_risk_count INTEGER NOT NULL, open_change_count INTEGER NOT NULL, health_score INTEGER NOT NULL, details_json TEXT NOT NULL DEFAULT '{}', UNIQUE(project_id,snapshot_date), FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS assistant_queries (
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, project_id INTEGER, question TEXT NOT NULL, answer TEXT NOT NULL, created_at TEXT NOT NULL, FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS health_checks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    check_name TEXT NOT NULL,
    status TEXT NOT NULL,
    details TEXT NOT NULL DEFAULT '',
    checked_at TEXT NOT NULL
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS user_preferences_v41 (
    user_id INTEGER PRIMARY KEY,
    start_page TEXT NOT NULL DEFAULT 'home',
    compact_mode INTEGER NOT NULL DEFAULT 0,
    favorite_project_id INTEGER,
    FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS project_calendars (
    project_id INTEGER PRIMARY KEY,
    work_week TEXT NOT NULL DEFAULT '1,2,3,4,5',
    hours_per_day REAL NOT NULL DEFAULT 8,
    holiday_json TEXT NOT NULL DEFAULT '[]',
    FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS risk_controls_v44 (
    risk_id INTEGER PRIMARY KEY,
    response_strategy TEXT NOT NULL DEFAULT '',
    mitigation TEXT NOT NULL DEFAULT '',
    contingency TEXT NOT NULL DEFAULT '',
    residual_probability INTEGER NOT NULL DEFAULT 1,
    residual_impact INTEGER NOT NULL DEFAULT 1,
    review_date TEXT NOT NULL DEFAULT '',
    FOREIGN KEY(risk_id) REFERENCES risks(id) ON DELETE CASCADE
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS notification_channels_v46 (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    channel_type TEXT NOT NULL,
    name TEXT NOT NULL,
    config_json TEXT NOT NULL DEFAULT '{}',
    enabled INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS automation_queue_v46 (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    rule_id INTEGER,
    project_id INTEGER,
    payload_json TEXT NOT NULL DEFAULT '{}',
    status TEXT NOT NULL DEFAULT 'Queued',
    attempts INTEGER NOT NULL DEFAULT 0,
    last_error TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    processed_at TEXT NOT NULL DEFAULT ''
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS enterprise_roles_v47 (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL UNIQUE,
    permissions_json TEXT NOT NULL DEFAULT '[]',
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS retention_policies_v47 (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    data_type TEXT NOT NULL UNIQUE,
    retention_days INTEGER NOT NULL DEFAULT 365,
    enabled INTEGER NOT NULL DEFAULT 1
);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS intelligence_notes_v50 (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    project_id INTEGER NOT NULL,
    note_type TEXT NOT NULL,
    source_text TEXT NOT NULL DEFAULT '',
    generated_json TEXT NOT NULL DEFAULT '{}',
    created_by INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
    FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL
);
        """)



        # v7.0.1: make late-roadmap tables part of normal startup migration.
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS intake_requests(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_type TEXT NOT NULL,
            title TEXT NOT NULL,
            customer TEXT DEFAULT '',
            requested_by TEXT DEFAULT '',
            priority TEXT DEFAULT 'Normal',
            requested_date TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'New',
            description TEXT DEFAULT '',
            project_id INTEGER
        );
        CREATE TABLE IF NOT EXISTS visual_automation_rules(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            trigger_name TEXT NOT NULL,
            condition_name TEXT DEFAULT '',
            action_name TEXT NOT NULL,
            enabled INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS strategic_goals(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT DEFAULT '',
            target_value REAL DEFAULT 0,
            current_value REAL DEFAULT 0,
            unit TEXT DEFAULT '%',
            owner TEXT DEFAULT '',
            due_date TEXT DEFAULT '',
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS goal_projects(
            goal_id INTEGER NOT NULL,
            project_id INTEGER NOT NULL,
            weight INTEGER DEFAULT 100,
            PRIMARY KEY(goal_id,project_id),
            FOREIGN KEY(goal_id) REFERENCES strategic_goals(id) ON DELETE CASCADE,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        );

        CREATE INDEX IF NOT EXISTS idx_tasks_project_status ON tasks(project_id,status);
        CREATE INDEX IF NOT EXISTS idx_tasks_project_end ON tasks(project_id,end_date);
        CREATE INDEX IF NOT EXISTS idx_tasks_owner ON tasks(owner);
        CREATE INDEX IF NOT EXISTS idx_risks_project_status ON risks(project_id,status);
        CREATE INDEX IF NOT EXISTS idx_notifications_user_read ON notifications(user_id,is_read);
        CREATE INDEX IF NOT EXISTS idx_resource_alloc_project_week ON resource_allocations(project_id,week_start);
        CREATE INDEX IF NOT EXISTS idx_audit_project_id ON audit_log(project_id,id);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS form_definitions(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,entity_type TEXT NOT NULL DEFAULT 'Project Request',description TEXT DEFAULT '',fields_json TEXT NOT NULL DEFAULT '[]',active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS form_submissions(id INTEGER PRIMARY KEY AUTOINCREMENT,form_id INTEGER NOT NULL,title TEXT NOT NULL,payload_json TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'New',score REAL DEFAULT 0,created_by INTEGER,created_at TEXT NOT NULL,FOREIGN KEY(form_id) REFERENCES form_definitions(id) ON DELETE CASCADE);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS business_cases(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER,submission_id INTEGER,title TEXT NOT NULL,strategic_fit INTEGER DEFAULT 5,business_value INTEGER DEFAULT 5,urgency INTEGER DEFAULT 5,regulatory INTEGER DEFAULT 5,technical_risk INTEGER DEFAULT 5,resource_demand INTEGER DEFAULT 5,cost_score INTEGER DEFAULT 5,expected_benefit INTEGER DEFAULT 5,total_score REAL DEFAULT 0,estimated_cost REAL DEFAULT 0,expected_value REAL DEFAULT 0,status TEXT DEFAULT 'Draft',created_at TEXT NOT NULL);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS cross_project_dependencies(id INTEGER PRIMARY KEY AUTOINCREMENT,predecessor_project_id INTEGER NOT NULL,successor_project_id INTEGER NOT NULL,predecessor_task_id INTEGER,successor_task_id INTEGER,link_type TEXT NOT NULL DEFAULT 'FS',lag_days INTEGER DEFAULT 0,status TEXT DEFAULT 'Active',notes TEXT DEFAULT '',created_at TEXT NOT NULL);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS resource_directory(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER,display_name TEXT NOT NULL,role_name TEXT DEFAULT '',location TEXT DEFAULT '',capacity_pct INTEGER DEFAULT 100,skills_json TEXT NOT NULL DEFAULT '[]',available_from TEXT DEFAULT '',active INTEGER NOT NULL DEFAULT 1);
CREATE UNIQUE INDEX IF NOT EXISTS idx_resource_directory_user ON resource_directory(user_id) WHERE user_id IS NOT NULL;
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS managed_templates(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,version INTEGER NOT NULL DEFAULT 1,description TEXT DEFAULT '',tasks_json TEXT NOT NULL DEFAULT '[]',active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS project_template_instances(project_id INTEGER PRIMARY KEY,managed_template_id INTEGER NOT NULL,template_version INTEGER NOT NULL,applied_at TEXT NOT NULL);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS portfolio_scenarios(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,description TEXT DEFAULT '',changes_json TEXT NOT NULL DEFAULT '[]',created_by INTEGER,created_at TEXT NOT NULL);
        """)

        conn.executescript("""
CREATE TABLE IF NOT EXISTS project_benefits(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER NOT NULL,title TEXT NOT NULL,unit TEXT DEFAULT '%',baseline_value REAL DEFAULT 0,target_value REAL DEFAULT 0,actual_value REAL,measurement_date TEXT DEFAULT '',owner TEXT DEFAULT '',status TEXT DEFAULT 'Planned',created_at TEXT NOT NULL);
        """)

        # Migration from older versions
        ensure_column(conn, "users", "force_password_change", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "users", "failed_logins", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "users", "locked_until", "TEXT DEFAULT ''")
        ensure_column(conn, "users", "last_login", "TEXT DEFAULT ''")
        ensure_column(conn, "projects", "created_by", "INTEGER")
        ensure_column(conn, "project_members", "added_at", "TEXT NOT NULL DEFAULT ''")
        ensure_column(conn, "project_members", "added_by", "INTEGER")
        ensure_column(conn, "audit_log", "ip_address", "TEXT DEFAULT ''")
        ensure_column(conn, "tasks", "planned_hours", "REAL DEFAULT 0")
        ensure_column(conn, "tasks", "remaining_hours", "REAL DEFAULT 0")
        ensure_column(conn, "tasks", "parent_task_id", "INTEGER")
        ensure_column(conn, "tasks", "duration_days", "INTEGER DEFAULT 1")
        ensure_column(conn, "tasks", "owner_user_id", "INTEGER")

        ensure_column(conn, "change_requests", "reason", "TEXT DEFAULT ''")
        ensure_column(conn, "change_requests", "decided_by", "INTEGER")
        ensure_column(conn, "change_requests", "decided_at", "TEXT DEFAULT ''")
        ensure_column(conn, "decisions", "owner", "TEXT DEFAULT ''")
        ensure_column(conn, "decisions", "decided_at", "TEXT DEFAULT ''")

        ensure_column(conn, "api_keys", "scopes", "TEXT NOT NULL DEFAULT 'read'")
        ensure_column(conn, "api_keys", "expires_at", "TEXT DEFAULT ''")

        ensure_column(conn, "projects", "archived_at", "TEXT DEFAULT ''")
        ensure_column(conn, "projects", "deleted_at", "TEXT DEFAULT ''")
        ensure_column(conn, "tasks", "deleted_at", "TEXT DEFAULT ''")

        conn.execute("INSERT OR IGNORE INTO schema_migrations(version,applied_at) VALUES(?,?)",
                     (APP_VERSION, datetime.now().isoformat(timespec="seconds")))
        conn.execute("INSERT OR IGNORE INTO system_settings(key,value) VALUES('allow_project_creation','pm')")
        conn.execute("INSERT OR IGNORE INTO system_settings(key,value) VALUES('lockout_attempts','5')")
        conn.execute("INSERT OR IGNORE INTO system_settings(key,value) VALUES('lockout_minutes','15')")

        templates = {
            "Standardprojekt":[
                {"wbs":"1","title":"Initiering","milestone":0},
                {"wbs":"1.1","title":"Projektmål och scope","milestone":0},
                {"wbs":"1.2","title":"Kickoff","milestone":1},
                {"wbs":"2","title":"Planering","milestone":0},
                {"wbs":"2.1","title":"Detaljplan","milestone":0},
                {"wbs":"2.2","title":"Riskworkshop","milestone":0},
                {"wbs":"3","title":"Genomförande","milestone":0},
                {"wbs":"4","title":"Verifiering och acceptans","milestone":0},
                {"wbs":"5","title":"Överlämning","milestone":1}
            ],
            "IT-/integrationsprojekt":[
                {"wbs":"1","title":"Förstudie och krav","milestone":0},
                {"wbs":"2","title":"Arkitektur och design","milestone":1},
                {"wbs":"3","title":"Utveckling/konfiguration","milestone":0},
                {"wbs":"4","title":"Integrationstest","milestone":0},
                {"wbs":"5","title":"UAT","milestone":1},
                {"wbs":"6","title":"Driftsättning","milestone":1},
                {"wbs":"7","title":"Hypercare och avslut","milestone":0}
            ],
            "LIMS-projekt":[
                {"wbs":"1","title":"Process- och kravanalys","milestone":0},
                {"wbs":"2","title":"LIMS-design","milestone":1},
                {"wbs":"3","title":"Konfiguration och utveckling","milestone":0},
                {"wbs":"4","title":"Instrument/integrationer","milestone":0},
                {"wbs":"5","title":"Verifiering/validering","milestone":0},
                {"wbs":"6","title":"Utbildning","milestone":0},
                {"wbs":"7","title":"Go-live","milestone":1},
                {"wbs":"8","title":"Stabilisering","milestone":0}
            ]
        }
        for name,tasks in templates.items():
            conn.execute("""INSERT OR IGNORE INTO project_templates(name,description,tasks_json,created_at)
                            VALUES(?,?,?,?)""",
                         (name,f"Inbyggd mall: {name}",json.dumps(tasks,ensure_ascii=False),
                          datetime.now().isoformat(timespec="seconds")))

def setting(key, default=""):
    with db() as conn:
        row = conn.execute("SELECT value FROM system_settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default

def user_count():
    with db() as conn:
        return conn.execute("SELECT COUNT(*) c FROM users").fetchone()["c"]

def current_user():
    uid = session.get("user_id")
    if not uid: return None
    with db() as conn:
        return conn.execute("SELECT * FROM users WHERE id=? AND active=1",(uid,)).fetchone()

def csrf_token():
    token = session.get("_csrf")
    if not token:
        token = secrets.token_urlsafe(32)
        session["_csrf"] = token
    return token


def active_language():
    lang = session.get("lang", "sv")
    return lang if lang in LANGUAGES else "sv"

def tr(key):
    lang = active_language()
    return TRANSLATIONS.get(lang, TRANSLATIONS["sv"]).get(key, TRANSLATIONS["sv"].get(key, key))

@app.get("/language/<lang>")
def set_language(lang):
    if lang not in LANGUAGES:
        abort(404)
    session["lang"] = lang
    return redirect(request.referrer or url_for("index"))


V152_UI = {
    "sv": {
        "language":"Språk","today":"Idag","my_work":"Mitt arbete","projects":"Projekt","portfolio":"Portfolio",
        "search":"Sök","create":"Skapa","admin":"Admin","notifications":"Notiser","password":"Lösenord",
        "logout":"Logga ut","quality":"Kvalitetscenter","integrations":"Integrationer","enterprise":"Företagsinställningar",
        "app_version":"Appversion","system_status":"Systemstatus","go_to":"Gå till","work":"Arbeta",
        "reports":"Rapporter","resources":"Resurser","capacity":"Kapacitet","project_health":"Projekthälsa",
        "azure_devops":"Azure DevOps","all_search":"Sök allt","choose_language":"Välj språk"
    },
    "en": {
        "language":"Language","today":"Today","my_work":"My work","projects":"Projects","portfolio":"Portfolio",
        "search":"Search","create":"Create","admin":"Admin","notifications":"Notifications","password":"Password",
        "logout":"Sign out","quality":"Quality center","integrations":"Integrations","enterprise":"Enterprise settings",
        "app_version":"App version","system_status":"System status","go_to":"Go to","work":"Work",
        "reports":"Reports","resources":"Resources","capacity":"Capacity","project_health":"Project health",
        "azure_devops":"Azure DevOps","all_search":"Search everything","choose_language":"Choose language"
    },
    "de": {
        "language":"Sprache","today":"Heute","my_work":"Meine Arbeit","projects":"Projekte","portfolio":"Portfolio",
        "search":"Suchen","create":"Erstellen","admin":"Admin","notifications":"Benachrichtigungen","password":"Passwort",
        "logout":"Abmelden","quality":"Qualitätscenter","integrations":"Integrationen","enterprise":"Unternehmenseinstellungen",
        "app_version":"App-Version","system_status":"Systemstatus","go_to":"Gehe zu","work":"Arbeiten",
        "reports":"Berichte","resources":"Ressourcen","capacity":"Kapazität","project_health":"Projektstatus",
        "azure_devops":"Azure DevOps","all_search":"Alles durchsuchen","choose_language":"Sprache wählen"
    },
    "no": {
        "language":"Språk","today":"I dag","my_work":"Mitt arbeid","projects":"Prosjekter","portfolio":"Portefølje",
        "search":"Søk","create":"Opprett","admin":"Admin","notifications":"Varsler","password":"Passord",
        "logout":"Logg ut","quality":"Kvalitetssenter","integrations":"Integrasjoner","enterprise":"Virksomhetsinnstillinger",
        "app_version":"Appversjon","system_status":"Systemstatus","go_to":"Gå til","work":"Arbeid",
        "reports":"Rapporter","resources":"Ressurser","capacity":"Kapasitet","project_health":"Prosjekthelse",
        "azure_devops":"Azure DevOps","all_search":"Søk i alt","choose_language":"Velg språk"
    },
    "da": {
        "language":"Sprog","today":"I dag","my_work":"Mit arbejde","projects":"Projekter","portfolio":"Portefølje",
        "search":"Søg","create":"Opret","admin":"Admin","notifications":"Notifikationer","password":"Adgangskode",
        "logout":"Log ud","quality":"Kvalitetscenter","integrations":"Integrationer","enterprise":"Virksomhedsindstillinger",
        "app_version":"Appversion","system_status":"Systemstatus","go_to":"Gå til","work":"Arbejde",
        "reports":"Rapporter","resources":"Ressourcer","capacity":"Kapacitet","project_health":"Projektsundhed",
        "azure_devops":"Azure DevOps","all_search":"Søg i alt","choose_language":"Vælg sprog"
    },
    "fi": {
        "language":"Kieli","today":"Tänään","my_work":"Oma työ","projects":"Projektit","portfolio":"Portfolio",
        "search":"Haku","create":"Luo","admin":"Admin","notifications":"Ilmoitukset","password":"Salasana",
        "logout":"Kirjaudu ulos","quality":"Laatukeskus","integrations":"Integraatiot","enterprise":"Yritysasetukset",
        "app_version":"Sovellusversio","system_status":"Järjestelmän tila","go_to":"Siirry","work":"Työ",
        "reports":"Raportit","resources":"Resurssit","capacity":"Kapasiteetti","project_health":"Projektin tila",
        "azure_devops":"Azure DevOps","all_search":"Hae kaikesta","choose_language":"Valitse kieli"
    }
}
def ui152(key):
    lang=active_language()
    return V152_UI.get(lang,V152_UI["sv"]).get(key,V152_UI["sv"].get(key,key))


@app.after_request
def force_utf8_charset_v1521(response):
    content_type = response.headers.get("Content-Type","")
    if content_type.startswith("text/") and "charset=" not in content_type.lower():
        response.headers["Content-Type"] = content_type + "; charset=utf-8"
    return response

@app.context_processor
def inject_i18n():
    return dict(t=tr, ui=ui152, active_lang=active_language(), languages=LANGUAGES)

@app.context_processor
def inject_globals():
    return dict(current_user=current_user(), app_version=APP_VERSION, csrf_token=csrf_token())

@app.before_request
def bootstrap_and_security():
    init_db()
    if request.method in ("POST","PUT","PATCH","DELETE"):
        sent = request.form.get("_csrf") or request.headers.get("X-CSRF-Token","")
        expected = session.get("_csrf","")
        if not expected or not sent or not hmac.compare_digest(sent, expected):
            abort(400, "Ogiltig eller saknad CSRF-token.")
    if session.get("user_id"):
        session.permanent = True
        u = current_user()
        if u and u["force_password_change"] and request.endpoint not in ("change_password","logout","static"):
            return redirect(url_for("change_password"))

def login_required(fn):
    @wraps(fn)
    def wrapper(*args,**kwargs):
        if user_count()==0: return redirect(url_for("setup"))
        if not current_user(): return redirect(url_for("login",next=request.path))
        return fn(*args,**kwargs)
    return wrapper

def role_required(*roles):
    def deco(fn):
        @wraps(fn)
        def wrapper(*args,**kwargs):
            u=current_user()
            if not u or u["role"] not in roles: abort(403)
            return fn(*args,**kwargs)
        return wrapper
    return deco

def project_access(project_id, write=False, manager=False):
    u=current_user()
    if not u: abort(401)
    if u["role"]=="admin": return "admin"
    with db() as conn:
        row=conn.execute("""SELECT pm.project_role FROM project_members pm
                            JOIN projects p ON p.id=pm.project_id
                            WHERE pm.project_id=? AND pm.user_id=?""",(project_id,u["id"])).fetchone()
    if not row: abort(403)
    pr=row["project_role"]
    if manager and pr!="pm": abort(403)
    if write and pr not in ("pm","member"): abort(403)
    return pr

def audit(project_id, entity_type, entity_id, action, details=""):
    u=current_user()
    ip=request.headers.get("X-Forwarded-For",request.remote_addr or "").split(",")[0].strip()
    with db() as conn:
        conn.execute("""INSERT INTO audit_log(project_id,user_id,entity_type,entity_id,action,details,created_at,ip_address)
                        VALUES(?,?,?,?,?,?,?,?)""",
                     (project_id,u["id"] if u else None,entity_type,entity_id,action,details,
                      datetime.now().isoformat(timespec="seconds"),ip))

def project_or_404(project_id, write=False, manager=False):
    project_access(project_id,write=write,manager=manager)
    with db() as conn:
        row=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
    if not row: abort(404)
    return row

def task_or_404(project_id,task_id,write=False):
    project_access(project_id,write=write)
    with db() as conn:
        row=conn.execute("SELECT * FROM tasks WHERE id=? AND project_id=?",(task_id,project_id)).fetchone()
    if not row: abort(404)
    return row

def parse_date(s):
    if not s:return None
    try:return datetime.strptime(s,"%Y-%m-%d").date()
    except:return None

def derived_status(task):
    if task["status"]=="Klar" or int(task["progress"] or 0)>=100:return "Klar"
    today=date.today(); end=parse_date(task["end_date"]); start=parse_date(task["start_date"])
    if task["status"]=="Blockerad":return "Blockerad"
    if end and end<today:return "Försenad"
    if start and start>today:return "Kommande"
    if int(task["progress"] or 0)>0 or (start and start<=today):return "Pågår"
    return task["status"] or "Ej startad"

def task_health(task):
    return {"Försenad":"danger","Blockerad":"warning","Klar":"success","Pågår":"info"}.get(derived_status(task),"muted")

@app.route("/setup",methods=["GET","POST"])
def setup():
    if user_count()>0:return redirect(url_for("login"))
    if request.method=="POST":
        username=request.form["username"].strip()
        password=request.form["password"]
        if len(password)<10:
            flash("Lösenordet måste vara minst 10 tecken.","danger")
            return render_template("setup.html")
        with db() as conn:
            conn.execute("""INSERT INTO users(username,display_name,password_hash,role,active,force_password_change,created_at)
                            VALUES(?,?,?,?,1,0,?)""",
                         (username,request.form.get("display_name","").strip() or username,
                          generate_password_hash(password),"admin",datetime.now().isoformat(timespec="seconds")))
        flash("Administratör skapad. Logga in.","success")
        return redirect(url_for("login"))
    return render_template("setup.html")

@app.route("/login",methods=["GET","POST"])
def login():
    if user_count()==0:return redirect(url_for("setup"))
    if request.method=="POST":
        username=request.form["username"].strip()
        with db() as conn:
            u=conn.execute("SELECT * FROM users WHERE username=?",(username,)).fetchone()
            if u:
                locked_until=None
                if u["locked_until"]:
                    try: locked_until=datetime.fromisoformat(u["locked_until"])
                    except: pass
                if not u["active"]:
                    flash("Kontot är inaktiverat.","danger"); return render_template("login.html")
                if locked_until and locked_until>datetime.now():
                    flash(f"Kontot är tillfälligt låst till {locked_until:%H:%M}.","danger")
                    return render_template("login.html")
                if check_password_hash(u["password_hash"],request.form["password"]):
                    conn.execute("UPDATE users SET failed_logins=0,locked_until='',last_login=? WHERE id=?",
                                 (datetime.now().isoformat(timespec="seconds"),u["id"]))
                    session.clear(); session["user_id"]=u["id"]; session["_csrf"]=secrets.token_urlsafe(32); session.permanent=True
                    return redirect(request.args.get("next") or url_for("ultimate_home_v140"))
                attempts=int(u["failed_logins"] or 0)+1
                limit=int(setting("lockout_attempts","5"))
                lock_until=""
                if attempts>=limit:
                    lock_until=(datetime.now()+timedelta(minutes=int(setting("lockout_minutes","15")))).isoformat(timespec="seconds")
                    attempts=0
                conn.execute("UPDATE users SET failed_logins=?,locked_until=? WHERE id=?",(attempts,lock_until,u["id"]))
        flash("Fel användarnamn eller lösenord.","danger")
    return render_template("login.html")

@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/change-password",methods=["GET","POST"])
@login_required
def change_password():
    u=current_user()
    if request.method=="POST":
        new=request.form["new_password"]
        if len(new)<10:
            flash("Det nya lösenordet måste vara minst 10 tecken.","danger")
        else:
            with db() as conn:
                dbu=conn.execute("SELECT * FROM users WHERE id=?",(u["id"],)).fetchone()
                if not check_password_hash(dbu["password_hash"],request.form["current_password"]):
                    flash("Nuvarande lösenord är fel.","danger")
                else:
                    conn.execute("UPDATE users SET password_hash=?,force_password_change=0 WHERE id=?",
                                 (generate_password_hash(new),u["id"]))
                    flash("Lösenordet är ändrat.","success")
                    return redirect(url_for("ultimate_home_v140"))
    return render_template("change_password.html")

@app.get("/health")
def health():
    try:
        with db() as conn: conn.execute("SELECT 1").fetchone()
        return jsonify(status="ok",version=APP_VERSION,database="ok"),200
    except Exception as e:
        return jsonify(status="error",version=APP_VERSION,error=str(e)),503

@app.route("/")
@login_required
def index():
    u=current_user()
    with db() as conn:
        if u["role"]=="admin":
            project_rows=conn.execute("""SELECT p.*,COUNT(t.id) task_count,COALESCE(ROUND(AVG(t.progress)),0) avg_progress,
                SUM(CASE WHEN t.status='Klar' OR t.progress>=100 THEN 1 ELSE 0 END) done_count,
                SUM(CASE WHEN t.status='Blockerad' THEN 1 ELSE 0 END) blocked_count
                FROM projects p LEFT JOIN tasks t ON t.project_id=p.id
                WHERE COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                GROUP BY p.id ORDER BY p.id DESC""").fetchall()
            task_rows=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
                WHERE t.end_date<>'' AND t.progress<100 AND COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                ORDER BY t.end_date LIMIT 50""").fetchall()
            open_risks=conn.execute("""SELECT r.*,p.name project_name FROM risks r JOIN projects p ON p.id=r.project_id
                WHERE r.status<>'Stängd' AND COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                ORDER BY (r.probability*r.impact) DESC LIMIT 20""").fetchall()
        else:
            project_rows=conn.execute("""SELECT p.*,COUNT(t.id) task_count,COALESCE(ROUND(AVG(t.progress)),0) avg_progress,
                SUM(CASE WHEN t.status='Klar' OR t.progress>=100 THEN 1 ELSE 0 END) done_count,
                SUM(CASE WHEN t.status='Blockerad' THEN 1 ELSE 0 END) blocked_count
                FROM projects p JOIN project_members pm ON pm.project_id=p.id AND pm.user_id=?
                LEFT JOIN tasks t ON t.project_id=p.id
                WHERE COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                GROUP BY p.id ORDER BY p.id DESC""",(u["id"],)).fetchall()
            task_rows=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
                JOIN project_members pm ON pm.project_id=p.id AND pm.user_id=?
                WHERE t.end_date<>'' AND t.progress<100 AND COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                ORDER BY t.end_date LIMIT 50""",(u["id"],)).fetchall()
            open_risks=conn.execute("""SELECT r.*,p.name project_name FROM risks r JOIN projects p ON p.id=r.project_id
                JOIN project_members pm ON pm.project_id=p.id AND pm.user_id=?
                WHERE r.status<>'Stängd' AND COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                ORDER BY (r.probability*r.impact) DESC LIMIT 20""",(u["id"],)).fetchall()

    today=date.today()
    projects=[]
    for row in project_rows:
        p=dict(row)
        p["progress"]=int(p.get("avg_progress") or 0)
        p["health"]=project_visual_health(p["id"])
        projects.append(p)

    overdue=[dict(t) for t in task_rows if parse_date(t["end_date"]) and parse_date(t["end_date"])<today]
    upcoming=[dict(t) for t in task_rows if parse_date(t["end_date"]) and today<=parse_date(t["end_date"])<=today+timedelta(days=21)]
    high_risks=[dict(r) for r in open_risks if int(r["probability"] or 0)*int(r["impact"] or 0)>=15]

    summary={
        "green":sum(1 for p in projects if p["health"]["rag"]=="green"),
        "amber":sum(1 for p in projects if p["health"]["rag"]=="amber"),
        "red":sum(1 for p in projects if p["health"]["rag"]=="red"),
        "avg_health":round(sum(p["health"]["score"] for p in projects)/len(projects)) if projects else 0,
        "avg_progress":round(sum(p["progress"] for p in projects)/len(projects)) if projects else 0,
    }
    return render_template("index.html",projects=projects,upcoming=upcoming,open_risks=open_risks,
                           high_risks=high_risks,overdue=overdue,summary=summary)

@app.route("/projects/new",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def new_project():
    u=current_user()
    with db() as conn: templates=conn.execute("SELECT * FROM project_templates ORDER BY name").fetchall()
    if request.method=="POST":
        with db() as conn:
            cur=conn.execute("""INSERT INTO projects(name,customer,project_manager,description,start_date,end_date,template_name,created_by,created_at)
                                VALUES(?,?,?,?,?,?,?,?,?)""",
                             (request.form["name"].strip(),request.form.get("customer","").strip(),
                              request.form.get("project_manager","").strip(),request.form.get("description","").strip(),
                              request.form.get("start_date",""),request.form.get("end_date",""),
                              request.form.get("template_name",""),u["id"],datetime.now().isoformat(timespec="seconds")))
            pid=cur.lastrowid
            conn.execute("""INSERT INTO project_members(project_id,user_id,project_role,added_at,added_by)
                            VALUES(?,?,?,?,?)""",(pid,u["id"],"pm",datetime.now().isoformat(timespec="seconds"),u["id"]))
            tmpl=request.form.get("template_name","")
            if tmpl:
                row=conn.execute("SELECT tasks_json FROM project_templates WHERE name=?",(tmpl,)).fetchone()
                if row:
                    for i,item in enumerate(json.loads(row["tasks_json"])):
                        conn.execute("""INSERT INTO tasks(project_id,wbs,title,status,priority,progress,milestone,sort_order)
                                        VALUES(?,?,?,?,?,?,?,?)""",(pid,item.get("wbs",""),item["title"],"Ej startad","Normal",0,item.get("milestone",0),i))
        audit(pid,"project",pid,"created","Projekt skapat")
        return redirect(url_for("project",project_id=pid))
    return render_template("project_form.html",project=None,templates=templates)

@app.route("/projects/<int:project_id>")
@login_required
def project(project_id):
    p=project_or_404(project_id)
    q=request.args.get("q","").strip(); status=request.args.get("status",""); owner=request.args.get("owner",""); priority=request.args.get("priority",""); sort=request.args.get("sort","wbs")
    sql="SELECT * FROM tasks WHERE project_id=?"; args=[project_id]
    if q:
        like=f"%{q}%"; sql+=" AND (title LIKE ? OR wbs LIKE ? OR owner LIKE ? OR notes LIKE ?)"; args += [like,like,like,like]
    if status: sql+=" AND status=?"; args.append(status)
    if owner: sql+=" AND owner=?"; args.append(owner)
    if priority: sql+=" AND priority=?"; args.append(priority)
    allowed={"wbs":"wbs COLLATE NOCASE","end":"end_date","start":"start_date","priority":"priority","owner":"owner COLLATE NOCASE"}
    sql+=" ORDER BY "+allowed.get(sort,allowed["wbs"])+",sort_order,id"
    with db() as conn:
        tasks=conn.execute(sql,args).fetchall()
        all_tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? ORDER BY (probability*impact) DESC,id DESC",(project_id,)).fetchall()
        comments=conn.execute("""SELECT c.*,u.display_name FROM comments c LEFT JOIN users u ON u.id=c.user_id
                                 WHERE c.project_id=? ORDER BY c.id DESC LIMIT 20""",(project_id,)).fetchall()
        history=conn.execute("""SELECT a.*,u.display_name FROM audit_log a LEFT JOIN users u ON u.id=a.user_id
                                WHERE a.project_id=? ORDER BY a.id DESC LIMIT 40""",(project_id,)).fetchall()
        baselines=conn.execute("SELECT * FROM baselines WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
        owners=conn.execute("SELECT DISTINCT owner FROM tasks WHERE project_id=? AND owner<>'' ORDER BY owner",(project_id,)).fetchall()
        members=conn.execute("""SELECT pm.*,u.username,u.display_name,u.role global_role FROM project_members pm
                                JOIN users u ON u.id=pm.user_id WHERE pm.project_id=? ORDER BY pm.project_role,u.display_name""",(project_id,)).fetchall()
        candidates=conn.execute("""SELECT u.* FROM users u WHERE u.active=1 AND u.id NOT IN
                                   (SELECT user_id FROM project_members WHERE project_id=?)
                                   ORDER BY u.display_name""",(project_id,)).fetchall()
    avg=round(sum(int(t["progress"] or 0) for t in all_tasks)/len(all_tasks)) if all_tasks else 0
    enriched=[dict(t,derived_status=derived_status(t),health=task_health(t)) for t in tasks]
    u=current_user()
    project_role="admin" if u["role"]=="admin" else next((m["project_role"] for m in members if m["user_id"]==u["id"]),"viewer")
    return render_template("project.html",project=p,tasks=enriched,all_tasks=all_tasks,risks=risks,comments=comments,
                           history=history,baselines=baselines,owners=owners,members=members,candidates=candidates,
                           project_role=project_role,statuses=STATUSES,priorities=PRIORITIES,risk_statuses=RISK_STATUSES,
                           avg=avg,done=sum(1 for t in all_tasks if derived_status(t)=="Klar"),
                           blocked=sum(1 for t in all_tasks if derived_status(t)=="Blockerad"),
                           overdue=sum(1 for t in all_tasks if derived_status(t)=="Försenad"),
                           upcoming=sum(1 for t in all_tasks if parse_date(t["end_date"]) and date.today()<=parse_date(t["end_date"])<=date.today()+timedelta(days=7) and int(t["progress"] or 0)<100))

@app.route("/projects/<int:project_id>/edit",methods=["GET","POST"])
@login_required
def edit_project(project_id):
    p=project_or_404(project_id,manager=True)
    with db() as conn: templates=conn.execute("SELECT * FROM project_templates ORDER BY name").fetchall()
    if request.method=="POST":
        with db() as conn:
            conn.execute("""UPDATE projects SET name=?,customer=?,project_manager=?,description=?,start_date=?,end_date=? WHERE id=?""",
                         (request.form["name"].strip(),request.form.get("customer","").strip(),
                          request.form.get("project_manager","").strip(),request.form.get("description","").strip(),
                          request.form.get("start_date",""),request.form.get("end_date",""),project_id))
        audit(project_id,"project",project_id,"updated","Projektinformation uppdaterad")
        return redirect(url_for("project",project_id=project_id))
    return render_template("project_form.html",project=p,templates=templates)

@app.post("/projects/<int:project_id>/members/add")
@login_required
def add_project_member(project_id):
    project_or_404(project_id,manager=True)
    uid=int(request.form["user_id"]); prole=request.form.get("project_role","viewer")
    if prole not in PROJECT_ROLES: prole="viewer"
    u=current_user()
    with db() as conn:
        conn.execute("""INSERT OR REPLACE INTO project_members(project_id,user_id,project_role,added_at,added_by)
                        VALUES(?,?,?,?,?)""",(project_id,uid,prole,datetime.now().isoformat(timespec="seconds"),u["id"]))
    audit(project_id,"member",uid,"added",prole)
    return redirect(url_for("project",project_id=project_id)+"#members")

@app.post("/projects/<int:project_id>/members/<int:user_id>/role")
@login_required
def change_project_member_role(project_id,user_id):
    project_or_404(project_id,manager=True)
    prole=request.form.get("project_role","viewer")
    if prole not in PROJECT_ROLES: abort(400)
    with db() as conn:
        conn.execute("UPDATE project_members SET project_role=? WHERE project_id=? AND user_id=?",(prole,project_id,user_id))
    audit(project_id,"member",user_id,"role-changed",prole)
    return redirect(url_for("project",project_id=project_id)+"#members")

@app.post("/projects/<int:project_id>/members/<int:user_id>/remove")
@login_required
def remove_project_member(project_id,user_id):
    project_or_404(project_id,manager=True)
    u=current_user()
    if u["id"]==user_id and u["role"]!="admin":
        flash("Du kan inte ta bort dig själv som projektledare.","danger")
    else:
        with db() as conn:
            conn.execute("DELETE FROM project_members WHERE project_id=? AND user_id=?",(project_id,user_id))
        audit(project_id,"member",user_id,"removed","")
    return redirect(url_for("project",project_id=project_id)+"#members")

@app.route("/projects/<int:project_id>/tasks/new",methods=["GET","POST"])
@login_required
def new_task(project_id):
    p=project_or_404(project_id,write=True)
    with db() as conn: candidates=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,)).fetchall()
    if request.method=="POST":
        with db() as conn:
            cur=conn.execute("""INSERT INTO tasks(project_id,wbs,title,owner,start_date,end_date,actual_start,actual_end,status,priority,progress,milestone,notes,sort_order)
                                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                             (project_id,request.form.get("wbs","").strip(),request.form["title"].strip(),request.form.get("owner","").strip(),
                              request.form.get("start_date",""),request.form.get("end_date",""),request.form.get("actual_start",""),request.form.get("actual_end",""),
                              request.form.get("status","Ej startad"),request.form.get("priority","Normal"),
                              max(0,min(100,int(request.form.get("progress","0") or 0))),1 if request.form.get("milestone")=="on" else 0,
                              request.form.get("notes","").strip(),int(request.form.get("sort_order","0") or 0)))
            tid=cur.lastrowid
            for dep in request.form.getlist("dependency_ids"):
                if dep.isdigit() and int(dep)!=tid:
                    conn.execute("INSERT OR IGNORE INTO task_dependencies(task_id,depends_on_task_id) VALUES(?,?)",(tid,int(dep)))
        audit(project_id,"task",tid,"created",request.form["title"].strip())
        return redirect(url_for("project",project_id=project_id))
    return render_template("task_form.html",project=p,task=None,candidates=candidates,selected_deps=[],statuses=STATUSES,priorities=PRIORITIES)

@app.route("/projects/<int:project_id>/tasks/<int:task_id>/edit",methods=["GET","POST"])
@login_required
def edit_task(project_id,task_id):
    p=project_or_404(project_id,write=True); t=task_or_404(project_id,task_id,write=True)
    with db() as conn:
        candidates=conn.execute("SELECT * FROM tasks WHERE project_id=? AND id<>? ORDER BY wbs,id",(project_id,task_id)).fetchall()
        selected=[r["depends_on_task_id"] for r in conn.execute("SELECT depends_on_task_id FROM task_dependencies WHERE task_id=?",(task_id,))]
    if request.method=="POST":
        prog=max(0,min(100,int(request.form.get("progress","0") or 0))); stat=request.form.get("status","Ej startad")
        if stat=="Klar": prog=100
        with db() as conn:
            conn.execute("""UPDATE tasks SET wbs=?,title=?,owner=?,start_date=?,end_date=?,actual_start=?,actual_end=?,status=?,priority=?,progress=?,milestone=?,notes=?,sort_order=?
                            WHERE id=? AND project_id=?""",
                         (request.form.get("wbs","").strip(),request.form["title"].strip(),request.form.get("owner","").strip(),
                          request.form.get("start_date",""),request.form.get("end_date",""),request.form.get("actual_start",""),request.form.get("actual_end",""),
                          stat,request.form.get("priority","Normal"),prog,1 if request.form.get("milestone")=="on" else 0,
                          request.form.get("notes","").strip(),int(request.form.get("sort_order","0") or 0),task_id,project_id))
            conn.execute("DELETE FROM task_dependencies WHERE task_id=?",(task_id,))
            for dep in request.form.getlist("dependency_ids"):
                if dep.isdigit() and int(dep)!=task_id:
                    conn.execute("INSERT OR IGNORE INTO task_dependencies(task_id,depends_on_task_id) VALUES(?,?)",(task_id,int(dep)))
        audit(project_id,"task",task_id,"updated",request.form["title"].strip())
        return redirect(url_for("project",project_id=project_id))
    return render_template("task_form.html",project=p,task=t,candidates=candidates,selected_deps=selected,statuses=STATUSES,priorities=PRIORITIES)

@app.post("/projects/<int:project_id>/tasks/<int:task_id>/delete")
@login_required
def delete_task(project_id,task_id):
    project_or_404(project_id,manager=True)
    t=task_or_404(project_id,task_id)
    with db() as conn: conn.execute("DELETE FROM tasks WHERE id=? AND project_id=?",(task_id,project_id))
    audit(project_id,"task",task_id,"deleted",t["title"])
    return redirect(url_for("project",project_id=project_id))

@app.post("/projects/<int:project_id>/risks/new")
@login_required
def new_risk(project_id):
    project_or_404(project_id,write=True)
    with db() as conn:
        cur=conn.execute("""INSERT INTO risks(project_id,kind,title,description,probability,impact,owner,action,status,due_date,created_at)
                            VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                         (project_id,request.form.get("kind","Risk"),request.form["title"].strip(),request.form.get("description","").strip(),
                          int(request.form.get("probability","3")),int(request.form.get("impact","3")),request.form.get("owner","").strip(),
                          request.form.get("action","").strip(),request.form.get("status","Öppen"),request.form.get("due_date",""),
                          datetime.now().isoformat(timespec="seconds")))
        rid=cur.lastrowid
    audit(project_id,"risk",rid,"created",request.form["title"].strip())
    return redirect(url_for("project",project_id=project_id)+"#risks")

@app.post("/projects/<int:project_id>/risks/<int:risk_id>/delete")
@login_required
def delete_risk(project_id,risk_id):
    project_or_404(project_id,manager=True)
    with db() as conn: conn.execute("DELETE FROM risks WHERE id=? AND project_id=?",(risk_id,project_id))
    audit(project_id,"risk",risk_id,"deleted","")
    return redirect(url_for("project",project_id=project_id)+"#risks")

@app.post("/projects/<int:project_id>/comments/new")
@login_required
def new_comment(project_id):
    project_or_404(project_id)
    body=request.form.get("body","").strip()
    if body:
        u=current_user()
        with db() as conn: conn.execute("INSERT INTO comments(project_id,user_id,body,created_at) VALUES(?,?,?,?)",
                                        (project_id,u["id"],body,datetime.now().isoformat(timespec="seconds")))
        audit(project_id,"comment",None,"created",body[:120])
    return redirect(url_for("project",project_id=project_id)+"#comments")

@app.post("/projects/<int:project_id>/baseline")
@login_required
def create_baseline(project_id):
    p=project_or_404(project_id,manager=True)
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,))]
        snap=json.dumps({"project":dict(p),"tasks":tasks},ensure_ascii=False)
        u=current_user()
        conn.execute("INSERT INTO baselines(project_id,name,created_at,created_by,snapshot_json) VALUES(?,?,?,?,?)",
                     (project_id,request.form.get("name","").strip() or f"Baseline {datetime.now():%Y-%m-%d %H:%M}",
                      datetime.now().isoformat(timespec="seconds"),u["id"],snap))
    audit(project_id,"baseline",None,"created","Ny baseline")
    return redirect(url_for("project",project_id=project_id)+"#baseline")

@app.route("/projects/<int:project_id>/gantt")
@login_required
def gantt(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND start_date<>'' AND end_date<>'' ORDER BY start_date,wbs,id",(project_id,)).fetchall()
        deps=conn.execute("""SELECT d.task_id,d.depends_on_task_id,t.title dep_title FROM task_dependencies d
                             JOIN tasks t ON t.id=d.depends_on_task_id WHERE d.task_id IN
                             (SELECT id FROM tasks WHERE project_id=?)""",(project_id,)).fetchall()
    if not tasks:return render_template("gantt.html",project=p,rows=[],months=[],range_start=None,range_end=None,deps={})
    starts=[parse_date(t["start_date"]) for t in tasks]; ends=[parse_date(t["end_date"]) for t in tasks]
    rs=min(starts); re=max(ends); total=max((re-rs).days+1,1); rows=[]
    for t in tasks:
        s=parse_date(t["start_date"]); e=parse_date(t["end_date"])
        rows.append(dict(t,left=((s-rs).days/total)*100,width=max(((e-s).days+1)/total*100,.7),derived_status=derived_status(t)))
    months=[]; cur=rs.replace(day=1)
    while cur<=re:
        nextm=(cur.replace(day=28)+timedelta(days=4)).replace(day=1); seg_start=max(cur,rs); seg_end=min(nextm-timedelta(days=1),re)
        months.append({"label":cur.strftime("%Y-%m"),"left":((seg_start-rs).days/total)*100,"width":((seg_end-seg_start).days+1)/total*100}); cur=nextm
    depmap={}
    for d in deps:depmap.setdefault(d["task_id"],[]).append(d["dep_title"])
    return render_template("gantt.html",project=p,rows=rows,months=months,range_start=rs,range_end=re,deps=depmap)


def autosize(ws):
    for col_cells in ws.columns:
        max_len=0
        col=get_column_letter(col_cells[0].column)
        for cell in col_cells:
            max_len=max(max_len,len("" if cell.value is None else str(cell.value)))
        ws.column_dimensions[col].width=min(max(max_len+2,10),42)

def style_header(ws,row=1):
    fill=PatternFill("solid",fgColor="0F4C81")
    font=Font(color="FFFFFF",bold=True)
    thin=Side(style="thin",color="D8DEE8")
    for c in ws[row]:
        c.fill=fill
        c.font=font
        c.alignment=Alignment(vertical="center")
        c.border=Border(bottom=thin)
    ws.row_dimensions[row].height=24

EXCEL_SCHEMA_VERSION="1"
EXCEL_IMPORT_DIR=DATA_DIR/"excel-imports"

def excel_canon_value(value):
    if value is None:
        return ""
    if isinstance(value,(datetime,date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value,bool):
        return 1 if value else 0
    if isinstance(value,float) and value.is_integer():
        return int(value)
    return value

def excel_row_hash(values):
    clean={str(k):excel_canon_value(v) for k,v in values.items()}
    payload=json.dumps(clean,ensure_ascii=False,sort_keys=True,separators=(",",":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:20]

def excel_date(value):
    if value is None or value=="":
        return ""
    if isinstance(value,(datetime,date)):
        return value.strftime("%Y-%m-%d")
    text=str(value).strip()
    for fmt in ("%Y-%m-%d","%Y/%m/%d","%d/%m/%Y","%d-%m-%Y"):
        try:
            return datetime.strptime(text,fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return text

def excel_int(value,default=0):
    if value in (None,""):
        return default
    try:
        return int(float(value))
    except (TypeError,ValueError):
        return default

def excel_float(value,default=0.0):
    if value in (None,""):
        return default
    try:
        return float(value)
    except (TypeError,ValueError):
        return default

def excel_bool(value):
    text=str(value or "").strip().lower()
    return 1 if text in ("1","true","ja","yes","x") else 0

def excel_text(value):
    return "" if value is None else str(value).strip()

def excel_add_validation(ws,column,values,start=2,end=2000):
    dv=DataValidation(type="list",formula1='"'+",".join(values)+'"',allow_blank=True)
    dv.error="Välj ett värde från listan."
    dv.errorTitle="Ogiltigt värde"
    ws.add_data_validation(dv)
    dv.add(f"{column}{start}:{column}{end}")

def excel_prepare_sheet(ws,headers,table_filter=True):
    ws.append(headers)
    style_header(ws)
    ws.freeze_panes="A2"
    if table_filter:
        ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}1"
    ws.sheet_view.showGridLines=False

def excel_hide_meta_columns(ws):
    for idx,cell in enumerate(ws[1],start=1):
        if str(cell.value or "").startswith("_"):
            ws.column_dimensions[get_column_letter(idx)].hidden=True

def excel_project_snapshot(p):
    return {
        "name":p["name"] or "",
        "customer":p["customer"] or "",
        "project_manager":p["project_manager"] or "",
        "description":p["description"] or "",
        "start_date":p["start_date"] or "",
        "end_date":p["end_date"] or "",
    }

def excel_task_snapshot(t):
    return {
        "wbs":t["wbs"] or "","title":t["title"] or "","owner":t["owner"] or "",
        "start_date":t["start_date"] or "","end_date":t["end_date"] or "",
        "actual_start":t["actual_start"] or "","actual_end":t["actual_end"] or "",
        "status":t["status"] or "","priority":t["priority"] or "",
        "progress":excel_int(t["progress"]),"milestone":excel_int(t["milestone"]),
        "notes":t["notes"] or "",
    }

def excel_risk_snapshot(r):
    return {
        "kind":r["kind"] or "Risk","title":r["title"] or "","description":r["description"] or "",
        "probability":excel_int(r["probability"],3),"impact":excel_int(r["impact"],3),
        "owner":r["owner"] or "","action":r["action"] or "","status":r["status"] or "",
        "due_date":r["due_date"] or "",
    }

def excel_change_snapshot(c):
    return {
        "title":c["title"] or "","description":c["description"] or "",
        "impact_scope":c["impact_scope"] or "","impact_days":excel_int(c["impact_days"]),
        "impact_cost":excel_float(c["impact_cost"]),"status":c["status"] or "",
    }

def excel_resource_snapshot(r):
    return {
        "resource_name":r["resource_name"] or "","week_start":r["week_start"] or "",
        "allocation_pct":excel_int(r["allocation_pct"]),"planned_hours":excel_float(r["planned_hours"]),
    }

def excel_cost_snapshot(c):
    return {
        "category":c["category"] or "","description":c["description"] or "",
        "planned":excel_float(c["planned"]),"actual":excel_float(c["actual"]),
        "cost_date":c["cost_date"] or "",
    }

def excel_decision_snapshot(d):
    return {
        "title":d["title"] or "","decision":d["decision"] or "",
        "decided_by":d["decided_by"] or "","decision_date":d["decision_date"] or "",
        "owner":d["owner"] or "",
    }

def excel_meeting_snapshot(m):
    return {
        "title":m["title"] or "","meeting_date":m["meeting_date"] or "",
        "attendees":m["attendees"] or "","notes":m["notes"] or "",
    }

def excel_action_snapshot(a):
    return {
        "title":a["title"] or "","owner":a["owner"] or "",
        "due_date":a["due_date"] or "","status":a["status"] or "",
    }

def excel_benefit_snapshot(b):
    return {
        "title":b["title"] or "","unit":b["unit"] or "%",
        "baseline_value":excel_float(b["baseline_value"]),
        "target_value":excel_float(b["target_value"]),
        "actual_value":None if b["actual_value"] is None else excel_float(b["actual_value"]),
        "measurement_date":b["measurement_date"] or "",
        "owner":b["owner"] or "","status":b["status"] or "",
    }

def build_roundtrip_workbook(project,data,scope="all"):
    wb=Workbook()
    default=wb.active
    wb.remove(default)

    # Project information
    info=wb.create_sheet("Projektinformation")
    info.sheet_view.showGridLines=False
    info["A1"]="Project Planer – Excel Round-trip"
    info["A1"].font=Font(size=18,bold=True,color="0F4C81")
    info.merge_cells("A1:B1")
    info["A3"]="Fält"; info["B3"]="Värde"; style_header(info,3)
    labels=[
        ("Projektnamn","name"),("Kund","customer"),("Projektledare","project_manager"),
        ("Beskrivning","description"),("Plan start","start_date"),("Plan slut","end_date")
    ]
    ps=excel_project_snapshot(project)
    for label,key in labels:
        info.append([label,ps[key]])
    info["A11"]="Instruktion"
    info["B11"]="Redigera värden i kolumn B. Importen visar alltid en förhandsgranskning innan något skrivs tillbaka."
    info["B11"].alignment=Alignment(wrap_text=True,vertical="top")
    info.column_dimensions["A"].width=24
    info.column_dimensions["B"].width=70

    if scope in ("all","plan"):
        ws=wb.create_sheet("Uppgifter")
        headers=["WBS","Aktivitet","Ansvarig","Plan start","Plan slut","Faktisk start","Faktiskt slut","Status","Prioritet","Progress %","Milstolpe","Kommentar","_ID","_Hash"]
        excel_prepare_sheet(ws,headers)
        for t in data["tasks"]:
            snap=excel_task_snapshot(t)
            ws.append([snap["wbs"],snap["title"],snap["owner"],snap["start_date"],snap["end_date"],snap["actual_start"],snap["actual_end"],snap["status"],snap["priority"],snap["progress"],"Ja" if snap["milestone"] else "Nej",snap["notes"],t["id"],excel_row_hash(snap)])
        excel_add_validation(ws,"H",STATUSES)
        excel_add_validation(ws,"I",PRIORITIES)
        excel_add_validation(ws,"K",["Ja","Nej"])
        if ws.max_row>1:
            ws.conditional_formatting.add(f"J2:J{ws.max_row}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=100,color="0F4C81"))
        excel_hide_meta_columns(ws); autosize(ws)
        ws.column_dimensions["B"].width=34; ws.column_dimensions["L"].width=40

        ms=wb.create_sheet("Milstolpar")
        excel_prepare_sheet(ms,["WBS","Milstolpe","Planerat datum","Status","Progress %"])
        for t in data["tasks"]:
            if excel_int(t["milestone"]):
                ms.append([t["wbs"],t["title"],t["end_date"],t["status"],t["progress"]])
        ms["A1"].comment=None
        ms.sheet_properties.tabColor="94A3B8"
        autosize(ms)

        dep=wb.create_sheet("Beroenden")
        excel_prepare_sheet(dep,["Föregående WBS","Efterföljande WBS","Typ","Förskjutning dagar","_ID","_Hash"])
        task_by_id={int(t["id"]):t for t in data["tasks"]}
        for l in data["links"]:
            pred=task_by_id.get(int(l["predecessor_id"]))
            succ=task_by_id.get(int(l["successor_id"]))
            snap={"predecessor_wbs":pred["wbs"] if pred else "","successor_wbs":succ["wbs"] if succ else "","link_type":l["link_type"] or "FS","lag_days":excel_int(l["lag_days"])}
            dep.append([snap["predecessor_wbs"],snap["successor_wbs"],snap["link_type"],snap["lag_days"],l["id"],excel_row_hash(snap)])
        excel_add_validation(dep,"C",["FS","SS","FF","SF"])
        excel_hide_meta_columns(dep); autosize(dep)

    if scope in ("all","risk"):
        rw=wb.create_sheet("Risker")
        excel_prepare_sheet(rw,["Typ","Titel","Beskrivning","Sannolikhet","Konsekvens","Ansvarig","Åtgärd","Status","Förfallodatum","_ID","_Hash"])
        for r in data["risks"]:
            snap=excel_risk_snapshot(r)
            rw.append([snap["kind"],snap["title"],snap["description"],snap["probability"],snap["impact"],snap["owner"],snap["action"],snap["status"],snap["due_date"],r["id"],excel_row_hash(snap)])
        excel_add_validation(rw,"A",["Risk","Issue"])
        excel_add_validation(rw,"H",RISK_STATUSES)
        excel_hide_meta_columns(rw); autosize(rw); rw.column_dimensions["C"].width=38; rw.column_dimensions["G"].width=38

        cw=wb.create_sheet("Ändringsärenden")
        excel_prepare_sheet(cw,["Rubrik","Beskrivning","Omfattningspåverkan","Dagar","Kostnad","Status","_ID","_Hash"])
        for c in data["changes"]:
            snap=excel_change_snapshot(c)
            cw.append([snap["title"],snap["description"],snap["impact_scope"],snap["impact_days"],snap["impact_cost"],snap["status"],c["id"],excel_row_hash(snap)])
        excel_add_validation(cw,"F",["Proposed","Submitted","Pending","Approved","Rejected","Closed"])
        excel_hide_meta_columns(cw); autosize(cw); cw.column_dimensions["B"].width=38; cw.column_dimensions["C"].width=34

        dw=wb.create_sheet("Beslut")
        excel_prepare_sheet(dw,["Titel","Beslut","Beslutat av","Beslutsdatum","Ansvarig","_ID","_Hash"])
        for d in data["decisions"]:
            snap=excel_decision_snapshot(d)
            dw.append([snap["title"],snap["decision"],snap["decided_by"],snap["decision_date"],snap["owner"],d["id"],excel_row_hash(snap)])
        excel_hide_meta_columns(dw); autosize(dw); dw.column_dimensions["B"].width=48

    if scope in ("all","resources"):
        rs=wb.create_sheet("Resurser")
        excel_prepare_sheet(rs,["Resurs","Vecka","Allokering %","Planerade timmar","_ID","_Hash"])
        for r in data["resources"]:
            snap=excel_resource_snapshot(r)
            rs.append([snap["resource_name"],snap["week_start"],snap["allocation_pct"],snap["planned_hours"],r["id"],excel_row_hash(snap)])
        excel_hide_meta_columns(rs); autosize(rs)

    if scope in ("all","finance"):
        co=wb.create_sheet("Kostnader")
        excel_prepare_sheet(co,["Kategori","Beskrivning","Planerat","Utfall","Datum","_ID","_Hash"])
        for c in data["costs"]:
            snap=excel_cost_snapshot(c)
            co.append([snap["category"],snap["description"],snap["planned"],snap["actual"],snap["cost_date"],c["id"],excel_row_hash(snap)])
        excel_hide_meta_columns(co); autosize(co); co.column_dimensions["B"].width=40

        be=wb.create_sheet("Nyttor")
        excel_prepare_sheet(be,["Titel","Enhet","Baslinje","Mål","Utfall","Mätdatum","Ansvarig","Status","_ID","_Hash"])
        for b in data["benefits"]:
            snap=excel_benefit_snapshot(b)
            be.append([snap["title"],snap["unit"],snap["baseline_value"],snap["target_value"],snap["actual_value"],snap["measurement_date"],snap["owner"],snap["status"],b["id"],excel_row_hash(snap)])
        excel_add_validation(be,"H",["Planned","Measured","Realized","Closed"])
        excel_hide_meta_columns(be); autosize(be)

    if scope=="all":
        mw=wb.create_sheet("Möten")
        excel_prepare_sheet(mw,["Titel","Datum","Deltagare","Anteckningar","_ID","_Hash"])
        for m in data["meetings"]:
            snap=excel_meeting_snapshot(m)
            mw.append([snap["title"],snap["meeting_date"],snap["attendees"],snap["notes"],m["id"],excel_row_hash(snap)])
        excel_hide_meta_columns(mw); autosize(mw); mw.column_dimensions["D"].width=45

        aw=wb.create_sheet("Åtgärder")
        excel_prepare_sheet(aw,["Titel","Ansvarig","Förfallodatum","Status","_ID","_Hash"])
        for a in data["actions"]:
            snap=excel_action_snapshot(a)
            aw.append([snap["title"],snap["owner"],snap["due_date"],snap["status"],a["id"],excel_row_hash(snap)])
        excel_add_validation(aw,"D",["Open","In Progress","Blocked","Done","Closed"])
        excel_hide_meta_columns(aw); autosize(aw)

    meta=wb.create_sheet("_Metadata")
    meta.sheet_state="veryHidden"
    meta.append(["key","value"])
    meta.append(["schema_version",EXCEL_SCHEMA_VERSION])
    meta.append(["project_id",project["id"]])
    meta.append(["project_name",project["name"]])
    meta.append(["project_hash",excel_row_hash(ps)])
    meta.append(["app_version",APP_VERSION])
    meta.append(["exported_at",datetime.now().isoformat(timespec="seconds")])
    meta.append(["scope",scope])
    return excel_pro_finalize_workbook(wb,project,data,scope)

def excel_project_data(project_id):
    with db() as conn:
        return {
            "tasks":conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY wbs,id",(project_id,)).fetchall(),
            "links":conn.execute("SELECT * FROM task_links WHERE project_id=? ORDER BY id",(project_id,)).fetchall(),
            "risks":conn.execute("SELECT * FROM risks WHERE project_id=? ORDER BY id",(project_id,)).fetchall(),
            "changes":conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id",(project_id,)).fetchall(),
            "resources":conn.execute("SELECT * FROM resource_allocations WHERE project_id=? ORDER BY resource_name,week_start,id",(project_id,)).fetchall(),
            "costs":conn.execute("SELECT * FROM project_costs WHERE project_id=? ORDER BY cost_date,id",(project_id,)).fetchall(),
            "decisions":conn.execute("SELECT * FROM decisions WHERE project_id=? ORDER BY decision_date,id",(project_id,)).fetchall(),
            "meetings":conn.execute("SELECT * FROM meetings WHERE project_id=? ORDER BY meeting_date,id",(project_id,)).fetchall(),
            "actions":conn.execute("SELECT * FROM action_items WHERE project_id=? ORDER BY due_date,id",(project_id,)).fetchall(),
            "benefits":conn.execute("SELECT * FROM project_benefits WHERE project_id=? ORDER BY id",(project_id,)).fetchall(),
        }


EXCEL_PRO_HEADER="163A5F"
EXCEL_PRO_ACCENT="0F4C81"
EXCEL_PRO_INPUT="EAF5EA"
EXCEL_PRO_READONLY="F3F4F6"
EXCEL_PRO_RED="FEE2E2"
EXCEL_PRO_AMBER="FEF3C7"
EXCEL_PRO_GREEN="DCFCE7"

def excel_safe_table_name(name):
    clean=re.sub(r"[^A-Za-z0-9_]","_",name or "Table")
    if not clean or clean[0].isdigit(): clean="T_"+clean
    return clean[:200]

def excel_add_table(ws,name):
    if ws.max_row<2 or ws.max_column<1: return
    tab=Table(displayName=excel_safe_table_name(name),ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
    tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
    ws.add_table(tab)

def excel_page_setup(ws,landscape=True):
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape" if landscape else "portrait"
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.oddFooter.center.text=f"Project Planer v{APP_VERSION} · {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws.oddFooter.center.size=8

def excel_to_date(v):
    if not isinstance(v,str) or not v: return v
    try: return datetime.strptime(v[:10],"%Y-%m-%d").date()
    except: return v

def excel_pro_style_sheet(ws,editable_headers=None,readonly=False,table_name=None):
    editable_headers=set(editable_headers or [])
    ws.freeze_panes="A2"; ws.sheet_view.showGridLines=False; style_header(ws,1)
    headers={c.column:excel_text(c.value) for c in ws[1]}
    date_headers={"Plan start","Plan slut","Faktisk start","Faktiskt slut","Förfallodatum","Datum","Vecka","Beslutsdatum","Mätdatum","Slutdatum"}
    money_headers={"Kostnad","Planerat","Utfall","Baslinje","Mål"}
    for col,header in headers.items():
        letter=get_column_letter(col)
        if header.startswith("_"):
            ws.column_dimensions[letter].hidden=True; continue
        fill=EXCEL_PRO_READONLY if readonly or header not in editable_headers else EXCEL_PRO_INPUT
        for row in range(2,ws.max_row+1):
            cell=ws.cell(row,col); cell.fill=PatternFill("solid",fgColor=fill)
            cell.alignment=Alignment(vertical="top",wrap_text=header in {"Aktivitet","Kommentar","Beskrivning","Åtgärd","Beslut","Anteckningar"})
            if header in date_headers:
                cell.value=excel_to_date(cell.value); cell.number_format="yyyy-mm-dd"
            elif header in money_headers: cell.number_format='#,##0.00'
    if "Progress %" in [c.value for c in ws[1]] and ws.max_row>=2:
        col=[c.column for c in ws[1] if c.value=="Progress %"][0]; letter=get_column_letter(col)
        ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=100,color=EXCEL_PRO_ACCENT))
    if "Status" in [c.value for c in ws[1]] and ws.max_row>=2:
        col=[c.column for c in ws[1] if c.value=="Status"][0]; letter=get_column_letter(col); rng=f"{letter}2:{letter}{ws.max_row}"
        ws.conditional_formatting.add(rng,FormulaRule(formula=[f'OR(${letter}2="Blockerad",${letter}2="Blocked")'],fill=PatternFill("solid",fgColor=EXCEL_PRO_RED)))
        ws.conditional_formatting.add(rng,FormulaRule(formula=[f'OR(${letter}2="Klar",${letter}2="Done",${letter}2="Closed",${letter}2="Stängd")'],fill=PatternFill("solid",fgColor=EXCEL_PRO_GREEN)))
    autosize(ws)
    for c in ws[1]:
        h=excel_text(c.value); letter=get_column_letter(c.column)
        if h in {"Aktivitet","Titel","Rubrik","Milstolpe"}: ws.column_dimensions[letter].width=34
        elif h in {"Kommentar","Beskrivning","Åtgärd","Beslut","Anteckningar"}: ws.column_dimensions[letter].width=42
    if table_name: excel_add_table(ws,table_name)
    excel_page_setup(ws,True)

def excel_pro_overview(wb,project,data):
    ws=wb.create_sheet("Översikt",0); ws.sheet_view.showGridLines=False
    ws["A1"]="Project Planer"; ws["A1"].font=Font(size=11,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=EXCEL_PRO_HEADER)
    ws["B1"]=f"Excel Pro · v{APP_VERSION}"; ws["B1"].font=Font(size=11,color="FFFFFF"); ws["B1"].fill=PatternFill("solid",fgColor=EXCEL_PRO_HEADER); ws.merge_cells("B1:F1")
    ws["A3"]=project["name"]; ws["A3"].font=Font(size=22,bold=True,color=EXCEL_PRO_HEADER); ws.merge_cells("A3:F3")
    ws["A4"]=f"{project['customer'] or 'Ingen kund angiven'} · Projektledare: {project['project_manager'] or 'Ej satt'}"; ws["A4"].font=Font(color="667085"); ws.merge_cells("A4:F4")
    tasks=data.get("tasks",[]); risks=data.get("risks",[]); costs=data.get("costs",[])
    milestones=[t for t in tasks if excel_int(t["milestone"])]
    overdue=[t for t in tasks if t["end_date"] and excel_int(t["progress"])<100 and str(t["end_date"])<date.today().isoformat()]
    blocked=[t for t in tasks if (t["status"] or "").lower() in ("blocked","blockerad")]
    high=[r for r in risks if excel_int(r["probability"])*excel_int(r["impact"])>=15]
    progress=round(sum(excel_int(t["progress"]) for t in tasks)/len(tasks)) if tasks else 0
    planned=sum(excel_float(c["planned"]) for c in costs); actual=sum(excel_float(c["actual"]) for c in costs)
    rag="RÖD" if len(overdue)>=3 or len(high)>=2 else ("GUL" if overdue or high or blocked else "GRÖN")
    kpis=[("RAG",rag),("Framdrift",f"{progress}%"),("Försenade",len(overdue)),("Blockerade",len(blocked)),("Höga risker",len(high)),("Milstolpar",len(milestones)),("Planerad kostnad",planned),("Utfall",actual)]
    for i,(label,val) in enumerate(kpis):
        col=1+(i%4)*2; row=6+(i//4)*3
        ws.cell(row,col,label).font=Font(size=9,color="667085",bold=True); ws.cell(row+1,col,val).font=Font(size=16,bold=True,color=EXCEL_PRO_HEADER)
        ws.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+1); ws.merge_cells(start_row=row+1,start_column=col,end_row=row+1,end_column=col+1)
        for rr in (row,row+1):
            for cc in (col,col+1): ws.cell(rr,cc).fill=PatternFill("solid",fgColor="F8FAFC")
    fill={"RÖD":EXCEL_PRO_RED,"GUL":EXCEL_PRO_AMBER,"GRÖN":EXCEL_PRO_GREEN}[rag]
    ws["A7"].fill=PatternFill("solid",fgColor=fill); ws["B7"].fill=PatternFill("solid",fgColor=fill)
    ws["A13"]="Ledningsbild"; ws["A13"].font=Font(size=14,bold=True,color=EXCEL_PRO_HEADER); ws.merge_cells("A13:F13")
    ws["A14"]=f"Projektet är {progress}% klart. {len(overdue)} aktiviteter är försenade, {len(blocked)} blockerade och {len(high)} höga risker är öppna. Planerat slutdatum är {project['end_date'] or 'inte satt'}."; ws.merge_cells("A14:F16"); ws["A14"].alignment=Alignment(wrap_text=True,vertical="top"); ws["A14"].fill=PatternFill("solid",fgColor="F8FAFC")
    ws["A18"]="Kommande milstolpar"; ws["A18"].font=Font(size=13,bold=True,color=EXCEL_PRO_HEADER)
    for c,v in enumerate(["Datum","Milstolpe","Status","Progress"],1): ws.cell(19,c,v); ws.cell(19,c).fill=PatternFill("solid",fgColor=EXCEL_PRO_HEADER); ws.cell(19,c).font=Font(color="FFFFFF",bold=True)
    for i,m in enumerate(sorted(milestones,key=lambda x:(x["end_date"] or "9999",x["id"]))[:8],20):
        ws.cell(i,1,excel_to_date(m["end_date"])); ws.cell(i,1).number_format="yyyy-mm-dd"; ws.cell(i,2,m["title"]); ws.cell(i,3,m["status"]); ws.cell(i,4,excel_int(m["progress"]))
    ws["J2"]="Status"; ws["K2"]="Antal"; counts={}
    for t in tasks: counts[t["status"] or "Ej satt"]=counts.get(t["status"] or "Ej satt",0)+1
    row=3
    for key,val in counts.items(): ws.cell(row,10,key); ws.cell(row,11,val); row+=1
    if row>3:
        ch=DoughnutChart(); ch.title="Aktiviteter per status"; ch.add_data(Reference(ws,min_col=11,min_row=2,max_row=row-1),titles_from_data=True); ch.set_categories(Reference(ws,min_col=10,min_row=3,max_row=row-1)); ch.height=7; ch.width=10; ch.dataLabels=DataLabelList(); ch.dataLabels.showPercent=True; ws.add_chart(ch,"F18")
    ws["J15"]="Ekonomi"; ws["K15"]="Belopp"; ws["J16"]="Planerat"; ws["K16"]=planned; ws["J17"]="Utfall"; ws["K17"]=actual
    if planned or actual:
        ch=BarChart(); ch.type="col"; ch.title="Ekonomi"; ch.legend=None; ch.add_data(Reference(ws,min_col=11,min_row=15,max_row=17),titles_from_data=True); ch.set_categories(Reference(ws,min_col=10,min_row=16,max_row=17)); ch.height=6; ch.width=10; ws.add_chart(ch,"F32")
    ws.column_dimensions["J"].hidden=True; ws.column_dimensions["K"].hidden=True; ws.freeze_panes="A5"
    for col,w in {"A":18,"B":35,"C":18,"D":14,"E":14,"F":14,"G":14,"H":14}.items(): ws.column_dimensions[col].width=w
    excel_page_setup(ws,True); ws.print_area=f"A1:H{max(42,ws.max_row)}"

def excel_pro_finalize_workbook(wb,project,data,scope="all"):
    wb.properties.title=f"{project['name']} – Project Planer"; wb.properties.subject="Excel Pro / Round-trip"; wb.properties.creator="Project Planer"
    wb.properties.description=f"Project Planer v{APP_VERSION} · {datetime.now().isoformat(timespec='seconds')}"
    if "Översikt" not in wb.sheetnames: excel_pro_overview(wb,project,data)
    specs={
      "Uppgifter":({"WBS","Aktivitet","Ansvarig","Plan start","Plan slut","Faktisk start","Faktiskt slut","Status","Prioritet","Progress %","Milstolpe","Kommentar"},False,"TasksTable"),
      "Milstolpar":(set(),True,"MilestonesView"),
      "Risker":({"Typ","Titel","Beskrivning","Sannolikhet","Konsekvens","Ansvarig","Åtgärd","Status","Förfallodatum"},False,"RisksTable"),
      "Ändringsärenden":({"Rubrik","Beskrivning","Omfattningspåverkan","Dagar","Kostnad","Status"},False,"ChangesTable"),
      "Resurser":({"Resurs","Vecka","Allokering %","Planerade timmar"},False,"ResourcesTable"),
      "Kostnader":({"Kategori","Beskrivning","Planerat","Utfall","Datum"},False,"CostsTable"),
      "Beroenden":({"Föregående WBS","Efterföljande WBS","Typ","Förskjutning dagar"},False,"DependenciesTable"),
      "Beslut":({"Titel","Beslut","Beslutat av","Beslutsdatum","Ansvarig"},False,"DecisionsTable"),
      "Möten":({"Titel","Datum","Deltagare","Anteckningar"},False,"MeetingsTable"),
      "Åtgärder":({"Titel","Ansvarig","Förfallodatum","Status"},False,"ActionsTable"),
      "Nyttor":({"Titel","Enhet","Baslinje","Mål","Utfall","Mätdatum","Ansvarig","Status"},False,"BenefitsTable")}
    for name,(editable,readonly,table) in specs.items():
        if name in wb.sheetnames: excel_pro_style_sheet(wb[name],editable,readonly,table)
    if "Projektinformation" in wb.sheetnames:
        info=wb["Projektinformation"]; info.sheet_view.showGridLines=False
        for row in range(4,10): info.cell(row,2).fill=PatternFill("solid",fgColor=EXCEL_PRO_INPUT)
        info["A12"]="Färgförklaring"; info["B12"]="Grönt = redigerbart. Grått/blått = rapport eller beräknat. Dolda tekniska kolumner används för säker round-trip."; info["B12"].alignment=Alignment(wrap_text=True); excel_page_setup(info,False)
    if "_Metadata" in wb.sheetnames: wb["_Metadata"].sheet_state="veryHidden"
    return wb

def excel_pro_report_finalize_v902(wb,payload):
    ws=wb["Projektstatus"]; ws.sheet_view.showGridLines=False; ws.freeze_panes="A3"; ws["D3"]="Excel Pro"; ws["D3"].font=Font(bold=True,color=EXCEL_PRO_ACCENT); ws["D4"]=f"v{APP_VERSION}"; ws["D5"]=datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["A18"]="Användning"; ws["B18"]="Rapportexport. För redigering och återimport används Excel Round-trip."; ws["B18"].alignment=Alignment(wrap_text=True); excel_page_setup(ws,False)
    for name in ["Milstolpar","Höga risker","Försenade","Ändringar","Beslut"]:
        if name in wb.sheetnames: excel_pro_style_sheet(wb[name],set(),True,f"Report_{name}")
    dash=wb.create_sheet("Dashboard",0); dash.sheet_view.showGridLines=False; dash["A1"]="Project Planer – Statusdashboard"; dash["A1"].font=Font(size=20,bold=True,color=EXCEL_PRO_HEADER); dash.merge_cells("A1:F1"); dash["A3"]=payload["project"]["name"]; dash["A3"].font=Font(size=16,bold=True); dash.merge_cells("A3:F3")
    metrics=[("RAG",payload["rag"].upper()),("Framdrift",f"{payload['progress']}%"),("Försenade",len(payload["overdue"])),("Höga risker",len(payload["high"])),("Planerat slut",payload["project"]["end_date"] or "–"),("Prognos slut",payload["forecast"]["forecast_end"] or "–")]
    for i,(label,val) in enumerate(metrics):
        row=5+(i//3)*3; col=1+(i%3)*2; dash.cell(row,col,label).font=Font(color="667085",bold=True); dash.cell(row+1,col,val).font=Font(size=16,bold=True,color=EXCEL_PRO_HEADER); dash.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+1); dash.merge_cells(start_row=row+1,start_column=col,end_row=row+1,end_column=col+1)
    planned=float(payload["costs"]["planned"] or 0); actual=float(payload["costs"]["actual"] or 0); dash["H2"]="Ekonomi"; dash["I2"]="Belopp"; dash["H3"]="Planerat"; dash["I3"]=planned; dash["H4"]="Utfall"; dash["I4"]=actual
    ch=BarChart(); ch.type="col"; ch.title="Ekonomi"; ch.legend=None; ch.add_data(Reference(dash,min_col=9,min_row=2,max_row=4),titles_from_data=True); ch.set_categories(Reference(dash,min_col=8,min_row=3,max_row=4)); ch.height=7; ch.width=10; dash.add_chart(ch,"A12"); dash.column_dimensions["H"].hidden=True; dash.column_dimensions["I"].hidden=True; excel_page_setup(dash,True)
    return wb

def build_portfolio_excel_v902(projects):
    wb=Workbook(); ws=wb.active; ws.title="Portfolio"; ws.sheet_view.showGridLines=False
    ws["A1"]="Project Planer – Portfolio Excel"; ws["A1"].font=Font(size=20,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=EXCEL_PRO_HEADER); ws.merge_cells("A1:L1"); ws["A2"]=f"Exporterad {datetime.now().strftime('%Y-%m-%d %H:%M')} · v{APP_VERSION}"; ws.merge_cells("A2:L2")
    headers=["Projekt","Kund","Projektledare","RAG","Framdrift %","Planerat slut","Prognos slut","Försenade","Blockerade","Höga risker","Planerad kostnad","Utfall"]
    ws.append([]); ws.append(headers); style_header(ws,4)
    for p in projects:
        pid=p["id"]; f=forecast_project_v850(pid)
        with db() as conn:
            tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL",(pid,)).fetchall(); risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status NOT IN ('Closed','Stängd')",(pid,)).fetchall(); costs=conn.execute("SELECT COALESCE(SUM(planned),0) p,COALESCE(SUM(actual),0) a FROM project_costs WHERE project_id=?",(pid,)).fetchone()
        overdue=[t for t in tasks if t["end_date"] and excel_int(t["progress"])<100 and t["end_date"]<date.today().isoformat()]; blocked=[t for t in tasks if (t["status"] or "").lower() in ("blocked","blockerad")]; high=[r for r in risks if excel_int(r["probability"])*excel_int(r["impact"])>=15]; progress=round(sum(excel_int(t["progress"]) for t in tasks)/len(tasks)) if tasks else 0; rag="Röd" if len(overdue)>=3 or len(high)>=2 else ("Gul" if overdue or high or blocked else "Grön")
        ws.append([p["name"],p["customer"],p["project_manager"],rag,progress,excel_to_date(p["end_date"]),excel_to_date(f["forecast_end"]),len(overdue),len(blocked),len(high),float(costs["p"] or 0),float(costs["a"] or 0)])
    if ws.max_row>=5:
        excel_add_table(ws,"PortfolioTable"); ws.conditional_formatting.add(f"E5:E{ws.max_row}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=100,color=EXCEL_PRO_ACCENT)); ws.conditional_formatting.add(f"D5:D{ws.max_row}",FormulaRule(formula=['$D5="Röd"'],fill=PatternFill("solid",fgColor=EXCEL_PRO_RED))); ws.conditional_formatting.add(f"D5:D{ws.max_row}",FormulaRule(formula=['$D5="Gul"'],fill=PatternFill("solid",fgColor=EXCEL_PRO_AMBER))); ws.conditional_formatting.add(f"D5:D{ws.max_row}",FormulaRule(formula=['$D5="Grön"'],fill=PatternFill("solid",fgColor=EXCEL_PRO_GREEN)))
        for row in range(5,ws.max_row+1): ws.cell(row,6).number_format="yyyy-mm-dd"; ws.cell(row,7).number_format="yyyy-mm-dd"; ws.cell(row,11).number_format='#,##0.00'; ws.cell(row,12).number_format='#,##0.00'
    ws.freeze_panes="A5"; autosize(ws); ws.column_dimensions["A"].width=30; ws.column_dimensions["B"].width=24; ws.column_dimensions["C"].width=24; excel_page_setup(ws,True)
    dash=wb.create_sheet("Översikt",0); dash.sheet_view.showGridLines=False; dash["A1"]="Portfolioöversikt"; dash["A1"].font=Font(size=22,bold=True,color=EXCEL_PRO_HEADER); dash.merge_cells("A1:F1")
    total=len(projects); vals=list(ws.iter_rows(min_row=5,values_only=True)) if ws.max_row>=5 else []; red=sum(1 for r in vals if r[3]=="Röd"); amber=sum(1 for r in vals if r[3]=="Gul"); green=sum(1 for r in vals if r[3]=="Grön")
    for i,(label,val) in enumerate([("Projekt",total),("Röda",red),("Gula",amber),("Gröna",green)]): col=1+i*2; dash.cell(3,col,label).font=Font(color="667085",bold=True); dash.cell(4,col,val).font=Font(size=18,bold=True,color=EXCEL_PRO_HEADER)
    dash["H2"]="RAG"; dash["I2"]="Antal"
    for rr,(label,val) in enumerate([("Röd",red),("Gul",amber),("Grön",green)],3): dash.cell(rr,8,label); dash.cell(rr,9,val)
    if total:
        ch=DoughnutChart(); ch.title="Portfolio RAG"; ch.add_data(Reference(dash,min_col=9,min_row=2,max_row=5),titles_from_data=True); ch.set_categories(Reference(dash,min_col=8,min_row=3,max_row=5)); ch.height=8; ch.width=12; ch.dataLabels=DataLabelList(); ch.dataLabels.showPercent=True; dash.add_chart(ch,"A7")
    dash.column_dimensions["H"].hidden=True; dash.column_dimensions["I"].hidden=True; excel_page_setup(dash,True)
    return wb

@app.get("/excel/portfolio")
@login_required
def excel_portfolio_v902():
    wb=build_portfolio_excel_v902(visible_projects_for_user()); bio=BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=f"Project-Planer-Portfolio-{date.today().isoformat()}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/projects/<int:project_id>/excel")
@login_required
def excel_center_v810(project_id):
    p=project_or_404(project_id)
    return render_template("excel_center_v810.html",project=p)

@app.get("/projects/<int:project_id>/excel/export")
@login_required
def excel_export_v810(project_id):
    p=project_or_404(project_id)
    scope=request.args.get("scope","all")
    if scope not in ("all","plan","risk","resources","finance"):
        scope="all"
    wb=build_roundtrip_workbook(p,excel_project_data(project_id),scope)
    bio=BytesIO()
    wb.save(bio)
    bio.seek(0)
    suffix="" if scope=="all" else f"-{scope}"
    audit(project_id,"project",project_id,"excel_export",f"scope={scope}; app={APP_VERSION}")
    return send_file(
        bio,as_attachment=True,
        download_name=f"{p['name']}-project-plan{suffix}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# Legacy Excel export now produces the full round-trip workbook.
@app.get("/projects/<int:project_id>/export.xlsx")
@login_required
def export_project(project_id):
    p=project_or_404(project_id)
    wb=build_roundtrip_workbook(p,excel_project_data(project_id),"all")
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    audit(project_id,"project",project_id,"excel_export","scope=all; legacy-route")
    return send_file(bio,as_attachment=True,download_name=f"{p['name']}-project-plan.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

EXCEL_SPECS={
    "Uppgifter":{
        "entity":"Uppgift","table":"tasks","required":"title",
        "headers":{"WBS":"wbs","Aktivitet":"title","Ansvarig":"owner","Plan start":"start_date","Plan slut":"end_date","Faktisk start":"actual_start","Faktiskt slut":"actual_end","Status":"status","Prioritet":"priority","Progress %":"progress","Milstolpe":"milestone","Kommentar":"notes"},
        "snapshot":excel_task_snapshot,
        "converters":{"start_date":excel_date,"end_date":excel_date,"actual_start":excel_date,"actual_end":excel_date,"progress":lambda v: max(0,min(100,excel_int(v))),"milestone":excel_bool},
    },
    "Risker":{
        "entity":"Risk","table":"risks","required":"title",
        "headers":{"Typ":"kind","Titel":"title","Beskrivning":"description","Sannolikhet":"probability","Konsekvens":"impact","Ansvarig":"owner","Åtgärd":"action","Status":"status","Förfallodatum":"due_date"},
        "snapshot":excel_risk_snapshot,
        "converters":{"probability":lambda v:max(1,min(5,excel_int(v,3))),"impact":lambda v:max(1,min(5,excel_int(v,3))),"due_date":excel_date},
    },
    "Ändringsärenden":{
        "entity":"Ändringsärende","table":"change_requests","required":"title",
        "headers":{"Rubrik":"title","Beskrivning":"description","Omfattningspåverkan":"impact_scope","Dagar":"impact_days","Kostnad":"impact_cost","Status":"status"},
        "snapshot":excel_change_snapshot,
        "converters":{"impact_days":excel_int,"impact_cost":excel_float},
    },
    "Resurser":{
        "entity":"Resursallokering","table":"resource_allocations","required":"resource_name",
        "headers":{"Resurs":"resource_name","Vecka":"week_start","Allokering %":"allocation_pct","Planerade timmar":"planned_hours"},
        "snapshot":excel_resource_snapshot,
        "converters":{"week_start":excel_date,"allocation_pct":lambda v:max(0,min(300,excel_int(v))),"planned_hours":excel_float},
    },
    "Kostnader":{
        "entity":"Kostnad","table":"project_costs","required":"description",
        "headers":{"Kategori":"category","Beskrivning":"description","Planerat":"planned","Utfall":"actual","Datum":"cost_date"},
        "snapshot":excel_cost_snapshot,
        "converters":{"planned":excel_float,"actual":excel_float,"cost_date":excel_date},
    },
    "Beslut":{
        "entity":"Beslut","table":"decisions","required":"title",
        "headers":{"Titel":"title","Beslut":"decision","Beslutat av":"decided_by","Beslutsdatum":"decision_date","Ansvarig":"owner"},
        "snapshot":excel_decision_snapshot,
        "converters":{"decision_date":excel_date},
    },
    "Möten":{
        "entity":"Möte","table":"meetings","required":"title",
        "headers":{"Titel":"title","Datum":"meeting_date","Deltagare":"attendees","Anteckningar":"notes"},
        "snapshot":excel_meeting_snapshot,
        "converters":{"meeting_date":excel_date},
    },
    "Åtgärder":{
        "entity":"Åtgärd","table":"action_items","required":"title",
        "headers":{"Titel":"title","Ansvarig":"owner","Förfallodatum":"due_date","Status":"status"},
        "snapshot":excel_action_snapshot,
        "converters":{"due_date":excel_date},
    },
    "Nyttor":{
        "entity":"Nytta","table":"project_benefits","required":"title",
        "headers":{"Titel":"title","Enhet":"unit","Baslinje":"baseline_value","Mål":"target_value","Utfall":"actual_value","Mätdatum":"measurement_date","Ansvarig":"owner","Status":"status"},
        "snapshot":excel_benefit_snapshot,
        "converters":{"baseline_value":excel_float,"target_value":excel_float,"actual_value":lambda v:None if v in (None,"") else excel_float(v),"measurement_date":excel_date},
    },
}

def excel_normalize_row(spec,raw):
    result={}
    converters=spec.get("converters",{})
    for field,value in raw.items():
        conv=converters.get(field,excel_text)
        result[field]=conv(value)
    return result

def excel_headers(ws):
    return {excel_text(cell.value):idx for idx,cell in enumerate(ws[1],start=1) if cell.value is not None}

def excel_parse_metadata(wb):
    if "_Metadata" not in wb.sheetnames:
        return {}
    ws=wb["_Metadata"]
    return {excel_text(r[0].value):excel_text(r[1].value) for r in ws.iter_rows(min_row=2,max_col=2) if r[0].value}

def excel_current_rows(conn,table,project_id):
    return {int(r["id"]):r for r in conn.execute(f"SELECT * FROM {table} WHERE project_id=?",(project_id,)).fetchall()}

def excel_diff(old,new):
    changes={}
    for key in new:
        if excel_canon_value(old.get(key)) != excel_canon_value(new.get(key)):
            changes[key]={"old":excel_canon_value(old.get(key)),"new":excel_canon_value(new.get(key))}
    return changes

def excel_import_preview(project_id,file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError("Välj en Excel-fil.")
    if not file_storage.filename.lower().endswith(".xlsx"):
        raise ValueError("Endast .xlsx stöds.")
    payload=file_storage.read()
    if len(payload)>15*1024*1024:
        raise ValueError("Excel-filen är större än 15 MB.")
    try:
        wb=load_workbook(BytesIO(payload),data_only=False)
    except Exception as ex:
        raise ValueError(f"Kunde inte läsa Excel-filen: {ex}")

    meta=excel_parse_metadata(wb)
    if meta.get("schema_version")!=EXCEL_SCHEMA_VERSION:
        raise ValueError("Filen är inte en kompatibel Project Planer Excel Round-trip-fil.")
    if excel_int(meta.get("project_id"))!=project_id:
        raise ValueError("Excel-filen tillhör ett annat projekt.")

    preview={"project_id":project_id,"exported_at":meta.get("exported_at",""),"app_version":meta.get("app_version",""),"items":[],"errors":[],"counts":{"create":0,"update":0,"conflict":0,"error":0,"unchanged":0}}

    with db() as conn:
        project=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
        # Project information sheet.
        if "Projektinformation" in wb.sheetnames:
            ws=wb["Projektinformation"]
            label_to_field={"Projektnamn":"name","Kund":"customer","Projektledare":"project_manager","Beskrivning":"description","Plan start":"start_date","Plan slut":"end_date"}
            raw={}
            for row in ws.iter_rows(min_row=4,max_col=2):
                label=excel_text(row[0].value)
                if label in label_to_field:
                    field=label_to_field[label]
                    raw[field]=excel_date(row[1].value) if field in ("start_date","end_date") else excel_text(row[1].value)
            current=excel_project_snapshot(project)
            current_hash=excel_row_hash(current)
            exported_hash=meta.get("project_hash","")
            excel_hash=excel_row_hash(raw)
            changes=excel_diff(current,raw)
            if changes:
                kind="conflict" if exported_hash and current_hash!=exported_hash and excel_hash!=current_hash else "update"
                preview["items"].append({"sheet":"Projektinformation","entity":"Projekt","kind":kind,"id":project_id,"row":0,"title":raw.get("name") or project["name"],"changes":changes,"data":raw})
                preview["counts"][kind]+=1

        for sheet_name,spec in EXCEL_SPECS.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws=wb[sheet_name]
            hdr=excel_headers(ws)
            missing=[h for h in spec["headers"] if h not in hdr]
            if missing:
                preview["errors"].append(f"{sheet_name}: saknar kolumner {', '.join(missing)}")
                preview["counts"]["error"]+=1
                continue
            id_col=hdr.get("_ID")
            hash_col=hdr.get("_Hash")
            current_rows=excel_current_rows(conn,spec["table"],project_id)
            for row_no in range(2,ws.max_row+1):
                raw={field:ws.cell(row_no,col).value for label,field in spec["headers"].items() for col in [hdr[label]]}
                data=excel_normalize_row(spec,raw)
                row_id=excel_int(ws.cell(row_no,id_col).value) if id_col else 0
                exported_hash=excel_text(ws.cell(row_no,hash_col).value) if hash_col else ""
                required=excel_text(data.get(spec["required"]))
                if not row_id and not required and all(v in ("",None,0,0.0) for v in data.values()):
                    continue
                if not required:
                    preview["items"].append({"sheet":sheet_name,"entity":spec["entity"],"kind":"error","id":row_id or None,"row":row_no,"title":f"Rad {row_no}","changes":{},"data":data,"message":"Obligatoriskt namn/rubrik saknas."})
                    preview["counts"]["error"]+=1
                    continue
                if not row_id:
                    preview["items"].append({"sheet":sheet_name,"entity":spec["entity"],"kind":"create","id":None,"row":row_no,"title":required,"changes":{k:{"old":"","new":excel_canon_value(v)} for k,v in data.items() if v not in ("",None)},"data":data})
                    preview["counts"]["create"]+=1
                    continue
                current_row=current_rows.get(row_id)
                if not current_row:
                    preview["items"].append({"sheet":sheet_name,"entity":spec["entity"],"kind":"error","id":row_id,"row":row_no,"title":required,"changes":{},"data":data,"message":"ID finns inte längre i appen."})
                    preview["counts"]["error"]+=1
                    continue
                current=spec["snapshot"](current_row)
                current_hash=excel_row_hash(current)
                excel_hash=excel_row_hash(data)
                if excel_hash==exported_hash:
                    preview["counts"]["unchanged"]+=1
                    continue
                changes=excel_diff(current,data)
                if not changes:
                    preview["counts"]["unchanged"]+=1
                    continue
                kind="conflict" if exported_hash and current_hash!=exported_hash and excel_hash!=current_hash else "update"
                preview["items"].append({"sheet":sheet_name,"entity":spec["entity"],"kind":kind,"id":row_id,"row":row_no,"title":required,"changes":changes,"data":data})
                preview["counts"][kind]+=1

        # Task dependencies need WBS mapping.
        if "Beroenden" in wb.sheetnames:
            ws=wb["Beroenden"]; hdr=excel_headers(ws)
            needed=["Föregående WBS","Efterföljande WBS","Typ","Förskjutning dagar"]
            if all(x in hdr for x in needed):
                tasks=conn.execute("SELECT id,wbs FROM tasks WHERE project_id=? AND deleted_at IS NULL",(project_id,)).fetchall()
                id_by_wbs={excel_text(t["wbs"]):int(t["id"]) for t in tasks if excel_text(t["wbs"])}
                links={int(r["id"]):r for r in conn.execute("SELECT * FROM task_links WHERE project_id=?",(project_id,)).fetchall()}
                id_col=hdr.get("_ID"); hash_col=hdr.get("_Hash")
                for row_no in range(2,ws.max_row+1):
                    pred=excel_text(ws.cell(row_no,hdr["Föregående WBS"]).value)
                    succ=excel_text(ws.cell(row_no,hdr["Efterföljande WBS"]).value)
                    typ=excel_text(ws.cell(row_no,hdr["Typ"]).value) or "FS"
                    lag=excel_int(ws.cell(row_no,hdr["Förskjutning dagar"]).value)
                    row_id=excel_int(ws.cell(row_no,id_col).value) if id_col else 0
                    exported_hash=excel_text(ws.cell(row_no,hash_col).value) if hash_col else ""
                    if not pred and not succ and not row_id: continue
                    if pred not in id_by_wbs or succ not in id_by_wbs:
                        preview["items"].append({"sheet":"Beroenden","entity":"Beroende","kind":"error","id":row_id or None,"row":row_no,"title":f"{pred} → {succ}","changes":{},"data":{},"message":"WBS finns inte bland projektets uppgifter."})
                        preview["counts"]["error"]+=1; continue
                    data={"predecessor_wbs":pred,"successor_wbs":succ,"link_type":typ,"lag_days":lag,"predecessor_id":id_by_wbs[pred],"successor_id":id_by_wbs[succ]}
                    visible={"predecessor_wbs":pred,"successor_wbs":succ,"link_type":typ,"lag_days":lag}
                    if not row_id:
                        preview["items"].append({"sheet":"Beroenden","entity":"Beroende","kind":"create","id":None,"row":row_no,"title":f"{pred} → {succ}","changes":{},"data":data})
                        preview["counts"]["create"]+=1; continue
                    cur=links.get(row_id)
                    if not cur:
                        preview["items"].append({"sheet":"Beroenden","entity":"Beroende","kind":"error","id":row_id,"row":row_no,"title":f"{pred} → {succ}","changes":{},"data":data,"message":"Beroendet finns inte längre."})
                        preview["counts"]["error"]+=1; continue
                    pred_cur=next((k for k,v in id_by_wbs.items() if v==int(cur["predecessor_id"])),"")
                    succ_cur=next((k for k,v in id_by_wbs.items() if v==int(cur["successor_id"])),"")
                    current={"predecessor_wbs":pred_cur,"successor_wbs":succ_cur,"link_type":cur["link_type"] or "FS","lag_days":excel_int(cur["lag_days"])}
                    excel_hash=excel_row_hash(visible); current_hash=excel_row_hash(current)
                    if excel_hash==exported_hash or excel_hash==current_hash:
                        preview["counts"]["unchanged"]+=1; continue
                    kind="conflict" if exported_hash and current_hash!=exported_hash else "update"
                    preview["items"].append({"sheet":"Beroenden","entity":"Beroende","kind":kind,"id":row_id,"row":row_no,"title":f"{pred} → {succ}","changes":excel_diff(current,visible),"data":data})
                    preview["counts"][kind]+=1
            else:
                preview["errors"].append("Beroenden: obligatoriska kolumner saknas.")
                preview["counts"]["error"]+=1

    return preview

def excel_save_preview(preview):
    EXCEL_IMPORT_DIR.mkdir(parents=True,exist_ok=True)
    token=secrets.token_urlsafe(18)
    path=EXCEL_IMPORT_DIR/f"{token}.json"
    path.write_text(json.dumps(preview,ensure_ascii=False,default=str),encoding="utf-8")
    return token

def excel_load_preview(token,project_id):
    if not re.fullmatch(r"[A-Za-z0-9_-]{10,80}",token or ""):
        abort(400)
    path=EXCEL_IMPORT_DIR/f"{token}.json"
    if not path.exists():
        abort(404)
    preview=json.loads(path.read_text(encoding="utf-8"))
    if int(preview.get("project_id") or 0)!=project_id:
        abort(403)
    return preview,path

def excel_backup_before_import():
    BACKUP_DIR.mkdir(parents=True,exist_ok=True)
    name=f"projectplan-before-excel-import-{datetime.now().strftime('%Y%m%d-%H%M%S')}.db"
    dst=BACKUP_DIR/name
    src=sqlite3.connect(DB_PATH)
    backup=sqlite3.connect(dst)
    try:
        src.backup(backup)
    finally:
        backup.close(); src.close()
    return name

def excel_insert(conn,table,project_id,data):
    now=datetime.now().isoformat(timespec="seconds")
    if table=="tasks":
        cols=["project_id","wbs","title","owner","start_date","end_date","actual_start","actual_end","status","priority","progress","milestone","notes","sort_order"]
        vals=[project_id,data["wbs"],data["title"],data["owner"],data["start_date"],data["end_date"],data["actual_start"],data["actual_end"],data["status"] or "Ej startad",data["priority"] or "Normal",data["progress"],data["milestone"],data["notes"],999999]
    elif table=="risks":
        cols=["project_id","kind","title","description","probability","impact","owner","action","status","due_date","created_at"]
        vals=[project_id,data["kind"] or "Risk",data["title"],data["description"],data["probability"],data["impact"],data["owner"],data["action"],data["status"] or "Öppen",data["due_date"],now]
    elif table=="change_requests":
        cols=["project_id","title","description","impact_scope","impact_days","impact_cost","status","created_at"]
        vals=[project_id,data["title"],data["description"],data["impact_scope"],data["impact_days"],data["impact_cost"],data["status"] or "Proposed",now]
    elif table=="resource_allocations":
        cols=["project_id","resource_name","week_start","allocation_pct","planned_hours"]
        vals=[project_id,data["resource_name"],data["week_start"],data["allocation_pct"],data["planned_hours"]]
    elif table=="project_costs":
        cols=["project_id","category","description","planned","actual","cost_date"]
        vals=[project_id,data["category"] or "External",data["description"],data["planned"],data["actual"],data["cost_date"]]
    elif table=="decisions":
        cols=["project_id","title","decision","decided_by","decision_date","owner","decided_at"]
        vals=[project_id,data["title"],data["decision"],data["decided_by"],data["decision_date"] or date.today().isoformat(),data["owner"],now]
    elif table=="meetings":
        cols=["project_id","title","meeting_date","attendees","notes"]
        vals=[project_id,data["title"],data["meeting_date"] or date.today().isoformat(),data["attendees"],data["notes"]]
    elif table=="action_items":
        cols=["project_id","title","owner","due_date","status"]
        vals=[project_id,data["title"],data["owner"],data["due_date"],data["status"] or "Open"]
    elif table=="project_benefits":
        cols=["project_id","title","unit","baseline_value","target_value","actual_value","measurement_date","owner","status","created_at"]
        vals=[project_id,data["title"],data["unit"] or "%",data["baseline_value"],data["target_value"],data["actual_value"],data["measurement_date"],data["owner"],data["status"] or "Planned",now]
    else:
        raise ValueError("Unsupported import table")
    placeholders=",".join("?" for _ in cols)
    cur=conn.execute(f"INSERT INTO {table}({','.join(cols)}) VALUES({placeholders})",vals)
    return cur.lastrowid

def excel_update(conn,table,row_id,project_id,data):
    fields=list(data.keys())
    setters=",".join(f"{f}=?" for f in fields)
    vals=[data[f] for f in fields]+[row_id,project_id]
    conn.execute(f"UPDATE {table} SET {setters} WHERE id=? AND project_id=?",vals)


def excel_safe_rows_v1522(conn,sql,args=()):
    try:
        return conn.execute(sql,args).fetchall()
    except Exception:
        return []

def excel_complete_data_v1522(project_id):
    """All project-related content that is useful in a complete Excel workbook."""
    data=excel_project_data(project_id)
    with db() as conn:
        data.update({
            "members":excel_safe_rows_v1522(conn,"""SELECT pm.*,COALESCE(u.display_name,u.username,'') display_name,u.username
                FROM project_members pm LEFT JOIN users u ON u.id=pm.user_id
                WHERE pm.project_id=? ORDER BY display_name""",(project_id,)),
            "status_reports":excel_safe_rows_v1522(conn,"SELECT * FROM status_reports WHERE project_id=? ORDER BY report_date,id",(project_id,)),
            "raid":excel_safe_rows_v1522(conn,"SELECT * FROM raid_items WHERE project_id=? ORDER BY item_type,status,due_date,id",(project_id,)),
            "approvals":excel_safe_rows_v1522(conn,"SELECT * FROM approvals WHERE project_id=? ORDER BY id",(project_id,)),
            "time_entries":excel_safe_rows_v1522(conn,"""SELECT te.*,COALESCE(u.display_name,u.username,'') user_name,t.title task_title
                FROM time_entries te LEFT JOIN users u ON u.id=te.user_id LEFT JOIN tasks t ON t.id=te.task_id
                WHERE te.project_id=? ORDER BY work_date,id""",(project_id,)),
            "documents":excel_safe_rows_v1522(conn,"SELECT * FROM documents WHERE project_id=? ORDER BY updated_at,id",(project_id,)),
            "attachments":excel_safe_rows_v1522(conn,"SELECT * FROM attachments WHERE project_id=? ORDER BY uploaded_at,id",(project_id,)),
            "comments":excel_safe_rows_v1522(conn,"""SELECT c.*,COALESCE(u.display_name,u.username,'') user_name
                FROM comments c LEFT JOIN users u ON u.id=c.user_id WHERE c.project_id=? ORDER BY created_at,id""",(project_id,)),
            "project_comments":excel_safe_rows_v1522(conn,"""SELECT c.*,COALESCE(u.display_name,u.username,'') user_name
                FROM project_comments c LEFT JOIN users u ON u.id=c.user_id WHERE c.project_id=? ORDER BY created_at,id""",(project_id,)),
            "baselines":excel_safe_rows_v1522(conn,"SELECT * FROM baselines WHERE project_id=? ORDER BY created_at,id",(project_id,)),
            "custom_fields":excel_safe_rows_v1522(conn,"""SELECT d.name,d.field_type,d.options_json,v.value_text,v.entity_type,v.entity_id
                FROM custom_field_values v JOIN custom_field_definitions d ON d.id=v.definition_id
                WHERE v.project_id=? ORDER BY d.name,v.entity_type,v.entity_id""",(project_id,)),
            "environments":excel_safe_rows_v1522(conn,"SELECT * FROM project_environments WHERE project_id=? ORDER BY id",(project_id,)),
            "deliverables":excel_safe_rows_v1522(conn,"SELECT * FROM deliverables WHERE project_id=? ORDER BY id",(project_id,)),
            "interfaces":excel_safe_rows_v1522(conn,"SELECT * FROM interface_register WHERE project_id=? ORDER BY id",(project_id,)),
            "test_cycles":excel_safe_rows_v1522(conn,"SELECT * FROM test_cycles WHERE project_id=? ORDER BY id",(project_id,)),
            "traceability":excel_safe_rows_v1522(conn,"SELECT * FROM requirements_traceability WHERE project_id=? ORDER BY id",(project_id,)),
            "cutover":excel_safe_rows_v1522(conn,"SELECT * FROM cutover_items WHERE project_id=? ORDER BY id",(project_id,)),
            "health_snapshots":excel_safe_rows_v1522(conn,"SELECT * FROM project_health_snapshots WHERE project_id=? ORDER BY id",(project_id,)),
            "cross_dependencies":excel_safe_rows_v1522(conn,"""SELECT x.*,
                pp.name predecessor_project,sp.name successor_project,
                pt.title predecessor_task,st.title successor_task
                FROM cross_project_dependencies x
                LEFT JOIN projects pp ON pp.id=x.predecessor_project_id
                LEFT JOIN projects sp ON sp.id=x.successor_project_id
                LEFT JOIN tasks pt ON pt.id=x.predecessor_task_id
                LEFT JOIN tasks st ON st.id=x.successor_task_id
                WHERE x.predecessor_project_id=? OR x.successor_project_id=?
                ORDER BY x.id""",(project_id,project_id)),
            "devops_links":excel_safe_rows_v1522(conn,"""SELECT l.*,c.name connection_name,c.project_name devops_project,t.title task_title
                FROM azure_devops_work_item_links l
                LEFT JOIN azure_devops_connections c ON c.id=l.connection_id
                LEFT JOIN tasks t ON t.id=l.task_id
                WHERE l.project_id=? ORDER BY l.id""",(project_id,)),
            "schedule_batches":excel_safe_rows_v1522(conn,"""SELECT b.*,COALESCE(u.display_name,u.username,'') created_name
                FROM schedule_change_batches b LEFT JOIN users u ON u.id=b.created_by
                WHERE b.project_id=? ORDER BY b.id""",(project_id,)),
            "schedule_items":excel_safe_rows_v1522(conn,"""SELECT i.*,t.title task_title
                FROM schedule_change_items i
                JOIN schedule_change_batches b ON b.id=i.batch_id
                LEFT JOIN tasks t ON t.id=i.task_id
                WHERE b.project_id=? ORDER BY i.batch_id,i.id""",(project_id,)),
            "project_finance":excel_safe_rows_v1522(conn,"SELECT * FROM project_finance WHERE project_id=?",(project_id,)),
        })
    return data

def excel_sheet_from_rows_v1522(wb,name,headers,rows,table_name=None,date_headers=None,money_headers=None,wide_headers=None):
    ws=wb.create_sheet(name)
    excel_prepare_sheet(ws,headers)
    for row in rows:
        ws.append([row.get(h,"") if isinstance(row,dict) else row[h] if h in row.keys() else "" for h in headers])
    style_header(ws,1)
    ws.freeze_panes="A2"
    ws.sheet_view.showGridLines=False
    date_headers=set(date_headers or [])
    money_headers=set(money_headers or [])
    wide_headers=set(wide_headers or [])
    for col,h in enumerate(headers,1):
        letter=get_column_letter(col)
        if h in wide_headers:
            ws.column_dimensions[letter].width=42
        for r in range(2,ws.max_row+1):
            cell=ws.cell(r,col)
            cell.alignment=Alignment(vertical="top",wrap_text=h in wide_headers)
            if h in date_headers:
                cell.value=excel_to_date(cell.value)
                cell.number_format="yyyy-mm-dd"
            if h in money_headers:
                cell.number_format='#,##0.00'
    autosize(ws)
    for col,h in enumerate(headers,1):
        if h in wide_headers:
            ws.column_dimensions[get_column_letter(col)].width=42
    if table_name and ws.max_row>=2:
        excel_add_table(ws,table_name)
    excel_page_setup(ws,True)
    return ws

def excel_complete_add_instructions_v1522(wb,project):
    ws=wb.create_sheet("LÄS MIG",0)
    ws.sheet_view.showGridLines=False
    ws["A1"]="Project Planer – Komplett Excel-mall"
    ws["A1"].font=Font(size=22,bold=True,color=EXCEL_PRO_HEADER)
    ws.merge_cells("A1:F1")
    ws["A3"]="Projekt"; ws["B3"]=project["name"]
    ws["A4"]="Appversion"; ws["B4"]=APP_VERSION
    ws["A5"]="Skapad"; ws["B5"]=datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["A7"]="Vad innehåller filen?"
    ws["A7"].font=Font(size=14,bold=True,color=EXCEL_PRO_HEADER)
    ws["A8"]="Den kompletta arbetsboken innehåller både round-trip-blad som kan redigeras/importeras tillbaka och referensblad med övrigt projektinnehåll."
    ws.merge_cells("A8:F9"); ws["A8"].alignment=Alignment(wrap_text=True,vertical="top")
    rows=[
        ("Gröna/redigerbara blad","Projektinformation, Uppgifter, Beroenden, Risker, Ändringsärenden, Resurser, Kostnader, Beslut, Möten, Åtgärder och Nyttor kan hanteras via den befintliga Excel-importen."),
        ("Referensblad","Medlemmar, statusrapporter, RAID, approvals, tid, dokument, bilagor, kommentarer, baselines, custom fields, leverabler, interfaces, test, cutover, DevOps och ändringshistorik följer med för komplett överblick."),
        ("Tekniska blad","_Metadata är dolt och krävs för säker round-trip. Ändra inte tekniska ID/hash-kolumner."),
        ("Säker import","Importen visar förhandsgranskning och konfliktkontroll innan commit. Databasbackup tas innan import.")
    ]
    ws["A11"]="Typ"; ws["B11"]="Beskrivning"; style_header(ws,11)
    for x in rows: ws.append(list(x))
    ws.column_dimensions["A"].width=24; ws.column_dimensions["B"].width=95
    for r in range(12,12+len(rows)): ws[f"B{r}"].alignment=Alignment(wrap_text=True,vertical="top")
    ws["A18"]="Importbara blad"
    ws["A18"].font=Font(size=13,bold=True,color=EXCEL_PRO_HEADER)
    importable=["Projektinformation","Uppgifter","Beroenden","Risker","Ändringsärenden","Resurser","Kostnader","Beslut","Möten","Åtgärder","Nyttor"]
    for i,name in enumerate(importable,19):
        ws.cell(i,1,"✓"); ws.cell(i,2,name)
    excel_page_setup(ws,False)

def build_complete_excel_template_v1522(project,data):
    # Start with full import-compatible Round-trip workbook.
    wb=build_roundtrip_workbook(project,data,"all")
    excel_complete_add_instructions_v1522(wb,project)

    # A compact data dictionary makes the workbook usable as a true template.
    dd=wb.create_sheet("Datamodell",1)
    dd_headers=["Blad","Syfte","Import tillbaka","Kommentar"]
    excel_prepare_sheet(dd,dd_headers)
    dictionary=[
      ("Projektinformation","Projektets grunddata","Ja","Namn, kund, projektledare, beskrivning och projektdatum."),
      ("Uppgifter","WBS och aktiviteter","Ja","Plan/faktiska datum, status, prioritet, progress och milstolpe."),
      ("Beroenden","Aktivitetsberoenden","Ja","FS/SS/FF/SF och lag i dagar."),
      ("Risker","Risker och issues","Ja","Sannolikhet, konsekvens, ansvarig och åtgärd."),
      ("Ändringsärenden","Change control","Ja","Omfattning, dagar, kostnad och status."),
      ("Resurser","Resursallokering","Ja","Vecka, allokering och timmar."),
      ("Kostnader","Projektkostnader","Ja","Planerat och utfall."),
      ("Nyttor","Benefits","Ja","Baslinje, mål och utfall."),
      ("Medlemmar","Projektteam","Nej","Projektroll och användare."),
      ("Statusrapporter","Historiska statusrapporter","Nej","RAG, sammanfattning, achievements och next steps."),
      ("RAID","Risk/assumption/issue/dependency","Nej","Komplett RAID-register."),
      ("Tid","Tidrapportering","Nej","Timmar per person/aktivitet/datum."),
      ("Leverabler","Leveransobjekt","Nej","Leverabler och status."),
      ("Interfaces","Integrations-/interface-register","Nej","Tekniska interfaces."),
      ("Testcykler","Testplanering","Nej","Testcykler och status."),
      ("Spårbarhet","Requirements traceability","Nej","Krav/spårbarhet."),
      ("Cutover","Cutover-plan","Nej","Go-live/cutover-aktiviteter."),
      ("Azure DevOps","Work Item-länkar","Nej","Synkstatus och länkning."),
      ("Omplaneringshistorik","Kontrollerade schemaändringar","Nej","Batchhistorik och old/new dates.")
    ]
    for row in dictionary: dd.append(list(row))
    excel_add_table(dd,"DataModelTable"); autosize(dd)
    dd.column_dimensions["B"].width=34; dd.column_dimensions["D"].width=55
    for r in range(2,dd.max_row+1): dd.cell(r,4).alignment=Alignment(wrap_text=True,vertical="top")

    def drows(rows):
        return [dict(r) for r in rows]

    excel_sheet_from_rows_v1522(wb,"Medlemmar",
        ["display_name","username","project_role","added_at","added_by"],
        drows(data.get("members",[])),"MembersTable",date_headers={"added_at"})

    excel_sheet_from_rows_v1522(wb,"Statusrapporter",
        ["report_date","overall_rag","scope_rag","schedule_rag","budget_rag","resources_rag","summary","achievements","next_steps","created_by","created_at"],
        drows(data.get("status_reports",[])),"StatusReportsTable",
        date_headers={"report_date","created_at"},wide_headers={"summary","achievements","next_steps"})

    excel_sheet_from_rows_v1522(wb,"RAID",
        ["item_type","title","owner","status","due_date","details"],
        drows(data.get("raid",[])),"RaidTable",date_headers={"due_date"},wide_headers={"details"})

    excel_sheet_from_rows_v1522(wb,"Godkännanden",
        ["entity_type","entity_id","status","requested_by","requested_at","decided_by","decided_at","comment"],
        drows(data.get("approvals",[])),"ApprovalsTable",
        date_headers={"requested_at","decided_at"},wide_headers={"comment"})

    excel_sheet_from_rows_v1522(wb,"Tid",
        ["work_date","user_name","task_title","hours","billable","note"],
        drows(data.get("time_entries",[])),"TimeEntriesTable",date_headers={"work_date"},wide_headers={"note"})

    excel_sheet_from_rows_v1522(wb,"Dokument",
        ["title","body","updated_by","updated_at"],
        drows(data.get("documents",[])),"DocumentsTable",date_headers={"updated_at"},wide_headers={"body"})

    excel_sheet_from_rows_v1522(wb,"Bilagor",
        ["filename","stored_name","content_type","size_bytes","uploaded_by","uploaded_at"],
        drows(data.get("attachments",[])),"AttachmentsTable",date_headers={"uploaded_at"})

    comments=[]
    for r in drows(data.get("comments",[])):
        comments.append({"källa":"Kommentar","user_name":r.get("user_name",""),"body":r.get("body",""),"created_at":r.get("created_at","")})
    for r in drows(data.get("project_comments",[])):
        comments.append({"källa":"Project Comment","user_name":r.get("user_name",""),"body":r.get("body",""),"created_at":r.get("created_at","")})
    excel_sheet_from_rows_v1522(wb,"Kommentarer",
        ["källa","user_name","body","created_at"],comments,"CommentsTable",date_headers={"created_at"},wide_headers={"body"})

    excel_sheet_from_rows_v1522(wb,"Baselines",
        ["name","created_at","created_by","snapshot_json"],
        drows(data.get("baselines",[])),"BaselinesTable",date_headers={"created_at"},wide_headers={"snapshot_json"})

    excel_sheet_from_rows_v1522(wb,"Anpassade fält",
        ["name","field_type","entity_type","entity_id","value_text","options_json"],
        drows(data.get("custom_fields",[])),"CustomFieldsTable",wide_headers={"value_text","options_json"})

    excel_sheet_from_rows_v1522(wb,"Miljöer",
        ["name","environment_type","url","owner","status","notes"],
        drows(data.get("environments",[])),"EnvironmentsTable",wide_headers={"notes"})

    excel_sheet_from_rows_v1522(wb,"Leverabler",
        ["title","owner","due_date","status","description"],
        drows(data.get("deliverables",[])),"DeliverablesTable",date_headers={"due_date"},wide_headers={"description"})

    excel_sheet_from_rows_v1522(wb,"Interfaces",
        ["name","source_system","target_system","owner","status","description"],
        drows(data.get("interfaces",[])),"InterfacesTable",wide_headers={"description"})

    excel_sheet_from_rows_v1522(wb,"Testcykler",
        ["name","start_date","end_date","owner","status","notes"],
        drows(data.get("test_cycles",[])),"TestCyclesTable",date_headers={"start_date","end_date"},wide_headers={"notes"})

    excel_sheet_from_rows_v1522(wb,"Spårbarhet",
        ["requirement_id","requirement_title","deliverable_id","test_reference","status","notes"],
        drows(data.get("traceability",[])),"TraceabilityTable",wide_headers={"notes"})

    excel_sheet_from_rows_v1522(wb,"Cutover",
        ["title","owner","planned_at","status","sequence_no","notes"],
        drows(data.get("cutover",[])),"CutoverTable",date_headers={"planned_at"},wide_headers={"notes"})

    excel_sheet_from_rows_v1522(wb,"Projekthälsa historik",
        list(drows(data.get("health_snapshots",[]))[0].keys()) if data.get("health_snapshots") else ["snapshot_date","health_score","rag","details_json"],
        drows(data.get("health_snapshots",[])),"HealthHistoryTable",wide_headers={"details_json"})

    excel_sheet_from_rows_v1522(wb,"Projektberoenden",
        ["predecessor_project","successor_project","predecessor_task","successor_task","link_type","lag_days","status","notes"],
        drows(data.get("cross_dependencies",[])),"CrossProjectDependenciesTable",wide_headers={"notes"})

    excel_sheet_from_rows_v1522(wb,"Azure DevOps",
        ["connection_name","devops_project","work_item_id","work_item_type","title","state","assigned_to","task_title","last_synced_at"],
        drows(data.get("devops_links",[])),"AzureDevOpsLinksTable",date_headers={"last_synced_at"})

    excel_sheet_from_rows_v1522(wb,"Omplaneringshistorik",
        ["id","title","reason","created_name","created_at","reverted_at","reverted_by"],
        drows(data.get("schedule_batches",[])),"ScheduleBatchesTable",
        date_headers={"created_at","reverted_at"},wide_headers={"reason"})

    excel_sheet_from_rows_v1522(wb,"Omplaneringsdetaljer",
        ["batch_id","task_title","old_start","old_end","new_start","new_end"],
        drows(data.get("schedule_items",[])),"ScheduleItemsTable",
        date_headers={"old_start","old_end","new_start","new_end"})

    excel_sheet_from_rows_v1522(wb,"Finanssammanfattning",
        list(drows(data.get("project_finance",[]))[0].keys()) if data.get("project_finance") else ["planned_budget","approved_budget","forecast","actual"],
        drows(data.get("project_finance",[])),"FinanceSummaryTable",
        money_headers={"planned_budget","approved_budget","forecast","actual"})

    # Ensure the user lands on the instructions sheet.
    wb.active=0
    wb.properties.title=f"{project['name']} – Komplett Project Planer Excel-mall"
    wb.properties.subject="Full project workbook / template"
    wb.properties.description=f"Project Planer v{APP_VERSION} · komplett Excel-mall med projektets round-trip-data och referensinnehåll."
    return wb

@app.get("/projects/<int:project_id>/excel/template/complete")
@login_required
def excel_complete_template_v1522(project_id):
    p=project_or_404(project_id)
    data=excel_complete_data_v1522(project_id)
    wb=build_complete_excel_template_v1522(p,data)
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    audit(project_id,"project",project_id,"excel_complete_template",f"app={APP_VERSION}")
    safe=re.sub(r"[^A-Za-z0-9ÅÄÖåäö _.-]","",p["name"] or "project").strip().replace(" ","-")
    return send_file(
        bio,as_attachment=True,
        download_name=f"{safe}-KOMPLETT-Excel-mall.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/projects/<int:project_id>/excel/import",methods=["GET","POST"])
@login_required
def excel_import_v810(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    if request.method=="POST":
        try:
            preview=excel_import_preview(project_id,request.files.get("file"))
            token=excel_save_preview(preview)
            return render_template("excel_import_preview_v810.html",project=p,preview=preview,token=token)
        except ValueError as ex:
            flash(str(ex),"error")
            return redirect(url_for("excel_import_v810",project_id=project_id))
    return render_template("excel_import_v810.html",project=p)

@app.post("/projects/<int:project_id>/excel/import/<token>/commit")
@login_required
def excel_import_commit_v810(project_id,token):
    p=project_or_404(project_id,write=True)
    preview,path=excel_load_preview(token,project_id)
    policy=request.form.get("conflict_policy","app")
    if policy not in ("app","excel"):
        policy="app"
    if preview.get("counts",{}).get("error",0):
        flash("Importen innehåller fel. Rätta Excel-filen och förhandsgranska igen.","error")
        return render_template("excel_import_preview_v810.html",project=p,preview=preview,token=token)

    backup_name=excel_backup_before_import()
    applied=created=updated=conflicts_skipped=0
    with db() as conn:
        conn.execute("BEGIN")
        try:
            for item in preview["items"]:
                kind=item["kind"]
                if kind=="error":
                    continue
                if kind=="conflict" and policy=="app":
                    conflicts_skipped+=1
                    continue
                if item["entity"]=="Projekt":
                    data=item["data"]
                    conn.execute("""UPDATE projects SET name=?,customer=?,project_manager=?,description=?,start_date=?,end_date=? WHERE id=?""",
                                 (data["name"],data["customer"],data["project_manager"],data["description"],data["start_date"],data["end_date"],project_id))
                    updated+=1; applied+=1
                    continue
                if item["sheet"]=="Beroenden":
                    d=item["data"]
                    if kind=="create":
                        conn.execute("""INSERT INTO task_links(project_id,predecessor_id,successor_id,link_type,lag_days) VALUES(?,?,?,?,?)""",
                                     (project_id,d["predecessor_id"],d["successor_id"],d["link_type"],d["lag_days"]))
                        created+=1
                    else:
                        conn.execute("""UPDATE task_links SET predecessor_id=?,successor_id=?,link_type=?,lag_days=? WHERE id=? AND project_id=?""",
                                     (d["predecessor_id"],d["successor_id"],d["link_type"],d["lag_days"],item["id"],project_id))
                        updated+=1
                    applied+=1
                    continue
                spec=EXCEL_SPECS[item["sheet"]]
                if kind=="create":
                    excel_insert(conn,spec["table"],project_id,item["data"])
                    created+=1
                else:
                    excel_update(conn,spec["table"],item["id"],project_id,item["data"])
                    updated+=1
                applied+=1
            conn.commit()
        except Exception:
            conn.rollback()
            raise

    audit(project_id,"project",project_id,"excel_import",
          f"applied={applied}; created={created}; updated={updated}; conflicts_skipped={conflicts_skipped}; backup={backup_name}")
    try:
        path.unlink()
    except OSError:
        pass
    flash(f"Excel-import klar: {applied} ändringar genomförda ({created} nya, {updated} uppdaterade). Backup: {backup_name}.","success")
    return redirect(url_for("project_workspace",project_id=project_id,tab="overview"))

@app.get("/projects/<int:project_id>/workspace-pro")
@login_required
def workspace_ux_pro_v820(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY wbs,id",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? ORDER BY probability*impact DESC,id",(project_id,)).fetchall()
    return render_template("workspace_ux_pro_v820.html",project=p,tasks=tasks,risks=risks)

@app.post("/projects/<int:project_id>/workspace-pro/tasks/<int:task_id>")
@login_required
def workspace_task_inline_v820(project_id,task_id):
    project_or_404(project_id,write=True)
    field=request.form.get("field","")
    allowed={"status","owner","start_date","end_date","priority","progress","notes"}
    if field not in allowed: abort(400)
    value=request.form.get("value","")
    if field=="progress": value=max(0,min(100,excel_int(value)))
    with db() as conn:
        row=conn.execute("SELECT id FROM tasks WHERE id=? AND project_id=?",(task_id,project_id)).fetchone()
        if not row: abort(404)
        conn.execute(f"UPDATE tasks SET {field}=? WHERE id=? AND project_id=?",(value,task_id,project_id))
        conn.commit()
    audit(project_id,"task",task_id,"workspace_inline_edit",field)
    return redirect(url_for("workspace_ux_pro_v820",project_id=project_id))

@app.get("/projects/<int:project_id>/views")
@login_required
def project_views_v830(project_id):
    p=project_or_404(project_id)
    view=request.args.get("view","list")
    if view not in ("list","board","timeline","calendar","milestones"): view="list"
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY end_date,wbs,id",(project_id,)).fetchall()
    return render_template("project_views_v830.html",project=p,tasks=tasks,view=view)

def gantt_impact_v840(project_id,task_id,new_end):
    with db() as conn:
        task=conn.execute("SELECT * FROM tasks WHERE id=? AND project_id=?",(task_id,project_id)).fetchone()
        if not task: abort(404)
        links=conn.execute("SELECT * FROM task_links WHERE project_id=? AND predecessor_id=?",(project_id,task_id)).fetchall()
        impacted=[]
        for l in links:
            succ=conn.execute("SELECT id,title,start_date,end_date FROM tasks WHERE id=?",(l["successor_id"],)).fetchone()
            if succ: impacted.append(dict(succ))
    delta=0
    try:
        if task["end_date"] and new_end: delta=(datetime.strptime(new_end,"%Y-%m-%d")-datetime.strptime(task["end_date"],"%Y-%m-%d")).days
    except ValueError: pass
    return {"delta":delta,"impacted":impacted}

@app.route("/projects/<int:project_id>/gantt-next",methods=["GET","POST"])
@login_required
def gantt_pro_v840(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    impact=None
    if request.method=="POST":
        task_id=excel_int(request.form.get("task_id")); new_end=excel_date(request.form.get("end_date"))
        impact=gantt_impact_v840(project_id,task_id,new_end)
        if request.form.get("apply")=="1":
            with db() as conn:
                conn.execute("UPDATE tasks SET end_date=? WHERE id=? AND project_id=?",(new_end,task_id,project_id)); conn.commit()
            flash("Planändringen är tillämpad.","success")
            return redirect(url_for("gantt_pro_v840",project_id=project_id))
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY start_date,wbs,id",(project_id,)).fetchall()
        links=conn.execute("SELECT * FROM task_links WHERE project_id=?",(project_id,)).fetchall()
    return render_template("gantt_pro_v840.html",project=p,tasks=tasks,links=links,impact=impact)

def forecast_project_v850(project_id):
    with db() as conn:
        p=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status NOT IN ('Closed','Stängd')",(project_id,)).fetchall()
        costs=conn.execute("SELECT COALESCE(SUM(planned),0) p,COALESCE(SUM(actual),0) a FROM project_costs WHERE project_id=?",(project_id,)).fetchone()
        alloc=conn.execute("SELECT COALESCE(MAX(allocation_pct),0) m FROM resource_allocations WHERE project_id=?",(project_id,)).fetchone()
    today=date.today(); overdue=[t for t in tasks if t["end_date"] and t["progress"]<100 and t["end_date"]<today.isoformat()]
    high=[r for r in risks if excel_int(r["probability"])*excel_int(r["impact"])>=15]
    slip=min(60,len(overdue)*3+len(high)*2+(5 if alloc["m"]>120 else 0))
    planned_end=p["end_date"] or ""
    forecast_end=planned_end
    try:
        forecast_end=(datetime.strptime(planned_end,"%Y-%m-%d")+timedelta(days=slip)).strftime("%Y-%m-%d")
    except: pass
    budget=float(costs["p"] or 0); actual=float(costs["a"] or 0)
    forecast_cost=max(actual,budget*(1+min(.30,(len(overdue)+len(high))*.025)))
    reasons=[]
    if overdue: reasons.append(f"{len(overdue)} försenade aktiviteter")
    if high: reasons.append(f"{len(high)} höga risker")
    if alloc["m"]>120: reasons.append(f"resursallokering {alloc['m']}%")
    return {"planned_end":planned_end,"forecast_end":forecast_end,"slip":slip,"budget":budget,"actual":actual,"forecast_cost":forecast_cost,"reasons":reasons,"overdue":overdue[:5],"high":high[:5]}

@app.get("/projects/<int:project_id>/forecast")
@login_required
def project_forecast_v850(project_id):
    p=project_or_404(project_id)
    return render_template("project_forecast_v850.html",project=p,f=forecast_project_v850(project_id))

@app.route("/projects/<int:project_id>/team-collaboration",methods=["GET","POST"])
@login_required
def collaboration_v860(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS project_conversations(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER NOT NULL,user_id INTEGER,body TEXT NOT NULL,created_at TEXT NOT NULL)""")
        if request.method=="POST":
            body=request.form.get("body","").strip()
            if body:
                conn.execute("INSERT INTO project_conversations(project_id,user_id,body,created_at) VALUES(?,?,?,?)",(project_id,current_user["id"],body,datetime.now().isoformat(timespec="seconds"))); conn.commit()
                audit(project_id,"project",project_id,"comment","collaboration")
                return redirect(url_for("collaboration_v860",project_id=project_id))
        rows=conn.execute("""SELECT c.*,COALESCE(NULLIF(u.display_name,''),u.username) user_name FROM project_conversations c LEFT JOIN users u ON u.id=c.user_id WHERE c.project_id=? ORDER BY c.id DESC LIMIT 100""",(project_id,)).fetchall()
        activity=conn.execute("SELECT * FROM audit_log WHERE project_id=? ORDER BY id DESC LIMIT 30",(project_id,)).fetchall()
    return render_template("collaboration_v860.html",project=p,rows=rows,activity=activity)

@app.route("/projects/<int:project_id>/stakeholder-share",methods=["GET","POST"])
@login_required
def stakeholder_share_v870(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS stakeholder_tokens(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER NOT NULL,token TEXT NOT NULL UNIQUE,active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL)""")
        if request.method=="POST":
            token=secrets.token_urlsafe(24)
            conn.execute("INSERT INTO stakeholder_tokens(project_id,token,active,created_at) VALUES(?,?,1,?)",(project_id,token,datetime.now().isoformat(timespec="seconds"))); conn.commit()
        tokens=conn.execute("SELECT * FROM stakeholder_tokens WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
    return render_template("stakeholder_share_v870.html",project=p,tokens=tokens)

@app.get("/share/project/<token>")
def stakeholder_public_v870(token):
    with db() as conn:
        row=conn.execute("SELECT * FROM stakeholder_tokens WHERE token=? AND active=1",(token,)).fetchone()
        if not row: abort(404)
        p=conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(row["project_id"],)).fetchone()
        if not p: abort(404)
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY end_date",(p["id"],)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND probability*impact>=15 ORDER BY probability*impact DESC",(p["id"],)).fetchall()
        decisions=conn.execute("SELECT * FROM decisions WHERE project_id=? ORDER BY decision_date DESC LIMIT 10",(p["id"],)).fetchall()
    return render_template("stakeholder_public_v870.html",project=p,tasks=tasks,risks=risks,decisions=decisions)

TEMPLATE_PRO_PRESETS={
 "LIMS Implementation":[
  ("1","Kickoff",1),("2","Krav & processkartläggning",0),("3","Konfiguration",0),("4","Integrationer",0),("5","Validering",0),("6","Utbildning",0),("7","Go-live",1)],
 "System Integration":[("1","Kickoff",1),("2","Interface design",0),("3","Utveckling",0),("4","SIT",0),("5","UAT",0),("6","Driftsättning",1)],
 "Upgrade":[("1","Planering",0),("2","Teknisk analys",0),("3","Uppgradering test",0),("4","Regressionstest",0),("5","Produktionssättning",1)]
}
@app.route("/projects/<int:project_id>/template-pro",methods=["GET","POST"])
@login_required
def template_pro_v880(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    if request.method=="POST":
        name=request.form.get("template","")
        preset=TEMPLATE_PRO_PRESETS.get(name)
        if not preset: abort(400)
        with db() as conn:
            existing=conn.execute("SELECT COUNT(*) c FROM tasks WHERE project_id=? AND deleted_at IS NULL",(project_id,)).fetchone()["c"]
            if existing and request.form.get("confirm")!="1":
                return render_template("template_pro_v880.html",project=p,presets=TEMPLATE_PRO_PRESETS,preview=preset,selected=name,needs_confirm=True)
            for i,(wbs,title,milestone) in enumerate(preset):
                conn.execute("""INSERT INTO tasks(project_id,wbs,title,status,priority,progress,milestone,sort_order) VALUES(?,?,?,?,?,?,?,?)""",(project_id,wbs,title,"Ej startad","Normal",0,milestone,i))
            conn.commit()
        audit(project_id,"project",project_id,"apply_template",name)
        flash(f"Mallen {name} har lagts till i projektet.","success")
        return redirect(url_for("project_workspace",project_id=project_id,tab="plan"))
    return render_template("template_pro_v880.html",project=p,presets=TEMPLATE_PRO_PRESETS,preview=None)

def pm_assistant_answer_v890(project_id,question):
    f=forecast_project_v850(project_id)
    q=(question or "").lower()
    with db() as conn:
        p=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status NOT IN ('Closed','Stängd')",(project_id,)).fetchall()
    overdue=[t for t in tasks if t["end_date"] and t["progress"]<100 and t["end_date"]<date.today().isoformat()]
    milestones=[t for t in tasks if t["milestone"] and t["progress"]<100]
    if "status" in q or "styrgrupp" in q:
        return f"{p['name']}: prognostiserat slut {f['forecast_end'] or 'saknas'}. {len(overdue)} försenade aktiviteter, {len([r for r in risks if excel_int(r['probability'])*excel_int(r['impact'])>=15])} höga risker. Kostnadsprognos {f['forecast_cost']:.0f}."
    if "milstolp" in q:
        return "Kommande öppna milstolpar: "+("; ".join(f"{t['title']} ({t['end_date']})" for t in milestones[:8]) or "inga öppna milstolpar")
    if "idag" in q or "fokus" in q:
        return "Fokusera på: "+("; ".join(t["title"] for t in overdue[:5]) or "inga försenade aktiviteter")+". "+("Orsaker: "+", ".join(f["reasons"]) if f["reasons"] else "Projektet saknar tydliga varningssignaler.")
    return f"Projektet har {len(tasks)} aktiviteter. Prognos: {f['forecast_end'] or 'okänd'}. Fråga gärna om fokus idag, milstolpar eller status inför styrgruppen."

@app.route("/projects/<int:project_id>/pm-assistant-2",methods=["GET","POST"])
@login_required
def pm_assistant_2_v890(project_id):
    p=project_or_404(project_id)
    answer=None; question=""
    if request.method=="POST":
        question=request.form.get("question","").strip()
        answer=pm_assistant_answer_v890(project_id,question)
        with db() as conn:
            conn.execute("INSERT INTO assistant_queries(user_id,project_id,question,answer,created_at) VALUES(?,?,?,?,?)",(current_user["id"],project_id,question,answer,datetime.now().isoformat(timespec="seconds"))); conn.commit()
    return render_template("pm_assistant_2_v890.html",project=p,question=question,answer=answer)

@app.get("/projects/<int:project_id>/next")
@login_required
def project_next_v900(project_id):
    p=project_or_404(project_id)
    f=forecast_project_v850(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY end_date,wbs,id",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status NOT IN ('Closed','Stängd') ORDER BY probability*impact DESC",(project_id,)).fetchall()
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? AND status IN ('Proposed','Submitted','Pending') ORDER BY id DESC",(project_id,)).fetchall()
    attention=[]
    for t in tasks:
        if t["end_date"] and t["progress"]<100 and t["end_date"]<date.today().isoformat(): attention.append(("Försenad aktivitet",t["title"],t["end_date"]))
    for r in risks:
        if excel_int(r["probability"])*excel_int(r["impact"])>=15: attention.append(("Hög risk",r["title"],f"Riskpoäng {excel_int(r['probability'])*excel_int(r['impact'])}"))
    for c in changes[:5]: attention.append(("Ändringsärende",c["title"],c["status"]))
    return render_template("project_next_v900.html",project=p,f=f,tasks=tasks,attention=attention[:8])

def report_studio_payload_v901(project_id):
    p=project_or_404(project_id)
    f=forecast_project_v850(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY end_date,wbs,id",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status NOT IN ('Closed','Stängd') ORDER BY probability*impact DESC,id",(project_id,)).fetchall()
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC LIMIT 12",(project_id,)).fetchall()
        decisions=conn.execute("SELECT * FROM decisions WHERE project_id=? ORDER BY decision_date DESC,id DESC LIMIT 12",(project_id,)).fetchall()
        costs=conn.execute("SELECT COALESCE(SUM(planned),0) planned,COALESCE(SUM(actual),0) actual FROM project_costs WHERE project_id=?",(project_id,)).fetchone()
        latest_status=conn.execute("SELECT * FROM status_reports WHERE project_id=? ORDER BY report_date DESC,id DESC LIMIT 1",(project_id,)).fetchone()
    milestones=[t for t in tasks if excel_int(t["milestone"])]
    overdue=[t for t in tasks if t["end_date"] and excel_int(t["progress"])<100 and t["end_date"]<date.today().isoformat()]
    blocked=[t for t in tasks if (t["status"] or "").lower() in ("blocked","blockerad")]
    high=[r for r in risks if excel_int(r["probability"])*excel_int(r["impact"])>=15]
    progress=round(sum(excel_int(t["progress"]) for t in tasks)/len(tasks)) if tasks else 0
    rag="red" if len(overdue)>=3 or len(high)>=2 else ("amber" if overdue or high or blocked else "green")
    summary=(
        f"{p['name']} är {progress}% klart. "
        f"{len(overdue)} aktiviteter är försenade, {len(blocked)} blockerade och {len(high)} höga risker är öppna. "
        f"Prognostiserat slut är {f['forecast_end'] or p['end_date'] or 'inte satt'}."
    )
    if latest_status and latest_status["summary"]:
        summary=latest_status["summary"]
    return {
        "project":p,"forecast":f,"tasks":tasks,"milestones":milestones,"risks":risks,"high":high,
        "changes":changes,"decisions":decisions,"overdue":overdue,"blocked":blocked,
        "progress":progress,"rag":rag,"costs":costs,"summary":summary
    }

def log_report_export_v901(project_id,fmt,report_type,file_name):
    with db() as conn:
        conn.execute(
            "INSERT INTO report_export_log(project_id,format,report_type,file_name,exported_by,created_at) VALUES(?,?,?,?,?,?)",
            (project_id,fmt,report_type,file_name,current_user()["id"],datetime.now().isoformat(timespec="seconds"))
        )
        conn.commit()
    audit(project_id,"project",project_id,"report_export",f"{report_type}:{fmt}:{file_name}")

def build_report_excel_v901(payload):
    wb=Workbook()
    ws=wb.active
    ws.title="Projektstatus"
    ws.sheet_view.showGridLines=False
    ws["A1"]="Project Planer – Projektstatus"
    ws["A1"].font=Font(size=18,bold=True,color="0F4C81")
    ws.merge_cells("A1:D1")
    project=payload["project"]
    rows=[
        ("Projekt",project["name"]),("Kund",project["customer"]),("Projektledare",project["project_manager"]),
        ("RAG",payload["rag"].upper()),("Framdrift %",payload["progress"]),
        ("Planerat slut",project["end_date"]),("Prognostiserat slut",payload["forecast"]["forecast_end"]),
        ("Försenade aktiviteter",len(payload["overdue"])),("Blockerade aktiviteter",len(payload["blocked"])),
        ("Höga risker",len(payload["high"])),("Planerad kostnad",float(payload["costs"]["planned"] or 0)),
        ("Utfall",float(payload["costs"]["actual"] or 0))
    ]
    ws.append([])
    for r in rows: ws.append(list(r))
    ws["A16"]="Ledningssammanfattning"; ws["A16"].font=Font(bold=True)
    ws["B16"]=payload["summary"]; ws["B16"].alignment=Alignment(wrap_text=True,vertical="top")
    ws.column_dimensions["A"].width=28; ws.column_dimensions["B"].width=70

    for title, headers, source, mapper in [
        ("Milstolpar",["WBS","Milstolpe","Slutdatum","Status","Progress %"],payload["milestones"],
         lambda x:[x["wbs"],x["title"],x["end_date"],x["status"],x["progress"]]),
        ("Höga risker",["Risk","Sannolikhet","Konsekvens","Poäng","Ansvarig","Status"],payload["high"],
         lambda x:[x["title"],x["probability"],x["impact"],excel_int(x["probability"])*excel_int(x["impact"]),x["owner"],x["status"]]),
        ("Försenade",["WBS","Aktivitet","Ansvarig","Slutdatum","Progress %"],payload["overdue"],
         lambda x:[x["wbs"],x["title"],x["owner"],x["end_date"],x["progress"]]),
        ("Ändringar",["Rubrik","Status","Dagar","Kostnad","Beskrivning"],payload["changes"],
         lambda x:[x["title"],x["status"],x["impact_days"],x["impact_cost"],x["description"]]),
        ("Beslut",["Titel","Beslut","Datum","Beslutat av"],payload["decisions"],
         lambda x:[x["title"],x["decision"],x["decision_date"],x["decided_by"]]),
    ]:
        sheet=wb.create_sheet(title)
        excel_prepare_sheet(sheet,headers)
        for item in source:
            sheet.append(mapper(item))
        autosize(sheet)
    return excel_pro_report_finalize_v902(wb,payload)

def build_report_pdf_v901(payload):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    b=BytesIO(); c=canvas.Canvas(b,pagesize=A4); w,h=A4
    c.setFillColor(HexColor("#0F4C81")); c.rect(0,h-92,w,92,fill=1,stroke=0)
    c.setFillColor(HexColor("#FFFFFF")); c.setFont("Helvetica-Bold",18)
    c.drawString(42,h-52,payload["project"]["name"][:58])
    c.setFont("Helvetica",9); c.drawString(42,h-72,f"Projektstatus · {date.today().isoformat()} · Project Planer v{APP_VERSION}")
    y=h-122
    c.setFillColor(HexColor("#111827")); c.setFont("Helvetica-Bold",13); c.drawString(42,y,"Ledningssammanfattning"); y-=20
    c.setFont("Helvetica",10)
    text=c.beginText(42,y)
    for line in textwrap.wrap(payload["summary"],92):
        text.textLine(line)
    c.drawText(text); y=text.getY()-18
    c.setFont("Helvetica-Bold",11)
    metrics=[
        f"RAG: {payload['rag'].upper()}",
        f"Framdrift: {payload['progress']}%",
        f"Försenade: {len(payload['overdue'])}",
        f"Blockerade: {len(payload['blocked'])}",
        f"Höga risker: {len(payload['high'])}",
        f"Prognos slut: {payload['forecast']['forecast_end'] or '–'}"
    ]
    for i,m in enumerate(metrics):
        col=i%2; row=i//2
        c.drawString(42+col*255,y-row*22,m)
    y-=78

    def section(title, rows):
        nonlocal y
        if y<120:
            c.showPage(); y=h-55
        c.setFont("Helvetica-Bold",12); c.drawString(42,y,title); y-=18
        c.setFont("Helvetica",9)
        for line in rows[:10]:
            if y<65:
                c.showPage(); y=h-55; c.setFont("Helvetica",9)
            c.drawString(52,y,line[:100]); y-=14
        y-=8

    section("Milstolpar",[f"{m['end_date'] or '–'}  {m['title']}  ({m['progress']}%)" for m in payload["milestones"]])
    section("Höga risker",[f"{r['title']} · poäng {excel_int(r['probability'])*excel_int(r['impact'])} · {r['owner'] or 'utan ansvarig'}" for r in payload["high"]])
    section("Behöver uppmärksamhet",[f"{t['title']} · slut {t['end_date']}" for t in payload["overdue"]])
    section("Senaste beslut",[f"{d['decision_date'] or '–'} · {d['title']}: {d['decision']}" for d in payload["decisions"]])
    c.save(); b.seek(0); return b

def build_report_pptx_v901(payload):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation()
    prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text=payload["project"]["name"]
    slide.placeholders[1].text=f"Projektstatus · {date.today().isoformat()} · Project Planer v{APP_VERSION}"

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text="Ledningssammanfattning"
    slide.placeholders[1].text=payload["summary"]+"\n\n"+(
        f"RAG: {payload['rag'].upper()}   |   Framdrift: {payload['progress']}%   |   "
        f"Försenade: {len(payload['overdue'])}   |   Höga risker: {len(payload['high'])}\n"
        f"Planerat slut: {payload['project']['end_date'] or '–'}   |   Prognos: {payload['forecast']['forecast_end'] or '–'}"
    )

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text="Milstolpar"
    slide.placeholders[1].text="\n".join(f"• {m['end_date'] or '–'} · {m['title']} · {m['progress']}%" for m in payload["milestones"][:10]) or "Inga milstolpar"

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text="Risk & uppmärksamhet"
    lines=[f"• RISK: {r['title']} · poäng {excel_int(r['probability'])*excel_int(r['impact'])}" for r in payload["high"][:6]]
    lines += [f"• FÖRSENAD: {t['title']} · {t['end_date']}" for t in payload["overdue"][:6]]
    slide.placeholders[1].text="\n".join(lines) or "Inga kritiska signaler"

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text="Ekonomi & prognos"
    planned=float(payload["costs"]["planned"] or 0); actual=float(payload["costs"]["actual"] or 0)
    slide.placeholders[1].text=(
        f"Planerat: {planned:,.0f}\nUtfall: {actual:,.0f}\n"
        f"Kostnadsprognos: {payload['forecast']['forecast_cost']:,.0f}\n"
        f"Prognostiserat slut: {payload['forecast']['forecast_end'] or '–'}"
    )

    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text="Beslut & nästa steg"
    decisions="\n".join(f"• {d['title']}: {d['decision']}" for d in payload["decisions"][:6]) or "Inga beslut registrerade"
    next_steps="\n".join(f"• {x}" for x in payload["forecast"]["reasons"]) or "• Fortsätt följa milstolpar, risker och resursbelastning."
    slide.placeholders[1].text=decisions+"\n\nNästa fokus:\n"+next_steps
    b=BytesIO(); prs.save(b); b.seek(0); return b

@app.get("/projects/<int:project_id>/report-studio")
@login_required
def report_studio_v901(project_id):
    payload=report_studio_payload_v901(project_id)
    with db() as conn:
        history=conn.execute(
            """SELECT l.*,COALESCE(NULLIF(u.display_name,''),u.username) exported_by_name
               FROM report_export_log l LEFT JOIN users u ON u.id=l.exported_by
               WHERE l.project_id=? ORDER BY l.id DESC LIMIT 20""",(project_id,)
        ).fetchall()
    return render_template("report_studio_v901.html",d=payload,history=history)

@app.get("/projects/<int:project_id>/report-preview")
@login_required
def report_preview_v901(project_id):
    return render_template("report_preview_v901.html",d=report_studio_payload_v901(project_id),generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"))

@app.get("/projects/<int:project_id>/report-studio.pdf")
@login_required
def report_pdf_v901(project_id):
    payload=report_studio_payload_v901(project_id)
    file_name=f"{secure_filename(payload['project']['name'])}-project-status.pdf"
    log_report_export_v901(project_id,"pdf","executive",file_name)
    return send_file(build_report_pdf_v901(payload),mimetype="application/pdf",as_attachment=True,download_name=file_name)

@app.get("/projects/<int:project_id>/report-studio.pptx")
@login_required
def report_pptx_v901(project_id):
    payload=report_studio_payload_v901(project_id)
    file_name=f"{secure_filename(payload['project']['name'])}-steering-group.pptx"
    log_report_export_v901(project_id,"pptx","steering",file_name)
    return send_file(build_report_pptx_v901(payload),mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",as_attachment=True,download_name=file_name)

@app.get("/projects/<int:project_id>/report-studio.xlsx")
@login_required
def report_excel_v901(project_id):
    payload=report_studio_payload_v901(project_id)
    wb=build_report_excel_v901(payload)
    b=BytesIO(); wb.save(b); b.seek(0)
    file_name=f"{secure_filename(payload['project']['name'])}-status-report.xlsx"
    log_report_export_v901(project_id,"xlsx","status",file_name)
    return send_file(b,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",as_attachment=True,download_name=file_name)

@app.get("/projects/<int:project_id>/report-pack-v901.zip")
@login_required
def report_pack_zip_v901(project_id):
    import zipfile
    payload=report_studio_payload_v901(project_id)
    base=secure_filename(payload["project"]["name"])
    package=BytesIO()
    with zipfile.ZipFile(package,"w",zipfile.ZIP_DEFLATED) as z:
        pdf=build_report_pdf_v901(payload); z.writestr(f"{base}-project-status.pdf",pdf.getvalue())
        ppt=build_report_pptx_v901(payload); z.writestr(f"{base}-steering-group.pptx",ppt.getvalue())
        wb=build_report_excel_v901(payload); x=BytesIO(); wb.save(x)
        z.writestr(f"{base}-status-report.xlsx",x.getvalue())
        z.writestr("README.txt",
            "Project Planer report pack\n"
            f"Project: {payload['project']['name']}\n"
            f"Generated: {datetime.now().isoformat(timespec='seconds')}\n"
            f"App version: {APP_VERSION}\n"
            "Contents: executive PDF, steering-group PowerPoint, status Excel workbook.\n")
    package.seek(0)
    file_name=f"{base}-report-pack.zip"
    log_report_export_v901(project_id,"zip","report_pack",file_name)
    return send_file(package,mimetype="application/zip",as_attachment=True,download_name=file_name)

@app.get("/everyday-ux")
@login_required
def everyday_ux_v910():
    return render_template("roadmap_910.html")

@app.get("/workspace-ux-3")
@login_required
def workspace_ux_v920():
    return render_template("roadmap_920.html")

@app.get("/planning-pro-3")
@login_required
def planning_pro_v930():
    return render_template("roadmap_930.html")

@app.get("/capacity-pro-3")
@login_required
def capacity_pro_v940():
    return render_template("roadmap_940.html")

@app.get("/project-control-center")
@login_required
def control_center_v950():
    return render_template("roadmap_950.html")

@app.get("/reporting-automation")
@login_required
def reporting_automation_v960():
    return render_template("roadmap_960.html")

@app.get("/collaboration-decisions")
@login_required
def collaboration_decisions_v970():
    return render_template("roadmap_970.html")

@app.get("/portfolio-programs")
@login_required
def portfolio_programs_v980():
    return render_template("roadmap_980.html")

@app.get("/pm-intelligence")
@login_required
def pm_intelligence_v990():
    return render_template("roadmap_990.html")

@app.get("/enterprise-next")
@login_required
def enterprise_ux_v1000():
    return render_template("roadmap_1000.html")

@app.get("/production-quality")
@login_required
def production_quality_v1001():
    return render_template("roadmap_1001.html")

def ensure_azure_devops_schema_v1010(conn):
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS azure_devops_connections(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL,
      organization_url TEXT NOT NULL,
      project_name TEXT NOT NULL,
      auth_mode TEXT NOT NULL DEFAULT 'PAT',
      pat_secret TEXT,
      enabled INTEGER NOT NULL DEFAULT 1,
      created_by INTEGER,
      created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
      updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS azure_devops_mappings(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      connection_id INTEGER NOT NULL,
      project_id INTEGER NOT NULL,
      devops_project TEXT NOT NULL,
      area_path TEXT,
      iteration_path TEXT,
      sync_direction TEXT NOT NULL DEFAULT 'DevOpsToProjectPlaner',
      enabled INTEGER NOT NULL DEFAULT 1,
      created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
      UNIQUE(connection_id,project_id)
    );
    CREATE TABLE IF NOT EXISTS azure_devops_work_item_links(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      connection_id INTEGER NOT NULL,
      project_id INTEGER NOT NULL,
      task_id INTEGER,
      work_item_id INTEGER NOT NULL,
      work_item_type TEXT,
      title TEXT,
      state TEXT,
      assigned_to TEXT,
      iteration_path TEXT,
      url TEXT,
      last_synced_at TEXT,
      UNIQUE(connection_id,work_item_id)
    );
    CREATE TABLE IF NOT EXISTS azure_devops_sync_log(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      connection_id INTEGER,
      project_id INTEGER,
      direction TEXT,
      status TEXT NOT NULL,
      message TEXT,
      items_read INTEGER NOT NULL DEFAULT 0,
      items_written INTEGER NOT NULL DEFAULT 0,
      created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    """)

@app.route("/integrations/azure-devops",methods=["GET","POST"])
@login_required
def azure_devops_v1010():
    if session.get("role") not in ("admin","pm"):
        abort(403)
    with db() as conn:
        ensure_azure_devops_schema_v1010(conn)
        if request.method=="POST":
            name=(request.form.get("name") or "").strip()
            org=(request.form.get("organization_url") or "").strip().rstrip("/")
            project=(request.form.get("project_name") or "").strip()
            pat=(request.form.get("pat_secret") or "").strip()
            if not name or not org or not project:
                flash("Namn, organisation och DevOps-projekt krävs.","error")
            elif not (org.startswith("https://dev.azure.com/") or org.startswith("https://") and "visualstudio.com" in org):
                flash("Ange en giltig Azure DevOps organisationsadress.","error")
            else:
                conn.execute("""INSERT INTO azure_devops_connections
                  (name,organization_url,project_name,auth_mode,pat_secret,created_by)
                  VALUES(?,?,?,?,?,?)""",(name,org,project,"PAT",pat or None,session.get("user_id")))
                conn.commit()
                flash("Azure DevOps-anslutningen har sparats. PAT visas inte i gränssnittet.","success")
        connections=conn.execute("""SELECT id,name,organization_url,project_name,auth_mode,enabled,created_at,updated_at
          FROM azure_devops_connections ORDER BY id DESC""").fetchall()
        syncs=conn.execute("""SELECT s.*,c.name connection_name FROM azure_devops_sync_log s
          LEFT JOIN azure_devops_connections c ON c.id=s.connection_id ORDER BY s.id DESC LIMIT 20""").fetchall()
    return render_template("azure_devops_v1010.html",connections=connections,syncs=syncs)

@app.get("/excel-reporting-pro-3")
@login_required
def excel_reporting_v1020():
    return render_template("roadmap_1020.html")

@app.get("/planning-engine-2")
@login_required
def planning_engine_v1030():
    return render_template("roadmap_1030.html")

@app.get("/resource-intelligence")
@login_required
def resource_intelligence_v1040():
    return render_template("roadmap_1040.html")

@app.get("/portfolio-control")
@login_required
def portfolio_control_v1050():
    return render_template("roadmap_1050.html")

@app.get("/automation-integrations")
@login_required
def automation_integrations_v1060():
    return render_template("roadmap_1060.html")

@app.get("/pm-intelligence-2")
@login_required
def pm_intelligence_v1070():
    return render_template("roadmap_1070.html")

@app.get("/ux")
@login_required
def ux_consolidation_v1080():
    return render_template("ux_1080.html")

@app.get("/workspace-pro")
@login_required
def workspace_pro_v1081():
    return render_template("ux_1081.html")

@app.get("/devops-ux")
@login_required
def devops_ux_v1082():
    return render_template("ux_1082.html")

@app.get("/excel-reporting-ux")
@login_required
def excel_reporting_ux_v1083():
    return render_template("ux_1083.html")

@app.get("/mobile-workspace")
@login_required
def mobile_ux_v1084():
    return render_template("ux_1084.html")

@app.get("/ux-edition")
@login_required
def ux_edition_v1100():
    return render_template("ux_1100.html")

@app.get("/production-audit")
@login_required
def production_audit_v1101():
    if session.get("role") not in ("admin","pm"):
        abort(403)
    checks = []
    with db() as conn:
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
        required = ["projects","tasks","risks","project_members","notifications","audit_log"]
        for name in required:
            checks.append({"name": f"Databastabell: {name}", "ok": name in tables})
        try:
            mode = conn.execute("PRAGMA journal_mode").fetchone()[0]
        except Exception:
            mode = "unknown"
        checks.append({"name":"SQLite journal mode", "ok": str(mode).lower() == "wal", "detail": str(mode)})
    return render_template("production_audit_v1101.html", checks=checks)

@app.get("/projects/<int:project_id>/home-2")
@login_required
def project_home_v1110(project_id):
    with db() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(project_id,)).fetchone()
        if not project: abort(404)
        tasks = conn.execute("""SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY end_date,sort_order LIMIT 200""",(project_id,)).fetchall()
        risks = conn.execute("""SELECT * FROM risks WHERE project_id=? AND status NOT IN ('Closed','Stängd') ORDER BY probability*impact DESC LIMIT 20""",(project_id,)).fetchall()
        costs = conn.execute("""SELECT COALESCE(SUM(planned),0) planned,COALESCE(SUM(actual),0) actual FROM project_costs WHERE project_id=?""",(project_id,)).fetchone()
    active=[t for t in tasks if not t["milestone"] and (t["status"] or "").lower() not in ("done","completed","closed","klar")]
    overdue=[t for t in active if t["end_date"] and t["end_date"] < datetime.utcnow().date().isoformat()]
    milestones=[t for t in tasks if t["milestone"]][:8]
    progress = round(sum((t["progress"] or 0) for t in tasks)/len(tasks)) if tasks else 0
    high_risks=[r for r in risks if (r["probability"] or 0)*(r["impact"] or 0) >= 12]
    return render_template("project_home_v1110.html",project=project,progress=progress,overdue=overdue,
                           milestones=milestones,high_risks=high_risks,costs=costs)

@app.get("/my-work-2")
@login_required
def my_work_v1120():
    u=current_user()
    uid=u["id"]
    today=date.today()
    projects=visible_projects_for_user()
    ids=[p["id"] for p in projects]
    tasks=[]; actions=[]; notifications=[]; approvals=[]; changes=[]
    with db() as conn:
        notifications=conn.execute("""SELECT * FROM notifications WHERE user_id=? AND is_read=0
                                      ORDER BY created_at DESC LIMIT 30""",(uid,)).fetchall()
        if ids:
            marks=",".join("?" for _ in ids)
            tasks=conn.execute(f"""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
              WHERE t.project_id IN ({marks}) AND COALESCE(t.deleted_at,'')=''
                AND (t.owner_user_id=? OR LOWER(COALESCE(t.owner,'')) IN (LOWER(?),LOWER(?)))
                AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')
              ORDER BY CASE WHEN t.end_date<>'' AND t.end_date<date('now') THEN 0 ELSE 1 END,
                       CASE WHEN LOWER(COALESCE(t.priority,'')) IN ('critical','kritisk','high','hög') THEN 0 ELSE 1 END,
                       CASE WHEN t.end_date IS NULL OR t.end_date='' THEN 1 ELSE 0 END,t.end_date LIMIT 150""",
              ids+[uid,u["display_name"],u["username"]]).fetchall()
            actions=conn.execute(f"""SELECT a.*,p.name project_name FROM action_items a JOIN projects p ON p.id=a.project_id
              WHERE a.project_id IN ({marks})
                AND LOWER(COALESCE(a.status,'')) NOT IN ('done','closed','klar')
                AND (a.owner='' OR LOWER(a.owner) IN (LOWER(?),LOWER(?)))
              ORDER BY CASE WHEN a.due_date<>'' AND a.due_date<date('now') THEN 0 ELSE 1 END,a.due_date LIMIT 80""",
              ids+[u["display_name"],u["username"]]).fetchall()
            try:
                changes=conn.execute(f"""SELECT c.*,p.name project_name FROM change_requests c JOIN projects p ON p.id=c.project_id
                  WHERE c.project_id IN ({marks}) AND LOWER(COALESCE(c.status,'')) IN ('proposed','submitted','pending')
                  ORDER BY c.id DESC LIMIT 40""",ids).fetchall()
            except Exception:
                changes=[]
    def decorate(row, field):
        d=_v15_date(row[field])
        return bool(d and d<today), bool(d and today<=d<=today+timedelta(days=7))
    task_rows=[]
    for r in tasks:
        d=dict(r); d["is_overdue"],d["is_week"]=decorate(d,"end_date"); task_rows.append(d)
    action_rows=[]
    for r in actions:
        d=dict(r); d["is_overdue"],d["is_week"]=decorate(d,"due_date"); action_rows.append(d)
    return render_template("my_work_v1500.html",tasks=task_rows,actions=action_rows,notifications=notifications,
                           changes=changes,user=u,today=today)

@app.post("/work/action")
@login_required
def work_action_v1500():
    kind=(request.form.get("kind") or "").strip()
    item_id=request.form.get("id",type=int)
    next_url=request.form.get("next") or url_for("my_work_v1120")
    if not item_id:
        return redirect(next_url)
    with db() as conn:
        if kind=="task":
            conn.execute("UPDATE tasks SET progress=100,status='Done' WHERE id=?",(item_id,))
        elif kind=="action":
            conn.execute("UPDATE action_items SET status='Done' WHERE id=?",(item_id,))
        elif kind=="notification":
            conn.execute("UPDATE notifications SET is_read=1 WHERE id=? AND user_id=?",(item_id,session.get("user_id")))
        conn.commit()
    return redirect(next_url)

@app.get("/projects/<int:project_id>/planning-pro-ux")
@login_required
def planning_ux_v1130(project_id):
    with db() as conn:
        project=conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(project_id,)).fetchone()
        if not project: abort(404)
        tasks=conn.execute("""SELECT id,wbs,title,owner,start_date,end_date,status,progress,milestone,parent_task_id
          FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY sort_order,id""",(project_id,)).fetchall()
        links=conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days FROM task_links WHERE project_id=?""",(project_id,)).fetchall()
    return render_template("planning_ux_v1130.html",project=project,tasks=tasks,links=links)


def _ado_patch_v1520(url,pat,operations):
    import urllib.request, base64, json as _json
    raw=(":"+pat).encode("utf-8")
    req=urllib.request.Request(url,method="PATCH")
    req.add_header("Authorization","Basic "+base64.b64encode(raw).decode("ascii"))
    req.add_header("Accept","application/json")
    req.add_header("Content-Type","application/json-patch+json")
    data=_json.dumps(operations).encode("utf-8")
    with urllib.request.urlopen(req,data=data,timeout=20) as resp:
        return _json.loads(resp.read().decode("utf-8"))

def _v152_devops_state(task_status):
    v=(task_status or "").strip().lower()
    if v in ("done","completed","closed","klar"): return "Closed"
    if v in ("blocked","blockerad"): return "Active"
    return "Active"

def _ado_json_v1140(method,url,pat,payload=None):
    import urllib.request, urllib.error, base64, json as _json
    raw=(":"+pat).encode("utf-8")
    req=urllib.request.Request(url,method=method)
    req.add_header("Authorization","Basic "+base64.b64encode(raw).decode("ascii"))
    req.add_header("Accept","application/json")
    if payload is not None:
        req.add_header("Content-Type","application/json")
        data=_json.dumps(payload).encode("utf-8")
    else:
        data=None
    with urllib.request.urlopen(req,data=data,timeout=20) as resp:
        return _json.loads(resp.read().decode("utf-8"))

@app.route("/integrations/azure-devops/live",methods=["GET","POST"])
@login_required
def devops_live_v1140():
    if session.get("role") not in ("admin","pm"): abort(403)
    message=None
    with db() as conn:
        ensure_azure_devops_schema_v1010(conn)
        connections=conn.execute("""SELECT id,name,organization_url,project_name,enabled FROM azure_devops_connections WHERE enabled=1 ORDER BY name""").fetchall()
        if request.method=="POST":
            cid=request.form.get("connection_id",type=int)
            c=conn.execute("SELECT * FROM azure_devops_connections WHERE id=? AND enabled=1",(cid,)).fetchone()
            if not c: abort(404)
            if not c["pat_secret"]:
                flash("Anslutningen saknar PAT.","error")
            else:
                try:
                    from urllib.parse import quote
                    base=c["organization_url"].rstrip("/")+"/"+quote(c["project_name"])
                    wiql=_ado_json_v1140("POST",base+"/_apis/wit/wiql?api-version=7.1",c["pat_secret"],
                        {"query":"SELECT [System.Id] FROM WorkItems WHERE [System.TeamProject] = @project ORDER BY [System.ChangedDate] DESC"})
                    ids=[x["id"] for x in wiql.get("workItems",[])[:200]]
                    written=0
                    if ids:
                        batch=_ado_json_v1140("POST",base+"/_apis/wit/workitemsbatch?api-version=7.1",c["pat_secret"],
                            {"ids":ids,"fields":["System.Id","System.WorkItemType","System.Title","System.State","System.AssignedTo","System.IterationPath","System.TeamProject"]})
                        for w in batch.get("value",[]):
                            f=w.get("fields",{})
                            assigned=f.get("System.AssignedTo")
                            if isinstance(assigned,dict): assigned=assigned.get("displayName") or assigned.get("uniqueName")
                            conn.execute("""INSERT INTO azure_devops_work_item_links
                              (connection_id,project_id,task_id,work_item_id,work_item_type,title,state,assigned_to,iteration_path,url,last_synced_at)
                              VALUES(?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP)
                              ON CONFLICT(connection_id,work_item_id) DO UPDATE SET
                              work_item_type=excluded.work_item_type,title=excluded.title,state=excluded.state,
                              assigned_to=excluded.assigned_to,iteration_path=excluded.iteration_path,url=excluded.url,last_synced_at=CURRENT_TIMESTAMP""",
                              (cid,0,None,w["id"],f.get("System.WorkItemType"),f.get("System.Title"),f.get("System.State"),
                               assigned,f.get("System.IterationPath"),w.get("url")))
                            written+=1
                    conn.execute("""INSERT INTO azure_devops_sync_log(connection_id,direction,status,message,items_read,items_written)
                      VALUES(?,?,?,?,?,?)""",(cid,"DevOpsToProjectPlaner","Success","Live Work Item sync",len(ids),written))
                    conn.commit()
                    flash(f"Azure DevOps-synk klar: {written} Work Items.","success")
                except Exception as ex:
                    conn.execute("""INSERT INTO azure_devops_sync_log(connection_id,direction,status,message)
                      VALUES(?,?,?,?)""",(cid,"DevOpsToProjectPlaner","Failed",str(ex)[:500]))
                    conn.commit()
                    flash("Azure DevOps-synk misslyckades: "+str(ex),"error")
        syncs=conn.execute("""SELECT s.*,c.name connection_name FROM azure_devops_sync_log s
          LEFT JOIN azure_devops_connections c ON c.id=s.connection_id ORDER BY s.id DESC LIMIT 30""").fetchall()
    return render_template("devops_live_v1140.html",connections=connections,syncs=syncs)

def ensure_reporting_presets_v1150(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS reporting_presets(
      id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,audience TEXT NOT NULL,format TEXT NOT NULL,
      sections_json TEXT NOT NULL DEFAULT '[]',created_by INTEGER,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""")

@app.route("/reporting-studio-2",methods=["GET","POST"])
@login_required
def reporting_studio_v1150():
    with db() as conn:
        ensure_reporting_presets_v1150(conn)
        if request.method=="POST":
            name=(request.form.get("name") or "").strip()
            audience=(request.form.get("audience") or "Styrgrupp").strip()
            fmt=(request.form.get("format") or "PPTX").strip().upper()
            sections=request.form.getlist("sections")
            if name:
                conn.execute("INSERT INTO reporting_presets(name,audience,format,sections_json,created_by) VALUES(?,?,?,?,?)",
                             (name,audience,fmt,json.dumps(sections,ensure_ascii=False),session.get("user_id")))
                conn.commit(); flash("Rapportprofil sparad.","success")
        presets=conn.execute("SELECT * FROM reporting_presets ORDER BY id DESC").fetchall()
    return render_template("reporting_studio_v1150.html",presets=presets)

@app.get("/resource-planner-2")
@login_required
def resource_planner_v1160():
    with db() as conn:
        rows=conn.execute("""SELECT COALESCE(NULLIF(resource_name,''),u.display_name,u.username,'Resurs') resource,
          week_start,SUM(COALESCE(allocation_pct,0)) allocation_pct,SUM(COALESCE(planned_hours,0)) planned_hours,
          COUNT(DISTINCT project_id) projects
          FROM resource_allocations ra LEFT JOIN users u ON u.id=ra.user_id
          GROUP BY resource,week_start ORDER BY week_start,resource LIMIT 500""").fetchall()
    overloaded=[r for r in rows if (r["allocation_pct"] or 0)>100]
    return render_template("resource_planner_v1160.html",rows=rows,overloaded=overloaded)

@app.get("/portfolio-cockpit")
@login_required
def portfolio_cockpit_v1170():
    with db() as conn:
        projects=[dict(r) for r in conn.execute("""SELECT p.*,
          COALESCE((SELECT AVG(COALESCE(t.progress,0)) FROM tasks t WHERE t.project_id=p.id AND COALESCE(t.deleted_at,'')=''),0) progress,
          COALESCE((SELECT COUNT(*) FROM tasks t WHERE t.project_id=p.id AND COALESCE(t.deleted_at,'')='' AND t.end_date<date('now')
                    AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')),0) overdue,
          COALESCE((SELECT COUNT(*) FROM tasks t WHERE t.project_id=p.id AND COALESCE(t.deleted_at,'')=''
                    AND LOWER(COALESCE(t.status,'')) IN ('blocked','blockerad')),0) blocked,
          COALESCE((SELECT COUNT(*) FROM risks r WHERE r.project_id=p.id AND COALESCE(r.probability,0)*COALESCE(r.impact,0)>=12
                    AND LOWER(COALESCE(r.status,'')) NOT IN ('closed','stängd')),0) high_risks,
          COALESCE((SELECT SUM(planned) FROM project_costs c WHERE c.project_id=p.id),0) planned_cost,
          COALESCE((SELECT SUM(actual) FROM project_costs c WHERE c.project_id=p.id),0) actual_cost
          FROM projects p WHERE COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')='' ORDER BY p.name""").fetchall()]
        try:
            resource_rows=conn.execute("""SELECT week_start,
                  SUM(CASE WHEN allocation_pct>100 THEN 1 ELSE 0 END) overload_rows
                  FROM resource_allocations
                  WHERE week_start>=date('now','-7 day')
                  GROUP BY week_start ORDER BY week_start LIMIT 8""").fetchall()
        except Exception:
            resource_rows=[]

    total_planned=total_actual=0
    red=amber=green=0
    for p in projects:
        h=_v15_health(p["progress"],p["overdue"],p["high_risks"],p["blocked"],p["planned_cost"],p["actual_cost"],p.get("end_date"))
        p["health"]=h
        total_planned+=_v15_money(p["planned_cost"])
        total_actual+=_v15_money(p["actual_cost"])
        if h["rag"]=="red": red+=1
        elif h["rag"]=="amber": amber+=1
        else: green+=1
    projects.sort(key=lambda p:({"red":0,"amber":1,"green":2}[p["health"]["rag"]],p["name"].lower()))
    portfolio_progress=round(sum(float(p["progress"] or 0) for p in projects)/len(projects)) if projects else 0
    return render_template("portfolio_v1500.html",projects=projects,red=red,amber=amber,green=green,
                           total_planned=total_planned,total_actual=total_actual,
                           portfolio_progress=portfolio_progress,resource_rows=resource_rows)

def ensure_collaboration_v1180(conn):
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS project_comments(
      id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER NOT NULL,user_id INTEGER,body TEXT NOT NULL,
      created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS project_watchers(
      project_id INTEGER NOT NULL,user_id INTEGER NOT NULL,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY(project_id,user_id));
    """)

@app.route("/projects/<int:project_id>/collaboration",methods=["GET","POST"])
@login_required
def collaboration_ux_v1180(project_id):
    with db() as conn:
        ensure_collaboration_v1180(conn)
        project=conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(project_id,)).fetchone()
        if not project: abort(404)
        if request.method=="POST":
            action=request.form.get("action")
            if action=="comment":
                body=(request.form.get("body") or "").strip()
                if body:
                    conn.execute("INSERT INTO project_comments(project_id,user_id,body) VALUES(?,?,?)",(project_id,session.get("user_id"),body))
                    conn.commit()
            elif action=="watch":
                conn.execute("INSERT OR IGNORE INTO project_watchers(project_id,user_id) VALUES(?,?)",(project_id,session.get("user_id"))); conn.commit()
            elif action=="unwatch":
                conn.execute("DELETE FROM project_watchers WHERE project_id=? AND user_id=?",(project_id,session.get("user_id"))); conn.commit()
        comments=conn.execute("""SELECT c.*,COALESCE(u.display_name,u.username,'Användare') author FROM project_comments c
          LEFT JOIN users u ON u.id=c.user_id WHERE c.project_id=? ORDER BY c.id DESC LIMIT 100""",(project_id,)).fetchall()
        watching=conn.execute("SELECT 1 FROM project_watchers WHERE project_id=? AND user_id=?",(project_id,session.get("user_id"))).fetchone() is not None
    return render_template("collaboration_ux_v1180.html",project=project,comments=comments,watching=watching)

def ensure_automation_runtime_v1190(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS automation_runtime_log(
      id INTEGER PRIMARY KEY AUTOINCREMENT,rule_name TEXT NOT NULL,status TEXT NOT NULL,affected INTEGER NOT NULL DEFAULT 0,
      message TEXT,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""")

@app.route("/automation-runtime",methods=["GET","POST"])
@login_required
def automation_runtime_v1190():
    if session.get("role") not in ("admin","pm"): abort(403)
    with db() as conn:
        ensure_automation_runtime_v1190(conn)
        if request.method=="POST":
            rows=conn.execute("""SELECT t.id,t.project_id,t.title,t.owner_user_id,t.end_date,p.name project_name
              FROM tasks t JOIN projects p ON p.id=t.project_id
              WHERE t.deleted_at IS NULL AND p.deleted_at IS NULL AND t.end_date<date('now')
              AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')""").fetchall()
            affected=0
            for t in rows:
                if t["owner_user_id"]:
                    exists=conn.execute("""SELECT 1 FROM notifications WHERE user_id=? AND project_id=? AND message=? AND is_read=0""",
                      (t["owner_user_id"],t["project_id"],f"Försenad aktivitet: {t['title']}")).fetchone()
                    if not exists:
                        conn.execute("""INSERT INTO notifications(user_id,project_id,message,link,is_read,created_at)
                          VALUES(?,?,?,?,0,CURRENT_TIMESTAMP)""",(t["owner_user_id"],t["project_id"],f"Försenad aktivitet: {t['title']}",f"/projects/{t['project_id']}/workspace/tasks"))
                        affected+=1
            conn.execute("""INSERT INTO automation_runtime_log(rule_name,status,affected,message)
              VALUES('Overdue task notification','Success',?,'Created notifications for overdue assigned tasks')""",(affected,))
            conn.commit(); flash(f"Automation körd. {affected} nya notifieringar skapades.","success")
        logs=conn.execute("SELECT * FROM automation_runtime_log ORDER BY id DESC LIMIT 50").fetchall()
    return render_template("automation_runtime_v1190.html",logs=logs)

@app.get("/today")
@login_required
def intelligent_pm_v1200():
    uid=session.get("user_id")
    today=datetime.utcnow().date().isoformat()
    with db() as conn:
        projects=conn.execute("""SELECT p.id,p.name,p.end_date,
          COALESCE((SELECT AVG(COALESCE(t.progress,0)) FROM tasks t WHERE t.project_id=p.id AND t.deleted_at IS NULL),0) progress,
          COALESCE((SELECT COUNT(*) FROM tasks t WHERE t.project_id=p.id AND t.deleted_at IS NULL AND t.end_date<date('now')
            AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')),0) overdue,
          COALESCE((SELECT COUNT(*) FROM risks r WHERE r.project_id=p.id AND COALESCE(r.probability,0)*COALESCE(r.impact,0)>=12
            AND LOWER(COALESCE(r.status,'')) NOT IN ('closed','stängd')),0) risks
          FROM projects p WHERE p.deleted_at IS NULL AND p.archived_at IS NULL ORDER BY p.name""").fetchall()
        my_tasks=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
          WHERE t.deleted_at IS NULL AND p.deleted_at IS NULL AND t.owner_user_id=? AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')
          ORDER BY t.end_date LIMIT 20""",(uid,)).fetchall()
        unread=conn.execute("SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0",(uid,)).fetchone()[0]
        try:
            devops_failed=conn.execute("SELECT COUNT(*) FROM azure_devops_sync_log WHERE status='Failed'").fetchone()[0]
        except Exception:
            devops_failed=0
    attention=sum(1 for p in projects if (p["overdue"] or 0)>0 or (p["risks"] or 0)>0)
    return render_template("intelligent_pm_v1200.html",projects=projects,my_tasks=my_tasks,unread=unread,
                           devops_failed=devops_failed,attention=attention,today=today)

@app.get("/ux-cleanup")
@login_required
def ux_cleanup_v1201():
    primary=[("Idag","/today"),("Mitt arbete","/my-work-2"),("Projekt","/"),("Portfolio","/portfolio-cockpit")]
    return render_template("ux_cleanup_v1201.html",primary=primary)

@app.get("/navigation-2")
@login_required
def navigation_v1210():
    return render_template("navigation_v1210.html")

@app.get("/projects/<int:project_id>/workspace-3")
@login_required
def workspace_v1220(project_id):
    with db() as conn:
        p=conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(project_id,)).fetchone()
        if not p: abort(404)
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? AND deleted_at IS NULL ORDER BY end_date,sort_order LIMIT 200",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND LOWER(COALESCE(status,'')) NOT IN ('closed','stängd') ORDER BY probability*impact DESC LIMIT 20",(project_id,)).fetchall()
    overdue=[t for t in tasks if t["end_date"] and t["end_date"]<datetime.utcnow().date().isoformat() and (t["status"] or "").lower() not in ("done","completed","closed","klar")]
    progress=round(sum((t["progress"] or 0) for t in tasks)/len(tasks)) if tasks else 0
    return render_template("workspace_v1220.html",project=p,tasks=tasks,risks=risks,overdue=overdue,progress=progress)

@app.route("/projects/<int:project_id>/quick-add",methods=["GET","POST"])
@login_required
def quick_actions_v1230(project_id):
    with db() as conn:
        project=conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(project_id,)).fetchone()
        if not project: abort(404)
        if request.method=="POST":
            kind=(request.form.get("kind") or "task").strip()
            title=(request.form.get("title") or "").strip()
            if not title:
                flash("Ange en rubrik.","error")
            elif kind in ("task","milestone"):
                conn.execute("""INSERT INTO tasks(project_id,title,status,priority,progress,milestone,sort_order)
                  VALUES(?,?, 'Not started','Normal',0,?,COALESCE((SELECT MAX(sort_order)+1 FROM tasks WHERE project_id=?),1))""",
                  (project_id,title,1 if kind=="milestone" else 0,project_id))
                conn.commit(); flash("Skapad.","success")
            elif kind=="risk":
                conn.execute("""INSERT INTO risks(project_id,kind,title,probability,impact,status,created_at)
                  VALUES(?, 'Risk', ?, 1, 1, 'Open', CURRENT_TIMESTAMP)""",(project_id,title)); conn.commit(); flash("Risk skapad.","success")
            elif kind=="decision":
                conn.execute("""INSERT INTO decisions(project_id,title,decision,decision_date) VALUES(?,?,?,date('now'))""",(project_id,title,"")); conn.commit(); flash("Beslut skapat.","success")
        return render_template("quick_actions_v1230.html",project=project)

@app.get("/search")
@login_required
def universal_search_v1240():
    q=(request.args.get("q") or "").strip()
    results=[]
    if q:
        like="%"+q+"%"
        with db() as conn:
            for r in conn.execute("SELECT id,name FROM projects WHERE deleted_at IS NULL AND name LIKE ? LIMIT 20",(like,)).fetchall():
                results.append({"kind":"Projekt","title":r["name"],"url":f"/projects/{r['id']}/workspace-3"})
            for r in conn.execute("SELECT id,project_id,title FROM tasks WHERE deleted_at IS NULL AND title LIKE ? LIMIT 30",(like,)).fetchall():
                results.append({"kind":"Aktivitet","title":r["title"],"url":f"/projects/{r['project_id']}/workspace-3"})
            for r in conn.execute("SELECT id,project_id,title FROM risks WHERE title LIKE ? LIMIT 20",(like,)).fetchall():
                results.append({"kind":"Risk","title":r["title"],"url":f"/projects/{r['project_id']}/workspace-3"})
            try:
                for r in conn.execute("SELECT work_item_id,title FROM azure_devops_work_item_links WHERE title LIKE ? LIMIT 20",(like,)).fetchall():
                    results.append({"kind":"DevOps","title":f"#{r['work_item_id']} {r['title']}","url":"/integrations/azure-devops/live"})
            except Exception: pass
    return render_template("universal_search_v1240.html",q=q,results=results)

@app.route("/projects/<int:project_id>/smart-task",methods=["GET","POST"])
@login_required
def smart_forms_v1250(project_id):
    with db() as conn:
        project=conn.execute("SELECT * FROM projects WHERE id=? AND deleted_at IS NULL",(project_id,)).fetchone()
        if not project: abort(404)
        if request.method=="POST":
            title=(request.form.get("title") or "").strip()
            owner=(request.form.get("owner") or "").strip()
            end_date=(request.form.get("end_date") or "").strip() or None
            priority=(request.form.get("priority") or "Normal").strip()
            notes=(request.form.get("notes") or "").strip()
            if title:
                conn.execute("""INSERT INTO tasks(project_id,title,owner,end_date,status,priority,progress,milestone,notes,sort_order)
                VALUES(?,?,?,?, 'Not started',?,0,0,?,COALESCE((SELECT MAX(sort_order)+1 FROM tasks WHERE project_id=?),1))""",
                (project_id,title,owner,end_date,priority,notes,project_id)); conn.commit(); flash("Aktivitet skapad.","success")
        return render_template("smart_forms_v1250.html",project=project)

@app.get("/smart-tables")
@login_required
def smart_tables_v1260():
    q=(request.args.get("q") or "").strip()
    status=(request.args.get("status") or "").strip()
    sql="""SELECT t.id,t.project_id,t.title,t.owner,t.status,t.priority,t.end_date,t.progress,p.name project_name
           FROM tasks t JOIN projects p ON p.id=t.project_id WHERE t.deleted_at IS NULL AND p.deleted_at IS NULL"""
    args=[]
    if q: sql+=" AND (t.title LIKE ? OR p.name LIKE ?)"; args += ["%"+q+"%","%"+q+"%"]
    if status: sql+=" AND t.status=?"; args.append(status)
    sql+=" ORDER BY t.end_date LIMIT 500"
    with db() as conn: rows=conn.execute(sql,args).fetchall()
    return render_template("smart_tables_v1260.html",rows=rows,q=q,status=status)

@app.get("/notifications-2")
@login_required
def notifications_v1270():
    uid=session.get("user_id")
    with db() as conn:
        rows=conn.execute("""SELECT * FROM notifications WHERE user_id=? ORDER BY is_read ASC,created_at DESC LIMIT 200""",(uid,)).fetchall()
    unread=[r for r in rows if not r["is_read"]]
    return render_template("notifications_v1270.html",rows=rows,unread=unread)

@app.route("/projects/new-guided",methods=["GET","POST"])
@login_required
def guided_setup_v1280():
    user=current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] not in ("admin","pm"):
        abort(403)

    form={
        "name":(request.form.get("name") or "").strip(),
        "customer":(request.form.get("customer") or "").strip(),
        "project_manager":(request.form.get("project_manager") or "").strip(),
        "description":(request.form.get("description") or "").strip(),
        "start_date":(request.form.get("start_date") or "").strip(),
        "end_date":(request.form.get("end_date") or "").strip(),
    }

    if request.method=="POST":
        errors=[]
        if not form["name"]:
            errors.append("Projektnamn krävs.")
        if len(form["name"])>200:
            errors.append("Projektnamnet är för långt.")
        start_d=_v15_date(form["start_date"]) if form["start_date"] else None
        end_d=_v15_date(form["end_date"]) if form["end_date"] else None
        if form["start_date"] and not start_d:
            errors.append("Ogiltigt startdatum.")
        if form["end_date"] and not end_d:
            errors.append("Ogiltigt slutdatum.")
        if start_d and end_d and end_d < start_d:
            errors.append("Slutdatum kan inte ligga före startdatum.")

        if errors:
            for msg in errors:
                flash(msg,"error")
        else:
            with db() as conn:
                cur=conn.execute("""INSERT INTO projects(
                    name,customer,project_manager,description,start_date,end_date,created_by,created_at
                ) VALUES(?,?,?,?,?,?,?,CURRENT_TIMESTAMP)""",(
                    form["name"],form["customer"],form["project_manager"],form["description"],
                    form["start_date"],form["end_date"],session.get("user_id")
                ))
                pid=cur.lastrowid
                conn.execute("""INSERT OR IGNORE INTO project_members(project_id,user_id,project_role,added_at,added_by)
                                VALUES(?,?,?,CURRENT_TIMESTAMP,?)""",
                             (pid,session.get("user_id"),"pm",session.get("user_id")))
                conn.commit()
            audit(pid,"project",pid,"create",f"Projekt skapat: {form['name']}")
            flash("Projektet skapades.","success")
            return redirect(url_for("ultimate_project_v140",project_id=pid))

    return render_template("guided_setup_v1280.html",form=form)

def ensure_personalization_v1290(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS user_favorites(
      user_id INTEGER NOT NULL,project_id INTEGER NOT NULL,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY(user_id,project_id))""")

@app.route("/personalization",methods=["GET","POST"])
@login_required
def personalization_v1290():
    uid=session.get("user_id")
    with db() as conn:
        ensure_personalization_v1290(conn)
        if request.method=="POST":
            pid=request.form.get("project_id",type=int); action=request.form.get("action")
            if pid and action=="favorite": conn.execute("INSERT OR IGNORE INTO user_favorites(user_id,project_id) VALUES(?,?)",(uid,pid))
            if pid and action=="remove": conn.execute("DELETE FROM user_favorites WHERE user_id=? AND project_id=?",(uid,pid))
            conn.commit()
        projects=conn.execute("""SELECT p.*,CASE WHEN f.project_id IS NULL THEN 0 ELSE 1 END favorite
          FROM projects p LEFT JOIN user_favorites f ON f.project_id=p.id AND f.user_id=?
          WHERE p.deleted_at IS NULL AND p.archived_at IS NULL ORDER BY favorite DESC,p.name""",(uid,)).fetchall()
    return render_template("personalization_v1290.html",projects=projects)

@app.get("/simple")
@login_required
def project_planer_simple_v1300():
    uid=session.get("user_id")
    with db() as conn:
        try: ensure_personalization_v1290(conn)
        except Exception: pass
        projects=conn.execute("""SELECT p.id,p.name,
          COALESCE((SELECT AVG(COALESCE(t.progress,0)) FROM tasks t WHERE t.project_id=p.id AND t.deleted_at IS NULL),0) progress,
          COALESCE((SELECT COUNT(*) FROM tasks t WHERE t.project_id=p.id AND t.deleted_at IS NULL AND t.end_date<date('now')
          AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')),0) overdue
          FROM projects p WHERE p.deleted_at IS NULL AND p.archived_at IS NULL ORDER BY p.name LIMIT 50""").fetchall()
        mine=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
          WHERE t.owner_user_id=? AND t.deleted_at IS NULL AND p.deleted_at IS NULL
          AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar') ORDER BY t.end_date LIMIT 12""",(uid,)).fetchall()
        unread=conn.execute("SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0",(uid,)).fetchone()[0]
    attention=sum(1 for p in projects if p["overdue"])
    return render_template("project_planer_simple_v1300.html",projects=projects,mine=mine,unread=unread,attention=attention)




def ensure_v152_schema(conn):
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS schedule_change_batches(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      project_id INTEGER NOT NULL,
      title TEXT NOT NULL,
      reason TEXT,
      created_by INTEGER,
      created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
      reverted_at TEXT,
      reverted_by INTEGER
    );
    CREATE TABLE IF NOT EXISTS schedule_change_items(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      batch_id INTEGER NOT NULL,
      task_id INTEGER NOT NULL,
      old_start TEXT, old_end TEXT,
      new_start TEXT, new_end TEXT,
      FOREIGN KEY(batch_id) REFERENCES schedule_change_batches(id) ON DELETE CASCADE
    );
    CREATE INDEX IF NOT EXISTS idx_sched_batch_project ON schedule_change_batches(project_id,created_at);
    CREATE INDEX IF NOT EXISTS idx_sched_item_batch ON schedule_change_items(batch_id);
    """)

def _v151_workday_add(d, days):
    """Add working days (Mon-Fri). Positive and negative values supported."""
    if not d:
        return None
    cur=d
    step=1 if days >= 0 else -1
    remaining=abs(int(days))
    while remaining:
        cur += timedelta(days=step)
        if cur.weekday() < 5:
            remaining -= 1
    return cur

def _v151_workdays_between(start, end):
    if not start or not end or end <= start:
        return 0
    cur=start
    count=0
    while cur < end:
        cur += timedelta(days=1)
        if cur.weekday() < 5:
            count += 1
    return count

def _v151_schedule(tasks, links):
    """Deterministic dependency forecast. Does not mutate stored project data."""
    items={int(t["id"]):dict(t) for t in tasks}
    incoming={tid:[] for tid in items}
    outgoing={tid:[] for tid in items}
    for raw in links:
        l=dict(raw)
        p=int(l["predecessor_id"]); q=int(l["successor_id"])
        if p in items and q in items:
            incoming[q].append(l)
            outgoing[p].append(l)

    indegree={tid:len(incoming[tid]) for tid in items}
    queue=[tid for tid,d in indegree.items() if d==0]
    order=[]
    while queue:
        tid=queue.pop(0); order.append(tid)
        for l in outgoing.get(tid,[]):
            q=int(l["successor_id"]); indegree[q]-=1
            if indegree[q]==0: queue.append(q)
    cyclic=len(order)!=len(items)
    if cyclic:
        order=list(items)

    result={}
    today=date.today()
    for tid in order:
        t=items[tid]
        orig_start=_v15_date(t.get("start_date"))
        orig_end=_v15_date(t.get("end_date"))
        duration=int(t.get("duration_days") or 0)
        if duration <= 0 and orig_start and orig_end:
            duration=max(1,_v151_workdays_between(orig_start,orig_end))
        if duration <= 0: duration=1

        start=orig_start or today
        end=orig_end or _v151_workday_add(start,duration-1)
        reasons=[]

        for l in incoming.get(tid,[]):
            pred=result.get(int(l["predecessor_id"]))
            if not pred: continue
            typ=(l.get("link_type") or "FS").upper()
            lag=int(l.get("lag_days") or 0)
            ps=pred["forecast_start"]; pe=pred["forecast_end"]
            if typ=="SS":
                candidate=_v151_workday_add(ps,lag)
                if candidate and candidate>start:
                    start=candidate; end=_v151_workday_add(start,duration-1); reasons.append(f"SS +{lag}")
            elif typ=="FF":
                candidate_end=_v151_workday_add(pe,lag)
                if candidate_end and candidate_end>end:
                    end=candidate_end; start=_v151_workday_add(end,-(duration-1)); reasons.append(f"FF +{lag}")
            elif typ=="SF":
                candidate_end=_v151_workday_add(ps,lag)
                if candidate_end and candidate_end>end:
                    end=candidate_end; start=_v151_workday_add(end,-(duration-1)); reasons.append(f"SF +{lag}")
            else: # FS
                candidate=_v151_workday_add(pe,lag+1)
                if candidate and candidate>start:
                    start=candidate; end=_v151_workday_add(start,duration-1); reasons.append(f"FS +{lag}")
        result[tid]={
            "id":tid,"title":t.get("title"),"owner":t.get("owner"),"status":t.get("status"),
            "duration":duration,"original_start":orig_start,"original_end":orig_end,
            "forecast_start":start,"forecast_end":end,"reasons":reasons
        }

    project_finish=max((r["forecast_end"] for r in result.values() if r["forecast_end"]), default=None)
    # Backward pass, approximate CPM using workdays.
    latest_finish={tid:project_finish for tid in result}
    for tid in reversed(order):
        successors=outgoing.get(tid,[])
        if successors:
            candidates=[]
            for l in successors:
                succ=result.get(int(l["successor_id"]))
                if not succ: continue
                typ=(l.get("link_type") or "FS").upper()
                lag=int(l.get("lag_days") or 0)
                if typ=="FS":
                    candidates.append(_v151_workday_add(succ["forecast_start"],-(lag+1)))
                else:
                    candidates.append(succ["forecast_end"])
            vals=[x for x in candidates if x]
            if vals: latest_finish[tid]=min(vals)
        r=result[tid]
        lf=latest_finish.get(tid) or r["forecast_end"]
        slack=_v151_workdays_between(r["forecast_end"],lf) if lf and r["forecast_end"] and lf>=r["forecast_end"] else 0
        r["slack_days"]=slack
        r["critical"]=slack<=0 and bool(outgoing.get(tid) or incoming.get(tid))
        r["delay_days"]=_v151_workdays_between(r["original_end"],r["forecast_end"]) if r["original_end"] and r["forecast_end"] and r["forecast_end"]>r["original_end"] else 0

    return {"tasks":result,"order":order,"cyclic":cyclic,"project_finish":project_finish}

def _v151_capacity(rows):
    grouped={}
    for raw in rows:
        r=dict(raw)
        key=(r.get("user_id") or 0, r.get("resource_name") or "Resurs", r.get("week_start") or "")
        g=grouped.setdefault(key,{"user_id":key[0],"resource":key[1],"week_start":key[2],"allocation_pct":0.0,"planned_hours":0.0})
        g["allocation_pct"] += float(r.get("allocation_pct") or 0)
        g["planned_hours"] += float(r.get("planned_hours") or 0)
    vals=list(grouped.values())
    for g in vals:
        g["status"]="over" if g["allocation_pct"]>100 else "high" if g["allocation_pct"]>=85 else "ok"
        g["available_pct"]=max(0,100-g["allocation_pct"])
    vals.sort(key=lambda x:(x["week_start"],-x["allocation_pct"],x["resource"].lower()))
    return vals

def _v15_is_closed(value):
    return (value or "").strip().lower() in ("done","completed","closed","klar","stängd")

def _v15_date(value):
    try:
        return date.fromisoformat((value or "")[:10])
    except Exception:
        return None

def _v15_money(value):
    try:
        return float(value or 0)
    except Exception:
        return 0.0

def _v15_health(progress=0, overdue=0, high_risks=0, blocked=0, planned=0, actual=0, end_date=None):
    score = 0
    score += min(40, int(overdue or 0) * 8)
    score += min(25, int(high_risks or 0) * 8)
    score += min(20, int(blocked or 0) * 10)
    if planned and actual > planned:
        score += 20
    end = _v15_date(end_date)
    if end and end < date.today() and float(progress or 0) < 100:
        score += 20
    score = min(100, score)
    rag = "green" if score < 20 else "amber" if score < 50 else "red"
    return {"score": score, "rag": rag}

def _v15_forecast(project, tasks):
    open_tasks = [t for t in tasks if not _v15_is_closed(t["status"])]
    overdue = [t for t in open_tasks if _v15_date(t["end_date"]) and _v15_date(t["end_date"]) < date.today()]
    if not overdue:
        return {"days": 0, "label": "Enligt plan", "projected_end": project["end_date"] or ""}
    max_delay = max((date.today() - _v15_date(t["end_date"])).days for t in overdue)
    end = _v15_date(project["end_date"])
    projected = end + timedelta(days=max_delay) if end else None
    return {
        "days": max_delay,
        "label": f"+{max_delay} dagar" if max_delay else "Enligt plan",
        "projected_end": projected.isoformat() if projected else (project["end_date"] or "")
    }

def _v15_project_access_sql():
    return """SELECT p.* FROM projects p
              LEFT JOIN project_members pm ON pm.project_id=p.id AND pm.user_id=?
              WHERE COALESCE(p.deleted_at,'')='' AND COALESCE(p.archived_at,'')=''
                AND (?='admin' OR p.created_by=? OR pm.user_id IS NOT NULL)
              GROUP BY p.id"""

def _ultimate_health_v140(project, tasks, risks, costs):
    open_tasks=[t for t in tasks if (t["status"] or "").lower() not in ("done","completed","closed","klar")]
    overdue=[t for t in open_tasks if t["end_date"] and t["end_date"] < date.today().isoformat()]
    blocked=[t for t in open_tasks if (t["status"] or "").lower() in ("blocked","blockerad")]
    high_risks=[r for r in risks if (r["probability"] or 0)*(r["impact"] or 0) >= 12 and (r["status"] or "").lower() not in ("closed","stängd")]
    planned=float(costs["planned"] or 0) if costs else 0
    actual=float(costs["actual"] or 0) if costs else 0
    budget_over=planned>0 and actual>planned
    score=0
    score += min(40,len(overdue)*8)
    score += min(25,len(high_risks)*8)
    score += min(20,len(blocked)*10)
    if budget_over: score += 20
    rag="green" if score<20 else "amber" if score<50 else "red"
    return {"rag":rag,"score":min(100,score),"overdue":overdue,"blocked":blocked,
            "high_risks":high_risks,"budget_over":budget_over,"planned":planned,"actual":actual}

@app.get("/home")
@login_required
def ultimate_home_v140():
    u=current_user()
    today=date.today().isoformat()
    with db() as conn:
        projects=[dict(r) for r in visible_projects_for_user()]
        ids=[p["id"] for p in projects]
        cards=[]
        attention_items=[]
        my_tasks=[]
        unread=conn.execute("SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0",(u["id"],)).fetchone()[0]
        failed_sync=0
        overloaded_people=0

        if ids:
            marks=",".join("?" for _ in ids)
            task_stats={r["project_id"]:dict(r) for r in conn.execute(f"""
                SELECT project_id,
                       COUNT(*) total,
                       COALESCE(AVG(COALESCE(progress,0)),0) progress,
                       SUM(CASE WHEN COALESCE(milestone,0)=0 AND LOWER(COALESCE(status,'')) NOT IN ('done','completed','closed','klar')
                                THEN 1 ELSE 0 END) open_count,
                       SUM(CASE WHEN end_date<>'' AND end_date<date('now')
                                 AND LOWER(COALESCE(status,'')) NOT IN ('done','completed','closed','klar')
                                THEN 1 ELSE 0 END) overdue,
                       SUM(CASE WHEN LOWER(COALESCE(status,'')) IN ('blocked','blockerad') THEN 1 ELSE 0 END) blocked
                FROM tasks
                WHERE project_id IN ({marks}) AND COALESCE(deleted_at,'')=''
                GROUP BY project_id
            """, ids).fetchall()}

            risk_stats={r["project_id"]:dict(r) for r in conn.execute(f"""
                SELECT project_id,
                       SUM(CASE WHEN COALESCE(probability,0)*COALESCE(impact,0)>=12
                                 AND LOWER(COALESCE(status,'')) NOT IN ('closed','stängd')
                                THEN 1 ELSE 0 END) high_risks
                FROM risks WHERE project_id IN ({marks})
                GROUP BY project_id
            """, ids).fetchall()}

            cost_stats={r["project_id"]:dict(r) for r in conn.execute(f"""
                SELECT project_id,COALESCE(SUM(planned),0) planned,COALESCE(SUM(actual),0) actual
                FROM project_costs WHERE project_id IN ({marks}) GROUP BY project_id
            """, ids).fetchall()}

            milestone_rows=conn.execute(f"""
                SELECT id,project_id,title,end_date,status
                FROM tasks
                WHERE project_id IN ({marks}) AND COALESCE(deleted_at,'')='' AND COALESCE(milestone,0)=1
                  AND end_date>=date('now')
                ORDER BY end_date LIMIT 120
            """, ids).fetchall()
            milestones_by={}
            for m in milestone_rows:
                milestones_by.setdefault(m["project_id"],[]).append(m)

            for p in projects:
                ts=task_stats.get(p["id"],{})
                rs=risk_stats.get(p["id"],{})
                cs=cost_stats.get(p["id"],{})
                progress=round(float(ts.get("progress") or 0))
                health=_v15_health(progress,ts.get("overdue",0),rs.get("high_risks",0),ts.get("blocked",0),
                                   _v15_money(cs.get("planned")),_v15_money(cs.get("actual")),p.get("end_date"))
                cards.append({
                    "project":p,
                    "progress":progress,
                    "open_count":int(ts.get("open_count") or 0),
                    "overdue":int(ts.get("overdue") or 0),
                    "blocked":int(ts.get("blocked") or 0),
                    "high_risks":int(rs.get("high_risks") or 0),
                    "planned":_v15_money(cs.get("planned")),
                    "actual":_v15_money(cs.get("actual")),
                    "health":health,
                    "upcoming":milestones_by.get(p["id"],[])[:2]
                })

            my_tasks=conn.execute(f"""
                SELECT t.id,t.project_id,t.title,t.status,t.priority,t.end_date,t.progress,p.name project_name
                FROM tasks t JOIN projects p ON p.id=t.project_id
                WHERE t.project_id IN ({marks}) AND COALESCE(t.deleted_at,'')=''
                  AND (t.owner_user_id=? OR LOWER(COALESCE(t.owner,'')) IN (LOWER(?),LOWER(?)))
                  AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')
                ORDER BY
                  CASE WHEN t.end_date<>'' AND t.end_date<date('now') THEN 0 ELSE 1 END,
                  CASE WHEN LOWER(COALESCE(t.priority,'')) IN ('high','hög','critical','kritisk') THEN 0 ELSE 1 END,
                  CASE WHEN t.end_date IS NULL OR t.end_date='' THEN 1 ELSE 0 END,
                  t.end_date
                LIMIT 40
            """, ids+[u["id"],u["display_name"],u["username"]]).fetchall()

            overdue_rows=conn.execute(f"""
                SELECT t.id,t.project_id,t.title,t.end_date,p.name project_name
                FROM tasks t JOIN projects p ON p.id=t.project_id
                WHERE t.project_id IN ({marks}) AND COALESCE(t.deleted_at,'')=''
                  AND t.end_date<>'' AND t.end_date<date('now')
                  AND LOWER(COALESCE(t.status,'')) NOT IN ('done','completed','closed','klar')
                ORDER BY t.end_date LIMIT 8
            """, ids).fetchall()
            for r in overdue_rows:
                days=(date.today()-_v15_date(r["end_date"])).days if _v15_date(r["end_date"]) else 0
                attention_items.append({"severity":"red","kind":"Försenad aktivitet","title":r["title"],
                                        "detail":f"{r['project_name']} · {days} dagar sen",
                                        "url":f"/projects/{r['project_id']}/ultimate"})

            risk_rows=conn.execute(f"""
                SELECT r.id,r.project_id,r.title,r.probability,r.impact,p.name project_name
                FROM risks r JOIN projects p ON p.id=r.project_id
                WHERE r.project_id IN ({marks})
                  AND COALESCE(r.probability,0)*COALESCE(r.impact,0)>=12
                  AND LOWER(COALESCE(r.status,'')) NOT IN ('closed','stängd')
                ORDER BY COALESCE(r.probability,0)*COALESCE(r.impact,0) DESC LIMIT 6
            """, ids).fetchall()
            for r in risk_rows:
                attention_items.append({"severity":"amber","kind":"Hög risk","title":r["title"],
                                        "detail":f"{r['project_name']} · score {(r['probability'] or 0)*(r['impact'] or 0)}",
                                        "url":f"/projects/{r['project_id']}/ultimate"})

            try:
                change_rows=conn.execute(f"""
                    SELECT c.id,c.project_id,c.title,c.status,p.name project_name
                    FROM change_requests c JOIN projects p ON p.id=c.project_id
                    WHERE c.project_id IN ({marks})
                      AND LOWER(COALESCE(c.status,'')) IN ('proposed','submitted','pending')
                    ORDER BY c.id DESC LIMIT 6
                """, ids).fetchall()
                for c in change_rows:
                    attention_items.append({"severity":"amber","kind":"Ändringsbegäran","title":c["title"],
                                            "detail":f"{c['project_name']} · väntar på hantering",
                                            "url":"/control-center"})
            except Exception:
                pass

            try:
                overload=conn.execute(f"""
                    SELECT COALESCE(NULLIF(resource_name,''),u.display_name,u.username,'Resurs') resource,
                           week_start,SUM(COALESCE(allocation_pct,0)) allocation_pct
                    FROM resource_allocations ra LEFT JOIN users u ON u.id=ra.user_id
                    WHERE ra.project_id IN ({marks}) AND week_start>=date('now','-7 day')
                    GROUP BY resource,week_start HAVING SUM(COALESCE(allocation_pct,0))>100
                    ORDER BY allocation_pct DESC LIMIT 10
                """, ids).fetchall()
                overloaded_people=len(overload)
                for r in overload[:3]:
                    attention_items.append({"severity":"amber","kind":"Överbelagd resurs","title":r["resource"],
                                            "detail":f"{r['allocation_pct']:.0f}% · vecka {r['week_start']}",
                                            "url":"/resource-planner-2"})
            except Exception:
                pass

        try:
            failed_sync=conn.execute("SELECT COUNT(*) FROM azure_devops_sync_log WHERE LOWER(status)='failed'").fetchone()[0]
        except Exception:
            failed_sync=0

    cards.sort(key=lambda c:({"red":0,"amber":1,"green":2}.get(c["health"]["rag"],3),c["project"]["name"].lower()))
    attention_items.sort(key=lambda x:{"red":0,"amber":1,"green":2}.get(x["severity"],3))
    attention=sum(1 for c in cards if c["health"]["rag"]!="green")
    return render_template("home_v1500.html",cards=cards,my_tasks=my_tasks,unread=unread,attention=attention,
                           attention_items=attention_items[:12],failed_sync=failed_sync,
                           overloaded_people=overloaded_people,user=u,today=today)

@app.get("/projects/<int:project_id>/ultimate")
@login_required
def ultimate_project_v140(project_id):
    project=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("""SELECT * FROM tasks WHERE project_id=? AND COALESCE(deleted_at,'')=''
                              ORDER BY CASE WHEN end_date IS NULL OR end_date='' THEN 1 ELSE 0 END,end_date,sort_order,id""",(project_id,)).fetchall()
        risks=conn.execute("""SELECT * FROM risks WHERE project_id=? ORDER BY COALESCE(probability,0)*COALESCE(impact,0) DESC,id DESC""",(project_id,)).fetchall()
        costs=conn.execute("""SELECT COALESCE(SUM(planned),0) planned,COALESCE(SUM(actual),0) actual FROM project_costs WHERE project_id=?""",(project_id,)).fetchone()
        members=conn.execute("""SELECT pm.*,COALESCE(u.display_name,u.username) display_name
             FROM project_members pm JOIN users u ON u.id=pm.user_id WHERE pm.project_id=? ORDER BY display_name""",(project_id,)).fetchall()
        actions=conn.execute("""SELECT * FROM action_items WHERE project_id=? AND LOWER(COALESCE(status,'')) NOT IN ('done','closed','klar')
                                ORDER BY CASE WHEN due_date='' THEN 1 ELSE 0 END,due_date LIMIT 12""",(project_id,)).fetchall()
        try:
            changes=conn.execute("""SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC LIMIT 8""",(project_id,)).fetchall()
        except Exception:
            changes=[]
        try:
            decisions=conn.execute("""SELECT * FROM decisions WHERE project_id=? ORDER BY id DESC LIMIT 8""",(project_id,)).fetchall()
        except Exception:
            decisions=[]
        try:
            devops=conn.execute("""SELECT state,COUNT(*) c FROM azure_devops_work_item_links
               WHERE project_id=? GROUP BY state ORDER BY c DESC""",(project_id,)).fetchall()
        except Exception:
            devops=[]
        try:
            comments=conn.execute("""SELECT c.*,COALESCE(u.display_name,u.username,'Användare') author
               FROM project_comments c LEFT JOIN users u ON u.id=c.user_id
               WHERE c.project_id=? ORDER BY c.id DESC LIMIT 6""",(project_id,)).fetchall()
        except Exception:
            comments=[]
        try:
            status_report=conn.execute("""SELECT * FROM status_reports WHERE project_id=? ORDER BY report_date DESC,id DESC LIMIT 1""",(project_id,)).fetchone()
        except Exception:
            status_report=None

    health=_ultimate_health_v140(project,tasks,risks,costs)
    progress=round(sum((t["progress"] or 0) for t in tasks)/len(tasks)) if tasks else 0
    milestones=sorted([t for t in tasks if t["milestone"] and t["end_date"]],key=lambda x:x["end_date"])[:10]
    open_tasks=[t for t in tasks if not _v15_is_closed(t["status"])]
    forecast=_v15_forecast(project,tasks)
    with db() as conn:
        schedule_links=conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days FROM task_links WHERE project_id=?""",(project_id,)).fetchall()
    schedule_plan=_v151_schedule(tasks,schedule_links)
    schedule_critical=sum(1 for x in schedule_plan["tasks"].values() if x["critical"])
    schedule_delayed=sum(1 for x in schedule_plan["tasks"].values() if x["delay_days"]>0)

    next_actions=[]
    for t in health["overdue"][:4]:
        days=(date.today()-_v15_date(t["end_date"])).days if _v15_date(t["end_date"]) else 0
        next_actions.append({"severity":"red","type":"Försenad","title":t["title"],"detail":f"{days} dagar sen","url":f"/projects/{project_id}/workspace-pro"})
    for t in health["blocked"][:3]:
        next_actions.append({"severity":"red","type":"Blockerad","title":t["title"],"detail":t["owner"] or "Ej tilldelad","url":f"/projects/{project_id}/workspace-pro"})
    for r in health["high_risks"][:3]:
        next_actions.append({"severity":"amber","type":"Hög risk","title":r["title"],"detail":f"P{r['probability']} × I{r['impact']}","url":f"/projects/{project_id}/risk-center"})
    for a in actions[:3]:
        due=_v15_date(a["due_date"])
        sev="red" if due and due < date.today() else "amber"
        next_actions.append({"severity":sev,"type":"Åtgärd","title":a["title"],"detail":a["due_date"] or "Utan datum","url":"/my-work-2"})
    if health["budget_over"]:
        next_actions.append({"severity":"red","type":"Budget","title":"Utfall över planerad kostnad",
                             "detail":f"{health['actual']-health['planned']:.0f} över plan","url":"/finance-control"})
    next_actions.sort(key=lambda x:{"red":0,"amber":1}.get(x["severity"],2))

    return render_template("project_cockpit_v1500.html",project=project,tasks=tasks,open_tasks=open_tasks,
                           risks=risks,costs=costs,members=members,devops=devops,comments=comments,
                           health=health,progress=progress,milestones=milestones,next_actions=next_actions,
                           actions=actions,changes=changes,decisions=decisions,status_report=status_report,
                           forecast=forecast,schedule_plan=schedule_plan,schedule_critical=schedule_critical,schedule_delayed=schedule_delayed)

@app.get("/projects/<int:project_id>/impact")
@login_required
def project_impact_v1500(project_id):
    project=project_or_404(project_id)
    shift=max(0,min(180,request.args.get("shift",7,type=int)))
    start_task=request.args.get("task_id",type=int)
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("""SELECT id,title,start_date,end_date,status,progress
                                                FROM tasks WHERE project_id=? AND COALESCE(deleted_at,'')=''
                                                ORDER BY sort_order,id""",(project_id,)).fetchall()]
        links=[dict(r) for r in conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days
                                                FROM task_links WHERE project_id=?""",(project_id,)).fetchall()]
    by_id={t["id"]:t for t in tasks}
    graph={}
    for l in links:
        graph.setdefault(l["predecessor_id"],[]).append(l["successor_id"])
    if not start_task and tasks:
        open_dated=[t for t in tasks if not _v15_is_closed(t["status"]) and _v15_date(t["end_date"])]
        start_task=(open_dated[0]["id"] if open_dated else tasks[0]["id"])
    affected=[]
    seen=set()
    queue=[start_task] if start_task else []
    while queue:
        tid=queue.pop(0)
        if tid in seen: continue
        seen.add(tid)
        t=by_id.get(tid)
        if t:
            ns=_v15_date(t["start_date"]); ne=_v15_date(t["end_date"])
            affected.append({
                **t,
                "new_start":(ns+timedelta(days=shift)).isoformat() if ns else "",
                "new_end":(ne+timedelta(days=shift)).isoformat() if ne else ""
            })
        queue.extend(graph.get(tid,[]))
    old_end=_v15_date(project["end_date"])
    projected=(old_end+timedelta(days=shift)).isoformat() if old_end and affected else (project["end_date"] or "")
    return render_template("impact_v1500.html",project=project,tasks=tasks,affected=affected,
                           shift=shift,start_task=start_task,projected_end=projected)

@app.get("/projects/<int:project_id>/schedule")
@login_required
def schedule_engine_v1510(project_id):
    project=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("""SELECT id,title,owner,status,start_date,end_date,duration_days,progress
                              FROM tasks WHERE project_id=? AND COALESCE(deleted_at,'')=''
                              ORDER BY sort_order,id""",(project_id,)).fetchall()
        links=conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days
                              FROM task_links WHERE project_id=?""",(project_id,)).fetchall()
    plan=_v151_schedule(tasks,links)
    rows=[plan["tasks"][tid] for tid in plan["order"] if tid in plan["tasks"]]
    critical=sum(1 for r in rows if r["critical"])
    delayed=sum(1 for r in rows if r["delay_days"]>0)
    return render_template("schedule_v1510.html",project=project,rows=rows,plan=plan,
                           critical=critical,delayed=delayed)

@app.get("/capacity")
@login_required
def capacity_v1510():
    with db() as conn:
        try:
            rows=conn.execute("""SELECT ra.user_id,
                       COALESCE(NULLIF(ra.resource_name,''),u.display_name,u.username,'Resurs') resource_name,
                       ra.week_start,ra.allocation_pct,ra.planned_hours,ra.project_id,p.name project_name
                FROM resource_allocations ra
                LEFT JOIN users u ON u.id=ra.user_id
                LEFT JOIN projects p ON p.id=ra.project_id
                WHERE ra.week_start>=date('now','-7 day')
                ORDER BY ra.week_start,resource_name""").fetchall()
        except Exception:
            rows=[]
    capacity=_v151_capacity(rows)
    weeks=sorted({r["week_start"] for r in capacity if r["week_start"]})[:12]
    people=sorted({r["resource"] for r in capacity})
    lookup={(r["resource"],r["week_start"]):r for r in capacity}
    over=sum(1 for r in capacity if r["status"]=="over")
    high=sum(1 for r in capacity if r["status"]=="high")
    return render_template("capacity_v1510.html",capacity=capacity,weeks=weeks,people=people,
                           lookup=lookup,over=over,high=high)

@app.get("/project-health")
@login_required
def project_health_v1510():
    projects=visible_projects_for_user()
    cards=[]
    with db() as conn:
        for p0 in projects:
            p=dict(p0)
            tasks=conn.execute("""SELECT * FROM tasks WHERE project_id=? AND COALESCE(deleted_at,'')=''""",(p["id"],)).fetchall()
            risks=conn.execute("""SELECT * FROM risks WHERE project_id=?""",(p["id"],)).fetchall()
            costs=conn.execute("""SELECT COALESCE(SUM(planned),0) planned,COALESCE(SUM(actual),0) actual
                                  FROM project_costs WHERE project_id=?""",(p["id"],)).fetchone()
            links=conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days FROM task_links WHERE project_id=?""",(p["id"],)).fetchall()
            health=_ultimate_health_v140(p,tasks,risks,costs)
            sched=_v151_schedule(tasks,links)
            cards.append({"project":p,"health":health,"finish":sched["project_finish"],
                          "critical":sum(1 for x in sched["tasks"].values() if x["critical"]),
                          "delayed":sum(1 for x in sched["tasks"].values() if x["delay_days"]>0)})
    cards.sort(key=lambda x:({"red":0,"amber":1,"green":2}.get(x["health"]["rag"],3),x["project"]["name"].lower()))
    return render_template("project_health_v1510.html",cards=cards)

@app.route("/projects/<int:project_id>/reschedule",methods=["GET"])
@login_required
def reschedule_preview_v1520(project_id):
    project=project_or_404(project_id)
    with db() as conn:
        ensure_v152_schema(conn)
        tasks=conn.execute("""SELECT id,title,owner,status,start_date,end_date,duration_days,progress
                              FROM tasks WHERE project_id=? AND COALESCE(deleted_at,'')=''
                              ORDER BY sort_order,id""",(project_id,)).fetchall()
        links=conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days
                              FROM task_links WHERE project_id=?""",(project_id,)).fetchall()
        history=conn.execute("""SELECT b.*,COALESCE(u.display_name,u.username,'System') created_name,
                               (SELECT COUNT(*) FROM schedule_change_items i WHERE i.batch_id=b.id) item_count
                               FROM schedule_change_batches b LEFT JOIN users u ON u.id=b.created_by
                               WHERE b.project_id=? ORDER BY b.id DESC LIMIT 20""",(project_id,)).fetchall()
    plan=_v151_schedule(tasks,links)
    changed=[]
    for tid in plan["order"]:
        r=plan["tasks"].get(tid)
        if not r: continue
        os=r["original_start"].isoformat() if r["original_start"] else ""
        oe=r["original_end"].isoformat() if r["original_end"] else ""
        ns=r["forecast_start"].isoformat() if r["forecast_start"] else ""
        ne=r["forecast_end"].isoformat() if r["forecast_end"] else ""
        if os!=ns or oe!=ne:
            changed.append({**r,"old_start":os,"old_end":oe,"new_start":ns,"new_end":ne})
    return render_template("reschedule_v1520.html",project=project,changed=changed,plan=plan,history=history)

@app.post("/projects/<int:project_id>/reschedule/apply")
@login_required
def reschedule_apply_v1520(project_id):
    project=project_or_404(project_id,write=True)
    selected={int(x) for x in request.form.getlist("task_id") if str(x).isdigit()}
    reason=(request.form.get("reason") or "Godkänd omplanering").strip()[:500]
    if not selected:
        flash("Välj minst en aktivitet att omplanera.","warning")
        return redirect(url_for("reschedule_preview_v1520",project_id=project_id))
    with db() as conn:
        ensure_v152_schema(conn)
        tasks=conn.execute("""SELECT id,title,owner,status,start_date,end_date,duration_days,progress
                              FROM tasks WHERE project_id=? AND COALESCE(deleted_at,'')=''""",(project_id,)).fetchall()
        links=conn.execute("""SELECT predecessor_id,successor_id,link_type,lag_days FROM task_links WHERE project_id=?""",(project_id,)).fetchall()
        plan=_v151_schedule(tasks,links)
        cur=conn.execute("""INSERT INTO schedule_change_batches(project_id,title,reason,created_by)
                            VALUES(?,?,?,?)""",(project_id,"Kontrollerad omplanering",reason,session.get("user_id")))
        batch_id=cur.lastrowid
        written=0
        for t in tasks:
            tid=int(t["id"])
            if tid not in selected or tid not in plan["tasks"]: continue
            p=plan["tasks"][tid]
            ns=p["forecast_start"].isoformat() if p["forecast_start"] else (t["start_date"] or "")
            ne=p["forecast_end"].isoformat() if p["forecast_end"] else (t["end_date"] or "")
            if ns==(t["start_date"] or "") and ne==(t["end_date"] or ""): continue
            conn.execute("""INSERT INTO schedule_change_items(batch_id,task_id,old_start,old_end,new_start,new_end)
                            VALUES(?,?,?,?,?,?)""",(batch_id,tid,t["start_date"] or "",t["end_date"] or "",ns,ne))
            conn.execute("UPDATE tasks SET start_date=?,end_date=? WHERE id=? AND project_id=?",(ns,ne,tid,project_id))
            written+=1
        if written==0:
            conn.execute("DELETE FROM schedule_change_batches WHERE id=?",(batch_id,))
        conn.commit()
    if written:
        audit(project_id,"schedule_batch",batch_id,"apply",f"{written} aktiviteter · {reason}")
        flash(f"Omplaneringen godkändes och {written} aktiviteter uppdaterades.","success")
    else:
        flash("Inga valda aktiviteter behövde ändras.","info")
    return redirect(url_for("reschedule_preview_v1520",project_id=project_id))

@app.post("/projects/<int:project_id>/reschedule/<int:batch_id>/undo")
@login_required
def reschedule_undo_v1520(project_id,batch_id):
    project=project_or_404(project_id,write=True)
    with db() as conn:
        ensure_v152_schema(conn)
        batch=conn.execute("SELECT * FROM schedule_change_batches WHERE id=? AND project_id=?",(batch_id,project_id)).fetchone()
        if not batch: abort(404)
        if batch["reverted_at"]:
            flash("Den här omplaneringen är redan återställd.","warning")
            return redirect(url_for("reschedule_preview_v1520",project_id=project_id))
        items=conn.execute("SELECT * FROM schedule_change_items WHERE batch_id=? ORDER BY id DESC",(batch_id,)).fetchall()
        for i in items:
            conn.execute("UPDATE tasks SET start_date=?,end_date=? WHERE id=? AND project_id=?",
                         (i["old_start"] or "",i["old_end"] or "",i["task_id"],project_id))
        conn.execute("""UPDATE schedule_change_batches SET reverted_at=CURRENT_TIMESTAMP,reverted_by=? WHERE id=?""",
                     (session.get("user_id"),batch_id))
        conn.commit()
    audit(project_id,"schedule_batch",batch_id,"undo",f"{len(items)} aktiviteter återställda")
    flash(f"Omplaneringen återställdes ({len(items)} aktiviteter).","success")
    return redirect(url_for("reschedule_preview_v1520",project_id=project_id))

@app.get("/capacity/level")
@login_required
def capacity_level_v1520():
    with db() as conn:
        rows=conn.execute("""SELECT ra.id,ra.user_id,
                    COALESCE(NULLIF(ra.resource_name,''),u.display_name,u.username,'Resurs') resource,
                    ra.week_start,ra.allocation_pct,ra.planned_hours,ra.project_id,p.name project_name
             FROM resource_allocations ra
             LEFT JOIN users u ON u.id=ra.user_id
             LEFT JOIN projects p ON p.id=ra.project_id
             WHERE ra.week_start>=date('now','-7 day')
             ORDER BY resource,week_start,ra.allocation_pct DESC""").fetchall()
    by={}
    for raw in rows:
        r=dict(raw); by.setdefault(r["resource"],[]).append(r)
    suggestions=[]
    for resource,items in by.items():
        totals={}
        for r in items: totals[r["week_start"]]=totals.get(r["week_start"],0)+float(r["allocation_pct"] or 0)
        under=[(w,v) for w,v in totals.items() if v<85]
        for r in items:
            total=totals.get(r["week_start"],0)
            if total<=100: continue
            needed=min(float(r["allocation_pct"] or 0),total-100)
            targets=sorted([(w,85-v) for w,v in under if w>r["week_start"] and 85-v>0],key=lambda x:x[0])
            if targets:
                w,room=targets[0]; move=min(needed,room)
                if move>0:
                    suggestions.append({"allocation_id":r["id"],"resource":resource,"project":r["project_name"],
                                        "from_week":r["week_start"],"to_week":w,"move_pct":round(move,1),
                                        "current_pct":float(r["allocation_pct"] or 0)})
    return render_template("resource_level_v1520.html",suggestions=suggestions)

@app.post("/capacity/level/apply")
@login_required
def capacity_level_apply_v1520():
    if current_user()["role"] not in ("admin","pm"): abort(403)
    allocation_id=request.form.get("allocation_id",type=int)
    to_week=(request.form.get("to_week") or "").strip()
    move_pct=request.form.get("move_pct",type=float)
    if not allocation_id or not to_week or not move_pct or move_pct<=0:
        flash("Ogiltigt utjämningsförslag.","error"); return redirect(url_for("capacity_level_v1520"))
    with db() as conn:
        row=conn.execute("SELECT * FROM resource_allocations WHERE id=?",(allocation_id,)).fetchone()
        if not row: abort(404)
        move=min(float(row["allocation_pct"] or 0),float(move_pct))
        old_hours=float(row["planned_hours"] or 0)
        old_pct=float(row["allocation_pct"] or 0)
        moved_hours=(old_hours*(move/old_pct)) if old_pct>0 else 0
        remain_pct=max(0,old_pct-move)
        remain_hours=max(0,old_hours-moved_hours)
        conn.execute("UPDATE resource_allocations SET allocation_pct=?,planned_hours=? WHERE id=?",
                     (remain_pct,remain_hours,allocation_id))
        existing=conn.execute("""SELECT id,allocation_pct,planned_hours FROM resource_allocations
                                 WHERE project_id=? AND COALESCE(user_id,0)=COALESCE(?,0)
                                   AND COALESCE(resource_name,'')=COALESCE(?, '') AND week_start=? LIMIT 1""",
                              (row["project_id"],row["user_id"],row["resource_name"],to_week)).fetchone()
        if existing:
            conn.execute("UPDATE resource_allocations SET allocation_pct=?,planned_hours=? WHERE id=?",
                         (float(existing["allocation_pct"] or 0)+move,float(existing["planned_hours"] or 0)+moved_hours,existing["id"]))
        else:
            conn.execute("""INSERT INTO resource_allocations(project_id,user_id,resource_name,week_start,allocation_pct,planned_hours)
                            VALUES(?,?,?,?,?,?)""",(row["project_id"],row["user_id"],row["resource_name"],to_week,move,moved_hours))
        conn.commit()
    audit(row["project_id"],"resource_allocation",allocation_id,"level",
          f"{move:.1f}% flyttat från {row['week_start']} till {to_week}")
    flash(f"{move:.0f}% flyttades till veckan {to_week}.","success")
    return redirect(url_for("capacity_level_v1520"))

@app.route("/integrations/azure-devops/bidirectional",methods=["GET","POST"])
@login_required
def devops_bidirectional_v1520():
    if current_user()["role"] not in ("admin","pm"): abort(403)
    with db() as conn:
        ensure_azure_devops_schema_v1010(conn)
        if request.method=="POST" and request.form.get("action")=="link":
            link_id=request.form.get("link_id",type=int)
            task_id=request.form.get("task_id",type=int)
            link=conn.execute("SELECT * FROM azure_devops_work_item_links WHERE id=?",(link_id,)).fetchone()
            task=conn.execute("SELECT * FROM tasks WHERE id=?",(task_id,)).fetchone()
            if not link or not task: abort(404)
            project_access(task["project_id"],write=True)
            conn.execute("UPDATE azure_devops_work_item_links SET task_id=?,project_id=? WHERE id=?",
                         (task_id,task["project_id"],link_id))
            conn.commit()
            flash(f"Work Item #{link['work_item_id']} länkades till {task['title']}.","success")
        connections=conn.execute("""SELECT id,name,organization_url,project_name,enabled FROM azure_devops_connections
                                    WHERE enabled=1 ORDER BY name""").fetchall()
        links=conn.execute("""SELECT l.*,c.name connection_name,t.title task_title,t.status task_status,p.name project_name
                              FROM azure_devops_work_item_links l
                              JOIN azure_devops_connections c ON c.id=l.connection_id
                              LEFT JOIN tasks t ON t.id=l.task_id
                              LEFT JOIN projects p ON p.id=l.project_id
                              ORDER BY l.last_synced_at DESC,l.id DESC LIMIT 300""").fetchall()
        tasks=conn.execute("""SELECT t.id,t.title,t.project_id,p.name project_name
                              FROM tasks t JOIN projects p ON p.id=t.project_id
                              WHERE COALESCE(t.deleted_at,'')='' AND COALESCE(p.deleted_at,'')=''
                              ORDER BY p.name,t.title LIMIT 1000""").fetchall()
    return render_template("devops_bidir_v1520.html",connections=connections,links=links,tasks=tasks)

@app.post("/integrations/azure-devops/bidirectional/<int:link_id>/push")
@login_required
def devops_push_v1520(link_id):
    if current_user()["role"] not in ("admin","pm"): abort(403)
    with db() as conn:
        ensure_azure_devops_schema_v1010(conn)
        link=conn.execute("""SELECT l.*,c.organization_url,c.project_name,c.pat_secret,t.title task_title,t.status task_status,t.project_id task_project
                             FROM azure_devops_work_item_links l
                             JOIN azure_devops_connections c ON c.id=l.connection_id
                             JOIN tasks t ON t.id=l.task_id WHERE l.id=?""",(link_id,)).fetchone()
        if not link: abort(404)
        project_access(link["task_project"],write=True)
        if not link["pat_secret"]:
            flash("Azure DevOps-anslutningen saknar PAT.","error"); return redirect(url_for("devops_bidirectional_v1520"))
        try:
            from urllib.parse import quote
            base=link["organization_url"].rstrip("/")+"/"+quote(link["project_name"])
            url=f"{base}/_apis/wit/workitems/{link['work_item_id']}?api-version=7.1"
            operations=[
                {"op":"add","path":"/fields/System.Title","value":link["task_title"]},
                {"op":"add","path":"/fields/System.State","value":_v152_devops_state(link["task_status"])}
            ]
            result=_ado_patch_v1520(url,link["pat_secret"],operations)
            fields=result.get("fields",{})
            conn.execute("""UPDATE azure_devops_work_item_links SET title=?,state=?,last_synced_at=CURRENT_TIMESTAMP WHERE id=?""",
                         (fields.get("System.Title",link["task_title"]),fields.get("System.State"),link_id))
            conn.execute("""INSERT INTO azure_devops_sync_log(connection_id,project_id,direction,status,message,items_read,items_written)
                            VALUES(?,?,?,?,?,?,?)""",
                         (link["connection_id"],link["task_project"],"ProjectPlanerToDevOps","Success",
                          f"Pushed Work Item #{link['work_item_id']}",0,1))
            conn.commit()
            flash(f"Work Item #{link['work_item_id']} uppdaterades i Azure DevOps.","success")
        except Exception as ex:
            conn.execute("""INSERT INTO azure_devops_sync_log(connection_id,project_id,direction,status,message)
                            VALUES(?,?,?,?,?)""",(link["connection_id"],link["task_project"],"ProjectPlanerToDevOps","Failed",str(ex)[:500]))
            conn.commit()
            flash("Push till Azure DevOps misslyckades: "+str(ex),"error")
    return redirect(url_for("devops_bidirectional_v1520"))

@app.post("/integrations/azure-devops/bidirectional/<int:link_id>/pull")
@login_required
def devops_pull_link_v1520(link_id):
    if current_user()["role"] not in ("admin","pm"): abort(403)
    with db() as conn:
        ensure_azure_devops_schema_v1010(conn)
        link=conn.execute("""SELECT l.*,c.organization_url,c.project_name,c.pat_secret,t.project_id task_project
                             FROM azure_devops_work_item_links l
                             JOIN azure_devops_connections c ON c.id=l.connection_id
                             JOIN tasks t ON t.id=l.task_id WHERE l.id=?""",(link_id,)).fetchone()
        if not link: abort(404)
        project_access(link["task_project"],write=True)
        try:
            from urllib.parse import quote
            base=link["organization_url"].rstrip("/")+"/"+quote(link["project_name"])
            wi=_ado_json_v1140("GET",f"{base}/_apis/wit/workitems/{link['work_item_id']}?api-version=7.1",link["pat_secret"])
            f=wi.get("fields",{})
            title=f.get("System.Title") or link["title"]
            state=f.get("System.State") or link["state"]
            mapped_status="Done" if str(state).lower() in ("done","closed","resolved","completed") else "In progress"
            conn.execute("UPDATE tasks SET title=?,status=? WHERE id=?",(title,mapped_status,link["task_id"]))
            conn.execute("""UPDATE azure_devops_work_item_links SET title=?,state=?,last_synced_at=CURRENT_TIMESTAMP WHERE id=?""",
                         (title,state,link_id))
            conn.execute("""INSERT INTO azure_devops_sync_log(connection_id,project_id,direction,status,message,items_read,items_written)
                            VALUES(?,?,?,?,?,?,?)""",
                         (link["connection_id"],link["task_project"],"DevOpsToProjectPlaner","Success",
                          f"Pulled Work Item #{link['work_item_id']}",1,1))
            conn.commit()
            audit(link["task_project"],"task",link["task_id"],"devops_pull",f"Work Item #{link['work_item_id']}")
            flash(f"Work Item #{link['work_item_id']} hämtades till Project Planer.","success")
        except Exception as ex:
            flash("Pull från Azure DevOps misslyckades: "+str(ex),"error")
    return redirect(url_for("devops_bidirectional_v1520"))

@app.get("/projects/<int:project_id>/excel")
@login_required
def project_excel_hub_v1524(project_id):
    project_or_404(project_id)
    return redirect(url_for("excel_center_v810",project_id=project_id))

def excel_xlsx_integrity_check_v1527(payload):
    """Validate the XLSX container and the XML parts Excel is strict about."""
    import zipfile as _zipfile
    import xml.etree.ElementTree as _ET
    from io import BytesIO as _BytesIO

    if not payload:
        raise ValueError("Tom XLSX-fil.")
    try:
        with _zipfile.ZipFile(_BytesIO(payload), "r") as zf:
            bad=zf.testzip()
            if bad:
                raise ValueError(f"Skadad ZIP-post i XLSX: {bad}")
            names=set(zf.namelist())
            required={"[Content_Types].xml","xl/workbook.xml","xl/styles.xml"}
            missing=required-names
            if missing:
                raise ValueError("XLSX saknar obligatoriska delar: "+", ".join(sorted(missing)))

            # Parse every XML/rels part. This catches malformed package content.
            for name in names:
                if name.endswith(".xml") or name.endswith(".rels"):
                    try:
                        _ET.fromstring(zf.read(name))
                    except Exception as ex:
                        raise ValueError(f"Ogiltig XML i {name}: {ex}")

            # Validate table references. A table must include header + at least one data row.
            ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            for name in names:
                if name.startswith("xl/tables/table") and name.endswith(".xml"):
                    root=_ET.fromstring(zf.read(name))
                    ref=root.attrib.get("ref","")
                    m=re.fullmatch(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)",ref)
                    if not m:
                        raise ValueError(f"Ogiltigt tabellområde i {name}: {ref}")
                    r1=int(m.group(2)); r2=int(m.group(4))
                    if r2 <= r1:
                        raise ValueError(f"Tabell utan datarad i {name}: {ref}")
    except _zipfile.BadZipFile as ex:
        raise ValueError(f"Ogiltig XLSX-container: {ex}")
    return True

def excel_serialize_workbook_v1527(wb):
    """Normalize through a save/reload/save cycle and verify the final XLSX package."""
    first=BytesIO()
    wb.save(first)
    payload=first.getvalue()
    excel_xlsx_integrity_check_v1527(payload)

    # Re-open and save once more. openpyxl normalizes relationships, names,
    # data validations and drawing/table package parts during this pass.
    check=load_workbook(BytesIO(payload),data_only=False)
    second=BytesIO()
    check.save(second)
    final_payload=second.getvalue()
    excel_xlsx_integrity_check_v1527(final_payload)
    return final_payload

def excel_blank_complete_workbook_v1525():
    """v15.2.8: build the blank template from a fresh Workbook.

    This intentionally does NOT reuse build_roundtrip_workbook() or
    build_complete_excel_template_v1522().  Microsoft Excel is stricter than
    openpyxl about some combinations of empty tables, filters, names and
    relationships.  The global blank template therefore uses only standard
    cells/styles and no tables/charts/defined names.
    """
    wb=Workbook()
    default=wb.active
    wb.remove(default)

    navy="0F4C81"; blue="DCEAF7"; green="E9F6EE"; grey="F4F6F8"; border="D0D7DE"; text="172B4D"; muted="667085"
    thin=Side(style="thin",color=border)

    def title(ws,title,subtitle=None):
        ws.sheet_view.showGridLines=False
        ws["A1"]="Project Planer"
        ws["A1"].font=Font(size=11,bold=True,color="FFFFFF")
        ws["A1"].fill=PatternFill("solid",fgColor=navy)
        ws["B1"]=f"Komplett projektmall · v{APP_VERSION}"
        ws["B1"].font=Font(size=11,color="FFFFFF")
        ws["B1"].fill=PatternFill("solid",fgColor=navy)
        ws.merge_cells("B1:F1")
        ws["A3"]=title
        ws["A3"].font=Font(size=20,bold=True,color=text)
        ws.merge_cells("A3:F3")
        if subtitle:
            ws["A4"]=subtitle
            ws["A4"].font=Font(size=10,color=muted)
            ws["A4"].alignment=Alignment(wrap_text=True,vertical="top")
            ws.merge_cells("A4:F5")
        for col,w in {"A":24,"B":28,"C":20,"D":20,"E":20,"F":22}.items(): ws.column_dimensions[col].width=w

    def plain_sheet(name,headers,description="",editable=True,rows=120,date_headers=None,money_headers=None,percent_headers=None):
        ws=wb.create_sheet(name)
        ws.sheet_view.showGridLines=False
        ws.freeze_panes="A2"
        ws.append(headers)
        for c in ws[1]:
            c.fill=PatternFill("solid",fgColor=navy); c.font=Font(bold=True,color="FFFFFF")
            c.alignment=Alignment(vertical="center",wrap_text=True); c.border=Border(bottom=thin)
        ws.row_dimensions[1].height=28
        date_headers=set(date_headers or [])
        money_headers=set(money_headers or [])
        percent_headers=set(percent_headers or [])
        for idx,h in enumerate(headers,1):
            letter=get_column_letter(idx)
            width=18
            if h in {"Aktivitet","Titel","Rubrik","Beskrivning","Åtgärd","Kommentar","Anteckningar","notes","body","details","details_json","snapshot_json"}: width=38
            elif h in {"WBS","Typ","Status","Prioritet","Enhet"}: width=16
            elif "datum" in h.lower() or "date" in h.lower() or h in {"Plan start","Plan slut","Faktisk start","Faktiskt slut","Vecka","Datum","Mätdatum","Förfallodatum"}: width=16
            ws.column_dimensions[letter].width=width
            for r in range(2,rows+2):
                cell=ws.cell(r,idx)
                cell.fill=PatternFill("solid",fgColor=(green if editable else grey))
                cell.border=Border(bottom=Side(style="hair",color="E8ECF0"))
                cell.alignment=Alignment(vertical="top",wrap_text=h in {"Aktivitet","Titel","Rubrik","Beskrivning","Åtgärd","Kommentar","Anteckningar","notes","body","details","details_json","snapshot_json"})
                if h in date_headers: cell.number_format="yyyy-mm-dd"
                elif h in money_headers: cell.number_format='#,##0.00'
                elif h in percent_headers: cell.number_format='0%'
        # A plain AutoFilter on a populated entry range is valid and useful.
        # Do not create an Excel Table; this avoids Excel repair warnings on blank templates.
        ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{rows+1}"
        ws.freeze_panes="A2"
        if description:
            ws.sheet_properties.pageSetUpPr.fitToPage=True
        return ws

    # Instructions first.
    readme=wb.create_sheet("LÄS MIG")
    title(readme,"Komplett Project Planer-mall","Denna arbetsbok är skapad från grunden för att vara stabil i Microsoft Excel. Fyll först i Projektinformation och sedan de blad du behöver.")
    readme["A7"]="1"
    readme["B7"]="Fyll i Projektinformation"
    readme["A8"]="2"
    readme["B8"]="Lägg in WBS/aktiviteter, risker, resurser, kostnader och övrigt innehåll"
    readme["A9"]="3"
    readme["B9"]="I Project Planer: Excel → Importera Excel och skapa projekt"
    readme["A10"]="4"
    readme["B10"]="Projektet öppnas automatiskt i Project Cockpit"
    readme["A12"]="Viktigt"
    readme["B12"]="Bladen Uppgifter, Beroenden, Risker, Ändringsärenden, Resurser, Kostnader, Beslut, Möten, Åtgärder och Nyttor importeras när ett nytt projekt skapas. Övriga blad finns med som komplett planerings- och referensstruktur."
    readme["B12"].alignment=Alignment(wrap_text=True,vertical="top")
    readme.column_dimensions["A"].width=12; readme.column_dimensions["B"].width=95

    info=wb.create_sheet("Projektinformation")
    title(info,"Projektinformation","Projektnamn är obligatoriskt när Excel-filen ska skapa ett nytt projekt.")
    info_rows=[("Projektnamn",""),("Kund",""),("Projektledare",""),("Beskrivning",""),("Plan start",""),("Plan slut","")]
    # Rows 4-5 are merged by title(..., subtitle=...). Start editable project
    # information at row 7 so no write can target an openpyxl MergedCell.
    for r,(label,value) in enumerate(info_rows,7):
        info.cell(r,1,label).font=Font(bold=True,color=text)
        info.cell(r,2,value).fill=PatternFill("solid",fgColor=green)
        info.cell(r,2).border=Border(bottom=thin)
        if label in ("Plan start","Plan slut"): info.cell(r,2).number_format="yyyy-mm-dd"
    info.column_dimensions["A"].width=24; info.column_dimensions["B"].width=60

    # Import-compatible sheets: exact headers expected by EXCEL_SPECS / dependency importer.
    core=[
      ("Uppgifter",list(EXCEL_SPECS["Uppgifter"]["headers"].keys()),{"Plan start","Plan slut","Faktisk start","Faktiskt slut"},set(),set()),
      ("Beroenden",["Föregående WBS","Efterföljande WBS","Typ","Förskjutning dagar"],set(),set(),set()),
      ("Risker",list(EXCEL_SPECS["Risker"]["headers"].keys()),{"Förfallodatum"},set(),set()),
      ("Ändringsärenden",list(EXCEL_SPECS["Ändringsärenden"]["headers"].keys()),set(),{"Kostnad"},set()),
      ("Resurser",list(EXCEL_SPECS["Resurser"]["headers"].keys()),{"Vecka"},set(),set()),
      ("Kostnader",list(EXCEL_SPECS["Kostnader"]["headers"].keys()),{"Datum"},{"Planerat","Utfall"},set()),
      ("Beslut",list(EXCEL_SPECS["Beslut"]["headers"].keys()),{"Beslutsdatum"},set(),set()),
      ("Möten",list(EXCEL_SPECS["Möten"]["headers"].keys()),{"Datum"},set(),set()),
      ("Åtgärder",list(EXCEL_SPECS["Åtgärder"]["headers"].keys()),{"Förfallodatum"},set(),set()),
      ("Nyttor",list(EXCEL_SPECS["Nyttor"]["headers"].keys()),{"Mätdatum"},{"Baslinje","Mål","Utfall"},set()),
    ]
    for name,headers,dates,money,pct in core:
        plain_sheet(name,headers,editable=True,date_headers=dates,money_headers=money,percent_headers=pct)

    # Helpful data validation only on simple literal lists. No defined names or cross-sheet formulas.
    validations={
      "Uppgifter":{"Status":["Ej påbörjad","Pågår","Blockerad","Klar"],"Prioritet":["Låg","Medium","Hög","Kritisk"],"Milstolpe":["Nej","Ja"]},
      "Beroenden":{"Typ":["FS","SS","FF","SF"]},
      "Risker":{"Typ":["Risk","Issue"],"Status":["Open","Closed"]},
      "Ändringsärenden":{"Status":["Open","Approved","Rejected","Closed"]},
      "Åtgärder":{"Status":["Open","Pågår","Klar","Closed"]},
      "Nyttor":{"Status":["Open","Pågår","Realiserad","Closed"]},
    }
    for sheet_name,cols in validations.items():
        ws=wb[sheet_name]
        header_map={excel_text(c.value):c.column for c in ws[1]}
        for header,values in cols.items():
            if header not in header_map: continue
            col=get_column_letter(header_map[header])
            formula='"'+','.join(values)+'"'
            dv=DataValidation(type="list",formula1=formula,allow_blank=True)
            dv.error="Välj ett värde i listan."; dv.errorTitle="Ogiltigt värde"
            ws.add_data_validation(dv); dv.add(f"{col}2:{col}121")

    # A read-only milestone planning sheet; milestones themselves are entered in Uppgifter.
    plain_sheet("Milstolpar",["WBS","Milstolpe","Ansvarig","Plan slut","Status","Progress %"],editable=False,date_headers={"Plan slut"})

    # Complete reference structure. These sheets are intentionally plain so Excel has nothing to repair.
    reference_specs=[
      ("Medlemmar",["display_name","username","project_role","added_at","added_by"]),
      ("Statusrapporter",["report_date","overall_rag","scope_rag","schedule_rag","budget_rag","resources_rag","summary","achievements","next_steps","created_by","created_at"]),
      ("RAID",["item_type","title","owner","status","due_date","details"]),
      ("Godkännanden",["entity_type","entity_id","status","requested_by","requested_at","decided_by","decided_at","comment"]),
      ("Tid",["work_date","user_name","task_title","hours","billable","note"]),
      ("Dokument",["title","body","updated_by","updated_at"]),
      ("Bilagor",["filename","stored_name","content_type","size_bytes","uploaded_by","uploaded_at"]),
      ("Kommentarer",["källa","user_name","body","created_at"]),
      ("Baselines",["name","created_at","created_by","snapshot_json"]),
      ("Anpassade fält",["name","field_type","entity_type","entity_id","value_text","options_json"]),
      ("Miljöer",["name","environment_type","url","owner","status","notes"]),
      ("Leverabler",["title","owner","due_date","status","description"]),
      ("Interfaces",["name","source_system","target_system","owner","status","description"]),
      ("Testcykler",["name","start_date","end_date","owner","status","notes"]),
      ("Spårbarhet",["requirement_id","requirement_title","deliverable_id","test_reference","status","notes"]),
      ("Cutover",["title","owner","planned_at","status","sequence_no","notes"]),
      ("Projekthälsa historik",["snapshot_date","health_score","rag","details_json"]),
      ("Projektberoenden",["predecessor_project","successor_project","predecessor_task","successor_task","link_type","lag_days","status","notes"]),
      ("Azure DevOps",["connection_name","devops_project","work_item_id","work_item_type","title","state","assigned_to","task_title","last_synced_at"]),
      ("Omplaneringshistorik",["id","title","reason","created_name","created_at","reverted_at","reverted_by"]),
      ("Omplaneringsdetaljer",["batch_id","task_title","old_start","old_end","new_start","new_end"]),
      ("Finanssammanfattning",["planned_budget","approved_budget","forecast","actual"]),
    ]
    for name,headers in reference_specs:
        plain_sheet(name,headers,editable=False,rows=20)

    # Data dictionary.
    dd=wb.create_sheet("Datamodell",2)
    dd.sheet_view.showGridLines=False
    dd.append(["Blad","Syfte","Import till nytt projekt","Kommentar"])
    for c in dd[1]: c.fill=PatternFill("solid",fgColor=navy); c.font=Font(bold=True,color="FFFFFF")
    purposes={
      "Projektinformation":"Projektets grunddata","Uppgifter":"WBS och aktiviteter","Beroenden":"Aktivitetsberoenden","Risker":"Risker och issues","Ändringsärenden":"Change control","Resurser":"Resursallokering","Kostnader":"Projektkostnader","Beslut":"Beslut","Möten":"Möten","Åtgärder":"Actions","Nyttor":"Benefits"
    }
    importable=set(purposes)
    for ws in wb.worksheets:
        if ws.title in {"LÄS MIG","Datamodell"}: continue
        dd.append([ws.title,purposes.get(ws.title,"Komplett projekt-/referensinnehåll"),"Ja" if ws.title in importable else "Nej","Fyll i efter behov."])
    dd.freeze_panes="A2"; dd.auto_filter.ref=f"A1:D{dd.max_row}"
    dd.column_dimensions["A"].width=28; dd.column_dimensions["B"].width=38; dd.column_dimensions["C"].width=24; dd.column_dimensions["D"].width=38

    # Metadata required by the importer. Keep simple key/value cells and use hidden (not veryHidden).
    meta=wb.create_sheet("_Metadata")
    meta.append(["key","value"])
    for k,v in [
      ("schema_version",EXCEL_SCHEMA_VERSION),("project_id","0"),("project_name","NEW_PROJECT_TEMPLATE"),
      ("project_hash",""),("app_version",APP_VERSION),("exported_at",datetime.now().isoformat(timespec="seconds")),
      ("template_kind","new_project_complete")
    ]: meta.append([k,v])
    meta.sheet_state="hidden"

    wb.active=0
    wb.properties.creator="Project Planer"
    wb.properties.title="Project Planer – Komplett projektmall"
    wb.properties.subject="Ny projektmall"
    wb.properties.description=f"Project Planer v{APP_VERSION} · ren Microsoft Excel-kompatibel mall byggd från grunden."
    # v15.2.9 defensive layout validation: Projectinformation inputs must
    # never overlap merged title/subtitle cells.
    if "Projektinformation" in wb.sheetnames:
        _info=wb["Projektinformation"]
        for _row in range(7,13):
            _cell=_info.cell(_row,2)
            if _cell.__class__.__name__ == "MergedCell":
                raise RuntimeError(f"Excel template layout error: B{_row} is merged/read-only")
    return wb

def excel_new_project_from_workbook_v1525(file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError("Välj en Excel-fil.")
    if not file_storage.filename.lower().endswith(".xlsx"):
        raise ValueError("Endast .xlsx stöds.")
    payload=file_storage.read()
    if len(payload)>15*1024*1024:
        raise ValueError("Excel-filen är större än 15 MB.")
    try:
        wb=load_workbook(BytesIO(payload),data_only=False)
    except Exception as ex:
        raise ValueError(f"Kunde inte läsa Excel-filen: {ex}")

    meta=excel_parse_metadata(wb)
    if meta.get("schema_version")!=EXCEL_SCHEMA_VERSION:
        raise ValueError("Filen är inte en kompatibel Project Planer Excel-fil.")
    if excel_int(meta.get("project_id")) not in (0,):
        raise ValueError("Filen tillhör redan ett befintligt projekt. Använd projektets vanliga Excel-import.")

    raw={"name":"","customer":"","project_manager":"","description":"","start_date":"","end_date":""}
    if "Projektinformation" in wb.sheetnames:
        ws=wb["Projektinformation"]
        mapping={"Projektnamn":"name","Kund":"customer","Projektledare":"project_manager","Beskrivning":"description","Plan start":"start_date","Plan slut":"end_date"}
        for row in ws.iter_rows(min_row=4,max_col=2):
            label=excel_text(row[0].value)
            if label in mapping:
                field=mapping[label]
                raw[field]=excel_date(row[1].value) if field in ("start_date","end_date") else excel_text(row[1].value)
    if not raw["name"]:
        raise ValueError("Fyll i Projektnamn på bladet Projektinformation innan import.")

    start_d=_v15_date(raw["start_date"]) if raw["start_date"] else None
    end_d=_v15_date(raw["end_date"]) if raw["end_date"] else None
    if raw["start_date"] and not start_d: raise ValueError("Ogiltigt startdatum i Projektinformation.")
    if raw["end_date"] and not end_d: raise ValueError("Ogiltigt slutdatum i Projektinformation.")
    if start_d and end_d and end_d<start_d: raise ValueError("Slutdatum kan inte ligga före startdatum.")

    with db() as conn:
        conn.execute("BEGIN")
        try:
            cur=conn.execute("""INSERT INTO projects(name,customer,project_manager,description,start_date,end_date,created_by,created_at)
                                VALUES(?,?,?,?,?,?,?,CURRENT_TIMESTAMP)""",
                             (raw["name"],raw["customer"],raw["project_manager"],raw["description"],
                              raw["start_date"],raw["end_date"],session.get("user_id")))
            pid=cur.lastrowid
            conn.execute("""INSERT OR IGNORE INTO project_members(project_id,user_id,project_role,added_at,added_by)
                            VALUES(?,?,?,CURRENT_TIMESTAMP,?)""",
                         (pid,session.get("user_id"),"pm",session.get("user_id")))

            created=0
            # Import normal editable sheets first. IDs/hashes are intentionally ignored for a new project.
            for sheet_name,spec in EXCEL_SPECS.items():
                if sheet_name not in wb.sheetnames: continue
                ws=wb[sheet_name]; hdr=excel_headers(ws)
                missing=[h for h in spec["headers"] if h not in hdr]
                if missing: continue
                for row_no in range(2,ws.max_row+1):
                    rawrow={field:ws.cell(row_no,hdr[label]).value for label,field in spec["headers"].items()}
                    data=excel_normalize_row(spec,rawrow)
                    required=excel_text(data.get(spec["required"]))
                    # Ignore the example rows if user left them untouched.
                    if required.startswith("Exempel") or required=="Kickoff":
                        continue
                    if not required and all(v in ("",None,0,0.0) for v in data.values()):
                        continue
                    if not required:
                        raise ValueError(f"{sheet_name}, rad {row_no}: obligatoriskt namn/rubrik saknas.")
                    excel_insert(conn,spec["table"],pid,data)
                    created+=1

            # Dependencies are imported after tasks so WBS can be resolved.
            if "Beroenden" in wb.sheetnames:
                ws=wb["Beroenden"]; hdr=excel_headers(ws)
                needed=["Föregående WBS","Efterföljande WBS","Typ","Förskjutning dagar"]
                if all(h in hdr for h in needed):
                    tasks=conn.execute("SELECT id,wbs FROM tasks WHERE project_id=? AND deleted_at IS NULL",(pid,)).fetchall()
                    id_by_wbs={excel_text(t["wbs"]):int(t["id"]) for t in tasks if excel_text(t["wbs"])}
                    for row_no in range(2,ws.max_row+1):
                        pred=excel_text(ws.cell(row_no,hdr["Föregående WBS"]).value)
                        succ=excel_text(ws.cell(row_no,hdr["Efterföljande WBS"]).value)
                        if not pred and not succ: continue
                        if pred not in id_by_wbs or succ not in id_by_wbs:
                            raise ValueError(f"Beroenden, rad {row_no}: WBS {pred} eller {succ} finns inte i Uppgifter.")
                        typ=excel_text(ws.cell(row_no,hdr["Typ"]).value) or "FS"
                        lag=excel_int(ws.cell(row_no,hdr["Förskjutning dagar"]).value)
                        conn.execute("""INSERT INTO task_links(project_id,predecessor_id,successor_id,link_type,lag_days)
                                      VALUES(?,?,?,?,?)""",(pid,id_by_wbs[pred],id_by_wbs[succ],typ,lag))
                        created+=1

            conn.commit()
        except Exception:
            conn.rollback()
            raise
    audit(pid,"project",pid,"excel_create_project",f"created_items={created}")
    return pid,created

@app.get("/excel")
@login_required
def excel_start_center_v1525():
    return render_template("excel_start_v1525.html")

@app.get("/excel/template")
@login_required
def excel_blank_template_v1525():
    wb=excel_blank_complete_workbook_v1525()
    payload=excel_serialize_workbook_v1527(wb)
    return send_file(BytesIO(payload),as_attachment=True,
        download_name="Project-Planer-KOMPLETT-projektmall.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/excel/create-project",methods=["GET","POST"])
@login_required
def excel_create_project_v1525():
    user=current_user()
    if user["role"] not in ("admin","pm"): abort(403)
    if request.method=="POST":
        try:
            pid,count=excel_new_project_from_workbook_v1525(request.files.get("file"))
            flash(f"Projektet skapades från Excel. {count} projektposter importerades.","success")
            return redirect(url_for("ultimate_project_v140",project_id=pid))
        except ValueError as ex:
            flash(str(ex),"error")
        except Exception as ex:
            flash("Kunde inte skapa projekt från Excel: "+str(ex),"error")
    return render_template("excel_create_project_v1525.html")

@app.get("/health/ui")
@login_required
def health_ui_v1526():
    base_path = os.path.join(app.template_folder or "templates", "base.html")
    try:
        text = open(base_path, "r", encoding="utf-8").read()
        excel_nav = 'data-nav-excel="v15.2.6"' in text
    except Exception:
        excel_nav = False
    return jsonify({
        "status": "ok" if excel_nav else "degraded",
        "version": APP_VERSION,
        "excel_nav": excel_nav,
        "excel_start_center": True
    }), (200 if excel_nav else 503)

@app.get("/ultimate/compare")
@login_required
def ultimate_capabilities_v140():
    return render_template("ultimate_capabilities_v140.html")

_UI_SV_V1401 = {
    "Not started": "Ej påbörjad", "Not Started": "Ej påbörjad",
    "In progress": "Pågår", "In Progress": "Pågår",
    "Done": "Klar", "Completed": "Klar", "Closed": "Stängd",
    "Blocked": "Blockerad", "Open": "Öppen",
    "High": "Hög", "Normal": "Normal", "Low": "Låg",
    "Success": "Lyckades", "Failed": "Misslyckades",
    "Risk": "Risk", "Issue": "Problem",
    "Green": "Grön", "Amber": "Gul", "Red": "Röd"
}

@app.template_filter("sv")
def ui_sv_v1401(value):
    if value is None:
        return ""
    return _UI_SV_V1401.get(str(value), str(value))

@app.route("/admin")
@login_required
@role_required("admin")
def admin_center():
    with db() as conn:
        counts={
            "users":conn.execute("SELECT COUNT(*) c FROM users").fetchone()["c"],
            "active_users":conn.execute("SELECT COUNT(*) c FROM users WHERE active=1").fetchone()["c"],
            "projects":conn.execute("SELECT COUNT(*) c FROM projects").fetchone()["c"],
            "memberships":conn.execute("SELECT COUNT(*) c FROM project_members").fetchone()["c"],
        }
        recent=conn.execute("""SELECT a.*,u.display_name,p.name project_name FROM audit_log a
                               LEFT JOIN users u ON u.id=a.user_id LEFT JOIN projects p ON p.id=a.project_id
                               ORDER BY a.id DESC LIMIT 50""").fetchall()
    return render_template("admin_center.html",counts=counts,recent=recent)

@app.route("/admin/users",methods=["GET","POST"])
@login_required
@role_required("admin")
def admin_users():
    if request.method=="POST":
        username=request.form["username"].strip(); password=request.form["password"]; role=request.form.get("role","member")
        if role not in ROLES:role="member"
        if len(password)<10:
            flash("Lösenord måste vara minst 10 tecken.","danger"); return redirect(url_for("admin_users"))
        try:
            with db() as conn:
                conn.execute("""INSERT INTO users(username,display_name,password_hash,role,active,force_password_change,created_at)
                                VALUES(?,?,?,?,1,1,?)""",
                             (username,request.form.get("display_name","").strip() or username,
                              generate_password_hash(password),role,datetime.now().isoformat(timespec="seconds")))
            flash("Användaren skapades och måste byta lösenord vid första login.","success")
        except sqlite3.IntegrityError:
            flash("Användarnamnet finns redan.","danger")
        return redirect(url_for("admin_users"))
    with db() as conn: users=conn.execute("SELECT * FROM users ORDER BY username").fetchall()
    return render_template("users.html",users=users,roles=ROLES)

@app.post("/admin/users/<int:user_id>/toggle")
@login_required
@role_required("admin")
def toggle_user(user_id):
    u=current_user()
    if u["id"]==user_id:
        flash("Du kan inte inaktivera ditt eget konto.","danger")
    else:
        with db() as conn:
            row=conn.execute("SELECT active FROM users WHERE id=?",(user_id,)).fetchone()
            if row:conn.execute("UPDATE users SET active=? WHERE id=?",(0 if row["active"] else 1,user_id))
    return redirect(url_for("admin_users"))

@app.post("/admin/users/<int:user_id>/reset-password")
@login_required
@role_required("admin")
def admin_reset_password(user_id):
    new=request.form["password"]
    if len(new)<10:
        flash("Lösenord måste vara minst 10 tecken.","danger")
    else:
        with db() as conn:
            conn.execute("UPDATE users SET password_hash=?,force_password_change=1,failed_logins=0,locked_until='' WHERE id=?",
                         (generate_password_hash(new),user_id))
        flash("Lösenordet är återställt. Användaren måste byta det vid nästa login.","success")
    return redirect(url_for("admin_users"))

@app.post("/admin/users/<int:user_id>/role")
@login_required
@role_required("admin")
def admin_change_role(user_id):
    role=request.form.get("role","member")
    if role not in ROLES:abort(400)
    with db() as conn:conn.execute("UPDATE users SET role=? WHERE id=?",(role,user_id))
    return redirect(url_for("admin_users"))

@app.route("/admin/projects")
@login_required
@role_required("admin")
def admin_projects():
    with db() as conn:
        projects=conn.execute("""SELECT p.*,COUNT(pm.user_id) member_count FROM projects p
                                 LEFT JOIN project_members pm ON pm.project_id=p.id GROUP BY p.id ORDER BY p.id DESC""").fetchall()
    return render_template("admin_projects.html",projects=projects)

@app.post("/admin/projects/<int:project_id>/delete")
@login_required
@role_required("admin")
def admin_delete_project(project_id):
    with db() as conn:conn.execute("DELETE FROM projects WHERE id=?",(project_id,))
    flash("Projektet är borttaget.","success")
    return redirect(url_for("admin_projects"))

@app.route("/admin/settings",methods=["GET","POST"])
@login_required
@role_required("admin")
def admin_settings():
    if request.method=="POST":
        updates={
            "lockout_attempts":str(max(1,min(20,int(request.form.get("lockout_attempts","5"))))),
            "lockout_minutes":str(max(1,min(1440,int(request.form.get("lockout_minutes","15"))))),
        }
        with db() as conn:
            for k,v in updates.items():
                conn.execute("INSERT OR REPLACE INTO system_settings(key,value) VALUES(?,?)",(k,v))
        flash("Systeminställningar sparade.","success")
        return redirect(url_for("admin_settings"))
    return render_template("admin_settings.html",
                           lockout_attempts=setting("lockout_attempts","5"),
                           lockout_minutes=setting("lockout_minutes","15"))

@app.errorhandler(400)
def bad_request(e):return render_template("error.html",code=400,message=getattr(e,"description","Ogiltig begäran.")),400
@app.errorhandler(403)
def forbidden(e):return render_template("error.html",code=403,message="Du saknar behörighet för den här åtgärden."),403
@app.errorhandler(404)
def not_found(e):return render_template("error.html",code=404,message="Sidan kunde inte hittas."),404

@app.errorhandler(500)
def internal_error(e):
    request_id=secrets.token_hex(4)
    app.logger.exception("Unhandled error request_id=%s path=%s",request_id,request.path)
    return render_template("error.html",code=500,message=f"Ett internt fel inträffade. Referens: {request_id}"),500



def visible_projects_for_user():
    u=current_user()
    with db() as conn:
        if u["role"]=="admin":
            return conn.execute("""SELECT * FROM projects
                WHERE COALESCE(archived_at,'')='' AND COALESCE(deleted_at,'')=''
                ORDER BY name""").fetchall()
        return conn.execute("""SELECT p.* FROM projects p JOIN project_members pm ON pm.project_id=p.id
                              WHERE pm.user_id=? AND COALESCE(p.archived_at,'')='' AND COALESCE(p.deleted_at,'')=''
                              ORDER BY p.name""",(u["id"],)).fetchall()

@app.get("/my-work")
@login_required
def my_work():
    u=current_user()
    with db() as conn:
        if u["role"]=="admin":
            rows=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
                                 WHERE t.owner=? ORDER BY t.end_date,t.priority""",(u["display_name"],)).fetchall()
        else:
            rows=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
                                 JOIN project_members pm ON pm.project_id=p.id AND pm.user_id=?
                                 WHERE t.owner IN (?,?) ORDER BY t.end_date,t.priority""",
                              (u["id"],u["display_name"],u["username"])).fetchall()
    items=[dict(r,derived_status=derived_status(r),health=task_health(r)) for r in rows]
    return render_template("my_work.html",items=items)

@app.get("/projects/<int:project_id>/kanban")
@login_required
def kanban(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY sort_order,wbs,id",(project_id,)).fetchall()
    cols={s:[] for s in STATUSES}
    for t in tasks: cols.setdefault(t["status"],[]).append(t)
    return render_template("kanban.html",project=p,columns=cols,project_role=project_access(project_id))

@app.post("/projects/<int:project_id>/tasks/<int:task_id>/status")
@login_required
def set_task_status(project_id,task_id):
    project_or_404(project_id,write=True); task_or_404(project_id,task_id,write=True)
    status=request.form.get("status","Ej startad")
    if status not in STATUSES: abort(400)
    with db() as conn:
        conn.execute("UPDATE tasks SET status=?,progress=CASE WHEN ?='Klar' THEN 100 ELSE progress END WHERE id=?",
                     (status,status,task_id))
    audit(project_id,"task",task_id,"status",status)
    return redirect(request.referrer or url_for("kanban",project_id=project_id))

@app.get("/projects/<int:project_id>/calendar")
@login_required
def project_calendar(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("""SELECT * FROM tasks WHERE project_id=? AND (start_date<>'' OR end_date<>'')
                              ORDER BY COALESCE(NULLIF(start_date,''),end_date),wbs""",(project_id,)).fetchall()
    return render_template("calendar.html",project=p,tasks=tasks)

def compute_cpm(project_id):
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]
        links=[dict(r) for r in conn.execute("SELECT * FROM task_links WHERE project_id=? AND link_type='FS'",(project_id,))]
    if not tasks: return []
    ids={t["id"]:t for t in tasks}
    pred={i:[] for i in ids}; succ={i:[] for i in ids}
    for l in links:
        if l["predecessor_id"] in ids and l["successor_id"] in ids:
            pred[l["successor_id"]].append((l["predecessor_id"],int(l["lag_days"] or 0)))
            succ[l["predecessor_id"]].append((l["successor_id"],int(l["lag_days"] or 0)))
    indeg={i:len(pred[i]) for i in ids}; q=[i for i in ids if indeg[i]==0]; order=[]
    while q:
        i=q.pop(0); order.append(i)
        for j,_ in succ[i]:
            indeg[j]-=1
            if indeg[j]==0:q.append(j)
    if len(order)!=len(ids): return [{"error":"Cyclic dependency detected"}]
    dur={}
    for i,t in ids.items():
        s=parse_date(t["start_date"]); e=parse_date(t["end_date"])
        dur[i]=max(((e-s).days+1) if s and e else int(t.get("duration_days") or 1),1)
    es={i:0 for i in ids}; ef={}
    for i in order:
        es[i]=max([ef[p]+lag for p,lag in pred[i]] or [0]); ef[i]=es[i]+dur[i]
    finish=max(ef.values()); lf={i:finish for i in ids}; ls={}
    for i in reversed(order):
        if succ[i]: lf[i]=min(ls[j]-lag for j,lag in succ[i])
        ls[i]=lf[i]-dur[i]
    out=[]
    for i in order:
        t=ids[i]; slack=ls[i]-es[i]
        out.append(dict(t,early_start=es[i],early_finish=ef[i],late_start=ls[i],late_finish=lf[i],slack=slack,critical=(slack==0)))
    return out

@app.get("/projects/<int:project_id>/critical-path")
@login_required
def critical_path(project_id):
    p=project_or_404(project_id)
    return render_template("critical_path.html",project=p,rows=compute_cpm(project_id))

@app.post("/projects/<int:project_id>/links")
@login_required
def add_task_link(project_id):
    project_or_404(project_id,manager=True)
    pred=int(request.form["predecessor_id"]); succ=int(request.form["successor_id"])
    ltype=request.form.get("link_type","FS")
    if ltype not in ("FS","SS","FF","SF"): abort(400)
    if pred==succ: abort(400)
    with db() as conn:
        conn.execute("""INSERT INTO task_links(project_id,predecessor_id,successor_id,link_type,lag_days)
                        VALUES(?,?,?,?,?)""",(project_id,pred,succ,ltype,int(request.form.get("lag_days","0") or 0)))
    audit(project_id,"dependency",None,"created",f"{pred}->{succ} {ltype}")
    return redirect(url_for("critical_path",project_id=project_id))

@app.post("/projects/<int:project_id>/auto-schedule")
@login_required
def auto_schedule(project_id):
    project_or_404(project_id,manager=True)
    with db() as conn:
        tasks={r["id"]:dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))}
        links=[dict(r) for r in conn.execute("SELECT * FROM task_links WHERE project_id=? AND link_type='FS'",(project_id,))]
        changed=0
        for _ in range(max(len(tasks),1)):
            round_changed=0
            for l in links:
                a=tasks.get(l["predecessor_id"]); b=tasks.get(l["successor_id"])
                if not a or not b: continue
                ae=parse_date(a["end_date"]); bs=parse_date(b["start_date"]); be=parse_date(b["end_date"])
                if not ae or not bs: continue
                target=ae+timedelta(days=1+int(l["lag_days"] or 0))
                if bs<target:
                    dur=max(((be-bs).days if be else 0),0)
                    newe=target+timedelta(days=dur)
                    conn.execute("UPDATE tasks SET start_date=?,end_date=? WHERE id=?",(target.isoformat(),newe.isoformat(),b["id"]))
                    b["start_date"]=target.isoformat(); b["end_date"]=newe.isoformat()
                    round_changed+=1; changed+=1
            if not round_changed: break
    audit(project_id,"schedule",None,"auto-schedule",f"{changed} flyttningar")
    flash(f"Auto-scheduling klar: {changed} datumjusteringar.","success")
    return redirect(url_for("gantt",project_id=project_id))


@app.get("/resources")
@login_required
def resources():
    u=current_user()
    with db() as conn:
        if u["role"]=="admin":
            users=conn.execute("""SELECT u.*,COALESCE(r.weekly_capacity,40) weekly_capacity,COALESCE(r.hourly_cost,0) hourly_cost
                                  FROM users u LEFT JOIN resource_profiles r ON r.user_id=u.id WHERE u.active=1 ORDER BY u.display_name""").fetchall()
        else:
            users=conn.execute("""SELECT u.*,COALESCE(r.weekly_capacity,40) weekly_capacity,COALESCE(r.hourly_cost,0) hourly_cost
                                  FROM users u LEFT JOIN resource_profiles r ON r.user_id=u.id
                                  WHERE u.id IN (SELECT DISTINCT pm2.user_id FROM project_members pm1
                                    JOIN project_members pm2 ON pm2.project_id=pm1.project_id WHERE pm1.user_id=?)
                                  ORDER BY u.display_name""",(u["id"],)).fetchall()
        loads=[]
        for person in users:
            hrs=conn.execute("""SELECT COALESCE(SUM(t.planned_hours),0) h FROM tasks t
                                WHERE t.owner IN (?,?) AND t.status<>'Klar'""",(person["display_name"],person["username"])).fetchone()["h"]
            loads.append(dict(person,planned_open=float(hrs or 0),utilization=round(float(hrs or 0)/max(float(person["weekly_capacity"] or 40),1)*100)))
    return render_template("resources.html",resources=loads)

@app.post("/admin/resources/<int:user_id>")
@login_required
@role_required("admin")
def update_resource(user_id):
    with db() as conn:
        conn.execute("""INSERT INTO resource_profiles(user_id,weekly_capacity,hourly_cost,hourly_rate)
                        VALUES(?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET weekly_capacity=excluded.weekly_capacity,
                        hourly_cost=excluded.hourly_cost,hourly_rate=excluded.hourly_rate""",
                     (user_id,float(request.form.get("weekly_capacity","40")),float(request.form.get("hourly_cost","0")),
                      float(request.form.get("hourly_rate","0"))))
    return redirect(url_for("resources"))

@app.route("/projects/<int:project_id>/time",methods=["GET","POST"])
@login_required
def project_time(project_id):
    p=project_or_404(project_id,write=(request.method=="POST"))
    u=current_user()
    with db() as conn:
        tasks=conn.execute("SELECT id,wbs,title FROM tasks WHERE project_id=? ORDER BY wbs",(project_id,)).fetchall()
        if request.method=="POST":
            conn.execute("""INSERT INTO time_entries(project_id,task_id,user_id,work_date,hours,note,created_at)
                            VALUES(?,?,?,?,?,?,?)""",(project_id,request.form.get("task_id") or None,u["id"],
                            request.form.get("work_date") or date.today().isoformat(),float(request.form["hours"]),
                            request.form.get("note","").strip(),datetime.now().isoformat(timespec="seconds")))
            audit(project_id,"time",None,"logged",f"{request.form['hours']} h")
            flash("Tid rapporterad.","success")
            return redirect(url_for("project_time",project_id=project_id))
        entries=conn.execute("""SELECT te.*,u.display_name,t.title task_title FROM time_entries te
                                JOIN users u ON u.id=te.user_id LEFT JOIN tasks t ON t.id=te.task_id
                                WHERE te.project_id=? ORDER BY te.work_date DESC,te.id DESC""",(project_id,)).fetchall()
        total=conn.execute("SELECT COALESCE(SUM(hours),0) h FROM time_entries WHERE project_id=?",(project_id,)).fetchone()["h"]
    return render_template("time.html",project=p,tasks=tasks,entries=entries,total=total)

@app.route("/projects/<int:project_id>/financials",methods=["GET","POST"])
@login_required
def financials(project_id):
    p=project_or_404(project_id,manager=(request.method=="POST"))
    with db() as conn:
        if request.method=="POST":
            if request.form.get("action")=="budget":
                conn.execute("""INSERT INTO project_finance(project_id,budget,currency) VALUES(?,?,?)
                                ON CONFLICT(project_id) DO UPDATE SET budget=excluded.budget,currency=excluded.currency""",
                             (project_id,float(request.form.get("budget","0")),request.form.get("currency","SEK")))
            else:
                conn.execute("""INSERT INTO project_costs(project_id,category,description,planned,actual,cost_date)
                                VALUES(?,?,?,?,?,?)""",(project_id,request.form.get("category","External"),request.form["description"],
                                float(request.form.get("planned","0")),float(request.form.get("actual","0")),request.form.get("cost_date","")))
            audit(project_id,"finance",None,"updated","")
            return redirect(url_for("financials",project_id=project_id))
        finance=conn.execute("SELECT * FROM project_finance WHERE project_id=?",(project_id,)).fetchone()
        costs=conn.execute("SELECT * FROM project_costs WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
        time_cost=conn.execute("""SELECT COALESCE(SUM(te.hours*COALESCE(r.hourly_cost,0)),0) c FROM time_entries te
                                  LEFT JOIN resource_profiles r ON r.user_id=te.user_id WHERE te.project_id=?""",(project_id,)).fetchone()["c"]
    budget=float(finance["budget"] if finance else 0); external=sum(float(c["actual"] or 0) for c in costs)
    actual=external+float(time_cost or 0)
    return render_template("financials.html",project=p,finance=finance,costs=costs,time_cost=time_cost,actual=actual,
                           variance=budget-actual,budget=budget)

@app.get("/projects/<int:project_id>/earned-value")
@login_required
def earned_value(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,)).fetchall()
        fin=conn.execute("SELECT * FROM project_finance WHERE project_id=?",(project_id,)).fetchone()
        actual_external=conn.execute("SELECT COALESCE(SUM(actual),0) a FROM project_costs WHERE project_id=?",(project_id,)).fetchone()["a"]
        time_cost=conn.execute("""SELECT COALESCE(SUM(te.hours*COALESCE(r.hourly_cost,0)),0) c FROM time_entries te
                                  LEFT JOIN resource_profiles r ON r.user_id=te.user_id WHERE te.project_id=?""",(project_id,)).fetchone()["c"]
    bac=float(fin["budget"] if fin else 0)
    today=date.today(); weighted_plan=0; weighted_ev=0
    for t in tasks:
        s=parse_date(t["start_date"]); e=parse_date(t["end_date"])
        if s and e:
            if today>=e: planned=1
            elif today<=s: planned=0
            else: planned=(today-s).days/max((e-s).days,1)
        else: planned=0
        weighted_plan+=planned; weighted_ev+=float(t["progress"] or 0)/100
    n=max(len(tasks),1); pv=bac*(weighted_plan/n); ev=bac*(weighted_ev/n); ac=float(actual_external or 0)+float(time_cost or 0)
    spi=(ev/pv) if pv else 0; cpi=(ev/ac) if ac else 0
    return render_template("earned_value.html",project=p,bac=bac,pv=pv,ev=ev,ac=ac,spi=spi,cpi=cpi)


@app.route("/pmo",methods=["GET","POST"])
@login_required
def pmo():
    u=current_user()
    with db() as conn:
        if request.method=="POST":
            if u["role"] not in ("admin","pm"): abort(403)
            conn.execute("INSERT INTO portfolios(name,description,owner) VALUES(?,?,?)",
                         (request.form["name"],request.form.get("description",""),u["display_name"]))
            return redirect(url_for("pmo"))
        portfolios=conn.execute("""SELECT pf.*,COUNT(pp.project_id) project_count FROM portfolios pf
                                   LEFT JOIN portfolio_projects pp ON pp.portfolio_id=pf.id GROUP BY pf.id ORDER BY pf.name""").fetchall()
        latest=conn.execute("""SELECT sr.*,p.name project_name FROM status_reports sr JOIN projects p ON p.id=sr.project_id
                               WHERE sr.id IN (SELECT MAX(id) FROM status_reports GROUP BY project_id)
                               ORDER BY sr.report_date DESC""").fetchall()
    return render_template("pmo.html",portfolios=portfolios,latest=latest)

@app.route("/projects/<int:project_id>/status-report",methods=["GET","POST"])
@login_required
def status_report(project_id):
    p=project_or_404(project_id,manager=(request.method=="POST"))
    u=current_user()
    with db() as conn:
        if request.method=="POST":
            conn.execute("""INSERT INTO status_reports(project_id,report_date,overall_rag,scope_rag,schedule_rag,budget_rag,resources_rag,
                            summary,achievements,next_steps,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                         (project_id,date.today().isoformat(),request.form["overall_rag"],request.form["scope_rag"],
                          request.form["schedule_rag"],request.form["budget_rag"],request.form["resources_rag"],
                          request.form.get("summary",""),request.form.get("achievements",""),request.form.get("next_steps",""),
                          u["id"],datetime.now().isoformat(timespec="seconds")))
            return redirect(url_for("status_report",project_id=project_id))
        reports=conn.execute("SELECT * FROM status_reports WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
    return render_template("status_report.html",project=p,reports=reports)

@app.route("/projects/<int:project_id>/governance",methods=["GET","POST"])
@login_required
def governance(project_id):
    p=project_or_404(project_id,write=(request.method=="POST"))
    u=current_user()
    with db() as conn:
        if request.method=="POST":
            kind=request.form["kind"]
            if kind=="raid":
                conn.execute("""INSERT INTO raid_items(project_id,item_type,title,owner,status,due_date,details) VALUES(?,?,?,?,?,?,?)""",
                             (project_id,request.form["item_type"],request.form["title"],request.form.get("owner",""),
                              request.form.get("status","Open"),request.form.get("due_date",""),request.form.get("details","")))
            elif kind=="decision":
                conn.execute("""INSERT INTO decisions(project_id,title,decision,decided_by,decision_date) VALUES(?,?,?,?,?)""",
                             (project_id,request.form["title"],request.form["decision"],request.form.get("decided_by",u["display_name"]),
                              request.form.get("decision_date") or date.today().isoformat()))
            elif kind=="change":
                conn.execute("""INSERT INTO change_requests(project_id,title,description,impact_scope,impact_days,impact_cost,status,requested_by,created_at)
                                VALUES(?,?,?,?,?,?,?,?,?)""",(project_id,request.form["title"],request.form.get("description",""),
                                request.form.get("impact_scope",""),int(request.form.get("impact_days","0") or 0),
                                float(request.form.get("impact_cost","0") or 0),"Proposed",u["id"],datetime.now().isoformat(timespec="seconds")))
            audit(project_id,kind,None,"created",request.form.get("title",""))
            return redirect(url_for("governance",project_id=project_id))
        raid=conn.execute("SELECT * FROM raid_items WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
        decisions=conn.execute("SELECT * FROM decisions WHERE project_id=? ORDER BY decision_date DESC,id DESC",(project_id,)).fetchall()
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
        approvals=conn.execute("SELECT * FROM approvals WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
    return render_template("governance.html",project=p,raid=raid,decisions=decisions,changes=changes,approvals=approvals)

@app.post("/projects/<int:project_id>/changes/<int:change_id>/decision")
@login_required
def decide_change(project_id,change_id):
    project_or_404(project_id,manager=True)
    status=request.form.get("status")
    if status not in ("Approved","Rejected"): abort(400)
    with db() as conn: conn.execute("UPDATE change_requests SET status=? WHERE id=? AND project_id=?",(status,change_id,project_id))
    audit(project_id,"change",change_id,status,"")
    return redirect(url_for("governance",project_id=project_id))


UPLOAD_DIR=DATA_DIR/"uploads"

@app.get("/notifications")
@login_required
def notifications_page():
    u=current_user()
    with db() as conn:
        items=conn.execute("SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 100",(u["id"],)).fetchall()
        conn.execute("UPDATE notifications SET is_read=1 WHERE user_id=?",(u["id"],))
    return render_template("notifications.html",items=items)

@app.route("/projects/<int:project_id>/collaboration",methods=["GET","POST"])
@login_required
def collaboration(project_id):
    p=project_or_404(project_id,write=(request.method=="POST"))
    u=current_user()
    with db() as conn:
        if request.method=="POST":
            kind=request.form["kind"]
            if kind=="doc":
                conn.execute("INSERT INTO documents(project_id,title,body,updated_by,updated_at) VALUES(?,?,?,?,?)",
                             (project_id,request.form["title"],request.form.get("body",""),u["id"],datetime.now().isoformat(timespec="seconds")))
            elif kind=="meeting":
                conn.execute("INSERT INTO meetings(project_id,title,meeting_date,attendees,notes) VALUES(?,?,?,?,?)",
                             (project_id,request.form["title"],request.form.get("meeting_date") or date.today().isoformat(),
                              request.form.get("attendees",""),request.form.get("notes","")))
            elif kind=="action":
                conn.execute("INSERT INTO action_items(project_id,title,owner,due_date,status) VALUES(?,?,?,?,?)",
                             (project_id,request.form["title"],request.form.get("owner",""),request.form.get("due_date",""),"Open"))
            audit(project_id,kind,None,"created",request.form.get("title",""))
            return redirect(url_for("collaboration",project_id=project_id))
        docs=conn.execute("SELECT * FROM documents WHERE project_id=? ORDER BY updated_at DESC",(project_id,)).fetchall()
        meetings=conn.execute("SELECT * FROM meetings WHERE project_id=? ORDER BY meeting_date DESC,id DESC",(project_id,)).fetchall()
        actions=conn.execute("SELECT * FROM action_items WHERE project_id=? ORDER BY status,due_date",(project_id,)).fetchall()
        files=conn.execute("SELECT * FROM attachments WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
    return render_template("collaboration.html",project=p,docs=docs,meetings=meetings,actions=actions,files=files)

@app.post("/projects/<int:project_id>/upload")
@login_required
def upload_attachment(project_id):
    project_or_404(project_id,write=True)
    f=request.files.get("file")
    if not f or not f.filename: abort(400)
    UPLOAD_DIR.mkdir(parents=True,exist_ok=True)
    original=secure_filename(f.filename)
    stored=f"{project_id}-{secrets.token_hex(8)}-{original}"
    f.save(UPLOAD_DIR/stored)
    u=current_user()
    with db() as conn:
        conn.execute("""INSERT INTO attachments(project_id,entity_type,original_name,stored_name,uploaded_by,uploaded_at)
                        VALUES(?,?,?,?,?,?)""",(project_id,"project",original,stored,u["id"],datetime.now().isoformat(timespec="seconds")))
    audit(project_id,"attachment",None,"uploaded",original)
    return redirect(url_for("collaboration",project_id=project_id))

@app.get("/projects/<int:project_id>/files/<int:file_id>")
@login_required
def download_attachment(project_id,file_id):
    project_or_404(project_id)
    with db() as conn: row=conn.execute("SELECT * FROM attachments WHERE id=? AND project_id=?",(file_id,project_id)).fetchone()
    if not row: abort(404)
    return send_file(UPLOAD_DIR/row["stored_name"],as_attachment=True,download_name=row["original_name"])

@app.post("/projects/<int:project_id>/notify")
@login_required
def notify_project(project_id):
    project_or_404(project_id,manager=True)
    msg=request.form["message"].strip()
    with db() as conn:
        members=conn.execute("SELECT user_id FROM project_members WHERE project_id=?",(project_id,)).fetchall()
        for m in members:
            conn.execute("INSERT INTO notifications(user_id,project_id,message,link,is_read,created_at) VALUES(?,?,?,?,0,?)",
                         (m["user_id"],project_id,msg,url_for("project",project_id=project_id),datetime.now().isoformat(timespec="seconds")))
    return redirect(url_for("collaboration",project_id=project_id))


BACKUP_DIR=DATA_DIR/"admin-backups"

def api_user():
    auth=request.headers.get("Authorization","")
    if not auth.startswith("Bearer "): return None
    raw=auth[7:].strip(); digest=hashlib.sha256(raw.encode()).hexdigest()
    with db() as conn:
        row=conn.execute("""SELECT u.*,k.id key_id FROM api_keys k JOIN users u ON u.id=k.user_id
                            WHERE k.key_hash=? AND k.active=1 AND u.active=1""",(digest,)).fetchone()
        if row: conn.execute("UPDATE api_keys SET last_used=? WHERE id=?",(datetime.now().isoformat(timespec="seconds"),row["key_id"]))
    return row

def api_required(fn):
    @wraps(fn)
    def wrap(*args,**kwargs):
        u=api_user()
        if not u:return jsonify(error="unauthorized"),401
        request.api_user=u
        return fn(*args,**kwargs)
    return wrap

@app.get("/api/v1/projects")
@api_required
def api_projects():
    u=request.api_user
    with db() as conn:
        if u["role"]=="admin":
            rows=conn.execute("SELECT * FROM projects ORDER BY id").fetchall()
        else:
            rows=conn.execute("""SELECT p.* FROM projects p JOIN project_members pm ON pm.project_id=p.id
                                 WHERE pm.user_id=? ORDER BY p.id""",(u["id"],)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.get("/api/v1/projects/<int:project_id>/tasks")
@api_required
def api_tasks(project_id):
    u=request.api_user
    with db() as conn:
        if u["role"]!="admin":
            ok=conn.execute("SELECT 1 FROM project_members WHERE project_id=? AND user_id=?",(project_id,u["id"])).fetchone()
            if not ok:return jsonify(error="forbidden"),403
        rows=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/admin/api-keys",methods=["GET","POST"])
@login_required
@role_required("admin")
def admin_api_keys():
    generated=None
    with db() as conn:
        if request.method=="POST":
            raw="pp_"+secrets.token_urlsafe(30); digest=hashlib.sha256(raw.encode()).hexdigest()
            conn.execute("INSERT INTO api_keys(user_id,name,key_hash,prefix,created_at) VALUES(?,?,?,?,?)",
                         (int(request.form["user_id"]),request.form["name"],digest,raw[:10],datetime.now().isoformat(timespec="seconds")))
            generated=raw
        keys=conn.execute("""SELECT k.*,u.username FROM api_keys k JOIN users u ON u.id=k.user_id ORDER BY k.id DESC""").fetchall()
        users=conn.execute("SELECT * FROM users WHERE active=1 ORDER BY username").fetchall()
    return render_template("api_keys.html",keys=keys,users=users,generated=generated)

@app.route("/admin/custom-fields",methods=["GET","POST"])
@login_required
@role_required("admin")
def admin_custom_fields():
    with db() as conn:
        if request.method=="POST":
            conn.execute("""INSERT INTO custom_field_definitions(entity_type,name,field_type,options_json,required)
                            VALUES(?,?,?,?,?)""",(request.form["entity_type"],request.form["name"],request.form["field_type"],
                            json.dumps([x.strip() for x in request.form.get("options","").split(",") if x.strip()]),
                            1 if request.form.get("required")=="on" else 0))
            return redirect(url_for("admin_custom_fields"))
        fields=conn.execute("SELECT * FROM custom_field_definitions ORDER BY entity_type,name").fetchall()
    return render_template("custom_fields.html",fields=fields)

@app.route("/admin/backups",methods=["GET","POST"])
@login_required
@role_required("admin")
def admin_backups():
    BACKUP_DIR.mkdir(parents=True,exist_ok=True)
    if request.method=="POST":
        name=f"projectplan-{datetime.now():%Y%m%d-%H%M%S}.db"
        shutil.copy2(DB_PATH,BACKUP_DIR/name)
        flash("Backup skapad.","success")
        return redirect(url_for("admin_backups"))
    files=sorted(BACKUP_DIR.glob("*.db"),reverse=True)
    return render_template("backups.html",files=[{"name":f.name,"size":f.stat().st_size} for f in files])

@app.get("/admin/backups/<name>")
@login_required
@role_required("admin")
def download_backup(name):
    safe=Path(name).name
    return send_file(BACKUP_DIR/safe,as_attachment=True,download_name=safe)

@app.get("/admin/diagnostics")
@login_required
@role_required("admin")
def diagnostics():
    with db() as conn:
        db_size=DB_PATH.stat().st_size if DB_PATH.exists() else 0
        tables=conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name").fetchall()
        migrations=conn.execute("SELECT * FROM schema_migrations ORDER BY applied_at DESC").fetchall()
    info={"version":APP_VERSION,"python":platform.python_version(),"platform":platform.platform(),
          "db_path":str(DB_PATH),"db_size":db_size,"table_count":len(tables),"upload_dir":str(UPLOAD_DIR)}
    return render_template("diagnostics.html",info=info,migrations=migrations)

@app.get("/api")
@login_required
def api_docs():
    return render_template("api_docs.html")


@app.get("/about")
@login_required
def about():
    return render_template("about.html")

def roadmap_accessible_projects():
    u=current_user()
    with db() as conn:
        if u["role"]=="admin":
            return conn.execute("SELECT * FROM projects ORDER BY name").fetchall()
        return conn.execute("""SELECT p.* FROM projects p JOIN project_members pm ON pm.project_id=p.id
                              WHERE pm.user_id=? ORDER BY p.name""",(u["id"],)).fetchall()

def roadmap_task_tree(tasks):
    rows=[dict(t) for t in tasks]
    by_id={int(t["id"]):t for t in rows}
    children={}
    roots=[]
    for t in rows:
        pid=t.get("parent_task_id")
        if pid and int(pid) in by_id and int(pid)!=int(t["id"]):
            children.setdefault(int(pid),[]).append(t)
        else:
            roots.append(t)
    def key(t):
        raw=str(t.get("wbs") or "")
        parts=[]
        for p in raw.split("."):
            try: parts.append((0,int(p)))
            except: parts.append((1,p.lower()))
        return (parts,int(t.get("sort_order") or 0),int(t["id"]))
    for vals in children.values(): vals.sort(key=key)
    roots.sort(key=key)
    out=[]
    def walk(t,depth):
        r=dict(t); r["depth"]=depth; r["derived"]=derived_status(t); r["children_count"]=len(children.get(int(t["id"]),[]))
        out.append(r)
        for c in children.get(int(t["id"]),[]): walk(c,depth+1)
    for r in roots: walk(r,0)
    return out

@app.get("/planning")
@login_required
def planning_hub():
    projects=roadmap_accessible_projects()
    return render_template("planning_hub.html",projects=projects)

@app.get("/projects/<int:project_id>/cockpit")
@login_required
def project_cockpit(project_id):
    p=project_or_404(project_id)
    today=date.today()
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,))]
        risks=[dict(r) for r in conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC",(project_id,))]
        actions=[dict(r) for r in conn.execute("SELECT * FROM action_items WHERE project_id=? AND status<>'Done' ORDER BY due_date,id",(project_id,))]
        recent=conn.execute("""SELECT a.*,u.display_name FROM audit_log a LEFT JOIN users u ON u.id=a.user_id
                              WHERE a.project_id=? ORDER BY a.id DESC LIMIT 10""",(project_id,)).fetchall()
    overdue=[t for t in tasks if int(t.get("progress") or 0)<100 and parse_date(t.get("end_date")) and parse_date(t.get("end_date"))<today]
    blocked=[t for t in tasks if t.get("status")=="Blockerad" and int(t.get("progress") or 0)<100]
    high_risks=[r for r in risks if int(r.get("probability") or 0)*int(r.get("impact") or 0)>=12]
    milestones=[t for t in tasks if t.get("milestone") and parse_date(t.get("end_date")) and int(t.get("progress") or 0)<100]
    milestones.sort(key=lambda x:parse_date(x.get("end_date")))
    avg=round(sum(int(t.get("progress") or 0) for t in tasks)/len(tasks)) if tasks else 0
    rag="red" if overdue or len(high_risks)>=2 or len(blocked)>=2 else ("amber" if high_risks or blocked else "green")
    attention=[]
    for t in overdue[:5]: attention.append(("Försenad",t["title"],t.get("end_date") or "","danger"))
    for t in blocked[:5]: attention.append(("Blockerad",t["title"],t.get("owner") or "","warning"))
    for r in high_risks[:5]: attention.append(("Hög risk",r["title"],f"Score {int(r.get('probability') or 0)*int(r.get('impact') or 0)}","danger"))
    return render_template("cockpit.html",project=p,tasks=tasks,risks=risks,actions=actions,recent=recent,
                           overdue=overdue,blocked=blocked,high_risks=high_risks,milestones=milestones[:8],
                           avg=avg,rag=rag,attention=attention)

@app.get("/projects/<int:project_id>/wbs")
@login_required
def project_wbs(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,sort_order,id",(project_id,)).fetchall()
    return render_template("wbs.html",project=p,rows=roadmap_task_tree(tasks))

@app.post("/projects/<int:project_id>/tasks/<int:task_id>/quick")
@login_required
def task_quick_update(project_id,task_id):
    project_or_404(project_id,write=True)
    progress=max(0,min(100,int(request.form.get("progress","0") or 0)))
    status=request.form.get("status","Pågår")
    if progress==100: status="Klar"
    with db() as conn:
        conn.execute("UPDATE tasks SET progress=?,status=? WHERE id=? AND project_id=?",(progress,status,task_id,project_id))
    audit(project_id,"task",task_id,"quick_updated",f"{status} {progress}%")
    return redirect(request.referrer or url_for("project_wbs",project_id=project_id))

@app.get("/my-work-legacy")
@login_required
def my_work_2():
    u=current_user(); today=date.today()
    with db() as conn:
        projects=roadmap_accessible_projects()
        ids=[p["id"] for p in projects]
        if not ids: tasks=[]; actions=[]
        else:
            ph=",".join("?" for _ in ids)
            tasks=[dict(r) for r in conn.execute(f"""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
                WHERE t.project_id IN ({ph}) AND t.progress<100 AND
                (LOWER(t.owner)=LOWER(?) OR t.owner_user_id=?) ORDER BY t.end_date,t.priority""",(*ids,u["display_name"],u["id"]))]
            actions=[dict(r) for r in conn.execute(f"""SELECT a.*,p.name project_name FROM action_items a JOIN projects p ON p.id=a.project_id
                WHERE a.project_id IN ({ph}) AND a.status<>'Done' AND LOWER(a.owner)=LOWER(?) ORDER BY a.due_date""",(*ids,u["display_name"]))]
    for t in tasks:
        d=parse_date(t.get("end_date")); t["is_overdue"]=bool(d and d<today); t["is_week"]=bool(d and today<=d<=today+timedelta(days=7))
    return render_template("my_work_2.html",tasks=tasks,actions=actions,today=today)

@app.get("/projects/<int:project_id>/gantt-pro")
@login_required
def gantt_pro(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,))]
        links=[dict(r) for r in conn.execute("SELECT * FROM task_links WHERE project_id=? ORDER BY id",(project_id,))]
    dated=[t for t in tasks if parse_date(t.get("start_date")) and parse_date(t.get("end_date"))]
    if dated:
        rs=min(parse_date(t["start_date"]) for t in dated); re=max(parse_date(t["end_date"]) for t in dated)
        total=max(1,(re-rs).days+1)
        for t in dated:
            st=parse_date(t["start_date"]); en=parse_date(t["end_date"])
            t["left"]=round((st-rs).days/total*100,3); t["width"]=max(.8,round(((en-st).days+1)/total*100,3))
    else: rs=re=None
    return render_template("gantt_pro.html",project=p,tasks=roadmap_task_tree(tasks),dated=dated,links=links,rs=rs,re=re)

@app.route("/control-center",methods=["GET","POST"])
@login_required
def control_center():
    projects=roadmap_accessible_projects()
    project_id=request.args.get("project_id",type=int) or (projects[0]["id"] if projects else None)
    if project_id: project_or_404(project_id)
    with db() as conn:
        raid=conn.execute("SELECT * FROM raid_items WHERE project_id=? ORDER BY item_type,status,due_date",(project_id,)).fetchall() if project_id else []
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall() if project_id else []
        decisions=conn.execute("SELECT * FROM decisions WHERE project_id=? ORDER BY decided_at DESC,id DESC",(project_id,)).fetchall() if project_id else []
        approvals=conn.execute("SELECT * FROM approvals WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall() if project_id else []
    return render_template("control_center.html",projects=projects,project_id=project_id,raid=raid,changes=changes,decisions=decisions,approvals=approvals)

@app.post("/projects/<int:project_id>/raid/new")
@login_required
def roadmap_raid_new(project_id):
    project_or_404(project_id,write=True)
    with db() as conn:
        conn.execute("""INSERT INTO raid_items(project_id,item_type,title,owner,status,due_date,details)
                        VALUES(?,?,?,?,?,?,?)""",(project_id,request.form.get("kind","Risk"),request.form["title"].strip(),
                        request.form.get("owner","").strip(),request.form.get("status","Open"),request.form.get("due_date",""),
                        request.form.get("description","").strip()))
    audit(project_id,"raid",None,"created",request.form["title"].strip())
    return redirect(url_for("control_center",project_id=project_id))

@app.post("/projects/<int:project_id>/changes/new")
@login_required
def roadmap_change_new(project_id):
    project_or_404(project_id,write=True)
    u=current_user()
    with db() as conn:
        conn.execute("""INSERT INTO change_requests(project_id,title,description,reason,impact_cost,impact_days,status,requested_by,created_at)
                        VALUES(?,?,?,?,?,?,?,?,?)""",(project_id,request.form["title"].strip(),request.form.get("description","").strip(),
                        request.form.get("reason","").strip(),float(request.form.get("impact_cost","0") or 0),int(request.form.get("impact_days","0") or 0),
                        "Submitted",u["id"],datetime.now().isoformat(timespec="seconds")))
    audit(project_id,"change",None,"submitted",request.form["title"].strip())
    return redirect(url_for("control_center",project_id=project_id))

@app.post("/projects/<int:project_id>/changes/<int:change_id>/<decision>")
@login_required
def roadmap_change_decision(project_id,change_id,decision):
    project_or_404(project_id,manager=True)
    status="Approved" if decision=="approve" else "Rejected"
    with db() as conn:
        conn.execute("UPDATE change_requests SET status=?,decided_by=?,decided_at=? WHERE id=? AND project_id=?",
                     (status,current_user()["id"],datetime.now().isoformat(timespec="seconds"),change_id,project_id))
    audit(project_id,"change",change_id,status.lower(),"")
    return redirect(url_for("control_center",project_id=project_id))

@app.post("/projects/<int:project_id>/decisions/new")
@login_required
def roadmap_decision_new(project_id):
    project_or_404(project_id,write=True)
    with db() as conn:
        conn.execute("""INSERT INTO decisions(project_id,title,decision,decided_by,decision_date,owner,decided_at)
                        VALUES(?,?,?,?,?,?,?)""",(project_id,request.form["title"].strip(),request.form.get("decision","").strip(),
                        request.form.get("owner","").strip(),date.today().isoformat(),request.form.get("owner","").strip(),
                        datetime.now().isoformat(timespec="seconds")))
    return redirect(url_for("control_center",project_id=project_id))

@app.route("/resource-plan",methods=["GET","POST"])
@login_required
def resource_plan():
    projects=roadmap_accessible_projects()
    if request.method=="POST":
        pid=int(request.form["project_id"]); project_or_404(pid,manager=True)
        with db() as conn:
            conn.execute("""INSERT INTO resource_allocations(project_id,user_id,resource_name,week_start,allocation_pct,planned_hours)
                            VALUES(?,?,?,?,?,?)""",(pid,int(request.form["user_id"]) if request.form.get("user_id") else None,
                            request.form["resource_name"].strip(),request.form["week_start"],int(request.form.get("allocation_pct","0") or 0),
                            float(request.form.get("planned_hours","0") or 0)))
        return redirect(url_for("resource_plan"))
    ids=[p["id"] for p in projects]
    with db() as conn:
        users=conn.execute("SELECT * FROM users WHERE active=1 ORDER BY display_name").fetchall()
        if ids:
            ph=",".join("?" for _ in ids)
            allocations=conn.execute(f"""SELECT a.*,p.name project_name FROM resource_allocations a JOIN projects p ON p.id=a.project_id
                                        WHERE a.project_id IN ({ph}) ORDER BY a.week_start,a.resource_name""",ids).fetchall()
        else: allocations=[]
    totals={}
    for a in allocations:
        k=(a["resource_name"],a["week_start"])
        totals[k]=totals.get(k,0)+int(a["allocation_pct"] or 0)
    return render_template("resource_plan.html",projects=projects,users=users,allocations=allocations,totals=totals)

@app.get("/finance-control")
@login_required
def finance_control():
    projects=roadmap_accessible_projects(); ids=[p["id"] for p in projects]
    rows=[]
    with db() as conn:
        for p in projects:
            f=conn.execute("SELECT * FROM project_finance WHERE project_id=?",(p["id"],)).fetchone()
            ac=conn.execute("SELECT COALESCE(SUM(actual),0) v FROM project_costs WHERE project_id=?",(p["id"],)).fetchone()["v"]
            planned=conn.execute("SELECT COALESCE(SUM(planned),0) v FROM project_costs WHERE project_id=?",(p["id"],)).fetchone()["v"]
            budget=float(f["budget"] or 0) if f else 0
            actual=float(ac or 0); forecast=max(float(planned or 0),actual)
            rows.append({"project":p,"budget":budget,"actual":actual,"forecast":forecast,
                         "variance":budget-forecast,"used_pct":round(actual/budget*100) if budget else 0})
    return render_template("finance_control.html",rows=rows)

@app.route("/executive-pmo",methods=["GET","POST"])
@login_required
def executive_pmo():
    projects=roadmap_accessible_projects()
    if request.method=="POST" and current_user()["role"] in ("admin","pm"):
        with db() as conn:
            conn.execute("INSERT OR IGNORE INTO programs(name,owner,description) VALUES(?,?,?)",
                         (request.form["name"].strip(),request.form.get("owner","").strip(),request.form.get("description","").strip()))
        return redirect(url_for("executive_pmo"))
    data=[]
    today=date.today()
    with db() as conn:
        programs=conn.execute("SELECT * FROM programs ORDER BY name").fetchall()
        for p in projects:
            tasks=[dict(x) for x in conn.execute("SELECT * FROM tasks WHERE project_id=?",(p["id"],))]
            risks=[dict(x) for x in conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd'",(p["id"],))]
            overdue=sum(1 for t in tasks if int(t.get("progress") or 0)<100 and parse_date(t.get("end_date")) and parse_date(t.get("end_date"))<today)
            high=sum(1 for r in risks if int(r.get("probability") or 0)*int(r.get("impact") or 0)>=12)
            progress=round(sum(int(t.get("progress") or 0) for t in tasks)/len(tasks)) if tasks else 0
            score=max(0,100-overdue*12-high*10-(100-progress)//5)
            rag="red" if score<60 else ("amber" if score<80 else "green")
            data.append({"p":p,"progress":progress,"overdue":overdue,"high":high,"score":score,"rag":rag})
    return render_template("executive_pmo.html",data=data,programs=programs)

@app.get("/portfolio/<int:portfolio_id>/dashboard")
@login_required
def portfolio_dashboard_v2(portfolio_id):
    with db() as conn:
        portfolio=conn.execute("SELECT * FROM portfolios WHERE id=?",(portfolio_id,)).fetchone()
        if not portfolio: abort(404)
        projects=conn.execute("""SELECT p.* FROM projects p JOIN portfolio_projects pp ON pp.project_id=p.id
                                 WHERE pp.portfolio_id=? ORDER BY p.name""",(portfolio_id,)).fetchall()
    return render_template("portfolio_dashboard_v2.html",portfolio=portfolio,projects=projects)

@app.route("/collaboration-hub",methods=["GET"])
@login_required
def collaboration_hub():
    projects=roadmap_accessible_projects(); u=current_user()
    ids=[p["id"] for p in projects]
    with db() as conn:
        mentions=conn.execute("""SELECT m.*,p.name project_name FROM mentions m JOIN projects p ON p.id=m.project_id
                                 WHERE m.mentioned_user_id=? ORDER BY m.id DESC LIMIT 30""",(u["id"],)).fetchall()
        if ids:
            ph=",".join("?" for _ in ids)
            meetings=conn.execute(f"SELECT m.*,p.name project_name FROM meetings m JOIN projects p ON p.id=m.project_id WHERE m.project_id IN ({ph}) ORDER BY meeting_date DESC LIMIT 30",ids).fetchall()
            actions=conn.execute(f"SELECT a.*,p.name project_name FROM action_items a JOIN projects p ON p.id=a.project_id WHERE a.project_id IN ({ph}) AND a.status<>'Done' ORDER BY a.due_date LIMIT 50",ids).fetchall()
        else: meetings=[]; actions=[]
    return render_template("collaboration_hub.html",projects=projects,mentions=mentions,meetings=meetings,actions=actions)

@app.post("/projects/<int:project_id>/mention")
@login_required
def create_mention(project_id):
    project_or_404(project_id,write=True)
    username=request.form.get("username","").strip().lstrip("@")
    with db() as conn:
        user=conn.execute("SELECT * FROM users WHERE username=? AND active=1",(username,)).fetchone()
        if not user:
            flash("Användaren hittades inte.","error")
        else:
            conn.execute("""INSERT INTO mentions(project_id,mentioned_user_id,source_type,message,created_at)
                            VALUES(?,?,?,?,?)""",(project_id,user["id"],request.form.get("source_type","message"),
                            request.form.get("message","").strip(),datetime.now().isoformat(timespec="seconds")))
            conn.execute("""INSERT INTO notifications(user_id,project_id,message,link,created_at)
                            VALUES(?,?,?,?,?)""",(user["id"],project_id,request.form.get("message","").strip(),
                            url_for("collaboration",project_id=project_id),datetime.now().isoformat(timespec="seconds")))
    return redirect(url_for("collaboration",project_id=project_id))

@app.post("/documents/<int:document_id>/version")
@login_required
def document_new_version(document_id):
    with db() as conn:
        doc=conn.execute("SELECT * FROM documents WHERE id=?",(document_id,)).fetchone()
        if not doc: abort(404)
    project_or_404(doc["project_id"],write=True)
    with db() as conn:
        n=conn.execute("SELECT COALESCE(MAX(version_no),0)+1 n FROM document_versions WHERE document_id=?",(document_id,)).fetchone()["n"]
        conn.execute("""INSERT INTO document_versions(document_id,version_no,content,created_by,created_at)
                        VALUES(?,?,?,?,?)""",(document_id,n,request.form.get("content",""),current_user()["id"],datetime.now().isoformat(timespec="seconds")))
        conn.execute("UPDATE documents SET body=?,updated_at=? WHERE id=?",(request.form.get("content",""),datetime.now().isoformat(timespec="seconds"),document_id))
    return redirect(url_for("collaboration",project_id=doc["project_id"]))

def project_assistant_summary(project_id):
    p=project_or_404(project_id)
    today=date.today()
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]
        risks=[dict(r) for r in conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd'",(project_id,))]
        changes=[dict(r) for r in conn.execute("SELECT * FROM change_requests WHERE project_id=? AND status='Submitted'",(project_id,))]
        actions=[dict(r) for r in conn.execute("SELECT * FROM action_items WHERE project_id=? AND status<>'Done'",(project_id,))]
    overdue=[t for t in tasks if int(t.get("progress") or 0)<100 and parse_date(t.get("end_date")) and parse_date(t.get("end_date"))<today]
    blocked=[t for t in tasks if t.get("status")=="Blockerad"]
    high=[r for r in risks if int(r.get("probability") or 0)*int(r.get("impact") or 0)>=12]
    progress=round(sum(int(t.get("progress") or 0) for t in tasks)/len(tasks)) if tasks else 0
    lines=[f"{p['name']} är {progress}% färdigt."]
    if overdue: lines.append(f"{len(overdue)} aktivitet(er) är försenade.")
    if blocked: lines.append(f"{len(blocked)} aktivitet(er) är blockerade.")
    if high: lines.append(f"{len(high)} höga risker behöver uppmärksamhet.")
    if changes: lines.append(f"{len(changes)} change request(s) väntar på beslut.")
    if actions: lines.append(f"{len(actions)} öppna actions finns.")
    if not any((overdue,blocked,high,changes)): lines.append("Inga kritiska avvikelser identifierades av regelmotorn.")
    return {"project":p,"progress":progress,"overdue":overdue,"blocked":blocked,"high":high,"changes":changes,"actions":actions,"text":" ".join(lines)}

@app.route("/workspace")
@login_required
def workspace():
    projects=roadmap_accessible_projects()
    prefs={}
    with db() as conn:
        row=conn.execute("SELECT layout_json FROM dashboard_preferences WHERE user_id=?",(current_user()["id"],)).fetchone()
        if row:
            try: prefs=json.loads(row["layout_json"] or "{}")
            except: prefs={}
    return render_template("workspace.html",projects=projects,prefs=prefs)

@app.post("/workspace/preferences")
@login_required
def workspace_preferences():
    layout=request.form.get("layout_json","{}")
    try: json.loads(layout)
    except: layout="{}"
    with db() as conn:
        conn.execute("DELETE FROM dashboard_preferences WHERE user_id=?",(current_user()["id"],))
        conn.execute("INSERT INTO dashboard_preferences(user_id,layout_json) VALUES(?,?)",(current_user()["id"],layout))
    flash("Dashboard sparad.","success")
    return redirect(url_for("workspace"))

@app.get("/assistant/project/<int:project_id>")
@login_required
def assistant_project(project_id):
    summary=project_assistant_summary(project_id)
    return render_template("assistant_project.html",summary=summary)

@app.route("/report-center",methods=["GET","POST"])
@login_required
def report_center():
    if request.method=="POST":
        with db() as conn:
            conn.execute("INSERT INTO report_definitions(name,report_type,config_json,owner_user_id,created_at) VALUES(?,?,?,?,?)",
                         (request.form["name"].strip(),request.form.get("report_type","portfolio"),request.form.get("config_json","{}"),
                          current_user()["id"],datetime.now().isoformat(timespec="seconds")))
        return redirect(url_for("report_center"))
    with db() as conn:
        reports=conn.execute("SELECT * FROM report_definitions WHERE owner_user_id=? OR ?='admin' ORDER BY id DESC",(current_user()["id"],current_user()["role"])).fetchall()
    return render_template("report_center.html",reports=reports)

@app.route("/admin/integrations",methods=["GET","POST"])
@login_required
@role_required("admin")
def admin_integrations():
    with db() as conn:
        if request.method=="POST":
            conn.execute("UPDATE oidc_settings SET enabled=?,issuer=?,client_id=?,scopes=? WHERE id=1",
                         (1 if request.form.get("enabled")=="on" else 0,request.form.get("issuer","").strip(),
                          request.form.get("client_id","").strip(),request.form.get("scopes","openid profile email").strip()))
            flash("OIDC-konfiguration sparad. Aktivering kräver IdP/client secret i servermiljön.","success")
            return redirect(url_for("admin_integrations"))
        oidc=conn.execute("SELECT * FROM oidc_settings WHERE id=1").fetchone()
        webhooks=conn.execute("SELECT * FROM webhooks ORDER BY id DESC").fetchall()
        rules=conn.execute("SELECT * FROM automation_rules ORDER BY id DESC").fetchall()
    return render_template("admin_integrations.html",oidc=oidc,webhooks=webhooks,rules=rules)

@app.get("/api/v1/health/summary")
def api_health_summary():
    return jsonify(status="ok",version=APP_VERSION,features=["planning","control","resources","finance","pmo","collaboration","reports","assistant"])


def record_recent(entity_type,entity_id,label,url):
    u=current_user()
    if not u: return
    with db() as conn:
        conn.execute("DELETE FROM recent_items WHERE user_id=? AND entity_type=? AND entity_id=?",(u["id"],entity_type,entity_id))
        conn.execute("INSERT INTO recent_items(user_id,entity_type,entity_id,label,url,viewed_at) VALUES(?,?,?,?,?,?)",
                     (u["id"],entity_type,entity_id,label,url,datetime.now().isoformat(timespec="seconds")))
        conn.execute("DELETE FROM recent_items WHERE id IN (SELECT id FROM recent_items WHERE user_id=? ORDER BY viewed_at DESC LIMIT -1 OFFSET 20)",(u["id"],))

@app.get("/productivity")
@login_required
def productivity_hub():
    u=current_user()
    with db() as conn:
        fav=conn.execute("SELECT * FROM favorites WHERE user_id=? ORDER BY created_at DESC",(u["id"],)).fetchall()
        recent=conn.execute("SELECT * FROM recent_items WHERE user_id=? ORDER BY viewed_at DESC LIMIT 12",(u["id"],)).fetchall()
        views=conn.execute("SELECT * FROM saved_views WHERE user_id=? ORDER BY created_at DESC",(u["id"],)).fetchall()
    return render_template("productivity.html",favorites=fav,recent=recent,views=views)

@app.get("/global-search")
@login_required
def global_search():
    q=request.args.get("q","").strip(); results=[]
    if q:
        like=f"%{q}%"; projects=roadmap_accessible_projects(); ids=[p["id"] for p in projects]
        with db() as conn:
            for p in projects:
                if q.lower() in (p["name"]+" "+(p["customer"] or "")).lower():
                    results.append({"type":"Projekt","label":p["name"],"detail":p["customer"],"url":url_for("project_cockpit",project_id=p["id"])})
            if ids:
                ph=",".join("?" for _ in ids)
                rows=conn.execute(f"SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id WHERE t.project_id IN ({ph}) AND (t.title LIKE ? OR t.wbs LIKE ?) LIMIT 50",(*ids,like,like)).fetchall()
                for t in rows: results.append({"type":"Aktivitet","label":t["title"],"detail":f"{t['project_name']} · {t['wbs']}","url":url_for("edit_task",project_id=t["project_id"],task_id=t["id"])})
    return render_template("global_search.html",q=q,results=results)

@app.post("/favorites/toggle")
@login_required
def toggle_favorite():
    et=request.form["entity_type"]; eid=int(request.form["entity_id"]); u=current_user()
    with db() as conn:
        row=conn.execute("SELECT 1 FROM favorites WHERE user_id=? AND entity_type=? AND entity_id=?",(u["id"],et,eid)).fetchone()
        if row: conn.execute("DELETE FROM favorites WHERE user_id=? AND entity_type=? AND entity_id=?",(u["id"],et,eid))
        else: conn.execute("INSERT INTO favorites(user_id,entity_type,entity_id,created_at) VALUES(?,?,?,?)",(u["id"],et,eid,datetime.now().isoformat(timespec="seconds")))
    return redirect(request.referrer or url_for("productivity_hub"))

@app.post("/projects/<int:project_id>/tasks/bulk")
@login_required
def tasks_bulk_update(project_id):
    project_or_404(project_id,write=True)
    ids=[int(x) for x in request.form.getlist("task_id") if x.isdigit()]
    if ids:
        status=request.form.get("status",""); priority=request.form.get("priority","")
        with db() as conn:
            for tid in ids:
                if status: conn.execute("UPDATE tasks SET status=? WHERE id=? AND project_id=?",(status,tid,project_id))
                if priority: conn.execute("UPDATE tasks SET priority=? WHERE id=? AND project_id=?",(priority,tid,project_id))
        audit(project_id,"task",None,"bulk_updated",f"{len(ids)} activities")
    return redirect(request.referrer or url_for("project_wbs",project_id=project_id))

@app.get("/projects/<int:project_id>/timeline")
@login_required
def project_timeline(project_id):
    p=project_or_404(project_id); record_recent("project",project_id,p["name"],url_for("project_cockpit",project_id=project_id))
    with db() as conn:
        rows=conn.execute("SELECT a.*,u.display_name FROM audit_log a LEFT JOIN users u ON u.id=a.user_id WHERE a.project_id=? ORDER BY a.id DESC LIMIT 100",(project_id,)).fetchall()
    return render_template("project_timeline.html",project=p,rows=rows)


def planning_network(project_id):
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]
        links=[dict(r) for r in conn.execute("SELECT * FROM task_links WHERE project_id=?",(project_id,))]
    by={int(t["id"]):t for t in tasks}; preds={k:[] for k in by}; succ={k:[] for k in by}
    for l in links:
        a=int(l["predecessor_id"]); b=int(l["successor_id"])
        if a in by and b in by: preds[b].append(l); succ[a].append(l)
    indeg={k:len(preds[k]) for k in by}; q=[k for k,v in indeg.items() if v==0]; order=[]
    while q:
        n=q.pop(0); order.append(n)
        for l in succ[n]:
            b=int(l["successor_id"]); indeg[b]-=1
            if indeg[b]==0:q.append(b)
    cycle=len(order)!=len(by)
    if cycle: order=list(by)
    es={k:0 for k in by}; ef={}
    for k in order:
        dur=max(1,int(by[k].get("duration_days") or 1)); start=0
        for l in preds[k]:
            a=int(l["predecessor_id"]); lag=int(l.get("lag_days") or 0); typ=l.get("link_type") or "FS"
            pd=max(1,int(by[a].get("duration_days") or 1))
            if typ=="FS": cand=ef.get(a,pd)+lag
            elif typ=="SS": cand=es.get(a,0)+lag
            elif typ=="FF": cand=ef.get(a,pd)+lag-dur
            else: cand=es.get(a,0)+lag-dur
            start=max(start,cand)
        es[k]=max(0,start); ef[k]=es[k]+dur
    finish=max(ef.values(),default=0); lf={k:finish for k in by}; ls={}
    for k in reversed(order):
        dur=max(1,int(by[k].get("duration_days") or 1)); latest=finish-dur
        if succ[k]:
            vals=[]
            for l in succ[k]:
                b=int(l["successor_id"]); lag=int(l.get("lag_days") or 0); typ=l.get("link_type") or "FS"; bd=max(1,int(by[b].get("duration_days") or 1))
                if typ=="FS": vals.append(ls.get(b,finish-bd)-lag-dur)
                elif typ=="SS": vals.append(ls.get(b,finish-bd)-lag)
                elif typ=="FF": vals.append(lf.get(b,finish)-lag-dur)
                else: vals.append(lf.get(b,finish)-lag)
            latest=min(vals)
        ls[k]=max(0,latest); lf[k]=ls[k]+dur
    rows=[]
    for k in order:
        t=dict(by[k]); t.update(es=es[k],ef=ef[k],ls=ls[k],lf=lf[k],slack=max(0,ls[k]-es[k]),critical=(ls[k]-es[k])==0); rows.append(t)
    return rows,links,cycle,finish

@app.get("/projects/<int:project_id>/advanced-planning")
@login_required
def advanced_planning(project_id):
    p=project_or_404(project_id); rows,links,cycle,finish=planning_network(project_id)
    with db() as conn: scenarios=conn.execute("SELECT * FROM planning_scenarios WHERE project_id=? ORDER BY id DESC",(project_id,)).fetchall()
    return render_template("advanced_planning.html",project=p,rows=rows,links=links,cycle=cycle,finish=finish,scenarios=scenarios)

@app.post("/projects/<int:project_id>/dependencies/new")
@login_required
def advanced_dependency_new(project_id):
    project_or_404(project_id,write=True); a=int(request.form["predecessor_id"]); b=int(request.form["successor_id"])
    if a==b: abort(400,"En aktivitet kan inte bero på sig själv")
    with db() as conn:
        conn.execute("INSERT INTO task_links(project_id,predecessor_id,successor_id,link_type,lag_days) VALUES(?,?,?,?,?)",(project_id,a,b,request.form.get("link_type","FS"),int(request.form.get("lag_days","0") or 0)))
    audit(project_id,"dependency",None,"created",f"{a}->{b}")
    return redirect(url_for("advanced_planning",project_id=project_id))

@app.post("/projects/<int:project_id>/scenario")
@login_required
def planning_scenario_save(project_id):
    project_or_404(project_id,write=True)
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]; links=[dict(r) for r in conn.execute("SELECT * FROM task_links WHERE project_id=?",(project_id,))]
        conn.execute("INSERT INTO planning_scenarios(project_id,name,snapshot_json,created_by,created_at) VALUES(?,?,?,?,?)",(project_id,request.form["name"].strip(),json.dumps({"tasks":tasks,"links":links},ensure_ascii=False),current_user()["id"],datetime.now().isoformat(timespec="seconds")))
    return redirect(url_for("advanced_planning",project_id=project_id))

@app.get("/projects/<int:project_id>/baseline-variance")
@login_required
def baseline_variance(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        current=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,))]
        base=conn.execute("SELECT * FROM baselines WHERE project_id=? ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
    snap=json.loads(base["snapshot_json"]) if base else []
    if isinstance(snap,dict): snap=snap.get("tasks",[])
    bm={str(x.get("id")):x for x in snap if isinstance(x,dict)}
    rows=[]
    for t in current:
        b=bm.get(str(t["id"])) or next((x for x in snap if isinstance(x,dict) and x.get("wbs")==t.get("wbs")),{})
        delay=0
        if b and parse_date(b.get("end_date")) and parse_date(t.get("end_date")): delay=(parse_date(t["end_date"])-parse_date(b["end_date"])).days
        rows.append({"task":t,"baseline":b,"delay":delay})
    return render_template("baseline_variance.html",project=p,baseline=base,rows=rows)


def evaluate_automation_rule(rule):
    cfg=json.loads(rule["config_json"] or "{}")
    pid=rule["project_id"]; trigger=rule["trigger_type"]; action=rule["action_type"]; matched=[]
    with db() as conn:
        if trigger=="task.overdue" and pid:
            for t in conn.execute("SELECT * FROM tasks WHERE project_id=? AND progress<100 AND end_date<>''",(pid,)).fetchall():
                if parse_date(t["end_date"]) and parse_date(t["end_date"])<date.today(): matched.append(dict(t))
        elif trigger=="risk.high" and pid:
            matched=[dict(r) for r in conn.execute("SELECT * FROM risks WHERE project_id=? AND probability*impact>=? AND status<>'Stängd'",(pid,int(cfg.get("score",12)))).fetchall()]
        elif trigger=="milestone.due" and pid:
            days=int(cfg.get("days",7)); until=date.today()+timedelta(days=days)
            for t in conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 AND progress<100",(pid,)).fetchall():
                d=parse_date(t["end_date"]); matched += [dict(t)] if d and date.today()<=d<=until else []
        if matched and action=="notify.project.members" and pid:
            members=conn.execute("SELECT user_id FROM project_members WHERE project_id=?",(pid,)).fetchall()
            msg=cfg.get("message") or f"Automation: {rule['name']} matched {len(matched)} item(s)."
            for m in members: conn.execute("INSERT INTO notifications(user_id,project_id,message,link,created_at) VALUES(?,?,?,?,?)",(m["user_id"],pid,msg,url_for("project_cockpit",project_id=pid),datetime.now().isoformat(timespec="seconds")))
    return matched

def run_automation_rules(project_id=None):
    with db() as conn:
        if project_id: rules=conn.execute("SELECT * FROM automation_rules WHERE active=1 AND (project_id=? OR project_id IS NULL)",(project_id,)).fetchall()
        else: rules=conn.execute("SELECT * FROM automation_rules WHERE active=1").fetchall()
    results=[]
    for rule in rules:
        started=datetime.now().isoformat(timespec="seconds")
        try:
            matched=evaluate_automation_rule(rule); status="Matched" if matched else "No match"; detail=f"{len(matched)} item(s)"
        except Exception as ex: status="Failed"; detail=str(ex)
        with db() as conn: conn.execute("INSERT INTO automation_executions(rule_id,project_id,status,details,started_at,finished_at) VALUES(?,?,?,?,?,?)",(rule["id"],rule["project_id"],status,detail,started,datetime.now().isoformat(timespec="seconds")))
        results.append((rule,status,detail))
    return results

@app.route("/automation-center",methods=["GET","POST"])
@login_required
def automation_center():
    projects=roadmap_accessible_projects()
    if request.method=="POST":
        pid=int(request.form["project_id"]) if request.form.get("project_id") else None
        if pid: project_or_404(pid,manager=True)
        with db() as conn: conn.execute("INSERT INTO automation_rules(project_id,name,trigger_type,action_type,config_json,active) VALUES(?,?,?,?,?,1)",(pid,request.form["name"].strip(),request.form["trigger_type"],request.form["action_type"],request.form.get("config_json","{}")))
        return redirect(url_for("automation_center"))
    with db() as conn:
        rules=conn.execute("SELECT * FROM automation_rules ORDER BY id DESC").fetchall(); history=conn.execute("SELECT e.*,r.name rule_name FROM automation_executions e LEFT JOIN automation_rules r ON r.id=e.rule_id ORDER BY e.id DESC LIMIT 100").fetchall()
    return render_template("automation_center.html",projects=projects,rules=rules,history=history)

@app.post("/automation/run")
@login_required
@role_required("admin","pm")
def automation_run_now():
    results=run_automation_rules(); flash(f"Automation körd: {len(results)} regel/regler.","success"); return redirect(url_for("automation_center"))

@app.post("/internal/automation/run")
def automation_internal_run():
    token=request.headers.get("X-Automation-Token","")
    expected=os.getenv("PROJECT_PLAN_AUTOMATION_TOKEN","")
    if not expected or not hmac.compare_digest(token,expected): abort(401)
    return jsonify(results=[{"rule":r[0]["name"],"status":r[1],"details":r[2]} for r in run_automation_rules()])


def project_report_payload(project_id):
    s=project_assistant_summary(project_id); p=s["project"]
    with db() as conn:
        milestones=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 ORDER BY end_date",(project_id,))]
        finance=conn.execute("SELECT * FROM project_finance WHERE project_id=?",(project_id,)).fetchone()
        costs=conn.execute("SELECT COALESCE(SUM(actual),0) actual,COALESCE(SUM(planned),0) planned FROM project_costs WHERE project_id=?",(project_id,)).fetchone()
    return {"summary":s,"project":p,"milestones":milestones,"finance":dict(finance) if finance else {},"costs":dict(costs)}

@app.get("/projects/<int:project_id>/report.pdf")
@login_required
def project_report_pdf(project_id):
    payload=project_report_payload(project_id)
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    b=BytesIO(); c=canvas.Canvas(b,pagesize=A4); w,h=A4; y=h-55
    c.setFont("Helvetica-Bold",18); c.drawString(45,y,f"Project Status – {payload['project']['name']}"); y-=30
    c.setFont("Helvetica",10); c.drawString(45,y,f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · Project Planer {APP_VERSION}"); y-=28
    text=c.beginText(45,y); text.setFont("Helvetica",11)
    for line in [payload["summary"]["text"],f"Progress: {payload['summary']['progress']}%",f"Overdue: {len(payload['summary']['overdue'])}",f"High risks: {len(payload['summary']['high'])}"]:
        for part in [line[i:i+95] for i in range(0,len(line),95)]: text.textLine(part)
    c.drawText(text); y=text.getY()-20; c.setFont("Helvetica-Bold",12); c.drawString(45,y,"Milestones"); y-=18; c.setFont("Helvetica",10)
    for m in payload["milestones"][:18]: c.drawString(55,y,f"{m.get('end_date','')}  {m.get('title','')}  {m.get('progress',0)}%"); y-=15
    c.showPage(); c.save(); b.seek(0)
    return send_file(b,mimetype="application/pdf",as_attachment=True,download_name=f"{secure_filename(payload['project']['name'])}-status.pdf")

@app.get("/projects/<int:project_id>/report.pptx")
@login_required
def project_report_pptx(project_id):
    payload=project_report_payload(project_id); from pptx import Presentation
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text=payload["project"]["name"]+" – Project Status"; slide.placeholders[1].text=payload["summary"]["text"]
    slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text="Key metrics"; slide.placeholders[1].text=f"Progress: {payload['summary']['progress']}%\nOverdue: {len(payload['summary']['overdue'])}\nBlocked: {len(payload['summary']['blocked'])}\nHigh risks: {len(payload['summary']['high'])}"
    slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text="Milestones"; slide.placeholders[1].text="\n".join(f"{m.get('end_date','')} – {m.get('title','')} ({m.get('progress',0)}%)" for m in payload["milestones"][:12]) or "No milestones"
    b=BytesIO(); prs.save(b); b.seek(0); return send_file(b,mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",as_attachment=True,download_name=f"{secure_filename(payload['project']['name'])}-status.pptx")

@app.get("/executive-reporting")
@login_required
def executive_reporting():
    rows=[]
    for p in roadmap_accessible_projects():
        s=project_assistant_summary(p["id"]); rag="red" if s["overdue"] or len(s["high"])>=2 else ("amber" if s["blocked"] or s["high"] else "green")
        rows.append({"p":p,"s":s,"rag":rag})
    return render_template("executive_reporting.html",rows=rows)

@app.route("/scheduled-reports",methods=["GET","POST"])
@login_required
def scheduled_reports():
    if request.method=="POST":
        pid=int(request.form["project_id"]) if request.form.get("project_id") else None
        if pid: project_or_404(pid,manager=True)
        with db() as conn: conn.execute("INSERT INTO scheduled_reports(project_id,name,format,cadence,recipient,created_by,created_at) VALUES(?,?,?,?,?,?,?)",(pid,request.form["name"],request.form.get("format","pdf"),request.form.get("cadence","weekly"),request.form.get("recipient",""),current_user()["id"],datetime.now().isoformat(timespec="seconds")))
        return redirect(url_for("scheduled_reports"))
    with db() as conn: rows=conn.execute("SELECT * FROM scheduled_reports ORDER BY id DESC").fetchall()
    return render_template("scheduled_reports.html",rows=rows,projects=roadmap_accessible_projects())


def api_auth_scope(scope="read"):
    def deco(fn):
        @wraps(fn)
        def wrapper(*args,**kwargs):
            auth=request.headers.get("Authorization","")
            if not auth.startswith("Bearer "): return jsonify(error="unauthorized"),401
            raw=auth[7:]; h=hashlib.sha256(raw.encode()).hexdigest()
            with db() as conn: key=conn.execute("SELECT * FROM api_keys WHERE key_hash=? AND active=1",(h,)).fetchone()
            if not key: return jsonify(error="unauthorized"),401
            scopes=set((key["scopes"] or "read").split());
            if scope not in scopes and "admin" not in scopes: return jsonify(error="insufficient_scope"),403
            if key["expires_at"] and parse_date(key["expires_at"]) and parse_date(key["expires_at"])<date.today(): return jsonify(error="key_expired"),401
            request.api_user=key["user_id"]; conn.execute("UPDATE api_keys SET last_used=? WHERE id=?",(datetime.now().isoformat(timespec="seconds"),key["id"]))
            return fn(*args,**kwargs)
        return wrapper
    return deco

@app.route("/api/v2/projects",methods=["GET","POST"])
@api_auth_scope("read")
def api_v2_projects():
    uid=request.api_user
    with db() as conn:
        u=conn.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
        if request.method=="POST":
            scopes=request.headers.get("X-API-Scope-Check","")
            # write authorization is revalidated against token below
            auth=request.headers.get("Authorization","")[7:]; kh=hashlib.sha256(auth.encode()).hexdigest(); key=conn.execute("SELECT * FROM api_keys WHERE key_hash=?",(kh,)).fetchone()
            if "write" not in set((key["scopes"] or "").split()) and "admin" not in set((key["scopes"] or "").split()): return jsonify(error="insufficient_scope"),403
            data=request.get_json(silent=True) or {}; name=(data.get("name") or "").strip()
            if not name:return jsonify(error="name_required"),400
            cur=conn.execute("INSERT INTO projects(name,customer,project_manager,description,created_by,created_at) VALUES(?,?,?,?,?,?)",(name,data.get("customer",""),data.get("project_manager",""),data.get("description",""),uid,datetime.now().isoformat(timespec="seconds")))
            pid=cur.lastrowid; conn.execute("INSERT OR IGNORE INTO project_members(project_id,user_id,project_role,added_at,added_by) VALUES(?,?,?,?,?)",(pid,uid,"pm",datetime.now().isoformat(timespec="seconds"),uid)); return jsonify(id=pid,name=name),201
        if u["role"]=="admin": rows=conn.execute("SELECT * FROM projects ORDER BY id DESC").fetchall()
        else: rows=conn.execute("SELECT p.* FROM projects p JOIN project_members pm ON pm.project_id=p.id WHERE pm.user_id=? ORDER BY p.id DESC",(uid,)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/v2/projects/<int:project_id>/tasks",methods=["GET","POST"])
@api_auth_scope("read")
def api_v2_tasks(project_id):
    uid=request.api_user
    with db() as conn:
        u=conn.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone(); allowed=u["role"]=="admin" or conn.execute("SELECT 1 FROM project_members WHERE project_id=? AND user_id=?",(project_id,uid)).fetchone()
        if not allowed:return jsonify(error="forbidden"),403
        if request.method=="POST":
            auth=request.headers.get("Authorization","")[7:]; key=conn.execute("SELECT * FROM api_keys WHERE key_hash=?",(hashlib.sha256(auth.encode()).hexdigest(),)).fetchone(); scopes=set((key["scopes"] or "").split())
            if "write" not in scopes and "admin" not in scopes:return jsonify(error="insufficient_scope"),403
            d=request.get_json(silent=True) or {}; title=(d.get("title") or "").strip()
            if not title:return jsonify(error="title_required"),400
            cur=conn.execute("INSERT INTO tasks(project_id,wbs,title,owner,start_date,end_date,status,priority,progress) VALUES(?,?,?,?,?,?,?,?,?)",(project_id,d.get("wbs",""),title,d.get("owner",""),d.get("start_date",""),d.get("end_date",""),d.get("status","Ej startad"),d.get("priority","Normal"),int(d.get("progress",0))))
            return jsonify(id=cur.lastrowid,title=title),201
        rows=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/integration-center",methods=["GET","POST"])
@login_required
@role_required("admin")
def integration_center_v35():
    providers=["Azure DevOps","GitHub","Jira","Microsoft Teams","SMTP"]
    if request.method=="POST":
        provider=request.form["provider"]
        with db() as conn: conn.execute("INSERT INTO integration_settings(provider,enabled,base_url,config_json,updated_at) VALUES(?,?,?,?,?) ON CONFLICT(provider) DO UPDATE SET enabled=excluded.enabled,base_url=excluded.base_url,config_json=excluded.config_json,updated_at=excluded.updated_at",(provider,1 if request.form.get("enabled") else 0,request.form.get("base_url",""),request.form.get("config_json","{}"),datetime.now().isoformat(timespec="seconds")))
        return redirect(url_for("integration_center_v35"))
    with db() as conn: settings={r["provider"]:r for r in conn.execute("SELECT * FROM integration_settings").fetchall()}
    return render_template("integration_center_v35.html",providers=providers,settings=settings)


@app.get("/security-center")
@login_required
@role_required("admin")
def security_center():
    with db() as conn:
        logins=conn.execute("SELECT * FROM login_history ORDER BY id DESC LIMIT 100").fetchall(); events=conn.execute("SELECT * FROM security_events ORDER BY id DESC LIMIT 100").fetchall(); sessions=conn.execute("SELECT s.*,u.username,u.display_name FROM user_sessions s JOIN users u ON u.id=s.user_id WHERE s.revoked_at='' ORDER BY s.last_seen DESC").fetchall()
    return render_template("security_center.html",logins=logins,events=events,sessions=sessions)

@app.post("/projects/<int:project_id>/archive")
@login_required
def project_archive(project_id):
    project_or_404(project_id,manager=True)
    with db() as conn: conn.execute("UPDATE projects SET archived_at=? WHERE id=?",(datetime.now().isoformat(timespec="seconds"),project_id))
    audit(project_id,"project",project_id,"archived",""); return redirect(url_for("index"))

@app.post("/projects/<int:project_id>/trash")
@login_required
def project_trash(project_id):
    project_or_404(project_id,manager=True)
    with db() as conn: conn.execute("UPDATE projects SET deleted_at=? WHERE id=?",(datetime.now().isoformat(timespec="seconds"),project_id))
    audit(project_id,"project",project_id,"trashed",""); return redirect(url_for("index"))

@app.get("/admin/trash")
@login_required
@role_required("admin")
def admin_trash():
    with db() as conn: rows=conn.execute("SELECT * FROM projects WHERE deleted_at<>'' ORDER BY deleted_at DESC").fetchall()
    return render_template("trash.html",rows=rows)

@app.post("/admin/trash/<int:project_id>/restore")
@login_required
@role_required("admin")
def admin_trash_restore(project_id):
    with db() as conn: conn.execute("UPDATE projects SET deleted_at='' WHERE id=?",(project_id,))
    return redirect(url_for("admin_trash"))

@app.get("/admin/backup/download-latest")
@login_required
@role_required("admin")
def download_live_database():
    if not DB_PATH.exists(): abort(404)
    stamp=datetime.now().strftime("%Y%m%d-%H%M%S")
    return send_file(DB_PATH,as_attachment=True,download_name=f"projectplan-{stamp}.db")

@app.after_request
def enterprise_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options","nosniff"); response.headers.setdefault("X-Frame-Options","SAMEORIGIN"); response.headers.setdefault("Referrer-Policy","strict-origin-when-cross-origin"); response.headers.setdefault("Permissions-Policy","camera=(), microphone=(), geolocation=()")
    if request.is_secure: response.headers.setdefault("Strict-Transport-Security","max-age=31536000; includeSubDomains")
    return response


@app.route("/projects/<int:project_id>/delivery",methods=["GET","POST"])
@login_required
def delivery_center(project_id):
    p=project_or_404(project_id)
    if request.method=="POST":
        project_or_404(project_id,write=True); kind=request.form["kind"]
        with db() as conn:
            if kind=="environment": conn.execute("INSERT OR IGNORE INTO project_environments(project_id,name,base_url,owner,status,notes) VALUES(?,?,?,?,?,?)",(project_id,request.form["name"],request.form.get("base_url",""),request.form.get("owner",""),request.form.get("status","Ready"),request.form.get("notes","")))
            elif kind=="deliverable": conn.execute("INSERT INTO deliverables(project_id,name,owner,due_date,status,acceptance_criteria) VALUES(?,?,?,?,?,?)",(project_id,request.form["name"],request.form.get("owner",""),request.form.get("due_date",""),request.form.get("status","Planned"),request.form.get("acceptance_criteria","")))
            elif kind=="interface": conn.execute("INSERT INTO interface_register(project_id,name,source_system,target_system,protocol,owner,status,notes) VALUES(?,?,?,?,?,?,?,?)",(project_id,request.form["name"],request.form.get("source_system",""),request.form.get("target_system",""),request.form.get("protocol",""),request.form.get("owner",""),request.form.get("status","Design"),request.form.get("notes","")))
            elif kind=="test": conn.execute("INSERT INTO test_cycles(project_id,name,test_type,start_date,end_date,status) VALUES(?,?,?,?,?,?)",(project_id,request.form["name"],request.form.get("test_type","SIT"),request.form.get("start_date",""),request.form.get("end_date",""),request.form.get("status","Planned")))
            elif kind=="requirement": conn.execute("INSERT INTO requirements_traceability(project_id,requirement_id,requirement,status) VALUES(?,?,?,?)",(project_id,request.form["requirement_id"],request.form["requirement"],request.form.get("status","Open")))
            elif kind=="cutover": conn.execute("INSERT INTO cutover_items(project_id,sequence_no,title,owner,planned_at,status,rollback_step) VALUES(?,?,?,?,?,?,?)",(project_id,int(request.form.get("sequence_no","0") or 0),request.form["title"],request.form.get("owner",""),request.form.get("planned_at",""),request.form.get("status","Planned"),request.form.get("rollback_step","")))
        audit(project_id,"delivery",None,"created",kind); return redirect(url_for("delivery_center",project_id=project_id))
    with db() as conn:
        envs=conn.execute("SELECT * FROM project_environments WHERE project_id=? ORDER BY name",(project_id,)).fetchall(); deliverables=conn.execute("SELECT * FROM deliverables WHERE project_id=? ORDER BY due_date",(project_id,)).fetchall(); interfaces=conn.execute("SELECT * FROM interface_register WHERE project_id=? ORDER BY name",(project_id,)).fetchall(); tests=conn.execute("SELECT * FROM test_cycles WHERE project_id=? ORDER BY start_date",(project_id,)).fetchall(); reqs=conn.execute("SELECT * FROM requirements_traceability WHERE project_id=? ORDER BY requirement_id",(project_id,)).fetchall(); cutover=conn.execute("SELECT * FROM cutover_items WHERE project_id=? ORDER BY sequence_no,id",(project_id,)).fetchall()
    return render_template("delivery_center.html",project=p,envs=envs,deliverables=deliverables,interfaces=interfaces,tests=tests,reqs=reqs,cutover=cutover)


def health_snapshot(project_id):
    s=project_assistant_summary(project_id); score=max(0,min(100,100-len(s["overdue"])*8-len(s["blocked"])*7-len(s["high"])*10-len(s["changes"])*3))
    return {"progress":s["progress"],"overdue":len(s["overdue"]),"blocked":len(s["blocked"]),"high":len(s["high"]),"changes":len(s["changes"]),"score":score,"summary":s}

def assistant_answer(project_id,question):
    q=question.lower(); h=health_snapshot(project_id); s=h["summary"]; p=s["project"]
    if any(x in q for x in ["sen","försen","late","delay"]):
        items=s["overdue"]; return (f"{len(items)} aktivitet(er) är försenade: "+", ".join(x["title"] for x in items[:10])) if items else "Inga försenade aktiviteter hittades."
    if any(x in q for x in ["block","go-live","golive"]):
        items=s["blocked"]+s["overdue"]; return ("De viktigaste blockerarna är: "+", ".join(dict.fromkeys(x["title"] for x in items[:10]))) if items else "Inga direkta blockerare identifierades."
    if "risk" in q:
        items=s["high"]; return (f"{len(items)} höga risker: "+", ".join(x["title"] for x in items[:10])) if items else "Inga höga risker identifierades."
    if any(x in q for x in ["ändrat","changed","förra veckan","last week"]):
        with db() as conn: changes=conn.execute("SELECT * FROM audit_log WHERE project_id=? AND created_at>=? ORDER BY id DESC LIMIT 30",(project_id,(datetime.now()-timedelta(days=7)).isoformat(timespec="seconds"))).fetchall()
        return f"{len(changes)} ändringar senaste 7 dagarna. "+"; ".join(f"{x['entity_type']} {x['action']}" for x in changes[:12])
    if any(x in q for x in ["hälsa","status","health"]): return f"{p['name']} har health score {h['score']}/100 och progress {h['progress']}%. {s['text']}"
    return s["text"]+" Fråga gärna om risker, förseningar, blockerare, status eller vad som ändrats senaste veckan."

@app.route("/intelligence",methods=["GET","POST"])
@login_required
def intelligence_center():
    projects=roadmap_accessible_projects(); pid=request.values.get("project_id",type=int) or (projects[0]["id"] if projects else None); answer=None; question=""
    if pid: project_or_404(pid)
    if request.method=="POST" and pid:
        question=request.form.get("question","").strip(); answer=assistant_answer(pid,question)
        with db() as conn: conn.execute("INSERT INTO assistant_queries(user_id,project_id,question,answer,created_at) VALUES(?,?,?,?,?)",(current_user()["id"],pid,question,answer,datetime.now().isoformat(timespec="seconds")))
    hist=[]
    if pid:
        with db() as conn: hist=conn.execute("SELECT * FROM assistant_queries WHERE user_id=? AND project_id=? ORDER BY id DESC LIMIT 10",(current_user()["id"],pid)).fetchall()
    return render_template("intelligence.html",projects=projects,project_id=pid,answer=answer,question=question,history=hist)

@app.post("/projects/<int:project_id>/health-snapshot")
@login_required
def create_health_snapshot(project_id):
    project_or_404(project_id,write=True); h=health_snapshot(project_id)
    with db() as conn: conn.execute("INSERT INTO project_health_snapshots(project_id,snapshot_date,progress,overdue_count,blocked_count,high_risk_count,open_change_count,health_score,details_json) VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(project_id,snapshot_date) DO UPDATE SET progress=excluded.progress,overdue_count=excluded.overdue_count,blocked_count=excluded.blocked_count,high_risk_count=excluded.high_risk_count,open_change_count=excluded.open_change_count,health_score=excluded.health_score,details_json=excluded.details_json",(project_id,date.today().isoformat(),h["progress"],h["overdue"],h["blocked"],h["high"],h["changes"],h["score"],json.dumps({"version":APP_VERSION})))
    return redirect(url_for("intelligence_center",project_id=project_id))

@app.get("/projects/<int:project_id>/health-trend")
@login_required
def health_trend(project_id):
    p=project_or_404(project_id)
    with db() as conn: rows=conn.execute("SELECT * FROM project_health_snapshots WHERE project_id=? ORDER BY snapshot_date",(project_id,)).fetchall()
    return render_template("health_trend.html",project=p,rows=rows)

@app.post("/projects/<int:project_id>/meeting-to-actions")
@login_required
def meeting_to_actions(project_id):
    project_or_404(project_id,write=True); notes=request.form.get("notes",""); created=0
    with db() as conn:
        for line in notes.splitlines():
            raw=line.strip()
            if raw.upper().startswith("ACTION:"):
                body=raw.split(":",1)[1].strip(); owner=""
                if "@" in body:
                    body,owner=body.rsplit("@",1); body=body.strip(); owner=owner.strip()
                if body: conn.execute("INSERT INTO action_items(project_id,title,owner,status) VALUES(?,?,?,'Open')",(project_id,body,owner)); created+=1
            elif raw.upper().startswith("DECISION:"):
                body=raw.split(":",1)[1].strip()
                if body: conn.execute("INSERT INTO decisions(project_id,title,decision,decided_by,decision_date,owner,decided_at) VALUES(?,?,?,?,?,?,?)",(project_id,body,body,current_user()["display_name"],date.today().isoformat(),current_user()["display_name"],datetime.now().isoformat(timespec="seconds"))); created+=1
    flash(f"Skapade {created} actions/beslut från mötesanteckningar.","success"); return redirect(url_for("collaboration",project_id=project_id))


def stability_snapshot():
    import shutil as _shutil
    results={}
    try:
        with db() as conn:
            integrity=conn.execute("PRAGMA integrity_check").fetchone()[0]
            results["database"]={"ok": integrity=="ok","detail":integrity}
            results["projects"]={"ok":True,"detail":conn.execute("SELECT COUNT(*) FROM projects").fetchone()[0]}
    except Exception as ex:
        results["database"]={"ok":False,"detail":str(ex)}
    try:
        usage=_shutil.disk_usage(DATA_DIR)
        results["disk"]={"ok":usage.free > 100*1024*1024,"detail":{"free_mb":round(usage.free/1024/1024),"total_mb":round(usage.total/1024/1024)}}
    except Exception as ex:
        results["disk"]={"ok":False,"detail":str(ex)}
    results["version"]={"ok":True,"detail":APP_VERSION}
    return results



@app.get("/health/live")
def health_live_v701():
    return jsonify(status="alive",version=APP_VERSION),200

@app.get("/health/ready")
def health_ready():
    snap=stability_snapshot()
    ok=all(v.get("ok") for k,v in snap.items() if k!="version")
    return jsonify(status="ready" if ok else "degraded",checks=snap), (200 if ok else 503)

@app.get("/admin/quality")
@login_required
@role_required("admin")
def admin_quality():
    snap=stability_snapshot()
    with db() as conn:
        migrations=conn.execute("SELECT * FROM schema_migrations ORDER BY applied_at DESC").fetchall()
    return render_template("quality_center.html",snap=snap,migrations=migrations)


@app.get("/home")
@login_required
def home_v41():
    u=current_user(); today=date.today(); projects=roadmap_accessible_projects()
    ids=[p["id"] for p in projects]
    with db() as conn:
        if ids:
            ph=",".join("?" for _ in ids)
            tasks=[dict(r) for r in conn.execute(f"SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id WHERE t.project_id IN ({ph}) AND t.progress<100 ORDER BY t.end_date LIMIT 30",ids)]
            approvals=conn.execute(f"SELECT COUNT(*) c FROM approvals WHERE project_id IN ({ph}) AND status='Pending'",ids).fetchone()["c"]
            high=conn.execute(f"SELECT COUNT(*) c FROM risks WHERE project_id IN ({ph}) AND status<>'Stängd' AND probability*impact>=12",ids).fetchone()["c"]
        else: tasks=[]; approvals=0; high=0
    overdue=[t for t in tasks if parse_date(t.get("end_date")) and parse_date(t["end_date"])<today]
    dueweek=[t for t in tasks if parse_date(t.get("end_date")) and today<=parse_date(t["end_date"])<=today+timedelta(days=7)]
    return render_template("home_v41.html",projects=projects,tasks=tasks,overdue=overdue,dueweek=dueweek,approvals=approvals,high=high)

@app.get("/command")
@login_required
def command_palette():
    q=(request.args.get("q") or "").strip()
    projects=[]; tasks=[]
    if q:
        like=f"%{q}%"
        allowed=roadmap_accessible_projects(); ids=[p["id"] for p in allowed]
        with db() as conn:
            projects=[p for p in allowed if q.lower() in (p["name"] or "").lower() or q.lower() in (p["customer"] or "").lower()]
            if ids:
                ph=",".join("?" for _ in ids)
                tasks=conn.execute(f"SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id WHERE t.project_id IN ({ph}) AND (t.title LIKE ? OR t.wbs LIKE ? OR t.owner LIKE ?) ORDER BY t.end_date LIMIT 50",(*ids,like,like,like)).fetchall()
    return render_template("command_palette.html",q=q,projects=projects,tasks=tasks)


@app.get("/projects/<int:project_id>/workspace")
@app.get("/projects/<int:project_id>/workspace/<tab>")
@login_required
def project_workspace(project_id,tab="overview"):
    p=project_or_404(project_id)
    allowed={"overview","plan","tasks","risks","resources","finance","documents","reports","meetings"}
    if tab not in allowed:
        abort(404)

    h=project_visual_health(project_id)
    today=date.today()

    with db() as conn:
        tasks=conn.execute("""
            SELECT * FROM tasks
            WHERE project_id=? AND deleted_at IS NULL
            ORDER BY COALESCE(sort_order,999999), wbs, id
        """,(project_id,)).fetchall()

        milestones=conn.execute("""
            SELECT * FROM tasks
            WHERE project_id=? AND milestone=1 AND deleted_at IS NULL
            ORDER BY CASE WHEN end_date='' THEN 1 ELSE 0 END,end_date,id
        """,(project_id,)).fetchall()

        risks=conn.execute("""
            SELECT * FROM risks
            WHERE project_id=? AND status<>'Stängd'
            ORDER BY probability*impact DESC,id DESC
        """,(project_id,)).fetchall()

        changes=conn.execute("""
            SELECT * FROM change_requests
            WHERE project_id=?
            ORDER BY id DESC
            LIMIT 30
        """,(project_id,)).fetchall()

        docs=conn.execute("""
            SELECT * FROM documents
            WHERE project_id=?
            ORDER BY updated_at DESC,id DESC
        """,(project_id,)).fetchall()

        meetings=conn.execute("""
            SELECT * FROM meetings
            WHERE project_id=?
            ORDER BY meeting_date DESC,id DESC
        """,(project_id,)).fetchall()

        actions=conn.execute("""
            SELECT * FROM action_items
            WHERE project_id=?
            ORDER BY CASE WHEN status IN ('Done','Closed','Stängd') THEN 1 ELSE 0 END,
                     CASE WHEN due_date='' THEN 1 ELSE 0 END,due_date,id
        """,(project_id,)).fetchall()

        decisions=conn.execute("""
            SELECT * FROM decisions
            WHERE project_id=?
            ORDER BY COALESCE(decided_at,decision_date,'') DESC,id DESC
            LIMIT 30
        """,(project_id,)).fetchall()

        audit_rows=conn.execute("""
            SELECT * FROM audit_log
            WHERE project_id=?
            ORDER BY id DESC
            LIMIT 30
        """,(project_id,)).fetchall()

        costs=conn.execute("""
            SELECT * FROM project_costs
            WHERE project_id=?
            ORDER BY cost_date DESC,id DESC
        """,(project_id,)).fetchall()

        allocations=conn.execute("""
            SELECT * FROM resource_allocations
            WHERE project_id=?
            ORDER BY resource_name,week_start
        """,(project_id,)).fetchall()

        reports=conn.execute("""
            SELECT * FROM status_reports
            WHERE project_id=?
            ORDER BY report_date DESC,id DESC
            LIMIT 20
        """,(project_id,)).fetchall()

        benefits=conn.execute("""
            SELECT * FROM project_benefits
            WHERE project_id=?
            ORDER BY id DESC
        """,(project_id,)).fetchall()

        incoming_deps=conn.execute("""
            SELECT d.*,p.name AS predecessor_name
            FROM cross_project_dependencies d
            JOIN projects p ON p.id=d.predecessor_project_id
            WHERE d.successor_project_id=? AND d.status='Active'
            ORDER BY d.id DESC
        """,(project_id,)).fetchall()

        outgoing_deps=conn.execute("""
            SELECT d.*,p.name AS successor_name
            FROM cross_project_dependencies d
            JOIN projects p ON p.id=d.successor_project_id
            WHERE d.predecessor_project_id=? AND d.status='Active'
            ORDER BY d.id DESC
        """,(project_id,)).fetchall()

    open_tasks=[t for t in tasks if int(t["progress"] or 0)<100]
    overdue_tasks=[
        t for t in open_tasks
        if parse_date(t["end_date"]) and parse_date(t["end_date"]) < today
    ]
    blocked_tasks=[
        t for t in open_tasks
        if (t["status"] or "").strip().lower() in ("blocked","blockerad")
    ]
    high_risks=[r for r in risks if int(r["probability"] or 0)*int(r["impact"] or 0)>=15]
    pending_changes=[c for c in changes if (c["status"] or "") in ("Proposed","Submitted","Pending")]

    budget=sum(float(c["planned"] or 0) for c in costs)
    actual=sum(float(c["actual"] or 0) for c in costs)
    variance=actual-budget
    budget_pct=round((actual/budget)*100) if budget else 0

    # Resource summary: peak allocation by resource across weeks.
    resource_summary={}
    for a in allocations:
        name=(a["resource_name"] or "Unnamed").strip()
        item=resource_summary.setdefault(name,{"name":name,"peak":0,"hours":0.0})
        item["peak"]=max(item["peak"],int(a["allocation_pct"] or 0))
        item["hours"]+=float(a["planned_hours"] or 0)
    resource_summary=sorted(resource_summary.values(),key=lambda x:(-x["peak"],x["name"]))[:10]

    # Unified attention queue: one place for PM action.
    attention=[]
    for t in overdue_tasks[:8]:
        attention.append({
            "severity":"red" if (today-parse_date(t["end_date"])).days>=7 else "amber",
            "kind":"Task",
            "title":t["title"],
            "detail":f"Försenad sedan {t['end_date']}",
            "url":url_for("project_workspace",project_id=project_id,tab="tasks")
        })
    for t in blocked_tasks[:5]:
        attention.append({
            "severity":"red","kind":"Blocker",
            "title":t["title"],"detail":"Blockerad aktivitet",
            "url":url_for("project_workspace",project_id=project_id,tab="tasks")
        })
    for r in high_risks[:5]:
        attention.append({
            "severity":"red","kind":"Risk",
            "title":r["title"],
            "detail":f"Risk score {int(r['probability'] or 0)*int(r['impact'] or 0)}",
            "url":url_for("project_workspace",project_id=project_id,tab="risks")
        })
    for c in pending_changes[:5]:
        attention.append({
            "severity":"amber","kind":"Change",
            "title":c["title"],"detail":c["status"],
            "url":url_for("project_workspace",project_id=project_id,tab="risks")
        })
    if budget and actual>budget:
        attention.append({
            "severity":"amber" if actual <= budget*1.10 else "red",
            "kind":"Budget","title":"Budget över plan",
            "detail":f"{variance:,.0f} över budget",
            "url":url_for("project_workspace",project_id=project_id,tab="finance")
        })
    for r in resource_summary:
        if r["peak"]>120:
            attention.append({
                "severity":"red","kind":"Resource",
                "title":f"{r['name']} är överallokerad",
                "detail":f"Peak {r['peak']}%",
                "url":url_for("project_workspace",project_id=project_id,tab="resources")
            })
        elif r["peak"]>100:
            attention.append({
                "severity":"amber","kind":"Resource",
                "title":f"{r['name']} har hög belastning",
                "detail":f"Peak {r['peak']}%",
                "url":url_for("project_workspace",project_id=project_id,tab="resources")
            })
    attention=attention[:12]

    next_milestones=[
        m for m in milestones
        if int(m["progress"] or 0)<100 and parse_date(m["end_date"])
    ][:8]

    return render_template(
        "project_workspace_v802.html",
        project=p,tab=tab,health=h,
        tasks=tasks,open_tasks=open_tasks,overdue_tasks=overdue_tasks,blocked_tasks=blocked_tasks,
        milestones=milestones,next_milestones=next_milestones,
        risks=risks,high_risks=high_risks,changes=changes,pending_changes=pending_changes,
        docs=docs,meetings=meetings,actions=actions,decisions=decisions,audit_rows=audit_rows,
        costs=costs,budget=budget,actual=actual,variance=variance,budget_pct=budget_pct,
        allocations=allocations,resource_summary=resource_summary,reports=reports,benefits=benefits,
        incoming_deps=incoming_deps,outgoing_deps=outgoing_deps,attention=attention,today=today
    )


def dependency_diagnostics(project_id):
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]
        links=[dict(r) for r in conn.execute("SELECT * FROM task_links WHERE project_id=?",(project_id,))]
    ids={int(t["id"]) for t in tasks}; problems=[]
    for l in links:
        if int(l["predecessor_id"]) not in ids or int(l["successor_id"]) not in ids:
            problems.append(f"Broken link #{l['id']}")
        if int(l["predecessor_id"])==int(l["successor_id"]):
            problems.append(f"Self dependency on task {l['predecessor_id']}")
    return problems

@app.route("/projects/<int:project_id>/planning-pro",methods=["GET","POST"])
@login_required
def planning_pro_v43(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    if request.method=="POST":
        with db() as conn:
            conn.execute("""INSERT INTO project_calendars(project_id,work_week,hours_per_day,holiday_json)
                            VALUES(?,?,?,?) ON CONFLICT(project_id) DO UPDATE SET work_week=excluded.work_week,hours_per_day=excluded.hours_per_day,holiday_json=excluded.holiday_json""",
                         (project_id,request.form.get("work_week","1,2,3,4,5"),float(request.form.get("hours_per_day","8") or 8),request.form.get("holiday_json","[]")))
        return redirect(url_for("planning_pro_v43",project_id=project_id))
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,)).fetchall()
        links=conn.execute("SELECT * FROM task_links WHERE project_id=? ORDER BY id",(project_id,)).fetchall()
        cal=conn.execute("SELECT * FROM project_calendars WHERE project_id=?",(project_id,)).fetchone()
    return render_template("planning_pro_v43.html",project=p,tasks=tasks,links=links,cal=cal,problems=dependency_diagnostics(project_id))

@app.post("/projects/<int:project_id>/tasks/<int:task_id>/shift")
@login_required
def shift_task_v43(project_id,task_id):
    project_or_404(project_id,write=True)
    delta=int(request.form.get("days","0") or 0)
    with db() as conn:
        t=conn.execute("SELECT * FROM tasks WHERE id=? AND project_id=?",(task_id,project_id)).fetchone()
        if not t: abort(404)
        sd=parse_date(t["start_date"]); ed=parse_date(t["end_date"])
        if sd: sd=sd+timedelta(days=delta)
        if ed: ed=ed+timedelta(days=delta)
        conn.execute("UPDATE tasks SET start_date=?,end_date=? WHERE id=?",(sd.isoformat() if sd else t["start_date"],ed.isoformat() if ed else t["end_date"],task_id))
    audit(project_id,"task",task_id,"shifted",f"{delta} days")
    return redirect(url_for("planning_pro_v43",project_id=project_id))


@app.route("/projects/<int:project_id>/risk-center",methods=["GET"])
@login_required
def risk_center_v44(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        risks=conn.execute("""SELECT r.*,c.response_strategy,c.mitigation,c.contingency,c.residual_probability,c.residual_impact,c.review_date
                              FROM risks r LEFT JOIN risk_controls_v44 c ON c.risk_id=r.id
                              WHERE r.project_id=? ORDER BY r.probability*r.impact DESC,r.id DESC""",(project_id,)).fetchall()
    matrix={(i,j):0 for i in range(1,6) for j in range(1,6)}
    for r in risks:
        prob=max(1,min(5,int(r["probability"] or 1))); impact=max(1,min(5,int(r["impact"] or 1))); matrix[(prob,impact)]+=1
    return render_template("risk_center_v44.html",project=p,risks=risks,matrix=matrix)

@app.post("/projects/<int:project_id>/risks/<int:risk_id>/control")
@login_required
def risk_control_save_v44(project_id,risk_id):
    project_or_404(project_id,write=True)
    with db() as conn:
        r=conn.execute("SELECT id FROM risks WHERE id=? AND project_id=?",(risk_id,project_id)).fetchone()
        if not r: abort(404)
        conn.execute("""INSERT INTO risk_controls_v44(risk_id,response_strategy,mitigation,contingency,residual_probability,residual_impact,review_date)
                        VALUES(?,?,?,?,?,?,?) ON CONFLICT(risk_id) DO UPDATE SET response_strategy=excluded.response_strategy,mitigation=excluded.mitigation,
                        contingency=excluded.contingency,residual_probability=excluded.residual_probability,residual_impact=excluded.residual_impact,review_date=excluded.review_date""",
                     (risk_id,request.form.get("response_strategy",""),request.form.get("mitigation",""),request.form.get("contingency",""),
                      int(request.form.get("residual_probability","1")),int(request.form.get("residual_impact","1")),request.form.get("review_date","")))
    return redirect(url_for("risk_center_v44",project_id=project_id))


def status_report_payload(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]
        risks=[dict(r) for r in conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC",(project_id,))]
        changes=[dict(r) for r in conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC LIMIT 8",(project_id,))]
    progress=round(sum(int(t.get("progress") or 0) for t in tasks)/len(tasks)) if tasks else 0
    overdue=[t for t in tasks if int(t.get("progress") or 0)<100 and parse_date(t.get("end_date")) and parse_date(t["end_date"])<date.today()]
    return p,tasks,risks,changes,progress,overdue

@app.get("/projects/<int:project_id>/executive.pdf")
@login_required
def executive_pdf_v45(project_id):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    import io
    p,tasks,risks,changes,progress,overdue=status_report_payload(project_id)
    buf=io.BytesIO(); c=canvas.Canvas(buf,pagesize=A4); w,h=A4; y=h-55
    c.setFont("Helvetica-Bold",18); c.drawString(45,y,f"Project Status – {p['name']}"); y-=30
    c.setFont("Helvetica",10); c.drawString(45,y,f"Progress: {progress}%   Overdue: {len(overdue)}   Open risks: {len(risks)}"); y-=30
    c.setFont("Helvetica-Bold",12); c.drawString(45,y,"Top risks"); y-=18; c.setFont("Helvetica",9)
    for r in risks[:8]:
        c.drawString(55,y,f"{r['title']} (score {int(r['probability'] or 0)*int(r['impact'] or 0)})"); y-=15
        if y<70: c.showPage(); y=h-55
    c.setFont("Helvetica-Bold",12); c.drawString(45,y,"Overdue activities"); y-=18; c.setFont("Helvetica",9)
    for t in overdue[:12]:
        c.drawString(55,y,f"{t.get('wbs','')} {t['title']} – {t.get('end_date','')}"); y-=15
        if y<70: c.showPage(); y=h-55
    c.save(); buf.seek(0)
    return send_file(buf,mimetype="application/pdf",as_attachment=True,download_name=f"{p['name']}-status.pdf")

@app.get("/projects/<int:project_id>/executive.pptx")
@login_required
def executive_pptx_v45(project_id):
    from pptx import Presentation
    from pptx.util import Inches
    import io
    p,tasks,risks,changes,progress,overdue=status_report_payload(project_id)
    prs=Presentation()
    s=prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text=p["name"]; s.placeholders[1].text=f"Executive Project Status · {date.today().isoformat()}"
    s=prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text="Project health"; s.placeholders[1].text=f"Progress: {progress}%\nOverdue activities: {len(overdue)}\nOpen risks: {len(risks)}\nRecent change requests: {len(changes)}"
    s=prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text="Top risks"; s.placeholders[1].text="\n".join(f"• {r['title']} – score {int(r['probability'] or 0)*int(r['impact'] or 0)}" for r in risks[:8]) or "No open risks"
    s=prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text="Needs attention"; s.placeholders[1].text="\n".join(f"• {t['title']} – due {t.get('end_date','')}" for t in overdue[:10]) or "No overdue activities"
    buf=io.BytesIO(); prs.save(buf); buf.seek(0)
    return send_file(buf,mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",as_attachment=True,download_name=f"{p['name']}-status.pptx")

@app.get("/projects/<int:project_id>/report-pack")
@login_required
def report_pack_v45(project_id):
    p=project_or_404(project_id)
    return render_template("report_pack_v45.html",project=p)


def automation_watch_v46():
    today=date.today(); created=0
    with db() as conn:
        projects=conn.execute("SELECT * FROM projects WHERE COALESCE(archived_at,'')='' AND COALESCE(deleted_at,'')=''").fetchall() if table_has_column(conn,"projects","archived") else conn.execute("SELECT * FROM projects").fetchall()
        for p in projects:
            overdue=conn.execute("SELECT COUNT(*) c FROM tasks WHERE project_id=? AND progress<100 AND end_date<>'' AND end_date<?",(p["id"],today.isoformat())).fetchone()["c"]
            high=conn.execute("SELECT COUNT(*) c FROM risks WHERE project_id=? AND status<>'Stängd' AND probability*impact>=15",(p["id"],)).fetchone()["c"]
            if overdue or high:
                exists=conn.execute("SELECT id FROM automation_queue_v46 WHERE project_id=? AND status='Queued' AND created_at LIKE ?",(p["id"],today.isoformat()+"%")).fetchone()
                if not exists:
                    payload=json.dumps({"overdue":overdue,"high_risks":high})
                    conn.execute("INSERT INTO automation_queue_v46(project_id,payload_json,status,created_at) VALUES(?,?,?,?)",(p["id"],payload,"Queued",datetime.now().isoformat(timespec="seconds")))
                    created+=1
    return created

@app.get("/automation-ops")
@login_required
def automation_ops_v46():
    with db() as conn:
        q=conn.execute("""SELECT q.*,p.name project_name FROM automation_queue_v46 q LEFT JOIN projects p ON p.id=q.project_id ORDER BY q.id DESC LIMIT 100""").fetchall()
        channels=conn.execute("SELECT * FROM notification_channels_v46 ORDER BY id DESC").fetchall()
    return render_template("automation_ops_v46.html",queue=q,channels=channels)

@app.post("/automation-ops/run")
@login_required
def automation_run_v46():
    if current_user()["role"] not in ("admin","pm"): abort(403)
    n=automation_watch_v46(); flash(f"Automation scan klar: {n} nya köposter.","success")
    return redirect(url_for("automation_ops_v46"))

@app.post("/automation-ops/process")
@login_required
@role_required("admin")
def automation_process_v46():
    with db() as conn:
        rows=conn.execute("SELECT * FROM automation_queue_v46 WHERE status='Queued' ORDER BY id LIMIT 50").fetchall()
        for r in rows:
            payload=json.loads(r["payload_json"] or "{}")
            members=conn.execute("SELECT user_id FROM project_members WHERE project_id=?",(r["project_id"],)).fetchall()
            for m in members:
                conn.execute("INSERT INTO notifications(user_id,project_id,message,link,is_read,created_at) VALUES(?,?,?,?,0,?)",
                             (m["user_id"],r["project_id"],f"Project attention needed · Overdue: {payload.get('overdue',0)} · High risks: {payload.get('high_risks',0)}",url_for("project_cockpit",project_id=r["project_id"]),datetime.now().isoformat(timespec="seconds")))
            conn.execute("UPDATE automation_queue_v46 SET status='Done',attempts=attempts+1,processed_at=? WHERE id=?",(datetime.now().isoformat(timespec="seconds"),r["id"]))
    return redirect(url_for("automation_ops_v46"))


def table_has_column(conn,table,column):
    try:
        return column in {r[1] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    except Exception:
        return False


@app.route("/admin/enterprise",methods=["GET","POST"])
@login_required
@role_required("admin")
def enterprise_center_v47():
    with db() as conn:
        if request.method=="POST":
            action=request.form.get("action")
            if action=="role":
                perms=[x.strip() for x in request.form.get("permissions","").split(",") if x.strip()]
                conn.execute("INSERT OR IGNORE INTO enterprise_roles_v47(name,permissions_json,created_at) VALUES(?,?,?)",(request.form["name"].strip(),json.dumps(perms),datetime.now().isoformat(timespec="seconds")))
            elif action=="retention":
                conn.execute("""INSERT INTO retention_policies_v47(data_type,retention_days,enabled) VALUES(?,?,1)
                                ON CONFLICT(data_type) DO UPDATE SET retention_days=excluded.retention_days,enabled=1""",(request.form["data_type"],int(request.form["retention_days"])))
            return redirect(url_for("enterprise_center_v47"))
        roles=conn.execute("SELECT * FROM enterprise_roles_v47 ORDER BY name").fetchall()
        retention=conn.execute("SELECT * FROM retention_policies_v47 ORDER BY data_type").fetchall()
        sessions=conn.execute("SELECT * FROM user_sessions ORDER BY id DESC LIMIT 50").fetchall()
        logins=conn.execute("SELECT * FROM login_history ORDER BY id DESC LIMIT 50").fetchall()
    return render_template("enterprise_center_v47.html",roles=roles,retention=retention,sessions=sessions,logins=logins)

@app.post("/admin/backup/verified")
@login_required
@role_required("admin")
def verified_backup_v47():
    import sqlite3 as _sqlite3, hashlib as _hashlib
    BACKUP_DIR.mkdir(parents=True,exist_ok=True)
    name=f"verified-{datetime.now().strftime('%Y%m%d-%H%M%S')}.db"
    dst=BACKUP_DIR/name
    src=_sqlite3.connect(str(DB_PATH)); out=_sqlite3.connect(str(dst))
    try: src.backup(out)
    finally: out.close(); src.close()
    digest=_hashlib.sha256(dst.read_bytes()).hexdigest()
    return jsonify(status="ok",file=name,sha256=digest,size=dst.stat().st_size)


def intelligence_answer_v50(project_id,question):
    p,tasks,risks,changes,progress,overdue=status_report_payload(project_id)
    q=(question or "").lower()
    blocked=[t for t in tasks if (t.get("status") or "").lower()=="blockerad"]
    milestones=[t for t in tasks if t.get("milestone") and int(t.get("progress") or 0)<100]
    high=[r for r in risks if int(r.get("probability") or 0)*int(r.get("impact") or 0)>=15]
    if "go-live" in q or "golive" in q or "block" in q:
        items=blocked+overdue
        return "Följande kan påverka go-live: " + (", ".join(x["title"] for x in items[:10]) if items else "inga uppenbara blockerare eller försenade aktiviteter.")
    if "risk" in q:
        return "Högsta riskerna: " + (", ".join(f"{r['title']} ({int(r.get('probability') or 0)*int(r.get('impact') or 0)})" for r in high[:8]) if high else "inga risker med score ≥15.")
    if "ändrat" in q or "changed" in q or "vecka" in q:
        cutoff=(datetime.now()-timedelta(days=7)).isoformat(timespec="seconds")
        with db() as conn:
            rows=conn.execute("SELECT * FROM audit_log WHERE project_id=? AND created_at>=? ORDER BY id DESC LIMIT 20",(project_id,cutoff)).fetchall()
        return f"{len(rows)} ändringar registrerade senaste sju dagarna. " + " ".join(f"{r['action']} {r['entity_type']}." for r in rows[:8])
    if "milestone" in q or "milstolp" in q:
        return "Öppna milstolpar: " + (", ".join(f"{t['title']} ({t.get('end_date','')})" for t in milestones[:10]) if milestones else "inga öppna milstolpar.")
    return f"{p['name']} är {progress}% färdigt. {len(overdue)} aktiviteter är försenade, {len(blocked)} blockerade och {len(high)} höga risker finns."

def extract_meeting_actions_v50(text):
    actions=[]; decisions=[]
    for raw in (text or "").splitlines():
        line=raw.strip()
        low=line.lower()
        if low.startswith(("action:","åtgärd:","todo:")):
            actions.append(line.split(":",1)[1].strip())
        elif low.startswith(("decision:","beslut:")):
            decisions.append(line.split(":",1)[1].strip())
    return {"actions":actions,"decisions":decisions}

@app.route("/projects/<int:project_id>/intelligence",methods=["GET","POST"])
@login_required
def intelligence_center_v50(project_id):
    p=project_or_404(project_id,write=request.method=="POST")
    answer=None; extracted=None
    if request.method=="POST":
        mode=request.form.get("mode","ask")
        if mode=="ask":
            answer=intelligence_answer_v50(project_id,request.form.get("question",""))
            with db() as conn:
                conn.execute("INSERT INTO assistant_queries(project_id,user_id,question,answer,created_at) VALUES(?,?,?,?,?)",(project_id,current_user()["id"],request.form.get("question",""),answer,datetime.now().isoformat(timespec="seconds")))
        elif mode=="meeting":
            text=request.form.get("notes",""); extracted=extract_meeting_actions_v50(text)
            with db() as conn:
                conn.execute("INSERT INTO intelligence_notes_v50(project_id,note_type,source_text,generated_json,created_by,created_at) VALUES(?,?,?,?,?,?)",(project_id,"meeting",text,json.dumps(extracted,ensure_ascii=False),current_user()["id"],datetime.now().isoformat(timespec="seconds")))
    with db() as conn:
        history=conn.execute("SELECT * FROM assistant_queries WHERE project_id=? ORDER BY id DESC LIMIT 15",(project_id,)).fetchall()
    return render_template("intelligence_center_v50.html",project=p,answer=answer,extracted=extracted,history=history)

@app.post("/projects/<int:project_id>/intelligence/create-actions")
@login_required
def intelligence_create_actions_v50(project_id):
    project_or_404(project_id,write=True)
    items=json.loads(request.form.get("items","[]"))
    with db() as conn:
        for title in items:
            if str(title).strip():
                conn.execute("INSERT INTO action_items(project_id,title,owner,due_date,status) VALUES(?,?,?,?,?)",(project_id,str(title).strip(),current_user()["display_name"],"","Open"))
    flash("Actions skapade.","success")
    return redirect(url_for("intelligence_center_v50",project_id=project_id))


def project_visual_health(project_id):
    today=date.today()
    with db() as conn:
        tasks=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=?",(project_id,))]
        risks=[dict(r) for r in conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd'",(project_id,))]
        changes=[dict(r) for r in conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC LIMIT 25",(project_id,))]
        milestones=[dict(r) for r in conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 ORDER BY end_date",(project_id,))]
    progress=round(sum(int(t.get("progress") or 0) for t in tasks)/len(tasks)) if tasks else 0
    overdue=[t for t in tasks if int(t.get("progress") or 0)<100 and parse_date(t.get("end_date")) and parse_date(t["end_date"])<today]
    blocked=[t for t in tasks if (t.get("status") or "").strip().lower() in ("blockerad","blocked")]
    high=[r for r in risks if int(r.get("probability") or 0)*int(r.get("impact") or 0)>=15]
    amber=[r for r in risks if 8 <= int(r.get("probability") or 0)*int(r.get("impact") or 0) < 15]
    pending_changes=[c for c in changes if (c.get("status") or "") in ("Submitted","Pending")]
    due_milestones=[m for m in milestones if int(m.get("progress") or 0)<100 and parse_date(m.get("end_date")) and parse_date(m["end_date"])<=today+timedelta(days=14)]

    score=100
    score -= min(35, len(overdue)*6)
    score -= min(30, len(high)*10)
    score -= min(20, len(blocked)*8)
    score -= min(15, len(pending_changes)*3)
    score=max(0,score)
    if high or blocked or len(overdue)>=3 or score<60:
        rag="red"; label="Röd"
    elif amber or overdue or pending_changes or score<80:
        rag="amber"; label="Orange"
    else:
        rag="green"; label="Grön"
    trend="down" if (len(overdue)>=3 or len(high)>=2) else ("up" if (not overdue and not high and progress>=50) else "stable")

    reasons=[]
    if blocked: reasons.append(f"{len(blocked)} blockerad")
    if high: reasons.append(f"{len(high)} hög risk")
    if overdue: reasons.append(f"{len(overdue)} försenad")
    if pending_changes: reasons.append(f"{len(pending_changes)} väntande CR")
    if not reasons: reasons.append("Inga kritiska signaler")
    return {
        "score":score,"rag":rag,"label":label,"trend":trend,"progress":progress,
        "overdue_count":len(overdue),"blocked_count":len(blocked),"high_risk_count":len(high),
        "amber_risk_count":len(amber),"pending_change_count":len(pending_changes),
        "milestones_due_count":len(due_milestones),"reasons":reasons[:3],
        "overdue":overdue[:8],"blocked":blocked[:8],"high_risks":high[:8],"milestones_due":due_milestones[:8]
    }

@app.get("/visual-dashboard")
@login_required
def visual_dashboard_v503():
    rows=roadmap_accessible_projects()
    cards=[]
    for p in rows:
        item=dict(p)
        item["health"]=project_visual_health(p["id"])
        cards.append(item)

    status=(request.args.get("status") or "all").lower()
    sort=(request.args.get("sort") or "attention").lower()
    q=(request.args.get("q") or "").strip().lower()

    all_cards=list(cards)
    if status in ("green","amber","red"):
        cards=[p for p in cards if p["health"]["rag"]==status]
    if q:
        cards=[p for p in cards if q in (p.get("name") or "").lower() or q in (p.get("customer") or "").lower()]

    if sort=="name":
        cards.sort(key=lambda p:(p.get("name") or "").lower())
    elif sort=="progress":
        cards.sort(key=lambda p:p["health"]["progress"],reverse=True)
    elif sort=="health":
        cards.sort(key=lambda p:p["health"]["score"],reverse=True)
    else:
        weight={"red":0,"amber":1,"green":2}
        cards.sort(key=lambda p:(weight.get(p["health"]["rag"],9),p["health"]["score"]))

    summary={
        "green":sum(1 for p in all_cards if p["health"]["rag"]=="green"),
        "amber":sum(1 for p in all_cards if p["health"]["rag"]=="amber"),
        "red":sum(1 for p in all_cards if p["health"]["rag"]=="red"),
        "overdue":sum(p["health"]["overdue_count"] for p in all_cards),
        "high_risks":sum(p["health"]["high_risk_count"] for p in all_cards),
        "blocked":sum(p["health"]["blocked_count"] for p in all_cards),
        "total":len(all_cards)
    }
    return render_template("visual_dashboard_v503.html",projects=cards,summary=summary,status=status,sort=sort,q=q)

@app.get("/projects/<int:project_id>/visual")
@login_required
def project_visual_v503(project_id):
    p=project_or_404(project_id)
    h=project_visual_health(project_id)
    with db() as conn:
        recent=conn.execute("SELECT * FROM audit_log WHERE project_id=? ORDER BY id DESC LIMIT 12",(project_id,)).fetchall()
        milestones=conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 ORDER BY end_date",(project_id,)).fetchall()
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY wbs,id",(project_id,)).fetchall()
    status_counts={}
    for t in tasks:
        key=(t["status"] or "Ej satt")
        status_counts[key]=status_counts.get(key,0)+1
    return render_template("project_visual_v503.html",project=p,health=h,recent=recent,milestones=milestones,status_counts=status_counts)

@app.context_processor
def visual_context_v503():
    return {"today_iso": date.today().isoformat()}


@app.route("/quick-add",methods=["GET","POST"])
@login_required
def quick_add_v51():
    projects=roadmap_accessible_projects()
    if request.method=="POST":
        project_id=int(request.form.get("project_id") or 0)
        project_or_404(project_id,write=True)
        kind=(request.form.get("kind") or "task").lower()
        title=(request.form.get("title") or "").strip()
        owner=(request.form.get("owner") or "").strip()
        due=(request.form.get("due_date") or "").strip()
        if not title:
            flash("Titel krävs.","error"); return redirect(url_for("quick_add_v51"))
        with db() as conn:
            if kind=="risk":
                conn.execute("INSERT INTO risks(project_id,kind,title,owner,due_date,created_at) VALUES(?,?,?,?,?,?)",(project_id,"Risk",title,owner,due,datetime.now().isoformat(timespec="seconds")))
            elif kind=="milestone":
                conn.execute("INSERT INTO tasks(project_id,title,owner,end_date,status,priority,progress,milestone) VALUES(?,?,?,?,?,?,?,1)",(project_id,title,owner,due,"Ej startad","Normal",0))
            else:
                conn.execute("INSERT INTO tasks(project_id,title,owner,end_date,status,priority,progress,milestone) VALUES(?,?,?,?,?,?,?,0)",(project_id,title,owner,due,"Ej startad","Normal",0))
            conn.commit()
        flash("Skapad.","success"); return redirect(url_for("quick_add_v51"))
    return render_template("quick_add_v51.html",projects=projects)

@app.post("/tasks/<int:task_id>/quick-update")
@login_required
def task_quick_update_v51(task_id):
    with db() as conn:
        t=conn.execute("SELECT * FROM tasks WHERE id=?",(task_id,)).fetchone()
        if not t: abort(404)
        project_or_404(t["project_id"],write=True)
        status=request.form.get("status") or t["status"]
        progress=max(0,min(100,int(request.form.get("progress") or t["progress"] or 0)))
        conn.execute("UPDATE tasks SET status=?,progress=? WHERE id=?",(status,progress,task_id)); conn.commit()
    flash("Aktiviteten uppdaterades.","success")
    return redirect(request.referrer or url_for("my_work_2"))


@app.get("/projects/<int:project_id>/workspace-2")
@login_required
def workspace_2_v52(project_id):
    p=project_or_404(project_id); h=project_visual_health(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY end_date,id",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC",(project_id,)).fetchall()
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC LIMIT 10",(project_id,)).fetchall()
    return render_template("workspace_2_v52.html",project=p,health=h,tasks=tasks,risks=risks,changes=changes)


@app.get("/projects/<int:project_id>/board-2")
@login_required
def board_v53(project_id):
    p=project_or_404(project_id)
    with db() as conn: tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY sort_order,id",(project_id,)).fetchall()
    columns=["Ej startad","Pågår","Blockerad","Klar"]
    return render_template("board_v53.html",project=p,tasks=tasks,columns=columns)

@app.post("/projects/<int:project_id>/tasks/<int:task_id>/move")
@login_required
def board_move_v53(project_id,task_id):
    project_or_404(project_id,write=True); status=request.form.get("status") or "Ej startad"
    if status not in ("Ej startad","Pågår","Blockerad","Klar"): abort(400)
    progress=100 if status=="Klar" else (50 if status=="Pågår" else 0)
    with db() as conn:
        conn.execute("UPDATE tasks SET status=?,progress=CASE WHEN ?='Klar' THEN 100 WHEN progress=100 THEN ? ELSE progress END WHERE id=? AND project_id=?",(status,status,progress,task_id,project_id)); conn.commit()
    return redirect(url_for("board_v53",project_id=project_id))

@app.get("/projects/<int:project_id>/timeline-2")
@login_required
def timeline_v53(project_id):
    p=project_or_404(project_id)
    with db() as conn: tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY start_date,end_date,id",(project_id,)).fetchall()
    return render_template("timeline_v53.html",project=p,tasks=tasks)


@app.get("/projects/<int:project_id>/action-center")
@login_required
def action_center_v54(project_id):
    p=project_or_404(project_id); h=project_visual_health(project_id)
    with db() as conn:
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC",(project_id,)).fetchall()
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? AND status IN ('Proposed','Submitted','Pending') ORDER BY id DESC",(project_id,)).fetchall()
    matrix={(prob,impact):[] for prob in range(1,6) for impact in range(1,6)}
    for r in risks: matrix[(int(r["probability"]),int(r["impact"]))].append(r)
    return render_template("action_center_v54.html",project=p,health=h,risks=risks,changes=changes,matrix=matrix)


@app.get("/team-capacity")
@login_required
def team_capacity_v55():
    projects=roadmap_accessible_projects(); pids=[int(p["id"]) for p in projects]
    rows=[]
    if pids:
        marks=",".join("?"*len(pids))
        with db() as conn:
            rows=conn.execute(f"SELECT resource_name,week_start,SUM(allocation_pct) allocation,SUM(planned_hours) hours FROM resource_allocations WHERE project_id IN ({marks}) GROUP BY resource_name,week_start ORDER BY week_start,resource_name",pids).fetchall()
    weeks=sorted({r["week_start"] for r in rows})[:4]
    people=sorted({r["resource_name"] for r in rows})
    cap={(r["resource_name"],r["week_start"]):r for r in rows}
    return render_template("team_capacity_v55.html",weeks=weeks,people=people,cap=cap)


@app.get("/focus")
@login_required
def focus_v56():
    u=current_user(); today=date.today()
    with db() as conn:
        if u["role"]=="admin":
            tasks=conn.execute("SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id WHERE t.progress<100 ORDER BY t.end_date").fetchall()
        else:
            tasks=conn.execute("""SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id
                JOIN project_members pm ON pm.project_id=p.id AND pm.user_id=? WHERE t.progress<100 ORDER BY t.end_date""",(u["id"],)).fetchall()
        notes=conn.execute("SELECT * FROM notifications WHERE user_id=? AND is_read=0 ORDER BY id DESC LIMIT 20",(u["id"],)).fetchall()
    overdue=[]; today_items=[]; week=[]
    for t in tasks:
        d=parse_date(t["end_date"])
        if not d: continue
        if d<today: overdue.append(t)
        elif d==today: today_items.append(t)
        elif d<=today+timedelta(days=7): week.append(t)
    return render_template("focus_v56.html",overdue=overdue,today_items=today_items,week=week,notes=notes)


@app.route("/projects/<int:project_id>/status-report-2",methods=["GET","POST"])
@login_required
def status_report_v57(project_id):
    p=project_or_404(project_id,write=(request.method=="POST")); h=project_visual_health(project_id)
    with db() as conn:
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC LIMIT 5",(project_id,)).fetchall()
        milestones=conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 ORDER BY end_date LIMIT 8",(project_id,)).fetchall()
        changes=conn.execute("SELECT * FROM change_requests WHERE project_id=? ORDER BY id DESC LIMIT 5",(project_id,)).fetchall()
        if request.method=="POST":
            rag={"green":"Green","amber":"Amber","red":"Red"}[h["rag"]]
            conn.execute("""INSERT INTO status_reports(project_id,report_date,overall_rag,scope_rag,schedule_rag,budget_rag,resources_rag,summary,achievements,next_steps,created_by,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",(project_id,date.today().isoformat(),rag,rag,rag,"Green","Green",request.form.get("summary",""),request.form.get("achievements",""),request.form.get("next_steps",""),current_user()["id"],datetime.now().isoformat(timespec="seconds")))
            conn.commit(); flash("Statusrapport sparad.","success")
    suggested=f"Projektet är {h['label'].lower()} med health score {h['score']}/100 och {h['progress']}% progress. {h['overdue_count']} aktiviteter är försenade och {h['high_risk_count']} höga risker är öppna."
    return render_template("status_report_v57.html",project=p,health=h,risks=risks,milestones=milestones,changes=changes,suggested=suggested)


@app.route("/project-wizard",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def project_wizard_v58():
    presets={
      "IT Project":["Kickoff","Requirements","Design","Build","System Test","UAT","Go-live"],
      "Integration Project":["Kickoff","Interface specification","Development","SIT","FAT","SAT","Cutover","Go-live"],
      "LIMS Implementation":["Kickoff","URS","Configuration","Interfaces","Validation","UAT","Training","Go-live"],
      "Upgrade Project":["Assessment","Backup","DEV upgrade","Regression test","UAT","Production upgrade","Hypercare"],
      "General Project":["Kickoff","Planning","Execution","Review","Handover"]
    }
    if request.method=="POST":
        name=(request.form.get("name") or "").strip(); template=request.form.get("template") or "General Project"
        if not name: flash("Projektnamn krävs.","error"); return redirect(url_for("project_wizard_v58"))
        start=request.form.get("start_date") or date.today().isoformat(); end=request.form.get("end_date") or ""
        with db() as conn:
            cur=conn.execute("INSERT INTO projects(name,customer,project_manager,description,start_date,end_date,template_name,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                (name,request.form.get("customer",""),request.form.get("project_manager",""),request.form.get("description",""),start,end,template,current_user()["id"],datetime.now().isoformat(timespec="seconds")))
            pid=cur.lastrowid
            conn.execute("INSERT OR IGNORE INTO project_members(project_id,user_id,project_role,added_at,added_by) VALUES(?,?,?,?,?)",(pid,current_user()["id"],"owner",datetime.now().isoformat(timespec="seconds"),current_user()["id"]))
            for i,title in enumerate(presets.get(template,presets["General Project"]),1):
                conn.execute("INSERT INTO tasks(project_id,wbs,title,status,priority,progress,sort_order,milestone) VALUES(?,?,?,?,?,?,?,?)",(pid,str(i),title,"Ej startad","Normal",0,i,1 if title in ("Go-live","Handover") else 0))
            conn.commit()
        flash("Projektet skapades från template.","success"); return redirect(url_for("workspace_2_v52",project_id=pid))
    return render_template("project_wizard_v58.html",presets=presets)


@app.route("/pm-copilot",methods=["GET","POST"])
@login_required
def pm_copilot_v60():
    projects=roadmap_accessible_projects(); answer=""; question=""; selected=None
    if request.method=="POST":
        question=(request.form.get("question") or "").strip()
        project_id=int(request.form.get("project_id") or 0)
        selected=project_or_404(project_id); h=project_visual_health(project_id)
        with db() as conn:
            risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC LIMIT 3",(project_id,)).fetchall()
            milestones=conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 AND progress<100 ORDER BY end_date LIMIT 3",(project_id,)).fetchall()
        lines=[f"{selected['name']}: {h['label']} · health {h['score']}/100 · progress {h['progress']}%."]
        q=question.lower()
        if "fokus" in q or "focus" in q or "idag" in q:
            if h["blocked"]: lines.append(f"1. Hantera blockerad aktivitet: {h['blocked'][0]['title']}.")
            if h["overdue"]: lines.append(f"2. Uppdatera försenad aktivitet: {h['overdue'][0]['title']} ({h['overdue'][0]['end_date']}).")
            if risks: lines.append(f"3. Följ upp risk: {risks[0]['title']} (score {risks[0]['probability']*risks[0]['impact']}).")
        elif "milstolp" in q or "milestone" in q:
            for i,m in enumerate(milestones,1): lines.append(f"{i}. {m['title']} – {m['end_date']} – {m['progress']}%.")
            if not milestones: lines.append("Inga öppna milstolpar hittades.")
        elif "risk" in q:
            for i,r in enumerate(risks,1): lines.append(f"{i}. {r['title']} – score {r['probability']*r['impact']} – owner {r['owner'] or 'ej satt'}.")
            if not risks: lines.append("Inga öppna risker hittades.")
        else:
            lines.append("Signaler: "+", ".join(h["reasons"])+".")
            lines.append("Fråga exempelvis: Vad ska jag fokusera på idag? Vilka risker är högst? Vilka milstolpar kommer härnäst?")
        answer="\n".join(lines)
        with db() as conn:
            conn.execute("INSERT INTO assistant_queries(user_id,project_id,question,answer,created_at) VALUES(?,?,?,?,?)",(current_user()["id"],project_id,question,answer,datetime.now().isoformat(timespec="seconds"))); conn.commit()
    return render_template("pm_copilot_v60.html",projects=projects,answer=answer,question=question,selected=selected)


@app.get("/task-experience")
@login_required
def task_experience_v61():
    projects=roadmap_accessible_projects()
    ids=[int(p["id"]) for p in projects]
    tasks=[]
    if ids:
        marks=",".join("?" for _ in ids)
        with db() as conn:
            tasks=conn.execute(f"""SELECT t.*,p.name project_name
                FROM tasks t JOIN projects p ON p.id=t.project_id
                WHERE t.project_id IN ({marks}) AND t.progress<100
                  AND COALESCE(t.deleted_at,'')=''
                ORDER BY CASE WHEN t.end_date='' THEN 1 ELSE 0 END,t.end_date,t.id LIMIT 30""",ids).fetchall()
    return render_template("task_experience_v61.html",projects=projects,tasks=tasks)

@app.post("/tasks/<int:task_id>/inline")
@login_required
def task_inline_v61(task_id):
    with db() as conn:
        t=conn.execute("SELECT * FROM tasks WHERE id=?",(task_id,)).fetchone()
        if not t: abort(404)
        project_or_404(t["project_id"])
        status=request.form.get("status") or t["status"]; owner=request.form.get("owner",t["owner"])
        end_date=request.form.get("end_date",t["end_date"]); priority=request.form.get("priority",t["priority"])
        progress=max(0,min(100,int(request.form.get("progress") or t["progress"] or 0)))
        conn.execute("UPDATE tasks SET status=?,owner=?,end_date=?,priority=?,progress=? WHERE id=?",(status,owner,end_date,priority,progress,task_id)); conn.commit()
    flash("Aktiviteten sparades.","success"); return redirect(request.referrer or url_for("task_experience_v61"))


@app.get("/projects/<int:project_id>/planning-engine")
@login_required
def planning_engine_v62(project_id):
    p=project_or_404(project_id)
    with db() as conn:
        tasks=conn.execute("SELECT * FROM tasks WHERE project_id=? ORDER BY start_date,end_date,id",(project_id,)).fetchall()
        links=conn.execute("SELECT * FROM task_links WHERE project_id=? ORDER BY id",(project_id,)).fetchall()
    critical={t["id"] for t in tasks if t["status"]=="Blockerad" or (parse_date(t["end_date"]) and parse_date(t["end_date"])<date.today() and t["progress"]<100)}
    return render_template("planning_engine_v62.html",project=p,tasks=tasks,links=links,critical=critical)

@app.post("/projects/<int:project_id>/tasks/<int:task_id>/schedule")
@login_required
def schedule_task_v62(project_id,task_id):
    project_or_404(project_id,write=True)
    with db() as conn:
        t=conn.execute("SELECT * FROM tasks WHERE id=? AND project_id=?",(task_id,project_id)).fetchone()
        if not t: abort(404)
        conn.execute("UPDATE tasks SET start_date=?,end_date=? WHERE id=?",(request.form.get("start_date",""),request.form.get("end_date",""),task_id)); conn.commit()
    return redirect(url_for("planning_engine_v62",project_id=project_id))


def ensure_intake_v63():
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS intake_requests(id INTEGER PRIMARY KEY AUTOINCREMENT,request_type TEXT NOT NULL,title TEXT NOT NULL,customer TEXT DEFAULT '',requested_by TEXT DEFAULT '',priority TEXT DEFAULT 'Normal',requested_date TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'New',description TEXT DEFAULT '',project_id INTEGER)"""); conn.commit()

@app.route("/intake",methods=["GET","POST"])
@login_required
def intake_v63():
    ensure_intake_v63()
    if request.method=="POST":
        title=(request.form.get("title") or "").strip()
        if not title:
            flash("Titel krävs.","error")
            return redirect(url_for("intake_v63"))
        with db() as conn:
            conn.execute("INSERT INTO intake_requests(request_type,title,customer,requested_by,priority,requested_date,status,description) VALUES(?,?,?,?,?,?,?,?)",
              (request.form.get("request_type","Project Request"),title,request.form.get("customer",""),request.form.get("requested_by",""),request.form.get("priority","Normal"),date.today().isoformat(),"New",request.form.get("description",""))); conn.commit()
        flash("Request mottagen.","success"); return redirect(url_for("intake_v63"))
    with db() as conn: rows=conn.execute("SELECT * FROM intake_requests ORDER BY id DESC").fetchall()
    return render_template("intake_v63.html",rows=rows)


def ensure_rules_v64():
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS visual_automation_rules(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,trigger_name TEXT NOT NULL,condition_name TEXT DEFAULT '',action_name TEXT NOT NULL,enabled INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL)"""); conn.commit()

@app.route("/automation-designer",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def automation_designer_v64():
    ensure_rules_v64()
    if request.method=="POST":
        with db() as conn:
            conn.execute("INSERT INTO visual_automation_rules(name,trigger_name,condition_name,action_name,enabled,created_at) VALUES(?,?,?,?,1,?)",
              (request.form.get("name","Rule"),request.form.get("trigger_name","Task overdue"),request.form.get("condition_name",""),request.form.get("action_name","Notify PM"),datetime.now().isoformat(timespec="seconds"))); conn.commit()
        flash("Automation sparad.","success"); return redirect(url_for("automation_designer_v64"))
    with db() as conn: rules=conn.execute("SELECT * FROM visual_automation_rules ORDER BY id DESC").fetchall()
    return render_template("automation_designer_v64.html",rules=rules)


@app.get("/resource-planning-pro")
@login_required
def resource_planning_v65():
    projects=roadmap_accessible_projects(); pids=[p["id"] for p in projects]
    rows=[]
    if pids:
        q=",".join("?"*len(pids))
        with db() as conn: rows=conn.execute(f"SELECT resource_name,week_start,SUM(allocation_pct) allocation,SUM(planned_hours) hours,GROUP_CONCAT(DISTINCT project_id) projects FROM resource_allocations WHERE project_id IN ({q}) GROUP BY resource_name,week_start ORDER BY week_start,resource_name",pids).fetchall()
    weeks=sorted({r["week_start"] for r in rows})[:8]; people=sorted({r["resource_name"] for r in rows}); cap={(r["resource_name"],r["week_start"]):r for r in rows}
    return render_template("resource_planning_v65.html",weeks=weeks,people=people,cap=cap)


def ensure_goals_v66():
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS strategic_goals(id INTEGER PRIMARY KEY AUTOINCREMENT,title TEXT NOT NULL,description TEXT DEFAULT '',target_value REAL DEFAULT 0,current_value REAL DEFAULT 0,unit TEXT DEFAULT '%',owner TEXT DEFAULT '',due_date TEXT DEFAULT '',created_at TEXT NOT NULL)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS goal_projects(goal_id INTEGER NOT NULL,project_id INTEGER NOT NULL,weight INTEGER DEFAULT 100,PRIMARY KEY(goal_id,project_id))"""); conn.commit()

@app.route("/goals",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def goals_v66():
    ensure_goals_v66()
    if request.method=="POST":
        with db() as conn:
            conn.execute("INSERT INTO strategic_goals(title,description,target_value,current_value,unit,owner,due_date,created_at) VALUES(?,?,?,?,?,?,?,?)",
             (request.form.get("title",""),request.form.get("description",""),float(request.form.get("target_value") or 100),float(request.form.get("current_value") or 0),request.form.get("unit","%"),request.form.get("owner",""),request.form.get("due_date",""),datetime.now().isoformat(timespec="seconds"))); conn.commit()
        return redirect(url_for("goals_v66"))
    with db() as conn: goals=conn.execute("SELECT * FROM strategic_goals ORDER BY id DESC").fetchall()
    return render_template("goals_v66.html",goals=goals)


@app.get("/financial-control-pro")
@login_required
def financial_control_v67():
    projects=roadmap_accessible_projects(); cards=[]
    with db() as conn:
        for p in projects:
            costs=conn.execute("SELECT COALESCE(SUM(planned),0) planned,COALESCE(SUM(actual),0) actual FROM project_costs WHERE project_id=?",(p["id"],)).fetchone()
            budget=float(costs["planned"] or 0); actual=float(costs["actual"] or 0)
            forecast=max(actual,budget); variance=forecast-budget
            cards.append({"project":p,"budget":budget,"actual":actual,"forecast":forecast,"variance":variance,"used":round(actual/budget*100) if budget else 0})
    return render_template("financial_control_v67.html",cards=cards)


@app.get("/stakeholder")
@login_required
def stakeholder_index_v68():
    return render_template("stakeholder_index_v68.html",projects=roadmap_accessible_projects())

@app.get("/stakeholder/<int:project_id>")
@login_required
def stakeholder_project_v68(project_id):
    p=project_or_404(project_id); h=project_visual_health(project_id)
    with db() as conn:
        milestones=conn.execute("SELECT * FROM tasks WHERE project_id=? AND milestone=1 ORDER BY end_date",(project_id,)).fetchall()
        risks=conn.execute("SELECT * FROM risks WHERE project_id=? AND status<>'Stängd' ORDER BY probability*impact DESC LIMIT 5",(project_id,)).fetchall()
        decisions=conn.execute("SELECT * FROM decisions WHERE project_id=? ORDER BY decision_date DESC LIMIT 8",(project_id,)).fetchall()
    return render_template("stakeholder_project_v68.html",project=p,health=h,milestones=milestones,risks=risks,decisions=decisions)


@app.route("/dashboard-designer-2",methods=["GET","POST"])
@login_required
def dashboard_designer_v69():
    u=current_user()
    defaults=["Portfolio Health","Progress","Needs Attention","Risk Summary","Budget","Resource Capacity"]
    with db() as conn:
        row=conn.execute("SELECT layout_json FROM dashboard_preferences WHERE user_id=?",(u["id"],)).fetchone()
        if request.method=="POST":
            widgets=request.form.getlist("widgets") or defaults
            payload=json.dumps({"widgets":widgets})
            conn.execute("INSERT INTO dashboard_preferences(user_id,layout_json) VALUES(?,?) ON CONFLICT(user_id) DO UPDATE SET layout_json=excluded.layout_json",(u["id"],payload)); conn.commit()
            flash("Dashboard sparad.","success"); return redirect(url_for("dashboard_designer_v69"))
    selected=defaults
    if row:
        try: selected=json.loads(row["layout_json"]).get("widgets",defaults)
        except: pass
    return render_template("dashboard_designer_v69.html",selected=selected,all_widgets=defaults)


@app.get("/v7")
@login_required
def enterprise_home_v70():
    u=current_user(); projects=roadmap_accessible_projects(); cards=[]
    for p in projects:
        item=dict(p); item["health"]=project_visual_health(p["id"]); cards.append(item)
    red=sum(1 for p in cards if p["health"]["rag"]=="red"); amber=sum(1 for p in cards if p["health"]["rag"]=="amber"); green=sum(1 for p in cards if p["health"]["rag"]=="green")
    overdue=sum(p["health"]["overdue_count"] for p in cards); risks=sum(p["health"]["high_risk_count"] for p in cards)
    return render_template("enterprise_home_v70.html",user=u,projects=cards,green=green,amber=amber,red=red,overdue=overdue,risks=risks)


@app.route("/form-builder",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def form_builder_v71():
    if request.method=="POST":
        name=(request.form.get("name") or "").strip()
        if not name: flash("Form name required.","error"); return redirect(url_for("form_builder_v71"))
        fields=[x.strip() for x in (request.form.get("fields") or "").splitlines() if x.strip()]
        with db() as conn:
            conn.execute("INSERT INTO form_definitions(name,entity_type,description,fields_json,created_at) VALUES(?,?,?,?,?)",(name,request.form.get("entity_type","Project Request"),request.form.get("description",""),json.dumps(fields),datetime.now().isoformat(timespec="seconds"))); conn.commit()
        return redirect(url_for("form_builder_v71"))
    with db() as conn: forms=conn.execute("SELECT * FROM form_definitions ORDER BY id DESC").fetchall()
    return render_template("form_builder_v71.html",forms=forms)

@app.route("/forms/<int:form_id>/submit",methods=["GET","POST"])
@login_required
def form_submit_v71(form_id):
    with db() as conn: form=conn.execute("SELECT * FROM form_definitions WHERE id=? AND active=1",(form_id,)).fetchone()
    if not form: abort(404)
    fields=json.loads(form["fields_json"] or "[]")
    if request.method=="POST":
        title=(request.form.get("title") or "").strip()
        if not title: flash("Title required.","error"); return redirect(url_for("form_submit_v71",form_id=form_id))
        payload={f:request.form.get(f,"") for f in fields}
        with db() as conn:
            conn.execute("INSERT INTO form_submissions(form_id,title,payload_json,created_by,created_at) VALUES(?,?,?,?,?)",(form_id,title,json.dumps(payload),current_user()["id"],datetime.now().isoformat(timespec="seconds"))); conn.commit()
        flash("Request submitted.","success"); return redirect(url_for("intake_v63"))
    return render_template("form_submit_v71.html",form=form,fields=fields)


@app.route("/business-cases",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def business_cases_v72():
    if request.method=="POST":
        vals=[max(0,min(10,int(request.form.get(k) or 0))) for k in ("strategic_fit","business_value","urgency","regulatory","technical_risk","resource_demand","cost_score","expected_benefit")]
        # Positive value dimensions + inverse risk/demand/cost dimensions.
        score=round((vals[0]+vals[1]+vals[2]+vals[3]+(10-vals[4])+(10-vals[5])+(10-vals[6])+vals[7])/80*100,1)
        with db() as conn:
            conn.execute("""INSERT INTO business_cases(title,strategic_fit,business_value,urgency,regulatory,technical_risk,resource_demand,cost_score,expected_benefit,total_score,estimated_cost,expected_value,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (request.form.get("title","").strip(),vals[0],vals[1],vals[2],vals[3],vals[4],vals[5],vals[6],vals[7],score,float(request.form.get("estimated_cost") or 0),float(request.form.get("expected_value") or 0),"Draft",datetime.now().isoformat(timespec="seconds"))); conn.commit()
        return redirect(url_for("business_cases_v72"))
    with db() as conn: cases=conn.execute("SELECT * FROM business_cases ORDER BY total_score DESC,id DESC").fetchall()
    return render_template("business_cases_v72.html",cases=cases)


@app.route("/portfolio-dependencies",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def portfolio_dependencies_v73():
    projects=roadmap_accessible_projects()
    if request.method=="POST":
        a=int(request.form.get("predecessor_project_id") or 0); b=int(request.form.get("successor_project_id") or 0)
        if not a or not b or a==b: flash("Choose two different projects.","error"); return redirect(url_for("portfolio_dependencies_v73"))
        project_or_404(a); project_or_404(b)
        with db() as conn:
            conn.execute("INSERT INTO cross_project_dependencies(predecessor_project_id,successor_project_id,link_type,lag_days,notes,created_at) VALUES(?,?,?,?,?,?)",(a,b,request.form.get("link_type","FS"),int(request.form.get("lag_days") or 0),request.form.get("notes",""),datetime.now().isoformat(timespec="seconds"))); conn.commit()
        return redirect(url_for("portfolio_dependencies_v73"))
    with db() as conn:
        deps=conn.execute("""SELECT d.*,a.name predecessor,b.name successor FROM cross_project_dependencies d JOIN projects a ON a.id=d.predecessor_project_id JOIN projects b ON b.id=d.successor_project_id ORDER BY d.id DESC""").fetchall()
    return render_template("portfolio_dependencies_v73.html",projects=projects,deps=deps)


@app.route("/resource-directory",methods=["GET","POST"])
@login_required
def resource_directory_v74():
    if request.method=="POST":
        if current_user()["role"] not in ("admin","pm"): abort(403)
        skills=[x.strip() for x in (request.form.get("skills") or "").split(",") if x.strip()]
        with db() as conn:
            conn.execute("INSERT INTO resource_directory(display_name,role_name,location,capacity_pct,skills_json,available_from) VALUES(?,?,?,?,?,?)",(request.form.get("display_name","").strip(),request.form.get("role_name",""),request.form.get("location",""),int(request.form.get("capacity_pct") or 100),json.dumps(skills),request.form.get("available_from",""))); conn.commit()
        return redirect(url_for("resource_directory_v74"))
    with db() as conn: resources=conn.execute("SELECT * FROM resource_directory WHERE active=1 ORDER BY display_name").fetchall()
    cards=[dict(r,skills=json.loads(r["skills_json"] or "[]")) for r in resources]
    return render_template("resource_directory_v74.html",resources=cards)


@app.route("/managed-templates",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def managed_templates_v75():
    if request.method=="POST":
        tasks=[x.strip() for x in (request.form.get("tasks") or "").splitlines() if x.strip()]
        with db() as conn:
            conn.execute("INSERT INTO managed_templates(name,description,tasks_json,created_at) VALUES(?,?,?,?)",(request.form.get("name","").strip(),request.form.get("description",""),json.dumps(tasks),datetime.now().isoformat(timespec="seconds"))); conn.commit()
        return redirect(url_for("managed_templates_v75"))
    with db() as conn: rows=conn.execute("SELECT * FROM managed_templates WHERE active=1 ORDER BY name,version DESC").fetchall()
    return render_template("managed_templates_v75.html",templates=rows)


@app.route("/scenario-planner",methods=["GET","POST"])
@login_required
@role_required("admin","pm")
def scenario_planner_v76():
    projects=roadmap_accessible_projects()
    if request.method=="POST":
        changes=[]
        for p in projects:
            decision=request.form.get(f"decision_{p['id']}","keep")
            shift=int(request.form.get(f"shift_{p['id']}") or 0)
            changes.append({"project_id":p["id"],"decision":decision,"shift_days":shift})
        with db() as conn:
            conn.execute("INSERT INTO portfolio_scenarios(name,description,changes_json,created_by,created_at) VALUES(?,?,?,?,?)",(request.form.get("name","Scenario"),request.form.get("description",""),json.dumps(changes),current_user()["id"],datetime.now().isoformat(timespec="seconds"))); conn.commit()
        flash("Scenario saved. No live project data was changed.","success"); return redirect(url_for("scenario_planner_v76"))
    with db() as conn: scenarios=conn.execute("SELECT * FROM portfolio_scenarios ORDER BY id DESC LIMIT 20").fetchall()
    return render_template("scenario_planner_v76.html",projects=projects,scenarios=scenarios)


@app.route("/benefits",methods=["GET","POST"])
@login_required
def benefits_v77():
    projects=roadmap_accessible_projects()
    if request.method=="POST":
        pid=int(request.form.get("project_id") or 0); project_or_404(pid,write=True)
        with db() as conn:
            conn.execute("INSERT INTO project_benefits(project_id,title,unit,baseline_value,target_value,actual_value,measurement_date,owner,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",(pid,request.form.get("title","").strip(),request.form.get("unit","%"),float(request.form.get("baseline_value") or 0),float(request.form.get("target_value") or 0),float(request.form["actual_value"]) if request.form.get("actual_value") else None,request.form.get("measurement_date",""),request.form.get("owner",""),"Measured" if request.form.get("actual_value") else "Planned",datetime.now().isoformat(timespec="seconds"))); conn.commit()
        return redirect(url_for("benefits_v77"))
    ids=[p["id"] for p in projects]; rows=[]
    if ids:
        marks=",".join("?"*len(ids))
        with db() as conn: rows=conn.execute(f"SELECT b.*,p.name project_name FROM project_benefits b JOIN projects p ON p.id=b.project_id WHERE b.project_id IN ({marks}) ORDER BY b.id DESC",ids).fetchall()
    return render_template("benefits_v77.html",projects=projects,benefits=rows)


@app.get("/action-inbox")
@login_required
def action_inbox_v78():
    u=current_user(); projects=roadmap_accessible_projects(); ids=[p["id"] for p in projects]
    overdue=[]; changes=[]; risks=[]
    if ids:
        marks=",".join("?"*len(ids))
        with db() as conn:
            overdue=conn.execute(f"SELECT t.*,p.name project_name FROM tasks t JOIN projects p ON p.id=t.project_id WHERE t.project_id IN ({marks}) AND t.progress<100 AND t.end_date<>'' AND t.end_date<? ORDER BY t.end_date LIMIT 30",ids+[date.today().isoformat()]).fetchall()
            changes=conn.execute(f"SELECT c.*,p.name project_name FROM change_requests c JOIN projects p ON p.id=c.project_id WHERE c.project_id IN ({marks}) AND c.status IN ('Proposed','Submitted','Pending') ORDER BY c.id DESC LIMIT 20",ids).fetchall()
            risks=conn.execute(f"SELECT r.*,p.name project_name FROM risks r JOIN projects p ON p.id=r.project_id WHERE r.project_id IN ({marks}) AND r.status<>'Stängd' AND r.probability*r.impact>=15 ORDER BY r.probability*r.impact DESC LIMIT 20",ids).fetchall()
        with db() as conn: notes=conn.execute("SELECT * FROM notifications WHERE user_id=? AND is_read=0 ORDER BY id DESC LIMIT 30",(u["id"],)).fetchall()
    else: notes=[]
    return render_template("action_inbox_v78.html",overdue=overdue,changes=changes,risks=risks,notes=notes)


@app.get("/control-engine")
@login_required
def control_engine_v79():
    projects=roadmap_accessible_projects(); signals=[]
    with db() as conn:
        for p in projects:
            h=project_visual_health(p["id"])
            alloc=conn.execute("SELECT COALESCE(MAX(allocation_pct),0) mx FROM resource_allocations WHERE project_id=?",(p["id"],)).fetchone()["mx"]
            deps=conn.execute("SELECT COUNT(*) c FROM cross_project_dependencies WHERE successor_project_id=? AND status='Active'",(p["id"],)).fetchone()["c"]
            costs=conn.execute("SELECT COALESCE(SUM(planned),0) b,COALESCE(SUM(actual),0) a FROM project_costs WHERE project_id=?",(p["id"],)).fetchone()
            score=100-h["overdue_count"]*8-h["high_risk_count"]*10-(15 if alloc>120 else 0)-(5 if deps else 0)-(15 if costs["b"] and costs["a"]>costs["b"] else 0)
            signals.append({"project":p,"health":h,"control_score":max(0,min(100,score)),"allocation":alloc,"dependencies":deps,"budget_over":bool(costs["b"] and costs["a"]>costs["b"])})
    signals.sort(key=lambda x:x["control_score"])
    return render_template("control_engine_v79.html",signals=signals)


@app.get("/v8")
@login_required
def intelligent_ppm_v80():
    projects=roadmap_accessible_projects(); cards=[]
    with db() as conn:
        intake_count=conn.execute("SELECT COUNT(*) c FROM form_submissions WHERE status='New'").fetchone()["c"]
        scenario_count=conn.execute("SELECT COUNT(*) c FROM portfolio_scenarios").fetchone()["c"]
        goals_count=conn.execute("SELECT COUNT(*) c FROM strategic_goals").fetchone()["c"]
        dep_count=conn.execute("SELECT COUNT(*) c FROM cross_project_dependencies WHERE status='Active'").fetchone()["c"]
    for p in projects:
        h=project_visual_health(p["id"]); cards.append({"project":p,"health":h})
    red=sum(1 for x in cards if x["health"]["rag"]=="red"); amber=sum(1 for x in cards if x["health"]["rag"]=="amber")
    return render_template("intelligent_ppm_v80.html",projects=cards,red=red,amber=amber,intake_count=intake_count,scenario_count=scenario_count,goals_count=goals_count,dep_count=dep_count)

if __name__=="__main__":
    init_db()
    app.run(host="0.0.0.0",port=int(os.getenv("PORT","8080")),debug=False)
